# -*- coding: utf-8 -*-
"""📊 «Costeo de órdenes de trabajo» — el tablero gerencial de Servicio Técnico (Daniel, 2026-10-25).

Pedido de Daniel: «cómo vamos semanalmente, cuánto ganamos, cuánto perdimos, cuánto estamos generando de garantías… cuánta
plata está rodando, es rentable el departamento, a nombre de quién la estamos haciendo, cuánto nos están cobrando nuestros
técnicos, los externos están yendo o están siendo reincidentes… con nuestros propios datos ordenados y clasificados no nos
puedan jugar a robarnos un poco de rol o protagonismo». Es EL informe para decir qué tan rentable es el departamento.

MÓDULO PURO (REGLA #9/#18): no importa Flask, no toca la base, no habla con Random ni con Check. Recibe FILAS ya
normalizadas (`armar_fila`) con la cuenta de cada OT que calcula `_ot_finanzas` en app.py, y devuelve números, listas y SVG.
La suma de la plata NO se reinventa: `armar_tablero` recibe `agregar` (en producción `_ot_fin_agregado`, la MISMA suma
de Vida del cliente, el panel de costos por técnico y el Excel) y de ahí salen Cobré, Me cobraron y Neto.

Las cuatro preguntas de la atribución son SEPARADAS e inmutables (cada una sale de su propia fuente guardada):
  · ORIGINÓ    → el ticket (o la OT) con quién y cuándo se creó.
  · AUTORIZÓ   → la autorización aprobada de `mant_ot_autorizaciones` (quién la resolvió, cuándo y su argumento).
  · LLEVA      → el responsable de la OT (REGLA #36).
  · EJECUTÓ    → el técnico asignado (y su empresa, si es externo).
Nadie recibe el crédito de lo que no hizo: cada persona suma por rol, nunca se mezclan los roles.

Confidencialidad (REGLAS #19 y #26): todo lo de este módulo es plata y proveedores; solo el superadmin lo ve (la ruta
lo exige) y cada sección lleva la etiqueta `CLASIFICACION`.
"""
from __future__ import annotations

import io
from datetime import date, datetime, timedelta, timezone

try:
    from zoneinfo import ZoneInfo
    TZ_CL = ZoneInfo("America/Santiago")
except Exception:                                    # pragma: no cover - sin tzdata (mismo respaldo que app.py)
    TZ_CL = timezone(timedelta(hours=-3))

NOMBRE_PANTALLA = "Costeo de órdenes de trabajo"
CLASIFICACION = "Confidencial · solo gerencia"

UMBRAL_BAJO = 10                    # % — igual que _OT_FIN_UMBRAL_BAJO («Margen bajo, menos de 10 %»)
COBERTURAS_REGALADAS = ("garantia", "regalia", "arriendo_leasing", "sin_costo")   # plata que se puso sin cobrar
COBERTURAS_OTRAS = ("cubierta", "interno", "contrato")                           # no son un regalo: van aparte
COBERTURA_ORDEN = COBERTURAS_REGALADAS + COBERTURAS_OTRAS
COBERTURA_TXT = {
    "garantia": "Garantía", "regalia": "Regalía", "arriendo_leasing": "Arriendo o leasing", "sin_costo": "Cortesía",
    "cubierta": "Desinstalación cubierta", "interno": "Trabajo interno", "contrato": "Mantención de contrato",
    "cobra": "Se cobra",
}
APARTE = ("interno", "contrato")    # mismo criterio que _OT_FIN_AGG_APARTE: no son negocio con nadie
MAX_RANGO_DIAS = 400

TIPO_EJECUTOR_TXT = {"interno": "Técnico propio", "elevado": "Técnico elevado", "externo": "Técnico externo",
                     "sin_asignar": "Sin técnico asignado"}
TIPO_EJECUTOR_ORDEN = ("externo", "elevado", "interno", "sin_asignar")

# Reincidencia (definición por defecto, visible en pantalla): una OT correctiva o de garantía sobre el MISMO equipo (o,
# si no hay equipos registrados, el mismo cliente) que se abre dentro de N días después de que un externo cerró una OT.
DIAS_REINCIDENCIA = (30, 60, 90)
DIAS_REINCIDENCIA_DEFECTO = 60
VENTANA_REINCIDENCIA_DIAS = 180     # se miran las OT que los externos cerraron en los últimos 6 meses
TIPOS_CORRECTIVOS = ("correctiva", "visita_correctiva", "garantia")
ESTADOS_FUERA = ("cancelada", "anulada")
ESTADOS_HECHOS = ("completada", "cerrada", "firmada_tecnico", "pendiente_aprobacion")
TASA_REVISAR = 30                   # % de reincidencia desde el que se sugiere revisar (ámbar, nunca rojo)
MIN_CASOS_REVISAR = 3               # con menos OT evaluadas no se saca conclusión

TIPOS_AUT_COBERTURA = ("cobro_cero", "crear_sin_documento", "cerrar_sin_documento", "desinstalacion_cubierta")
TIPOS_AUT_PERDIDA = ("cerrar_con_perdida", "exceder_saldo")
TIPO_AUT_TXT = {
    "cobro_cero": "Cobro en $0", "crear_sin_documento": "Crear sin documento", "cerrar_sin_documento": "Cerrar sin documento",
    "cerrar_con_perdida": "Cerrar con pérdida", "exceder_saldo": "Cobrar sobre el saldo",
    "documento_observado": "Documento observado", "desinstalacion_cubierta": "Desinstalación cubierta",
}
SIN_CLASIFICAR = "Sin clasificar"

MESES_ES = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
            "noviembre", "diciembre")
MESES_CORTO = ("ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic")


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Números y textos
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def _f(x, defecto=0.0):
    """float seguro (Decimal, str, None)."""
    if x is None or x == "":
        return defecto
    try:
        return float(x)
    except (TypeError, ValueError):
        return defecto


def clp(n, signo=False):
    """$1.234.567 (negativo: -$1.234). Sin decimales; None → «—»."""
    if n is None:
        return "—"
    v = int(round(_f(n)))
    s = "$" + f"{abs(v):,}".replace(",", ".")
    if v < 0:
        return "-" + s
    return ("+" + s) if (signo and v > 0) else s


def pct_txt(p, decimales=1):
    """12,5 % (coma decimal); None → «—»."""
    if p is None:
        return "—"
    txt = f"{_f(p):.{decimales}f}".replace(".", ",")
    if decimales and txt.endswith("," + "0" * decimales):
        txt = txt[: -(decimales + 1)]
    return txt + " %"


def fecha_txt(d):
    """dd/mm/aaaa (nunca ISO crudo, REGLA #6). Acepta date o datetime ya en hora Chile."""
    if not d:
        return ""
    try:
        return d.strftime("%d/%m/%Y")
    except Exception:
        return ""


def a_chile(dt):
    """datetime NAIVE en UTC (así guarda MySQL con NOW()) → datetime con hora Chile. None → None."""
    if dt is None:
        return None
    if isinstance(dt, date) and not isinstance(dt, datetime):
        return datetime(dt.year, dt.month, dt.day, tzinfo=TZ_CL)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(TZ_CL)


def fecha_hora_chile_txt(dt):
    """dd/mm/aaaa hh:mm en hora Chile para un datetime naive-UTC; vacío si no hay."""
    c = a_chile(dt)
    return c.strftime("%d/%m/%Y %H:%M") if c else ""


def inicio_dia_utc(d):
    """Primer instante del día `d` en hora Chile, como datetime NAIVE en UTC (para comparar con columnas DATETIME).
    El fin de un día y el inicio del siguiente son EL MISMO instante (rangos contiguos, también en el cambio de hora)."""
    return datetime(d.year, d.month, d.day, tzinfo=TZ_CL).astimezone(timezone.utc).replace(tzinfo=None)


def lunes_de(d):
    return d - timedelta(days=d.weekday())


def semana_de(dt_utc):
    """Lunes (hora Chile) de la semana en que ocurrió `dt_utc` (naive UTC)."""
    c = a_chile(dt_utc)
    return lunes_de(c.date()) if c else None


def mes_de(dt_utc):
    c = a_chile(dt_utc)
    return f"{c.year}-{c.month:02d}" if c else None


def _etiqueta_semana(lunes):
    dom = lunes + timedelta(days=6)
    if lunes.month == dom.month:
        return f"{lunes.day} al {dom.day} de {MESES_ES[dom.month - 1]}"
    return f"{lunes.day} de {MESES_CORTO[lunes.month - 1]} al {dom.day} de {MESES_CORTO[dom.month - 1]}"


def _mes_mas(anio, mes, k):
    i = anio * 12 + (mes - 1) + k
    return i // 12, i % 12 + 1


def _ultimo_dia_mes(anio, mes):
    a2, m2 = _mes_mas(anio, mes, 1)
    return date(a2, m2, 1) - timedelta(days=1)


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Períodos (semana lunes–domingo en hora Chile por defecto)
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def _periodo_base(modo, desde, hasta):
    p = {"modo": modo, "desde": desde, "hasta": hasta, "dias": (hasta - desde).days + 1,
         "ini_utc": inicio_dia_utc(desde), "fin_utc": inicio_dia_utc(hasta + timedelta(days=1))}
    if modo == "semana":
        p["etiqueta"] = "Semana del " + _etiqueta_semana(desde) + f" de {hasta.year}"
    elif modo == "mes":
        p["etiqueta"] = f"{MESES_ES[desde.month - 1].capitalize()} de {desde.year}"
    elif modo == "trimestre":
        p["etiqueta"] = f"Trimestre {MESES_CORTO[desde.month - 1]}–{MESES_CORTO[hasta.month - 1]} de {desde.year}"
    else:
        p["etiqueta"] = f"Del {fecha_txt(desde)} al {fecha_txt(hasta)}"
    return p


def rango_periodo(hoy, modo="semana", desde=None, hasta=None, ancla=None):
    """El período a mirar y el anterior, con sus límites en UTC naive (exclusivo al final).
    modo: semana (lunes–domingo) | mes | trimestre | rango (desde/hasta inclusive). `ancla`: un día cualquiera del período
    (por defecto hoy). Un rango mal dado (falta una punta, al revés, de más de 400 días) se corrige, nunca revienta."""
    modo = (modo or "semana").strip().lower()
    if modo not in ("semana", "mes", "trimestre", "rango"):
        modo = "semana"
    ancla = ancla or hoy
    if modo == "rango":
        if not (desde and hasta):
            modo = "semana"
        else:
            if hasta < desde:
                desde, hasta = hasta, desde
            if (hasta - desde).days + 1 > MAX_RANGO_DIAS:
                desde = hasta - timedelta(days=MAX_RANGO_DIAS - 1)
    if modo == "semana":
        desde = lunes_de(ancla)
        hasta = desde + timedelta(days=6)
        prev_d, prev_h = desde - timedelta(days=7), desde - timedelta(days=1)
    elif modo == "mes":
        desde = date(ancla.year, ancla.month, 1)
        hasta = _ultimo_dia_mes(ancla.year, ancla.month)
        pa, pm = _mes_mas(ancla.year, ancla.month, -1)
        prev_d, prev_h = date(pa, pm, 1), _ultimo_dia_mes(pa, pm)
    elif modo == "trimestre":
        m0 = ((ancla.month - 1) // 3) * 3 + 1
        desde = date(ancla.year, m0, 1)
        hasta = _ultimo_dia_mes(ancla.year, m0 + 2)
        pa, pm = _mes_mas(ancla.year, m0, -3)
        prev_d, prev_h = date(pa, pm, 1), _ultimo_dia_mes(pa, pm + 2)
    else:
        n = (hasta - desde).days + 1
        prev_h = desde - timedelta(days=1)
        prev_d = prev_h - timedelta(days=n - 1)
    p = _periodo_base(modo, desde, hasta)
    p["previo"] = _periodo_base(modo, prev_d, prev_h)
    p["hoy"] = hoy
    p["en_curso"] = desde <= hoy <= hasta
    p["transcurridos"] = max(0, min(p["dias"], (hoy - desde).days + 1)) if desde <= hoy else 0
    p["ancla_previa"] = prev_d if modo != "rango" else None
    sig = hasta + timedelta(days=1)
    p["ancla_siguiente"] = sig if (modo != "rango" and sig <= hoy) else None
    return p


def ventana_carga(periodo, semanas=12, meses=6):
    """(desde_date, hasta_date): lo que hay que leer de la base para el período, el anterior, las últimas `semanas` y los
    últimos `meses` meses. Una sola lectura sirve a la tendencia, a las series y a las garantías por mes."""
    hasta = periodo["hasta"]
    d1 = periodo["previo"]["desde"]
    d2 = lunes_de(hasta) - timedelta(days=7 * (semanas - 1))
    a, m = _mes_mas(hasta.year, hasta.month, -(meses - 1))
    d3 = date(a, m, 1)
    return min(d1, d2, d3), hasta


def ultimas_semanas(hasta, n=12):
    """Lunes de las últimas n semanas que terminan en la semana de `hasta` (la más vieja primero)."""
    fin = lunes_de(hasta)
    return [fin - timedelta(days=7 * i) for i in range(n - 1, -1, -1)]


def ultimos_meses(hasta, n=6):
    """[(año, mes)] de los últimos n meses que terminan en el de `hasta` (el más viejo primero)."""
    return [_mes_mas(hasta.year, hasta.month, -i) for i in range(n - 1, -1, -1)]


def _fecha(s):
    try:
        return datetime.strptime((s or "").strip()[:10], "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def leer_parametros(args, hoy):
    """Lee ?modo=&ancla=&desde=&hasta=&dias= sin confiar en nada: un valor raro vuelve al defecto (semana actual, 60 días).
    `ancla` en el futuro se corrige a hoy. `args`: cualquier mapa con .get (request.args)."""
    modo = ((args.get("modo") or "semana").strip().lower())
    if modo not in ("semana", "mes", "trimestre", "rango"):
        modo = "semana"
    ancla = _fecha(args.get("ancla"))
    if ancla and ancla > hoy:
        ancla = hoy
    desde, hasta = _fecha(args.get("desde")), _fecha(args.get("hasta"))
    try:
        dias = int(args.get("dias") or DIAS_REINCIDENCIA_DEFECTO)
    except (TypeError, ValueError):
        dias = DIAS_REINCIDENCIA_DEFECTO
    if dias not in DIAS_REINCIDENCIA:
        dias = DIAS_REINCIDENCIA_DEFECTO
    return {"modo": modo, "ancla": ancla, "desde": desde, "hasta": hasta, "dias": dias}


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Una fila = una OT cerrada, normalizada
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def tipo_ejecutor(rol, proveedor_tipo, ficha_externa_id, tecnico_user_id):
    """interno | elevado | externo | sin_asignar según el rol del usuario asignado y la ficha de proveedor."""
    r = (rol or "").strip().lower()
    if r.startswith("tecnico_externo") or ficha_externa_id or (proveedor_tipo or "").strip().lower() == "externo":
        return "externo"
    if r == "tecnico_ejecutivo":
        return "elevado"
    if r.startswith("tecnico") or tecnico_user_id:
        return "interno"
    return "sin_asignar"


def armar_ejecutor(v):
    """Quién EJECUTÓ la OT: el técnico asignado y, si es externo, su empresa. La clave agrupa a los externos por empresa
    (una empresa puede mandar a distintas personas) y a los propios por usuario."""
    tipo = tipo_ejecutor(v.get("tecnico_role"), v.get("proveedor_tipo"), v.get("te_id"), v.get("tecnico_user_id"))
    persona = (v.get("tecnico_nombre") or "").strip() or None
    empresa = ((v.get("te_empresa") or "").strip() or (v.get("proveedor_nombre") or "").strip() or None)
    if tipo == "externo":
        nombre = empresa or persona or "Proveedor sin nombre"
        clave = "emp:" + nombre.lower()
    elif tipo == "sin_asignar":
        nombre, clave = "Sin técnico asignado", "sin"
    else:
        nombre = persona or "Técnico sin nombre"
        clave = "usr:" + str(v.get("tecnico_user_id") or nombre.lower())
    return {"clave": clave, "nombre": nombre, "persona": persona, "empresa": empresa if tipo == "externo" else None,
            "tipo": tipo, "tipo_txt": TIPO_EJECUTOR_TXT[tipo], "user_id": v.get("tecnico_user_id")}


def armar_origen(v, nombres=None):
    """Quién ORIGINÓ la OT: el ticket que la pidió (con quién y cuándo se creó) o, si nació sin ticket, quien la creó."""
    nombres = nombres or {}

    def _n(u):
        u = (u or "").strip()
        return nombres.get(u.lower()) or u or None

    if v.get("tk_id"):
        origen = (v.get("tk_origen") or "").strip().lower()
        num = v.get("numero_ticket") or f"#{v.get('tk_id')}"
        if origen == "form":
            quien = "Cliente (formulario web)"
        elif origen == "erp":
            quien = "ERP (documento de venta)"
        else:
            quien = _n(v.get("tk_creado_por"))
        return {"nombre": quien, "at": v.get("tk_creado_at"), "fuente": f"Ticket {num}", "ticket": num,
                "ticket_id": v.get("tk_id")}
    return {"nombre": _n(v.get("created_by")), "at": v.get("created_at"), "fuente": "OT creada directamente (sin ticket)",
            "ticket": None, "ticket_id": None}


def armar_lleva(v):
    """Quién LLEVA la OT (REGLA #36). Las OT anteriores a esa regla no se rellenan: «sin responsable» es un dato."""
    nombre = (v.get("responsable_nombre") or "").strip() or None
    return {"nombre": nombre, "at": v.get("responsable_at"),
            "fuente": "Responsable de la OT" if nombre else "Sin responsable declarado"}


def elegir_autorizacion(auts, tipos):
    """De las autorizaciones APROBADAS de una OT, la que corresponde a lo que se atribuye: la más reciente del primer tipo
    de `tipos` que exista. Una autorización de OTRO asunto (p. ej. un documento observado) NO se le acredita a un $0: queda
    en `todas` para consultarla, pero `principal` es None si nadie autorizó justo esto. «Solicitó» y «autorizó» son personas
    distintas: aquí manda quien RESOLVIÓ."""
    ok = [a for a in (auts or []) if str(a.get("estado") or "aprobada").lower() == "aprobada"
          and (a.get("resuelto_por_nombre") or "").strip()]
    ok.sort(key=lambda a: (a.get("resuelto_at") or datetime.min, a.get("id") or 0))
    elegida = None
    for t in tipos:
        c = [a for a in ok if a.get("tipo") == t]
        if c:
            elegida = c[-1]
            break

    def _fmt(a):
        return {"id": a.get("id"), "tipo": a.get("tipo"), "tipo_txt": TIPO_AUT_TXT.get(a.get("tipo"), a.get("tipo") or ""),
                "nombre": (a.get("resuelto_por_nombre") or "").strip(), "at": a.get("resuelto_at"),
                "solicito": (a.get("solicitado_por_nombre") or "").strip() or None, "solicitado_at": a.get("solicitado_at"),
                "argumento": (a.get("argumento") or "").strip() or None, "motivo": a.get("motivo"),
                "fuente": f"Autorización #{a.get('id')} · {TIPO_AUT_TXT.get(a.get('tipo'), a.get('tipo') or '')}"}

    return {"principal": _fmt(elegida) if elegida else None, "todas": [_fmt(a) for a in ok]}


def armar_autorizo(cobertura, auts, clasif=None):
    """Quién AUTORIZÓ (y su argumento guardado). Primero la autorización aprobada; si la OT la creó directamente un
    superadmin, queda la constancia de la clasificación guardada al crear. Sin ninguna: `sin_autorizacion`=True (dato, no
    error: las OT anteriores a la regla del $0 no la tienen)."""
    tipos = TIPOS_AUT_COBERTURA if cobertura != "cobra" else TIPOS_AUT_PERDIDA
    r = elegir_autorizacion(auts, tipos)
    p = r["principal"]
    if p:
        return {"nombre": p["nombre"], "at": p["at"], "at_txt": fecha_hora_chile_txt(p["at"]), "argumento": p["argumento"],
                "fuente": p["fuente"], "tipo_txt": p["tipo_txt"], "solicito": p["solicito"], "todas": r["todas"],
                "sin_autorizacion": False}
    c = clasif or {}
    quien = (c.get("autorizado_por") or "").strip()
    if quien:
        return {"nombre": quien, "at": None, "at_txt": (c.get("autorizado_at") or ""),
                "argumento": (c.get("argumento") or "").strip() or None,
                "fuente": "Clasificación guardada al crear la OT", "tipo_txt": c.get("texto") or "",
                "solicito": (c.get("pedido_por") or None), "todas": [], "sin_autorizacion": False}
    return {"nombre": None, "at": None, "at_txt": "", "argumento": None, "fuente": "Sin autorización registrada", "tipo_txt": "",
            "solicito": None, "todas": [], "sin_autorizacion": True}


def respaldo_de(codigos, cobra):
    """Con qué documento se respalda lo cobrado, mirando TODAS las marcas de la clasificación (no solo la principal):
    factura (o boleta) · nota de venta/cotización/guía (falta la factura) · sin documento (autorizado) · otro · na (la OT
    no cobra). Una factura siempre gana: si el documento es una factura, lo cobrado tiene respaldo aunque además esté
    marcado «documento de más de 3 meses»."""
    if not cobra:
        return "na"
    cs = set(codigos or ())
    if "factura" in cs:
        return "factura"
    if cs & {"nota_venta", "cotizacion", "otro_documento"}:
        return "nota_venta"
    if "sin_documento" in cs:
        return "sin_documento"
    return "otro"


def clase_neto(f):
    """En qué parte de la cuenta cae una OT (espejo de _ot_fin_agregado, que es quien suma):
    aparte (interno/contrato) · sin_medir (falta un costo) · ganada · perdida (se cobra y queda < 0) · cubierta ·
    cobertura (garantía, regalía, arriendo/leasing, cortesía)."""
    if f["cobertura"] in APARTE:
        return "aparte"
    if not f["mide"]:
        return "sin_medir"
    if f["cobra"]:
        return "perdida" if f["queda"] < 0 else "ganada"
    return "cubierta" if f["cobertura"] == "cubierta" else "cobertura"


def armar_fila(v, fin, ctx=None):
    """Una OT cerrada lista para el tablero. `v`: la fila de mant_visitas con sus joins; `fin`: _ot_finanzas(v, repuestos);
    `ctx`: lo que no está en la fila — nombres de usuario, autorizaciones aprobadas, clasificación, equipos, intención/área."""
    ctx = ctx or {}
    cob = fin.get("cobertura") or "cobra"
    me = fin.get("me_cobraron") or {}
    cobre_d = fin.get("cobre") or {}
    queda_d = fin.get("queda") or {}
    val = (fin.get("valorizado") or {}).get("monto")
    cl = ctx.get("clasif") or {}
    ejecutor = armar_ejecutor(v)
    cerr = v.get("cerrada_at")
    f = {
        "id": v.get("id"), "numero": v.get("numero_ot") or f"VS-{int(v.get('id') or 0):05d}",
        "titulo": (v.get("titulo") or "").strip(), "estado": (v.get("estado") or "").lower(),
        "cliente": (v.get("cliente") or "").strip() or "Trabajo interno", "cliente_id": v.get("cliente_id"),
        "tipo": (v.get("tipo") or "").lower(), "tipo_txt": ctx.get("tipo_txt") or (v.get("tipo") or "").replace("_", " ").capitalize(),
        "centro_costo": (v.get("centro_costo") or "").strip().lower() or None,
        "centro_txt": ctx.get("centro_txt") or ((v.get("centro_costo") or "").strip() or "Sin centro de costo"),
        "cerrada_at": cerr, "semana": semana_de(cerr), "mes": mes_de(cerr),
        "cobertura": cob, "cobertura_txt": COBERTURA_TXT.get(cob, cob), "cobra": bool(fin.get("cobra")),
        "mide": bool(queda_d.get("mostrar")), "falta_tecnico": bool(me.get("falta_tecnico")),
        "cobre": _f(cobre_d.get("total")), "me_total": _f(me.get("total")), "a_pagar": _f(fin.get("a_pagar_proveedor")),
        "repuestos": _f(me.get("repuestos")), "queda": _f(queda_d.get("total")), "pct": queda_d.get("pct"),
        "valorizado": (_f(val) if val is not None else None),
        "clasif_codigo": cl.get("codigo"), "clasif_texto": cl.get("texto") or "",
        "respaldo": respaldo_de(cl.get("codigos") or ([cl["codigo"]] if cl.get("codigo") else []), bool(fin.get("cobra"))),
        "ejecutor": ejecutor, "origino": armar_origen(v, ctx.get("nombres")), "lleva": armar_lleva(v),
        "autorizo": armar_autorizo(cob, ctx.get("autorizaciones"), cl),
        "intencion": (ctx.get("intencion") or "").strip() or None, "asume_area": (ctx.get("asume_area") or "").strip() or None,
        "maquinas": sorted({int(m) for m in (ctx.get("maquinas") or []) if m}),
        "ot_madre_id": v.get("ot_madre_id"), "ot_relacion": (v.get("ot_relacion") or "").lower() or None,
        "fin": fin,
    }
    f["clase_neto"] = clase_neto(f)
    f["costo_cobertura"] = None if f["falta_tecnico"] else f["me_total"]
    f["perdida"] = round(-f["queda"], 2) if f["clase_neto"] == "perdida" else 0.0
    return f


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Cortes de tiempo
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def dentro(f, periodo):
    c = f.get("cerrada_at")
    return c is not None and periodo["ini_utc"] <= c < periodo["fin_utc"]


def dividir(filas, periodo):
    """(del período, del período anterior), por fecha de cierre en hora Chile."""
    return ([f for f in filas if dentro(f, periodo)], [f for f in filas if dentro(f, periodo["previo"])])


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  1 · Resumen: Cobré · Me cobraron · Gané · Perdí · Neto
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def resumen(filas, agregar):
    """Las cifras del período. Cobré / Me cobraron / Neto salen de `agregar` (la suma única del proyecto); Gané, Perdí y
    lo regalado se reparten de la MISMA lista de OT, así que Neto = Gané − Perdí − garantías y coberturas siempre."""
    fins = [f["fin"] for f in filas]
    ag = agregar(fins)
    ganadas = [f for f in filas if f["clase_neto"] == "ganada"]
    perdidas = [f for f in filas if f["clase_neto"] == "perdida"]
    cober = [f for f in filas if f["clase_neto"] == "cobertura"]
    cubiertas = [f for f in filas if f["clase_neto"] == "cubierta"]
    medidas = [f for f in filas if f["clase_neto"] in ("ganada", "perdida", "cobertura", "cubierta")]
    sin_medir = [f for f in filas if f["clase_neto"] == "sin_medir"]
    aparte = {}
    for f in filas:
        if f["clase_neto"] == "aparte":
            a = aparte.setdefault(f["cobertura"], {"n": 0, "costo": 0.0, "txt": COBERTURA_TXT.get(f["cobertura"], f["cobertura"])})
            a["n"] += 1
            a["costo"] = round(a["costo"] + f["me_total"], 2)
    respaldo = {"factura": 0.0, "nota_venta": 0.0, "sin_documento": 0.0, "otro": 0.0}
    for f in ganadas + perdidas:
        respaldo[f["respaldo"] if f["respaldo"] in respaldo else "otro"] += f["cobre"]
    pagos = round(sum(f["a_pagar"] for f in medidas), 2)
    reps = round(sum(f["repuestos"] for f in medidas), 2)
    r = {
        "n_ot": len(filas), "n_medidas": len(medidas), "n_sin_medir": len(sin_medir),
        "cobre": round(_f(ag.get("cobre")), 2), "me_cobraron": round(_f(ag.get("me_cobraron")), 2),
        "neto": round(_f(ag.get("queda")), 2), "pct": ag.get("pct"),
        "ganamos": round(sum(f["queda"] for f in ganadas), 2), "n_ganadas": len(ganadas),
        "perdimos": round(sum(f["perdida"] for f in perdidas), 2), "n_perdidas": len(perdidas),
        "regalado_n": len(cober), "regalado_costo": round(sum(f["me_total"] for f in cober), 2),
        "regalado_valor": round(sum(f["valorizado"] or 0 for f in cober), 2),
        "cubiertas_n": len(cubiertas), "cubiertas_costo": round(sum(f["me_total"] for f in cubiertas), 2),
        "pago_tecnicos": pagos, "repuestos": reps, "respaldo": {k: round(v, 2) for k, v in respaldo.items()},
        "aparte": aparte, "aparte_costo": round(sum(a["costo"] for a in aparte.values()), 2),
        "cobre_sin_medir": round(sum(f["cobre"] for f in sin_medir), 2),
        "perdidas_autorizadas": sum(1 for f in perdidas if not f["autorizo"]["sin_autorizacion"]),
        "ag": ag,
    }
    r["clase"] = clase_margen(r["pct"], r["neto"] if r["n_medidas"] else None)
    return r


def clase_margen(pct, queda):
    """Semáforo del margen (mismo criterio que Vida del cliente): rojo SOLO si se pierde plata de verdad."""
    if queda is None:
        return "gris"
    if queda < 0:
        return "rojo"
    if pct is None:
        return "gris"
    return "bajo" if pct < UMBRAL_BAJO else "ok"


def comparar(valor, previo, mejor="sube"):
    """Compara contra el período anterior. `mejor`: «sube» (más es mejor), «baja» (menos es mejor) o «neutro».
    Devuelve flecha (sube|baja|igual|sin_base), el cambio y el tono (bueno|malo|neutro)."""
    v, p = _f(valor), previo
    if p is None:
        return {"flecha": "sin_base", "delta": None, "delta_pct": None, "tono": "neutro", "previo": None}
    p = _f(p)
    delta = round(v - p, 2)
    if abs(delta) < 0.5:
        return {"flecha": "igual", "delta": 0.0, "delta_pct": 0.0, "tono": "neutro", "previo": p}
    dpct = round(delta / abs(p) * 100, 1) if abs(p) >= 1 else None
    flecha = "sube" if delta > 0 else "baja"
    if mejor == "neutro":
        tono = "neutro"
    else:
        tono = "bueno" if (flecha == mejor) else "malo"
    return {"flecha": flecha, "delta": delta, "delta_pct": dpct, "tono": tono, "previo": p}


def tarjetas(res, prev, hay_previo):
    """Las cinco tarjetas de arriba, EN EL ORDEN de Daniel: Cobré · Me cobraron · Gané · Perdí · Neto."""
    def cmp(clave, mejor):
        return comparar(res[clave], prev[clave] if hay_previo else None, mejor)

    rs = res["respaldo"]
    return [
        {"clave": "cobre", "titulo": "Cobré", "valor": res["cobre"], "tono": "base", "icono": "bi-cash-coin",
         "cmp": cmp("cobre", "sube"),
         "detalle": [("Con factura o boleta", rs["factura"], "$"), ("Con nota de venta (falta la factura)", rs["nota_venta"], "$"),
                     ("Sin documento, autorizado", rs["sin_documento"], "$"), ("Otro respaldo", rs["otro"], "$")],
         "ayuda": "Lo cobrado al cliente en las OT que se cobran y cerraron en el período (servicio + despacho). Se separa por el "
                  "documento de Random que lo respalda; es la suma única del proyecto (_ot_fin_agregado)."},
        {"clave": "me_cobraron", "titulo": "Me cobraron", "valor": res["me_cobraron"], "tono": "base", "icono": "bi-person-workspace",
         "cmp": cmp("me_cobraron", "neutro"),
         "detalle": [("Técnicos y proveedores (con despacho)", res["pago_tecnicos"], "$"), ("Repuestos instalados", res["repuestos"], "$")],
         "ayuda": "Lo que cobraron los técnicos y proveedores (incluye el despacho) más los repuestos instalados, en las OT medidas."},
        {"clave": "ganamos", "titulo": "Gané", "valor": res["ganamos"], "tono": "ok", "icono": "bi-graph-up-arrow",
         "cmp": cmp("ganamos", "sube"),
         "detalle": [("OT que dejaron margen", res["n_ganadas"], "n")],
         "ayuda": "La suma de lo que quedó (cobré − me cobraron) en las OT que se cobran y cerraron con margen de cero o más."},
        {"clave": "perdimos", "titulo": "Perdí", "valor": res["perdimos"], "tono": "perdida" if res["perdimos"] > 0 else "base",
         "icono": "bi-graph-down-arrow", "cmp": cmp("perdimos", "baja"),
         "detalle": [("OT cobradas a pérdida", res["n_perdidas"], "n"), ("Con autorización registrada", res["perdidas_autorizadas"], "n")],
         "ayuda": "La suma de lo que se perdió en las OT que se cobran y cerraron con la cuenta en negativo. Las garantías y "
                  "coberturas NO entran aquí: van aparte, más abajo."},
        {"clave": "neto", "titulo": "Neto", "valor": res["neto"], "tono": res["clase"], "icono": "bi-bullseye",
         "cmp": cmp("neto", "sube"), "pct": res["pct"],
         "detalle": [("Gané", res["ganamos"], "$"), ("Perdí", -res["perdimos"], "$"), ("Garantías y coberturas", -res["regalado_costo"], "$")]
                    + ([("Desinstalaciones cubiertas", -res["cubiertas_costo"], "$")] if res["cubiertas_n"] else []),
         "ayuda": "Gané − Perdí − garantías y coberturas (− desinstalaciones cubiertas). Es lo que de verdad dejó el período."},
    ]


def veredicto(res, prev, hay_previo, tendencia, periodo):
    """Una frase de gerente: ¿es rentable el departamento? Tono amable; rojo solo si se perdió plata de verdad."""
    cuando = {"semana": "esta semana", "mes": "este mes", "trimestre": "este trimestre"}.get(periodo["modo"], "en el período")
    if not periodo.get("en_curso") and periodo["modo"] != "rango":
        cuando = "en el período"
    if res["n_medidas"] == 0:
        return {"clase": "gris", "titulo": "Todavía no hay OT medidas en este período",
                "texto": "Cuando se cierren OT con su costo declarado, aquí verás si el departamento deja margen."
                         + (f" Hay {res['n_sin_medir']} OT cerradas sin medir (falta el costo del técnico o del despacho)."
                            if res["n_sin_medir"] else "")}
    q, pct = res["neto"], res["pct"]
    if q < 0:
        clase, titulo = "rojo", f"El departamento perdió {clp(-q)} {cuando}"
    elif pct is None:
        clase, titulo = "gris", f"Sin movimiento de cobro {cuando}"
    elif pct < UMBRAL_BAJO:
        clase, titulo = "bajo", f"Rentable, pero justo {cuando}: margen de {pct_txt(pct)}"
    else:
        clase, titulo = "ok", f"El departamento es rentable {cuando}: margen de {pct_txt(pct)}"
    partes = [f"Cobramos {clp(res['cobre'])}, nos cobraron {clp(res['me_cobraron'])} y quedaron {clp(q)}."]
    if hay_previo and prev["n_medidas"] and prev["pct"] is not None and pct is not None:
        d = round(pct - prev["pct"], 1)
        if abs(d) < 0.1:
            partes.append(f"El margen es igual al del período anterior ({pct_txt(prev['pct'])}).")
        else:
            puntos = f"{abs(d):.1f}".replace(".", ",").replace(",0", "")
            partes.append(f"El margen está {puntos} {'punto' if puntos == '1' else 'puntos'} {'sobre' if d > 0 else 'bajo'} el del período anterior ({pct_txt(prev['pct'])}).")
    if res["regalado_costo"] > 0 and res["cobre"] > 0:
        partes.append(f"Las garantías y coberturas costaron {clp(res['regalado_costo'])}, el "
                      f"{pct_txt(res['regalado_costo'] / res['cobre'] * 100)} de lo cobrado.")
    elif res["regalado_costo"] > 0:
        partes.append(f"Las garantías y coberturas costaron {clp(res['regalado_costo'])}.")
    con_datos = [s for s in (tendencia or []) if s["n"] and s["pct"] is not None][-4:]
    if len(con_datos) >= 3:
        dd = con_datos[-1]["pct"] - con_datos[0]["pct"]
        partes.append("En las últimas semanas con datos el margen va " + ("subiendo." if dd >= 2 else "bajando." if dd <= -2 else "estable."))
    if res["n_sin_medir"]:
        partes.append(f"{res['n_sin_medir']} OT quedan fuera porque falta el costo del técnico o del despacho.")
    return {"clase": clase, "titulo": titulo, "texto": " ".join(partes)}


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  2 · Tendencia de 12 semanas y «a dónde se va la plata»
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def tendencia(filas, hasta, agregar, n=12):
    """Margen y neto de las últimas `n` semanas (lunes–domingo en hora Chile), con la misma suma que el resumen."""
    por = {}
    for f in filas:
        if f.get("semana"):
            por.setdefault(f["semana"], []).append(f)
    out = []
    for lunes in ultimas_semanas(hasta, n):
        fs = por.get(lunes, [])
        ag = agregar([f["fin"] for f in fs])
        out.append({"lunes": lunes, "domingo": lunes + timedelta(days=6), "etiqueta": _etiqueta_semana(lunes),
                    "corta": f"{lunes.day}/{lunes.month}", "n": len(fs), "cobre": round(_f(ag.get("cobre")), 2),
                    "me_cobraron": round(_f(ag.get("me_cobraron")), 2), "neto": round(_f(ag.get("queda")), 2),
                    "pct": ag.get("pct"), "n_medidas": int(ag.get("n") or 0)})
    return out


def tendencia_pct(tend):
    """El margen % de cada semana como línea (las semanas sin cobro quedan como hueco, no como 0 %), con su mínimo, máximo y
    última cifra para leerlo de un vistazo."""
    pcts = [s["pct"] for s in tend]
    con = [p for p in pcts if p is not None]
    return {"serie": pcts, "svg": svg_sparkline(pcts, 420, 64, "Margen por semana, últimas %d semanas" % len(pcts)),
            "min": min(con) if con else None, "max": max(con) if con else None, "ultimo": con[-1] if con else None,
            "n_semanas_con_dato": len(con)}


def donde_se_va(filas):
    """Pareto del gasto del período: todo lo que nos cobraron (técnico, despacho y repuestos) por tipo de OT, centro de costo
    y proveedor. Cuenta TAMBIÉN el trabajo interno y las mantenciones de contrato: es plata que salió igual."""
    total = round(sum(f["me_total"] for f in filas), 2)

    def agrupar(clave_fn, texto_fn):
        g = {}
        for f in filas:
            k = clave_fn(f)
            e = g.setdefault(k, {"clave": k, "etiqueta": texto_fn(f), "n": 0, "costo": 0.0})
            e["n"] += 1
            e["costo"] = round(e["costo"] + f["me_total"], 2)
        lista = sorted(g.values(), key=lambda e: (-e["costo"], -e["n"], e["etiqueta"]))
        acum = 0.0
        for e in lista:
            acum += e["costo"]
            e["pct"] = round(e["costo"] / total * 100, 1) if total > 0 else 0.0
            e["acum"] = round(acum / total * 100, 1) if total > 0 else 0.0
        return lista

    def _prov(f):
        e = f["ejecutor"]
        return "propios" if e["tipo"] in ("interno", "elevado") else e["clave"]

    def _prov_txt(f):
        e = f["ejecutor"]
        return "Técnicos propios" if e["tipo"] in ("interno", "elevado") else e["nombre"]

    return {"total": total, "n": len(filas), "n_sin_costo": sum(1 for f in filas if f["falta_tecnico"]),
            "por_tipo": agrupar(lambda f: f["tipo"] or "sin_tipo", lambda f: f["tipo_txt"]),
            "por_centro": agrupar(lambda f: f["centro_costo"] or "sin_centro", lambda f: f["centro_txt"]),
            "por_proveedor": agrupar(_prov, _prov_txt)}


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  3 · Garantías y coberturas
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def _bucket(valor):
    return valor if valor else SIN_CLASIFICAR


def garantias(filas, filas_ventana, hasta, n_meses=6):
    """Todo lo que NO se cobra, por familia, por cliente y por mes. «Nos costó» es la cuenta de la OT; si el técnico aún no
    declara su costo se dice («sin costo declarado»), nunca se inventa un $0. Las familias que son plata regalada
    (garantía, regalía, arriendo/leasing, cortesía) van separadas de las que no lo son (desinstalación cubierta, interno,
    contrato). Los puntos de extensión `intencion` y `asume_area` agrupan cuando existan; mientras no, «Sin clasificar»."""
    todas = [f for f in filas if f["cobertura"] != "cobra"]

    def totales(fs):
        conocidas = [f for f in fs if not f["falta_tecnico"]]
        return {"n": len(fs), "costo": round(sum(f["me_total"] for f in conocidas), 2),
                "valor": round(sum(f["valorizado"] or 0 for f in fs), 2),
                "n_sin_costo": len(fs) - len(conocidas), "n_sin_valor": sum(1 for f in fs if not f["valorizado"]),
                "n_costo_mayor": sum(1 for f in conocidas if f["valorizado"] and f["me_total"] > f["valorizado"])}

    por_familia = []
    for c in COBERTURA_ORDEN:
        fs = [f for f in todas if f["cobertura"] == c]
        if fs:
            t = totales(fs)
            t.update({"cobertura": c, "txt": COBERTURA_TXT[c], "regalada": c in COBERTURAS_REGALADAS,
                      "pct_costo_valor": round(t["costo"] / t["valor"] * 100, 1) if t["valor"] > 0 else None})
            por_familia.append(t)
    regaladas = [f for f in todas if f["cobertura"] in COBERTURAS_REGALADAS]
    otras = [f for f in todas if f["cobertura"] in ("cubierta",)]

    cli = {}
    for f in regaladas + otras:
        e = cli.setdefault(f["cliente_id"] or f["cliente"], {"cliente_id": f["cliente_id"], "cliente": f["cliente"],
                                                             "n": 0, "costo": 0.0, "valor": 0.0, "n_sin_costo": 0, "familias": {},
                                                             "ots": []})
        e["n"] += 1
        if len(e["ots"]) < 12:
            e["ots"].append({"id": f["id"], "numero": f["numero"]})
        if f["falta_tecnico"]:
            e["n_sin_costo"] += 1
        else:
            e["costo"] = round(e["costo"] + f["me_total"], 2)
        e["valor"] = round(e["valor"] + (f["valorizado"] or 0), 2)
        e["familias"][COBERTURA_TXT[f["cobertura"]]] = e["familias"].get(COBERTURA_TXT[f["cobertura"]], 0) + 1
    por_cliente = sorted(cli.values(), key=lambda e: (-e["costo"], -e["n"], e["cliente"]))

    meses = []
    for anio, mes in ultimos_meses(hasta, n_meses):
        clave = f"{anio}-{mes:02d}"
        fs = [f for f in filas_ventana if f["mes"] == clave and f["cobertura"] in COBERTURAS_REGALADAS]
        t = totales(fs)
        t.update({"clave": clave, "etiqueta": f"{MESES_CORTO[mes - 1].capitalize()} {anio}"})
        meses.append(t)
    mx_mes = max([m["costo"] for m in meses] + [1.0])
    for m in meses:
        m["pct"] = round(m["costo"] / mx_mes * 100, 1)

    def por_extension(clave):
        g = {}
        for f in regaladas:
            k = _bucket(f.get(clave))
            e = g.setdefault(k, {"etiqueta": k, "n": 0, "costo": 0.0, "valor": 0.0, "sin_clasificar": k == SIN_CLASIFICAR})
            e["n"] += 1
            e["costo"] = round(e["costo"] + (f["me_total"] if not f["falta_tecnico"] else 0), 2)
            e["valor"] = round(e["valor"] + (f["valorizado"] or 0), 2)
        return sorted(g.values(), key=lambda e: (e["sin_clasificar"], -e["costo"], e["etiqueta"]))

    cc = {}
    for f in regaladas:
        e = cc.setdefault(f["centro_costo"] or "sin_centro", {"etiqueta": f["centro_txt"], "n": 0, "costo": 0.0, "valor": 0.0})
        e["n"] += 1
        e["costo"] = round(e["costo"] + (f["me_total"] if not f["falta_tecnico"] else 0), 2)
        e["valor"] = round(e["valor"] + (f["valorizado"] or 0), 2)
    por_centro = sorted(cc.values(), key=lambda e: (-e["costo"], -e["n"], e["etiqueta"]))
    tot_reg = totales(regaladas)
    return {"por_familia": por_familia, "total_regalado": tot_reg, "n_cubiertas": len(otras),
            "por_cliente": por_cliente, "por_mes": meses, "por_centro": por_centro,
            "por_intencion": por_extension("intencion"), "por_area": por_extension("asume_area"),
            "costo_mayor_valor": [f for f in regaladas if not f["falta_tecnico"] and f["valorizado"] and f["me_total"] > f["valorizado"]],
            "clasificadas": sum(1 for f in regaladas if f.get("asume_area") or f.get("intencion")),
            "regaladas": regaladas}


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  4 · A nombre de quién (atribución separada e inmutable)
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
ROLES = (("origino", "Originó", "Quién pidió el trabajo: el ticket (o la OT, si nació sin ticket)."),
         ("autorizo", "Autorizó", "Quién aprobó el $0 o la pérdida, con su argumento guardado."),
         ("lleva", "Lleva", "El responsable de la OT (REGLA #36)."),
         ("ejecuto", "Ejecutó", "El técnico asignado y, si es externo, su empresa."))


def _nombre_ejecuto(f):
    """Quién EJECUTÓ: el técnico asignado (o su empresa, si es externo); nadie si la OT no tiene técnico."""
    e = f["ejecutor"]
    return None if e["tipo"] == "sin_asignar" else e["nombre"]


def atribucion(filas):
    """Filas = cada garantía/cobertura y cada OT a pérdida del período, con los cuatro roles SEPARADOS y la fuente de cada
    uno. El resumen por persona suma por rol: quien ejecutó no recibe lo que otro autorizó. Sin dato = «sin registrar»."""
    casos = []
    for f in filas:
        if f["clase_neto"] == "perdida":
            tipo, monto = "perdida", f["perdida"]
        elif f["clase_neto"] in ("cobertura", "cubierta") or (f["cobertura"] in COBERTURAS_REGALADAS and not f["mide"]):
            tipo, monto = "cobertura", f["costo_cobertura"]
        else:
            continue
        casos.append({"id": f["id"], "numero": f["numero"], "cliente": f["cliente"], "tipo": tipo, "monto": monto,
                      "cobertura_txt": f["cobertura_txt"] if tipo == "cobertura" else "Cobrada a pérdida",
                      "cobre": f["cobre"], "me_total": f["me_total"], "valorizado": f["valorizado"],
                      "centro_txt": f["centro_txt"], "centro_costo": f["centro_costo"],
                      "intencion": f["intencion"], "asume_area": f["asume_area"],
                      "cerrada_at": f["cerrada_at"], "origino": f["origino"], "autorizo": f["autorizo"], "lleva": f["lleva"],
                      "ejecuto": {"nombre": _nombre_ejecuto(f), "tipo_txt": f["ejecutor"]["tipo_txt"],
                                  "persona": f["ejecutor"]["persona"], "empresa": f["ejecutor"]["empresa"],
                                  "fuente": "Técnico asignado a la OT"},
                      "argumento": f["autorizo"]["argumento"], "sin_autorizacion": f["autorizo"]["sin_autorizacion"]})
    casos.sort(key=lambda c: (-_f(c["monto"]), c["numero"]))
    personas = {}
    for c in casos:
        for rol, _t, _a in ROLES:
            nombre = (c["ejecuto"]["nombre"] if rol == "ejecuto" else (c[rol] or {}).get("nombre"))
            if not nombre:
                continue
            p = personas.setdefault(nombre.strip().lower(), {"nombre": nombre.strip(), **{r: {"n": 0, "costo": 0.0, "perdida": 0.0, "n_cob": 0, "n_per": 0}
                                                                                           for r, _x, _y in ROLES}})
            e = p[rol]
            e["n"] += 1
            if c["tipo"] == "perdida":
                e["n_per"] += 1
                e["perdida"] = round(e["perdida"] + _f(c["monto"]), 2)
            else:
                e["n_cob"] += 1
                e["costo"] = round(e["costo"] + _f(c["monto"]), 2)
    lista = sorted(personas.values(), key=lambda p: -sum(p[r]["costo"] + p[r]["perdida"] for r, _x, _y in ROLES))
    return {"casos": casos, "personas": lista, "roles": ROLES,
            "n_sin_autorizacion": sum(1 for c in casos if c["sin_autorizacion"]),
            "n_sin_argumento": sum(1 for c in casos if not c["argumento"]),
            "n_sin_clasificar": sum(1 for c in casos if c["tipo"] == "cobertura" and not (c["asume_area"] and c["intencion"]))}


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  5 · Técnicos y proveedores · cumplimiento · reincidencia
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def tecnicos(filas, filas_prev, filas_ventana, hasta, semanas_serie=8):
    """Cuánto nos cobra cada técnico/empresa: total, N° de OT, ticket promedio, % del costo total y evolución (contra el
    período anterior y semana a semana). Los externos se agrupan por empresa; los propios, por persona."""
    def cobran(fs):
        return round(sum(f["a_pagar"] for f in fs), 2)

    total = cobran(filas)
    prev_por = {}
    for f in filas_prev:
        prev_por[f["ejecutor"]["clave"]] = round(prev_por.get(f["ejecutor"]["clave"], 0.0) + f["a_pagar"], 2)
    sem = ultimas_semanas(hasta, semanas_serie)
    ven_por = {}
    for f in filas_ventana:
        if f.get("semana") in sem:
            d = ven_por.setdefault(f["ejecutor"]["clave"], {})
            d[f["semana"]] = round(d.get(f["semana"], 0.0) + f["a_pagar"], 2)
    grupos = {}
    for f in filas:
        e = f["ejecutor"]
        g = grupos.setdefault(e["clave"], {"clave": e["clave"], "nombre": e["nombre"], "tipo": e["tipo"], "tipo_txt": e["tipo_txt"],
                                            "personas": [], "filas": []})
        g["filas"].append(f)
        if e["persona"] and e["persona"] not in g["personas"]:
            g["personas"].append(e["persona"])
    lista = []
    for g in grupos.values():
        fs = g.pop("filas")
        c = cobran(fs)
        con_costo = [f for f in fs if not f["falta_tecnico"]]
        previo = prev_por.get(g["clave"])
        serie = [ven_por.get(g["clave"], {}).get(s, 0.0) for s in sem]
        g.update({"n": len(fs), "cobran": c, "n_sin_costo": len(fs) - len(con_costo),
                  "promedio": round(c / len(con_costo), 2) if con_costo else None,
                  "pct_costo": round(c / total * 100, 1) if total > 0 else 0.0,
                  "n_garantias": sum(1 for f in fs if f["cobertura"] in COBERTURAS_REGALADAS),
                  "costo_garantias": round(sum(f["me_total"] for f in fs if f["cobertura"] in COBERTURAS_REGALADAS
                                               and not f["falta_tecnico"]), 2),
                  "previo": previo, "cmp": comparar(c, previo if previo is not None else (0.0 if filas_prev else None), "neutro"),
                  "serie": serie, "ots": [{"id": f["id"], "numero": f["numero"], "cobran": f["a_pagar"]} for f in fs][:30]})
        lista.append(g)
    lista.sort(key=lambda g: (-g["cobran"], -g["n"], g["nombre"]))
    tipos = []
    for t in TIPO_EJECUTOR_ORDEN:
        gs = [g for g in lista if g["tipo"] == t]
        if gs:
            tipos.append({"tipo": t, "txt": TIPO_EJECUTOR_TXT[t], "n_grupos": len(gs), "n": sum(g["n"] for g in gs),
                          "cobran": round(sum(g["cobran"] for g in gs), 2)})
    return {"grupos": lista, "total": total, "tipos": tipos, "semanas": sem}


def cumplimiento(agendadas, periodo, hoy):
    """¿Están yendo? Por técnico/empresa: OT agendadas en el período, cuántas ya están hechas y cuántas vencieron sin
    cerrarse. `agendadas`: [{id, numero, ejecutor, fecha_programada (date), estado}]."""
    g = {}
    for a in agendadas:
        est = (a.get("estado") or "").lower()
        fp = a.get("fecha_programada")
        if est in ESTADOS_FUERA or est == "reagendada" or not fp:
            continue
        if not (periodo["desde"] <= fp <= periodo["hasta"]):
            continue
        e = a["ejecutor"]
        x = g.setdefault(e["clave"], {"clave": e["clave"], "nombre": e["nombre"], "tipo": e["tipo"], "agendadas": 0, "hechas": 0,
                                       "vencidas": 0, "casos_vencidos": []})
        x["agendadas"] += 1
        if est in ESTADOS_HECHOS:
            x["hechas"] += 1
        elif fp < hoy:
            x["vencidas"] += 1
            x["casos_vencidos"].append({"id": a.get("id"), "numero": a.get("numero")})
    for x in g.values():
        x["pct_hechas"] = round(x["hechas"] / x["agendadas"] * 100, 1) if x["agendadas"] else None
    return g


def _comparten_equipo(a, b):
    """'equipo' si comparten un equipo; 'cliente' si alguno no tiene equipos registrados (se compara solo por cliente);
    None si ambos tienen equipos y ninguno coincide."""
    ea, eb = set(a.get("maquinas") or ()), set(b.get("maquinas") or ())
    if ea and eb:
        return "equipo" if ea & eb else None
    return "cliente"


def detectar_reincidencias(cerradas, candidatas, hoy_utc, dias=DIAS_REINCIDENCIA_DEFECTO):
    """Reincidencia de los EXTERNOS. Definición (por defecto, y es solo un indicador para revisar, no una sentencia):
    una OT correctiva o de garantía sobre el mismo equipo (o el mismo cliente, si no hay equipos registrados) que se crea
    hasta `dias` después de que un externo cerró una OT.

    · No cuentan las OT que son continuación planificada de la misma (hijas: desinstalación, pendiente por stock,
      seguimiento de un hallazgo) ni las canceladas o anuladas.
    · Una OT cuya ventana de `dias` aún no termina y no tiene reincidencia se marca «en observación» y NO entra al
      denominador de la tasa (no se sabe todavía): se informa aparte.
    `cerradas`: [{id, numero, cliente_id, cliente, cerrada_at, ejecutor, maquinas}] · `candidatas`: [{id, numero, cliente_id,
    tipo, es_garantia, estado, created_at, ot_madre_id, ot_relacion, maquinas, ejecutor}].
    Devuelve {por_empresa, casos, ...} (tasa en %, None si no hay con qué medir)."""
    por_cli = {}
    for c in candidatas:
        por_cli.setdefault(c.get("cliente_id"), []).append(c)
    limite = timedelta(days=int(dias))
    emp = {}
    casos = []
    for x in cerradas:
        e = x["ejecutor"]
        if e["tipo"] != "externo" or not x.get("cerrada_at") or not x.get("cliente_id"):
            continue
        g = emp.setdefault(e["clave"], {"clave": e["clave"], "nombre": e["nombre"], "n_cerradas": 0, "n_reincidentes": 0,
                                        "n_en_observacion": 0, "n_repite_visita": 0, "casos": []})
        g["n_cerradas"] += 1
        encontrados = []
        for y in por_cli.get(x["cliente_id"], []):
            if y["id"] == x["id"] or (y.get("estado") or "").lower() in ESTADOS_FUERA:
                continue
            if y.get("ot_madre_id") == x["id"]:
                continue                                    # continuación planificada de la misma OT
            if not (y["tipo"] in TIPOS_CORRECTIVOS or y.get("es_garantia")):
                continue
            ya = y.get("created_at")
            if ya is None or not (x["cerrada_at"] < ya <= x["cerrada_at"] + limite):
                continue
            nivel = _comparten_equipo(x, y)
            if nivel is None:
                continue
            ey = y.get("ejecutor") or {}
            encontrados.append({"x_id": x["id"], "x_numero": x["numero"], "y_id": y["id"], "y_numero": y["numero"],
                                "cliente": x.get("cliente"), "cliente_id": x["cliente_id"], "nivel": nivel,
                                "dias_despues": max(0, (ya - x["cerrada_at"]).days), "y_tipo": y.get("tipo"),
                                "y_ejecutor": ey.get("nombre"), "misma_empresa": ey.get("clave") == e["clave"],
                                "cerrada_at": x["cerrada_at"], "empresa": e["nombre"]})
        if encontrados:
            g["n_reincidentes"] += 1
            if any(c["misma_empresa"] for c in encontrados):
                g["n_repite_visita"] += 1
            encontrados.sort(key=lambda c: c["dias_despues"])
            g["casos"].extend(encontrados)
            casos.extend(encontrados)
        elif x["cerrada_at"] + limite > hoy_utc:
            g["n_en_observacion"] += 1
    lista = []
    for g in emp.values():
        den = g["n_cerradas"] - g["n_en_observacion"]
        g["n_evaluadas"] = den
        g["tasa"] = round(g["n_reincidentes"] / den * 100, 1) if den > 0 else None
        g["revisar"] = bool(g["tasa"] is not None and g["tasa"] >= TASA_REVISAR and den >= MIN_CASOS_REVISAR)
        lista.append(g)
    lista.sort(key=lambda g: (-(g["tasa"] if g["tasa"] is not None else -1), -g["n_cerradas"], g["nombre"]))
    casos.sort(key=lambda c: (c["cerrada_at"], c["y_numero"]), reverse=True)
    den = sum(g["n_evaluadas"] for g in lista)
    num = sum(g["n_reincidentes"] for g in lista)
    return {"dias": int(dias), "por_empresa": lista, "casos": casos,
            "tasa_global": round(num / den * 100, 1) if den > 0 else None, "n_evaluadas": den, "n_reincidentes": num,
            "n_en_observacion": sum(g["n_en_observacion"] for g in lista),
            "n_repite_visita": sum(g["n_repite_visita"] for g in lista),
            "n_empresas": len(lista), "n_revisar": sum(1 for g in lista if g["revisar"])}


def resumen_abiertas(items):
    """Plata en camino: lo que se cobrará y lo que se pagará con las OT que siguen abiertas. `items`: [{cobra, mide, cobre,
    a_pagar}] ya calculados con _ot_finanzas."""
    por_cobrar = round(sum(i["cobre"] for i in items if i.get("cobra") and i.get("cobre")), 2)
    por_pagar = round(sum(i["a_pagar"] for i in items if i.get("a_pagar")), 2)
    return {"n": len(items), "por_cobrar": por_cobrar, "por_pagar": por_pagar,
            "n_sin_medir": sum(1 for i in items if not i.get("mide")),
            "n_con_cobro": sum(1 for i in items if i.get("cobra") and i.get("cobre"))}


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Gráficos SVG (inline, con clases: los colores viven en ot_gerencial.css)
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def _esc(s):
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;"))


def svg_barras_neto(semanas, ancho=760, alto=200):
    """Barras del neto por semana (verde si dejó margen, rojo si se perdió) sobre una grilla tenue; la última semana lleva su
    cifra. Accesible: role=img con una descripción y una tabla equivalente en la página."""
    if not semanas:
        return ""
    izq, der, arr, aba = 8, 8, 24, 30
    w, h = ancho - izq - der, alto - arr - aba
    vals = [s["neto"] for s in semanas]
    mx = max(max(vals), 0.0)
    mn = min(min(vals), 0.0)
    rango = (mx - mn) or 1.0
    y0 = arr + h * (mx / rango)
    paso = w / len(semanas)
    ancho_b = max(6.0, paso * 0.62)
    partes = [f'<svg class="gc-svg" viewBox="0 0 {ancho} {alto}" role="img" aria-label="Neto semanal de las últimas '
              f'{len(semanas)} semanas" preserveAspectRatio="xMidYMid meet">']
    for i in range(5):
        y = arr + h * i / 4
        partes.append(f'<line class="gc-grid" x1="{izq}" x2="{ancho - der}" y1="{y:.1f}" y2="{y:.1f}"/>')
    partes.append(f'<line class="gc-cero" x1="{izq}" x2="{ancho - der}" y1="{y0:.1f}" y2="{y0:.1f}"/>')
    for i, s in enumerate(semanas):
        cx = izq + paso * i + paso / 2
        alt = abs(s["neto"]) / rango * h
        y = y0 - alt if s["neto"] >= 0 else y0
        clase = "gc-b-vacia" if not s["n_medidas"] else ("gc-b-neg" if s["neto"] < 0 else "gc-b-pos")
        ultima = i == len(semanas) - 1
        partes.append(f'<rect class="{clase}{" gc-b-ult" if ultima else ""}" x="{cx - ancho_b / 2:.1f}" y="{y:.1f}" '
                      f'width="{ancho_b:.1f}" height="{max(alt, 1.5):.1f}" rx="3"><title>{_esc(s["etiqueta"])}: '
                      f'{_esc(clp(s["neto"]))} ({s["n_medidas"]} OT medidas)</title></rect>')
        if i % 2 == 0 or ultima:
            partes.append(f'<text class="gc-eje" x="{cx:.1f}" y="{alto - 10}" text-anchor="middle">{_esc(s["corta"])}</text>')
        if ultima and s["n_medidas"]:
            ty = (y - 6) if s["neto"] >= 0 else (y + alt + 14)
            partes.append(f'<circle class="gc-punto" cx="{cx:.1f}" cy="{(y if s["neto"] >= 0 else y + alt):.1f}" r="4"/>')
            partes.append(f'<text class="gc-valor" x="{cx:.1f}" y="{ty:.1f}" text-anchor="middle">{_esc(clp(s["neto"]))}</text>')
    partes.append("</svg>")
    return "".join(partes)


def svg_sparkline(valores, ancho=110, alto=30, etiqueta="Evolución"):
    """Línea con relleno y punto final resaltado. Los huecos (None) se saltan."""
    pts = [(i, v) for i, v in enumerate(valores or []) if v is not None]
    if len(pts) < 2:
        return ""
    xs = [p[0] for p in pts]
    vs = [p[1] for p in pts]
    mn, mx = min(vs), max(vs)
    rango = (mx - mn) or 1.0
    n = max(len(valores) - 1, 1)
    pad = 3

    def xy(i, v):
        return pad + (ancho - 2 * pad) * i / n, pad + (alto - 2 * pad) * (1 - (v - mn) / rango)

    coords = [xy(i, v) for i, v in pts]
    linea = " ".join(f"{x:.1f},{y:.1f}" for x, y in coords)
    area = f"{coords[0][0]:.1f},{alto - pad} " + linea + f" {coords[-1][0]:.1f},{alto - pad}"
    ux, uy = coords[-1]
    return (f'<svg class="gc-spark" viewBox="0 0 {ancho} {alto}" role="img" aria-label="{_esc(etiqueta)}">'
            f'<polygon class="gc-spark-area" points="{area}"/><polyline class="gc-spark-linea" points="{linea}"/>'
            f'<circle class="gc-punto" cx="{ux:.1f}" cy="{uy:.1f}" r="3"/></svg>')


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  El tablero completo
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def armar_tablero(periodo, filas_ventana, agregar, hoy, agendadas=None, reinc=None, abiertas=None, cxp=None, meta=None):
    """Junta todo. `filas_ventana`: las OT cerradas leídas (período, anterior, 12 semanas y 6 meses) ya normalizadas con
    `armar_fila`; `agregar`: la suma única (`_ot_fin_agregado`); `reinc`: {cerradas, candidatas, dias, hoy_utc}."""
    meta = dict(meta or {})
    filas, filas_prev = dividir(filas_ventana, periodo)
    hay_previo = bool(filas_prev)
    res = resumen(filas, agregar)
    prev = resumen(filas_prev, agregar)
    tend = tendencia(filas_ventana, min(periodo["hasta"], hoy), agregar)
    tecs = tecnicos(filas, filas_prev, filas_ventana, min(periodo["hasta"], hoy))
    cump = cumplimiento(agendadas or [], periodo, hoy) if agendadas is not None else {}
    for g in tecs["grupos"]:
        g["cumplimiento"] = cump.get(g["clave"])
    re_ = None
    if reinc is not None:
        re_ = detectar_reincidencias(reinc.get("cerradas") or [], reinc.get("candidatas") or [],
                                     reinc.get("hoy_utc") or datetime.utcnow(), reinc.get("dias") or DIAS_REINCIDENCIA_DEFECTO)
        por = {g["clave"]: g for g in re_["por_empresa"]}
        for g in tecs["grupos"]:
            g["reincidencia"] = por.get(g["clave"])
    gar = garantias(filas, filas_ventana, periodo["hasta"])
    atr = atribucion(filas)
    return {
        "nombre": NOMBRE_PANTALLA, "clasificacion": CLASIFICACION, "periodo": periodo, "hay_previo": hay_previo,
        "resumen": res, "previo": prev, "tarjetas": tarjetas(res, prev, hay_previo),
        "veredicto": veredicto(res, prev, hay_previo, tend, periodo),
        "tendencia": tend, "svg_tendencia": svg_barras_neto(tend), "tendencia_pct": tendencia_pct(tend),
        "donde_se_va": donde_se_va(filas),
        "garantias": gar, "atribucion": atr, "tecnicos": tecs, "reincidencia": re_,
        "perdidas": [f for f in filas if f["clase_neto"] == "perdida"],
        "sin_medir": [f for f in filas if f["clase_neto"] == "sin_medir"],
        "abiertas": resumen_abiertas(abiertas) if abiertas is not None else None, "cxp": cxp,
        "meta": meta,
    }


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Excel (solo superadmin): una hoja por sección
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def excel_bytes(d):
    """El período en un .xlsx — Resumen, Tendencia, Dónde se va la plata, Garantías y coberturas, A nombre de quién,
    Técnicos y proveedores y Reincidencia. Cada hoja dice «Confidencial · solo gerencia»."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    import xlsx_fechas as XF     # 📅 2026-10-29: fechas reales de Excel, hora Chile, formato español (un solo helper)

    wb = openpyxl.Workbook()
    NEGRO, GRIS = "0A0A0A", "F3F4F6"
    # Encabezados de las columnas de apoyo (Año, N° mes, Semana…): son números, pero NO plata ni cantidades con miles.
    _apoyo_cierre = XF.encabezados_apoyo("de la fecha de cierre")
    _apoyo_semana = XF.encabezados_apoyo("del lunes de la semana")
    _sin_miles = set(_apoyo_cierre) | set(_apoyo_semana)

    def hoja(titulo, encabezados, filas, anchos=None, nota=None, primera=False):
        ws = wb.active if primera else wb.create_sheet()
        ws.title = titulo[:31]
        ws.cell(row=1, column=1, value=f"{NOMBRE_PANTALLA} · {d['periodo']['etiqueta']} · {CLASIFICACION}").font = Font(bold=True, size=11)
        if nota:
            ws.cell(row=2, column=1, value=nota).font = Font(italic=True, size=9, color="6B7280")
        fila0 = 4
        for ci, h in enumerate(encabezados, 1):
            c = ws.cell(row=fila0, column=ci, value=h)
            c.font = Font(bold=True, color="FFFFFF", size=10)
            c.fill = PatternFill("solid", fgColor=NEGRO)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for ri, fila in enumerate(filas, fila0 + 1):
            for ci, val in enumerate(fila, 1):
                c = ws.cell(row=ri, column=ci, value=val)
                c.alignment = Alignment(vertical="top", wrap_text=isinstance(val, str) and len(val) > 40)
                if XF.formato(val):                  # 📅 2026-10-29: fecha real, dd/mm/yyyy [hh:mm]
                    c.number_format = XF.formato(val)
                elif ci <= len(encabezados) and encabezados[ci - 1] in _sin_miles:
                    c.number_format = "0" if isinstance(val, int) and not isinstance(val, bool) else "General"
                elif (isinstance(val, float) or (isinstance(val, int) and not isinstance(val, bool))) and ci <= len(encabezados):
                    c.number_format = '"$"#,##0' if "$" in encabezados[ci - 1] else "0.0" if "%" in encabezados[ci - 1] else "#,##0"
                if (ri - fila0) % 2 == 0:
                    c.fill = PatternFill("solid", fgColor=GRIS)
        for ci in range(1, len(encabezados) + 1):
            ws.column_dimensions[get_column_letter(ci)].width = (anchos or {}).get(ci, max(14, min(48, len(str(encabezados[ci - 1])) + 6)))
        ws.freeze_panes = ws.cell(row=fila0 + 1, column=1)
        # 📅 2026-10-29: filtro automático sobre la PRIMERA tabla de la hoja (hasta la primera fila vacía: «Por cliente», «Casos»…
        # son tablas aparte y un filtro que las mezcle las desordenaría).
        n_tabla = next((i for i, f in enumerate(filas) if not f), len(filas))
        ws.auto_filter.ref = f"A{fila0}:{get_column_letter(max(1, len(encabezados)))}{fila0 + n_tabla}"
        return ws

    r = d["resumen"]
    hoja("Resumen", ["Concepto", "Valor ($)", "Detalle"], [
        ["Cobré", r["cobre"], f"Factura/boleta {clp(r['respaldo']['factura'])} · nota de venta {clp(r['respaldo']['nota_venta'])} · sin documento {clp(r['respaldo']['sin_documento'])}"],
        ["Me cobraron", r["me_cobraron"], f"Técnicos y proveedores {clp(r['pago_tecnicos'])} · repuestos {clp(r['repuestos'])}"],
        ["Gané", r["ganamos"], f"{r['n_ganadas']} OT con margen"],
        ["Perdí", r["perdimos"], f"{r['n_perdidas']} OT cobradas a pérdida"],
        ["Garantías y coberturas (costo)", r["regalado_costo"], f"{r['regalado_n']} OT · valor del servicio {clp(r['regalado_valor'])}"],
        ["Desinstalaciones cubiertas (costo)", r["cubiertas_costo"], f"{r['cubiertas_n']} OT"],
        ["Neto", r["neto"], f"Margen {pct_txt(r['pct'])}"],
        ["OT sin medir (fuera de las cifras)", r["n_sin_medir"], "Falta el costo del técnico o del despacho"],
        ["Veredicto", d["veredicto"]["titulo"], d["veredicto"]["texto"]],
    ], {1: 38, 2: 22, 3: 90}, primera=True)
    # 📅 2026-10-29: al final, el lunes y el domingo de cada semana como fechas reales y Año · Mes · Semana… para filtrar por mes.
    hoja("Tendencia", ["Semana", "OT medidas", "Cobré ($)", "Me cobraron ($)", "Neto ($)", "Margen %",
                       "Desde (lunes)", "Hasta (domingo)"] + _apoyo_semana,
         [[s["etiqueta"], s["n_medidas"], s["cobre"], s["me_cobraron"], s["neto"], s["pct"],
           XF.a_fecha_excel(s.get("lunes")), XF.a_fecha_excel(s.get("domingo"))] + XF.columnas_apoyo(XF.a_fecha_excel(s.get("lunes")))
          for s in d["tendencia"]], {1: 34})
    ds = d["donde_se_va"]
    filas = []
    for corte, lista in (("Tipo de OT", ds["por_tipo"]), ("Centro de costo", ds["por_centro"]), ("Proveedor", ds["por_proveedor"])):
        filas += [[corte, e["etiqueta"], e["n"], e["costo"], e["pct"], e["acum"]] for e in lista]
    hoja("Dónde se va la plata", ["Corte", "Detalle", "OT", "Costo ($)", "% del total", "% acumulado"], filas, {2: 40},
         nota="Todo lo que nos cobraron (técnico, despacho y repuestos) en las OT cerradas del período.")
    g = d["garantias"]
    hoja("Garantías y coberturas",
         ["Familia", "OT", "Valor del servicio ($)", "Nos costó ($)", "Costo / valor %", "Sin costo declarado"],
         [[x["txt"], x["n"], x["valor"], x["costo"], x["pct_costo_valor"], x["n_sin_costo"]] for x in g["por_familia"]]
         + [[], ["Por cliente"]] + [[x["cliente"], x["n"], x["valor"], x["costo"], None, x["n_sin_costo"]] for x in g["por_cliente"]]
         + [[], ["Por mes"]] + [[x["etiqueta"], x["n"], x["valor"], x["costo"], None, x["n_sin_costo"]] for x in g["por_mes"]],
         {1: 44})
    a = d["atribucion"]
    # 📅 2026-10-29: «Cuándo autorizó» como fecha y hora real (hora Chile) y, al final, la fecha de cierre de la OT con Año · Mes ·
    # Semana… (todas las OT del costeo están cerradas: la fecha que manda es la de cierre).
    hoja("A nombre de quién",
         ["OT", "Cliente", "Qué es", "Monto ($)", "Área que asume", "Intención", "Centro de costo", "Originó", "Fuente (originó)",
          "Autorizó", "Cuándo autorizó", "Argumento guardado", "Lleva", "Ejecutó", "Fecha de cierre"] + _apoyo_cierre,
         [[c["numero"], c["cliente"], c["cobertura_txt"], c["monto"], c["asume_area"] or SIN_CLASIFICAR, c["intencion"] or SIN_CLASIFICAR,
           c["centro_txt"], (c["origino"]["nombre"] or "sin registrar"), c["origino"]["fuente"],
           c["autorizo"]["nombre"] or "Sin autorización registrada", XF.a_fecha_excel(c["autorizo"]["at"]),
           c["argumento"] or "", c["lleva"]["nombre"] or "Sin responsable", c["ejecuto"]["nombre"] or "Sin técnico",
           XF.a_fecha_excel(c.get("cerrada_at"))] + XF.columnas_apoyo(XF.a_fecha_excel(c.get("cerrada_at"))) for c in a["casos"]],
         {2: 30, 8: 24, 9: 30, 10: 24, 12: 60, 13: 24, 14: 26, 15: 17},
         nota="Cuatro roles separados: quién originó, quién autorizó, quién lleva y quién ejecutó. Cada uno sale de su propia fuente.")
    t = d["tecnicos"]
    hoja("Técnicos y proveedores",
         ["Técnico o empresa", "Tipo", "OT", "Nos cobran ($)", "Promedio por OT ($)", "% del costo", "Garantías", "Costo garantías ($)", "Cambio vs anterior %"],
         [[x["nombre"], x["tipo_txt"], x["n"], x["cobran"], x["promedio"], x["pct_costo"], x["n_garantias"], x["costo_garantias"],
           x["cmp"]["delta_pct"]] for x in t["grupos"]], {1: 40})
    if d.get("reincidencia"):
        ri = d["reincidencia"]
        hoja("Reincidencia",
             ["Empresa externa", "OT cerradas", "Evaluadas", "En observación", "Con reincidencia", "Tasa %", "Volvió la misma empresa"],
             [[x["nombre"], x["n_cerradas"], x["n_evaluadas"], x["n_en_observacion"], x["n_reincidentes"], x["tasa"], x["n_repite_visita"]]
              for x in ri["por_empresa"]]
             + [[], ["Casos"], ["OT cerrada", "OT posterior", "Cliente", "Días después", "Nivel", "Quién fue", "Misma empresa",
                                "Cerrada el"]]
             + [[c["x_numero"], c["y_numero"], c["cliente"], c["dias_despues"], "mismo equipo" if c["nivel"] == "equipo" else "mismo cliente",
                 c["y_ejecutor"] or "", "sí" if c["misma_empresa"] else "no",
                 XF.a_fecha_excel(c.get("cerrada_at"))] for c in ri["casos"]],     # 📅 2026-10-29: fecha real, hora Chile
             {1: 40},
             nota=f"Reincidencia = OT correctiva o de garantía sobre el mismo equipo (o cliente) dentro de {ri['dias']} días de que un externo cerró una OT. "
                  "Es un indicador para revisar, no una sentencia.")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
