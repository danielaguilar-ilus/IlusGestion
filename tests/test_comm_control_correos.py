# -*- coding: utf-8 -*-
"""CONTROL DE CORREOS (Comunicaciones → «Control de correos», Daniel, 20-oct-2026).

Daniel: «una pestaña donde el superadmin ve TODOS los correos que el sistema envía y puede APAGAR cada uno (con motivo)».

Qué se prueba aquí (todo con correo de MENTIRA — REGLA #22: jamás un correo real ni a clientes):
  1. El REGISTRO (comm_correos_registro.py): completo, sin claves repetidas, cada llave existe, cada plantilla está en el editor,
     y las tablas de traducción de Retiros / Tickets / Accesos / Técnicos / Transporte caen en claves registradas.
  2. COBERTURA: TODA llamada de envío del proyecto (todos los .py, salvo tests) lleva `correo=` con una clave registrada, o está en
     la lista de excepciones con su razón; el `modulo=` de cada llamada es la llave del registro.
  3. La PUERTA ÚNICA (`_send_ilus_email` de app.py, función real): general > llave del módulo > interruptor del correo > cola;
     sin fila = encendido; clave desconocida sale; un correo protegido nunca se frena; base caída = último valor conocido.
  4. La COLA: la clave viaja en su columna (nunca en kwargs_json), se vuelve a revisar al drenar, y sin la columna igual se encola.
  5. La MIGRACIÓN: barata (REGLA #18): la segunda corrida no ejecuta ningún DDL.
  6. REINTENTAR respeta el interruptor.
  7. La PANTALLA (fase 2): los 5 endpoints reales (solo superadmin, 401/403 en JSON), apagar/encender con motivo, versión vista
     (409 si otra persona ya lo cambió), auditoría, vista previa con datos de mentira, «Enviarme una prueba» solo a quien está
     conectado, cifras de 30 días, el Historial («Apagado: «Nombre»» y filtro por clave) y las reglas de UI del proyecto
     (sin popups nativos, tabla paginada, marcar todo, nombres completos, 44 px / 16 px en el celular).
  8. Los avisos que envían FedEx y SimpliRoute (fase 3): su interruptor y el general, sin la llave del módulo, fail-open.
  9. Un correo cuyo ENVÍO vive en otra rama («Falta de stock: aviso a bodega», `ot_que_paso_bodega`): el catálogo ya lo trae, la
     puerta lo resuelve por su evento aunque el envío aún no lleve `correo=`, y la regla de cobertura lo acepta tal como llegará.

Sin BD ni Flask: las funciones de app.py se extraen con `ast` (importar app.py exige BD y credenciales) y se ejecutan en un ambiente de
mentira. Correr con:  py -m unittest tests.test_comm_control_correos
"""
import ast
import datetime as _dtm
import html as _html
import json
import os
import re
import shutil
import subprocess
import types
import unittest
import warnings
from unittest import mock

import comm_correos_registro as R
import ot_correos_tecnicos as OTC
from tests import _fuente_codigo as FC
from tests import _mundo_correos as MC

RAIZ = FC.RAIZ


def _leer(rel):
    return FC.texto(rel)


def _extraer(amb, funcs=(), consts=()):
    """Ejecuta en `amb` las funciones y constantes de nivel módulo de app.py indicadas (sin decoradores)."""
    return FC.extraer(amb, "app.py", funcs, consts)


def _llaves_validas():
    """`_MODULOS_KILLSWITCH_VALIDOS` leído del código de app.py."""
    return tuple(FC.valor_literal("app.py", "_MODULOS_KILLSWITCH_VALIDOS"))

# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  1) EL REGISTRO
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestRegistro(unittest.TestCase):
    def test_el_registro_se_valida_solo(self):
        self.assertEqual(R.validar(_llaves_validas()), [])

    def test_claves_unicas_y_campos_completos(self):
        claves = [c["clave"] for c in R.REGISTRO]
        self.assertEqual(len(claves), len(set(claves)))
        for c in R.REGISTRO:
            for campo in ("clave", "nombre", "modulo_ui", "llave", "destinatario", "disparo", "modo", "tipo"):
                self.assertTrue(str(c[campo]).strip(), f"{c['clave']}: falta {campo}")

    def test_la_llave_de_cada_correo_existe_incluida_ot_tecnicos(self):
        validas = _llaves_validas()
        self.assertIn("ot_tecnicos", validas)
        for c in R.REGISTRO:
            self.assertIn(c["llave"], validas, c["clave"])
        self.assertEqual({c["llave"] for c in R.REGISTRO if c["clave"].startswith("ot_tec_")}, {"ot_tecnicos"})

    def test_el_destinatario_es_un_valor_conocido_y_responsable_ot_es_valido(self):
        self.assertIn("responsable_ot", R.DESTINATARIOS)
        self.assertIn("bodega", R.DESTINATARIOS)
        self.assertEqual(set(R.DESTINATARIOS_ORDEN), set(R.DESTINATARIOS))
        for c in R.REGISTRO:
            self.assertIn(c["destinatario"], R.DESTINATARIOS, c["clave"])
        # el correo a bodega de Retiros se identifica como «bodega»
        self.assertEqual(R.por_clave("ret_aviso_bodega")["destinatario"], "bodega")

    def test_las_ordenes_de_trabajo_forman_su_propio_grupo_separado(self):
        grupos = {c["modulo_ui"] for c in R.listar()}
        self.assertIn("ot", grupos)
        self.assertEqual({"ot", "tickets", "mantenciones", "retiros", "transporte", "cotizaciones", "prospectos",
                          "repuestos", "catalogo", "accesos"}, grupos)
        for k in ("ot_creada_cliente", "ot_tec_asignada", "ot_tec_agenda_manana", "ot_pendiente_aprobacion",
                  "ot_informe_servicio", "ot_anexo_envio", "ot_acta_recepcion"):
            self.assertEqual(R.por_clave(k)["modulo_ui"], "ot", k)
        for k in ("tk_creacion", "tk_asignacion"):
            self.assertEqual(R.por_clave(k)["modulo_ui"], "tickets", k)
        for k in ("mant_plan_propuesto", "mant_recordatorio_visita"):
            self.assertEqual(R.por_clave(k)["modulo_ui"], "mantenciones", k)

    def test_los_tres_correos_de_seguridad_son_protegidos_y_solo_ellos(self):
        self.assertEqual({c["clave"] for c in R.REGISTRO if c["protegido"]},
                         {"acceso_olvido_clave", "acceso_invitacion", "acceso_cambio_clave"})

    def test_las_autorizaciones_son_sensibles(self):
        self.assertTrue(R.por_clave("ot_autorizacion_pendiente")["sensible"])
        self.assertTrue(R.por_clave("ot_pendiente_aprobacion")["sensible"])

    def test_los_cuatro_avisos_de_terceros_y_las_dos_herramientas(self):
        self.assertEqual(set(R.claves(tipo="externo")),
                         {"ext_fedex_notif_envio", "ext_fedex_notif_retiro", "ext_fedex_export_email", "ext_simpliroute_contacto"})
        self.assertEqual(set(R.HERRAMIENTAS), {"comm_prueba_manual", "comm_reenviar_historial"})
        # las herramientas no son tarjetas
        self.assertFalse(set(R.HERRAMIENTAS) & {c["clave"] for c in R.listar()})

    def test_cada_plantilla_del_registro_existe_en_el_editor_de_comunicaciones(self):
        js = _leer("static/comunicaciones_index.js")
        listas = {"transporte": "TPL_ESTADOS_TRANSPORTE", "transporte_cotizaciones": "TPL_ESTADOS_COTIZACIONES_TRANSPORTE",
                  "retiros": "TPL_ESTADOS_RETIROS", "mantenciones": "TPL_ESTADOS_MANTENCIONES",
                  "prospectos": "TPL_ESTADOS_PROSPECTOS", "comunicacion_interna": "TPL_ESTADOS_COMUNICACION_INTERNA",
                  "tickets": "TPL_ESTADOS_TICKETS", "catalogo": "TPL_ESTADOS_CATALOGO"}
        bloques = {}
        for modulo, nombre in listas.items():
            i = js.index(f"const {nombre} = [")
            bloques[modulo] = js[i:js.index("\n];", i)]
        for c in R.REGISTRO:
            if c["plantilla"] is None:
                continue
            modulo, estado = c["plantilla"]
            self.assertIn(modulo, bloques, f"{c['clave']}: módulo de plantilla {modulo}")
            self.assertRegex(bloques[modulo], r"key:\s*'" + re.escape(estado) + "'",
                             f"{c['clave']}: la plantilla {modulo}/{estado} no está en el editor")

    def test_clave_invalida_o_repetida_se_detecta(self):
        malo = R.REGISTRO + (dict(R.REGISTRO[0]),)
        original = R.REGISTRO
        try:
            R.REGISTRO = malo
            problemas = R.validar()
        finally:
            R.REGISTRO = original
        self.assertTrue(any("repetida" in p for p in problemas))


class TestTablasDeTraduccion(unittest.TestCase):
    """Cada tabla «evento/estado/kind → clave» del código cae en claves registradas, con la llave y la plantilla correctas."""

    def test_avisos_a_tecnicos(self):
        self.assertEqual(set(OTC.CORREO_DE), set(OTC.EVENTOS))
        for ev, clave in OTC.CORREO_DE.items():
            c = R.por_clave(clave)
            self.assertIsNotNone(c, clave)
            self.assertEqual(c["llave"], OTC.MODULO_LLAVE)
            self.assertEqual(c["plantilla"], (OTC.MODULO_PLANTILLAS, OTC.PLANTILLA_DE[ev]), clave)
        agenda = R.por_clave(OTC.CORREO_AGENDA)
        self.assertEqual(agenda["plantilla"], (OTC.MODULO_PLANTILLAS, OTC.PLANTILLA_AGENDA))
        self.assertEqual(agenda["llave"], OTC.MODULO_LLAVE)
        self.assertEqual(R.por_clave(OTC.CORREO_VISITA_COPIA)["llave"], OTC.MODULO_LLAVE)
        # una plantilla apagada en el editor hace que NO salga (así lo hace _ot_tec_plantilla_apagada)
        for clave in list(OTC.CORREO_DE.values()) + [OTC.CORREO_AGENDA]:
            self.assertEqual(R.por_clave(clave)["plantilla_apagada"], "no_sale", clave)

    def test_retiros_kind_a_clave(self):
        est = FC.valor_literal_anidado("pickups_module.py", "_KIND_TO_ESTADO")
        cor = FC.valor_literal_anidado("pickups_module.py", "_KIND_TO_CORREO")
        self.assertEqual(set(est), set(cor), "todo kind de notify() necesita su clave")
        for kind, clave in cor.items():
            c = R.por_clave(clave)
            self.assertIsNotNone(c, f"{kind}: {clave}")
            self.assertEqual(c["llave"], "retiros", clave)
            if est[kind]:
                self.assertEqual(c["plantilla"], ("retiros", est[kind]), f"{kind}: {clave}")
            else:
                self.assertIsNone(c["plantilla"], f"{kind}: {clave}")

    def test_tickets_slug_a_clave(self):
        defaults = set(FC.valor_literal_anidado("tickets_module.py", "_TK_LIFECYCLE_DEFAULTS", solo_claves=True))
        corr = FC.valor_literal_anidado("tickets_module.py", "_TK_CORREO_DE_SLUG")
        self.assertEqual(defaults, set(corr), "todo estado del ciclo de vida necesita su clave")
        for slug, clave in corr.items():
            c = R.por_clave(clave)
            self.assertIsNotNone(c, clave)
            self.assertEqual(c["llave"], "tickets")
            self.assertEqual(c["plantilla"], ("tickets", slug), clave)

    def test_accesos_plantilla_a_clave(self):
        tabla = FC.valor_literal("app.py", "_ACCESO_CORREO_DE_PLANTILLA")
        self.assertEqual(set(tabla), {"usuario_nuevo", "cambio_clave", "olvido_contrasena"})
        for estado, clave in tabla.items():
            c = R.por_clave(clave)
            self.assertTrue(c["protegido"], clave)
            self.assertEqual(c["plantilla"], ("comunicacion_interna", estado), clave)
            self.assertEqual(c["llave"], "comunicacion_interna")

    def test_transporte_los_seis_estados_se_resuelven_por_el_evento(self):
        """`_tr_notificar_cliente` NO se toca (otras 3 pruebas la comparan con origin/main): su evento `tracking_<estado>` se
        resuelve a la clave con `inferir()`. Se cruza con el `_slug` (plantilla) que usa esa misma función."""
        fn = ast.parse(FC.funcion_top("app.py", "_tr_notificar_cliente")).body[0]
        slug = None
        for n in ast.walk(fn):
            if isinstance(n, ast.Assign) and isinstance(n.targets[0], ast.Name) and n.targets[0].id == "_slug":
                # {...}.get(estado, 'entregado')  → el dict de estado → plantilla
                slug = ast.literal_eval(n.value.func.value)
        # el evento se arma exactamente así en el código (si cambia, esta prueba avisa)
        self.assertIn("evento=f\"tracking_{estado.lower().replace(' ','_')}\"", FC.funcion_top("app.py", "_tr_notificar_cliente"))
        self.assertEqual(set(slug), {"En preparación", "Entregado a transporte", "En ruta", "Entregado", "Entrega fallida",
                                     "Problema"})
        esperado = {"En preparación": "tr_en_preparacion", "Entregado a transporte": "tr_despachado", "En ruta": "tr_en_ruta",
                    "Entregado": "tr_entregado", "Entrega fallida": "tr_entrega_fallida", "Problema": "tr_entrega_fallida"}
        for estado, plantilla in slug.items():
            evento = f"tracking_{estado.lower().replace(' ', '_')}"
            clave = R.inferir("transporte", evento)
            self.assertEqual(clave, esperado[estado], estado)
            c = R.por_clave(clave)
            self.assertEqual(c["llave"], "transporte")
            self.assertEqual(c["plantilla"], ("transporte", plantilla), estado)
    def test_inferir_no_adivina_lo_ambiguo(self):
        # Todo Retiros sale con evento «pickup_created» y «visita_agendada» sirve a tres correos distintos: sin clave explícita no
        # se adivina nada (sale con el general y la llave, como hoy).
        self.assertIsNone(R.inferir("retiros", "pickup_created"))
        self.assertIsNone(R.inferir("mantenciones", "visita_agendada"))
        self.assertIsNone(R.inferir("comunicacion_interna", "password_access"))
        self.assertIsNone(R.inferir("general", None))
        self.assertEqual(R.inferir("tickets", "retry_ticket_resuelto"), "tk_resuelto")       # reintento del Historial
        self.assertEqual(R.inferir("tickets", "prueba_tk_creacion"), "tk_creacion")          # prueba de la pantalla
        self.assertIsNone(R.inferir("tickets", "prueba_clave_que_no_existe"))


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  2) COBERTURA: toda llamada de envío lleva una clave del registro
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
# Funciones que ENVUELVEN el envío y la llave de paso (`modulo=`) con que salen. `_send_ilus_email` toma su `modulo=` del argumento.
ENVOLTORIOS = {"_send_ilus_email": None, "_mant_send_email_multi": "mantenciones", "_send_pickup_email_multi": "retiros",
               "_prospecto_correo_por_ticket": "tickets", "_ot_tec_enviar": "ot_tecnicos"}

# Llamadas cuya clave NO es una constante en el sitio de la llamada. Cada una dice por qué es segura y qué otra prueba la cubre.
EXCEPCIONES_DINAMICAS = {
    ("app.py", "_ot_tec_despachar"):
        "correo=_otc.CORREO_DE[ev]: tabla de ot_correos_tecnicos (TestTablasDeTraduccion.test_avisos_a_tecnicos)",
    ("app.py", "_ot_tec_barrido_agenda"): "correo=_otc.CORREO_AGENDA: constante de ot_correos_tecnicos",
    ("app.py", "mant_visita_enviar_email"): "correo=_otc.CORREO_VISITA_COPIA: constante de ot_correos_tecnicos",
    ("app.py", "_send_password_access_email"):
        "correo=_clave_correo sale de _ACCESO_CORREO_DE_PLANTILLA (TestTablasDeTraduccion.test_accesos_plantilla_a_clave)",
    ("app.py", "comm_log_reintentar"):
        "Reintentar sale con la clave del correo ORIGINAL (email_log.correo_clave) — prueba TestReintentar",
    ("pickups_module.py", "notify"): "correo=_KIND_TO_CORREO.get(kind) (TestTablasDeTraduccion.test_retiros_kind_a_clave)",
    ("tickets_module.py", "_tk_notificar_lifecycle"):
        "correo=_TK_CORREO_DE_SLUG.get(estado_slug) (TestTablasDeTraduccion.test_tickets_slug_a_clave)",
}

# Llamadas que NO llevan `correo=` a propósito, con su razón.
EXCEPCIONES_SIN_CLAVE = {
    ("app.py", "_tr_notificar_cliente"):
        "su evento tracking_<estado> se resuelve con inferir() (test_transporte_los_seis_estados_se_resuelven_por_el_evento); "
        "la función no se toca: test_correo_cliente_no_miente_al_subir, test_manifiesto_sin_lineas_zz_confirmado y "
        "test_pantalla_despacho_semaforo la comparan con origin/main",
}


# Correos cuyo ENVÍO vive en otra rama que todavía no se mezcla: el catálogo ya los trae (así el interruptor existe el día de la
# mezcla y nada falla) y por eso, mientras esta rama no los envía, no se exige que una llamada use su clave. Cuando el envío llegue
# aparece solo en el código y TestCobertura lo revisa como a cualquier otro; la entrada puede quitarse de aquí (la prueba pasa con o
# sin ella). Cada una dice dónde está su envío.
ENVIO_EN_OTRA_RAMA = {
    "ot_que_paso_bodega":
        "«¿Qué pasó?» de la OT (falta de stock → aviso interno a bodega): su envío está en la rama de esa función, aún sin mezclar",
}


def _archivos_de_codigo():
    out = []
    for nombre in sorted(os.listdir(RAIZ)):
        if nombre.endswith(".py") and not nombre.startswith("test_"):
            out.append(nombre)
    return out


def _constantes_de(nodo):
    """Todas las claves constantes que puede tomar `nodo` (str, o IfExp con str en ambas ramas); None si no es constante."""
    if isinstance(nodo, ast.Constant) and isinstance(nodo.value, str):
        return [nodo.value]
    if isinstance(nodo, ast.IfExp):
        a, b = _constantes_de(nodo.body), _constantes_de(nodo.orelse)
        return (a + b) if (a is not None and b is not None) else None
    return None


def _clave_por_evento(kw):
    """La clave que la puerta deduce SOLA cuando una llamada no declara `correo=` pero su `evento=` es una constante que
    `inferir()` conoce (`_EVENTO_A_CLAVE` del registro). None si no hay evento constante o no hay una sola lectura posible."""
    ev = kw.get("evento")
    if isinstance(ev, ast.Constant) and isinstance(ev.value, str):
        return R.inferir(None, ev.value)
    return None


def _problemas_de_claves(llamadas):
    """Qué llamadas de envío NO dejan clara su clave: ni `correo=` registrada, ni excepción con razón, ni evento constante que el
    registro resuelve solo. Cada llamada: (archivo, linea, envoltorio, funcion, params, kw)."""
    problemas = []
    for archivo, linea, envoltorio, fn, params, kw in llamadas:
        donde = f"{archivo}:{linea} {envoltorio}() en {fn}"
        if "correo" not in kw:
            if (archivo, fn) in EXCEPCIONES_SIN_CLAVE:
                continue
            if _clave_por_evento(kw):
                continue           # p. ej. el envío de otra rama que aún no pasa correo=: la puerta lo resuelve por su evento
            problemas.append(f"{donde}: falta correo= (agrégalo; REGLA #33 de CLAUDE.md)")
            continue
        valor = kw["correo"]
        claves = _constantes_de(valor)
        if claves is not None:
            for k in claves:
                if R.por_clave(k) is None:
                    problemas.append(f"{donde}: la clave {k!r} no está en el registro (comm_correos_registro.py)")
            continue
        if isinstance(valor, ast.Name) and valor.id in params:
            continue                                  # el envoltorio la recibe de quien lo llama
        if (archivo, fn) in EXCEPCIONES_DINAMICAS:
            continue
        problemas.append(f"{donde}: la clave no es una constante ({ast.unparse(valor)}); declárala en EXCEPCIONES_DINAMICAS "
                         f"con su razón o pásala como constante")
    return problemas


def _problemas_de_llave(llamadas):
    """Qué llamadas salen por una llave de paso distinta de la que el registro declara para su clave (la explícita `correo=` o la
    deducida del evento constante)."""
    problemas = []
    for archivo, linea, envoltorio, fn, params, kw in llamadas:
        if "correo" in kw:
            claves = _constantes_de(kw["correo"])
        else:
            deducida = _clave_por_evento(kw)
            claves = [deducida] if deducida else None
        if not claves:
            continue
        for k in claves:
            c = R.por_clave(k)
            if c is None:
                continue
            if ENVOLTORIOS[envoltorio] is not None:
                esperado = ENVOLTORIOS[envoltorio]
            elif "modulo" in kw and isinstance(kw["modulo"], ast.Constant):
                esperado = kw["modulo"].value
            elif "modulo" not in kw and "**" in kw:
                continue                         # el modulo viaja dentro de un **dict: no se ve en el sitio de la llamada
            elif "modulo" not in kw:
                esperado = "general"             # sin modulo= ni evento conocido, la puerta lo manda a «general»
            else:
                continue
            if c["llave"] != esperado:
                problemas.append(f"{archivo}:{linea} {envoltorio}(): {k} sale por la llave {esperado!r} pero el registro dice "
                                 f"{c['llave']!r}")
    return problemas


def _llamadas_de_envio():
    """[(archivo, linea, envoltorio, funcion_contenedora, params_de_la_funcion, kw {nombre: nodo})] de TODO el código.
    Se lee por TOKENS (tests/_fuente_codigo.py): el árbol completo de app.py pesa varios GB."""
    resultado = []
    for archivo in _archivos_de_codigo():
        for ll in FC.llamadas(archivo, set(ENVOLTORIOS)):
            kw = {}
            for nombre, valor in ll["kw"].items():
                try:
                    kw[nombre] = ast.parse(valor.strip(), mode="eval").body
                except SyntaxError:
                    kw[nombre] = ast.Name(id="<expresion no analizable: " + valor + ">", ctx=ast.Load())
            resultado.append((archivo, ll["linea"], ll["nombre"], ll["funcion"], ll["params"], kw))
    return resultado


class TestCobertura(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.llamadas = _llamadas_de_envio()

    def test_se_encontraron_las_llamadas_conocidas(self):
        # Si este número baja a casi cero el escáner dejó de ver el código y la prueba de abajo no estaría probando nada.
        self.assertGreater(len(self.llamadas), 55)
        archivos = {l[0] for l in self.llamadas}
        self.assertTrue({"app.py", "pickups_module.py", "tickets_module.py", "catalogo_module.py",
                         "logistica_cotizaciones.py"} <= archivos)

    def test_toda_llamada_lleva_una_clave_registrada_o_esta_en_las_excepciones(self):
        problemas = _problemas_de_claves(self.llamadas)
        self.assertEqual(problemas, [], "\n".join(problemas))

    def test_las_excepciones_siguen_existiendo(self):
        """Una excepción que ya no corresponde a ninguna llamada es una mentira en la lista."""
        vistas = {(a, fn) for a, _, _, fn, _, kw in self.llamadas}
        for (archivo, fn), razon in list(EXCEPCIONES_DINAMICAS.items()) + list(EXCEPCIONES_SIN_CLAVE.items()):
            self.assertIn((archivo, fn), vistas, f"{archivo}:{fn} ya no tiene llamadas de envío ({razon[:50]})")
            self.assertGreater(len(razon), 20)

    def test_el_modulo_de_cada_llamada_es_la_llave_del_registro(self):
        problemas = _problemas_de_llave(self.llamadas)
        self.assertEqual(problemas, [], "\n".join(problemas))

    def test_ninguna_clave_ilus_del_registro_queda_sin_usar(self):
        """Un correo registrado que nada envía es una tarjeta que miente (se vería un interruptor que no apaga nada).
        Excepción declarada: los correos de `ENVIO_EN_OTRA_RAMA` (su envío aún no está en esta rama)."""
        fuentes = "\n".join(_leer(a) for a in _archivos_de_codigo() if a != "comm_correos_registro.py")
        inferibles = {"tr_en_preparacion", "tr_despachado", "tr_en_ruta", "tr_entregado", "tr_entrega_fallida"}   # por evento
        for c in R.listar("ilus"):
            k = c["clave"]
            if k in inferibles or k in ENVIO_EN_OTRA_RAMA:
                continue
            self.assertRegex(fuentes, r'["\']' + re.escape(k) + r'["\']', f"{k}: ninguna llamada de envío usa esta clave")

    def test_el_envio_directo_no_se_hace_fuera_de_la_puerta(self):
        """Solo estos lugares llaman a `_send_ilus_email_real`: la puerta, el drenado de la cola y «Enviar prueba» (herramienta).
        Desde la fase 2 también «Enviarme una prueba» de Control de correos (`comm_correos_prueba`): un correo de muestra SOLO a quien está
        conectado; no puede pasar por la puerta porque la puerta lo frenaría con el interruptor del mismo correo que se está probando."""
        lugares = set()
        for archivo in _archivos_de_codigo():
            for ll in FC.llamadas(archivo, {"_send_ilus_email_real"}):
                lugares.add((archivo, ll["funcion"]))
        base = {("app.py", "_send_ilus_email"), ("app.py", "_email_queue_send_job"), ("app.py", "_bg_send_email")}
        self.assertLessEqual(base, lugares, lugares)
        self.assertLessEqual(lugares, base | {("app.py", "comm_correos_prueba")}, lugares)

# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  3) LA PUERTA ÚNICA (funciones REALES de app.py en un ambiente de mentira)
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
FUNCS_PUERTA = ("_correo_ctrl_invalidar", "_correo_ctrl_estado", "_correo_resolver_clave", "_correo_permitido",
                "_correo_bloqueo_previo", "_email_marcar_bloqueo", "_email_ultimo_bloqueo", "_email_motivo_bloqueo_txt",
                "_email_resultado", "_modulo_desde_evento", "_send_ilus_email", "_email_queue_send_job")
CONSTS_PUERTA = ("_CORREO_CTRL_CACHE", "_CORREO_CTRL_TTL", "_CORREO_AVISADOS", "_KS_MODULO_TXT")


class _Cursor:
    def __init__(self, mundo):
        self.m = mundo

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def execute(self, sql, params=()):
        self.m.sql.append((sql, params))
        if "INSERT INTO comm_email_jobs" in sql:
            if not self.m.columna_cola and "correo_clave" in sql:
                raise RuntimeError("(1054, \"Unknown column 'correo_clave' in 'field list'\")")
            if not self.m.columna_ref and "ref_json" in sql:
                raise RuntimeError("(1054, \"Unknown column 'ref_json' in 'field list'\")")
            self.m.cola.append((sql, params))


class _Conn:
    def __init__(self, mundo):
        self.m = mundo

    def cursor(self):
        return _Cursor(self.m)

    def commit(self):
        pass


class Mundo:
    """Base de mentira (comm_correo_control + cola) y correo de mentira. La puerta y la cola son las funciones REALES."""

    def __init__(self, filas=None, global_on=True, llaves_cerradas=(), columna_cola=True, columna_ref=True):
        self.filas = dict(filas or {})          # clave → activo (0/1), como en comm_correo_control
        self.global_on = global_on
        self.llaves_cerradas = set(llaves_cerradas)
        self.columna_cola = columna_cola
        self.columna_ref = columna_ref
        self.db_cae = False
        self.lecturas = 0
        self.salieron, self.logs, self.sql, self.cola, self.avisos = [], [], [], [], []
        self.jobs_estado = []
        amb = {"json": json, "print": lambda *a, **k: self.avisos.append(" ".join(str(x) for x in a)),
               "g": types.SimpleNamespace(), "_ccr": R,
               "mysql_fetchall": self._fetchall, "mysql_execute": lambda q, p=(): self.jobs_estado.append((q, p)),
               "comm_is_enabled": lambda canal: self.global_on,
               "_modulo_canal_bloqueado": lambda mod, canal: mod in self.llaves_cerradas,
               "_email_log": lambda *a, **k: self.logs.append((a, k)) or 1,
               "_email_log_norm_ref": lambda ref, asunto: (ref.get("tipo"), ref.get("id"), ref.get("codigo")),
               "get_db": lambda: _Conn(self), "_email_queue_kick": lambda: None,
               "_email_adjuntos_para_cola": lambda adj: None, "_email_adjuntos_de_cola": lambda adj: adj,
               "_send_ilus_email_real": self._real}
        _extraer(amb, FUNCS_PUERTA, CONSTS_PUERTA)
        self.amb = amb

    def _fetchall(self, sql, params=()):
        self.lecturas += 1
        if self.db_cae:
            raise RuntimeError("la base no responde")
        return [{"clave": k, "activo": a, "motivo": "", "actualizado_por": "Daniel", "actualizado_en": None}
                for k, a in self.filas.items()]

    def _real(self, to, asunto, html, **kw):
        self.salieron.append({"to": to, "asunto": asunto, **kw})
        return True

    def enviar(self, **kw):
        kw.setdefault("modulo", "retiros")
        return self.amb["_send_ilus_email"]("cliente@ejemplo.cl", "Asunto", "<p>x</p>", **kw)

    def estados_log(self):
        return [(k.get("correo"), a[3]) for a, k in self.logs]       # (clave, estado)


class TestOrdenDeRevision(unittest.TestCase):
    """general > llave del módulo > interruptor del correo > cola. Nada de esto llama jamás al envío real."""

    def test_correo_apagado_no_llama_al_envio_real_y_queda_bloqueado_con_su_motivo(self):
        m = Mundo(filas={"ret_cita_confirmada": 0})
        ok = m.enviar(evento="pickup_created", correo="ret_cita_confirmada")
        self.assertFalse(ok)
        self.assertEqual(m.salieron, [])
        (args, kw), = m.logs
        self.assertEqual(args[3], "bloqueado")
        self.assertEqual(kw["correo"], "ret_cita_confirmada")
        self.assertIn("Apagado en Control de correos", args[4] if len(args) > 4 else kw["error_msg"])
        self.assertIn("Cita de retiro confirmada", kw["error_msg"])
        self.assertEqual(kw["metadata"], {"modulo": "retiros", "bloqueo": "correo", "correo": "ret_cita_confirmada"})
        self.assertEqual(kw["html"], "<p>x</p>")            # el cuerpo se guarda, como con las llaves
        r = m.amb["_email_resultado"](ok)
        self.assertEqual((r["causa"], r["correo"]), ("apagado", "ret_cita_confirmada"))
        self.assertIn("«Cita de retiro confirmada» está apagado en Comunicaciones → Control de correos", r["mensaje"])

    def test_el_interruptor_general_manda_sobre_todo(self):
        m = Mundo(global_on=False, filas={"ret_cita_confirmada": 0})
        self.assertFalse(m.enviar(correo="ret_cita_confirmada"))
        self.assertEqual(m.salieron, [])
        (args, kw), = m.logs
        self.assertEqual(args[3], "bloqueado")
        self.assertIn("kill switch global", kw["error_msg"])
        self.assertEqual(kw["correo"], "ret_cita_confirmada")                 # igual cuenta como frenado de ese correo
        self.assertEqual(m.amb["_email_ultimo_bloqueo"]()["tipo"], "global")

    def test_la_llave_del_modulo_va_antes_que_el_interruptor_del_correo(self):
        m = Mundo(llaves_cerradas={"retiros"}, filas={"ret_cita_confirmada": 0})
        self.assertFalse(m.enviar(correo="ret_cita_confirmada"))
        self.assertEqual(m.salieron, [])
        self.assertEqual(m.amb["_email_ultimo_bloqueo"]()["tipo"], "modulo")
        self.assertIn("Llave de paso cerrada", m.logs[0][1]["error_msg"])

    def test_con_todo_abierto_sale_y_queda_enviado_con_su_clave(self):
        m = Mundo()
        self.assertTrue(m.enviar(correo="ret_cita_confirmada", evento="pickup_created"))
        self.assertEqual(len(m.salieron), 1)
        self.assertNotIn("correo", m.salieron[0])                 # la clave NUNCA llega al envío real (no la conoce)
        self.assertEqual(m.estados_log(), [("ret_cita_confirmada", "enviado")])
        self.assertIsNone(m.amb["_email_ultimo_bloqueo"]())

    def test_un_correo_apagado_no_frena_a_los_demas(self):
        m = Mundo(filas={"ret_cita_confirmada": 0})
        self.assertFalse(m.enviar(correo="ret_cita_confirmada"))
        self.assertTrue(m.enviar(correo="ret_propuesta"))
        self.assertTrue(m.enviar(correo="ret_en_preparacion"))
        self.assertEqual(len(m.salieron), 2)

    def test_el_bloqueo_de_un_correo_no_contamina_al_siguiente(self):
        m = Mundo(filas={"ret_cita_confirmada": 0})
        m.enviar(correo="ret_cita_confirmada")
        self.assertEqual(m.amb["_email_ultimo_bloqueo"]()["tipo"], "correo")
        m.enviar(correo="ret_propuesta")
        self.assertIsNone(m.amb["_email_ultimo_bloqueo"]())


class TestPorDefecto(unittest.TestCase):
    def test_sin_fila_en_la_tabla_todos_los_correos_salen(self):
        # Salvo los que NACEN APAGADOS (avisos internos nuevos, Daniel 26-oct-2026): esos se prueban abajo.
        m = Mundo()
        normales = [c for c in R.listar() if not c.get("nace_apagado")]
        for c in normales:
            self.assertTrue(m.enviar(modulo=c["llave"], correo=c["clave"]), c["clave"])
        self.assertEqual(len(m.salieron), len(normales))
        self.assertEqual({l[1]["correo"] for l in m.logs}, {c["clave"] for c in normales})

    def test_los_que_nacen_apagados_no_salen_sin_fila_y_se_encienden_desde_la_pantalla(self):
        # Daniel 26-oct: «agrégalos apagados para encenderlos desde el front».
        nuevos = [c for c in R.listar() if c.get("nace_apagado")]
        self.assertEqual({c["clave"] for c in nuevos}, {"inq_nueva_gestion", "inq_asignada_responsable", "inq_sla_vencido",
                                                       "hall_nuevos_gestion", "hall_tomado_responsable",
                                                       # 27-oct (Tickets 2.0): recordatorios del SLA de respuesta.
                                                       "tk_sla_por_vencer", "tk_sla_vencido"})
        for c in nuevos:
            self.assertNotEqual(c["destinatario"], "cliente", "un correo a clientes jamás se enciende desde la pantalla (REGLA #22)")
            self.assertEqual(c["llave"], "comunicacion_interna", c["clave"])
            self.assertFalse(Mundo().enviar(modulo=c["llave"], correo=c["clave"]), "sin fila: apagado · " + c["clave"])
            self.assertTrue(Mundo(filas={c["clave"]: 1}).enviar(modulo=c["llave"], correo=c["clave"]), "encendido en la pantalla · " + c["clave"])
            self.assertFalse(Mundo(filas={c["clave"]: 0}).enviar(modulo=c["llave"], correo=c["clave"]), c["clave"])
            st = R.estado_efectivo(c, switch_activo=R.switch_de_fila(c, None))
            self.assertEqual(st["efectivo"], "correo", "la pantalla lo muestra apagado por su interruptor (ámbar), no por configuración")
            self.assertTrue(st["puede_encender"], "se puede encender desde la pantalla")
        self.assertTrue(R.switch_de_fila("ot_que_paso_bodega", None), "los de siempre: sin fila, encendidos")
        self.assertFalse(R.switch_de_fila("inq_nueva_gestion", None))
        self.assertTrue(R.switch_de_fila("inq_nueva_gestion", {"activo": 1}))

    def test_clave_desconocida_sale_y_deja_un_aviso(self):
        m = Mundo(filas={"clave_que_no_existe": 0})
        self.assertTrue(m.enviar(correo="clave_que_no_existe", evento="algo"))
        self.assertTrue(any("clave desconocida" in a and "clave_que_no_existe" in a for a in m.avisos), m.avisos)
        self.assertIsNone(m.logs[0][1]["correo"])                  # no se registra una clave que no existe

    def test_sin_clave_y_sin_evento_conocido_sale_con_las_llaves_de_siempre_y_avisa_una_vez(self):
        m = Mundo()
        self.assertTrue(m.enviar(evento="pickup_created"))
        self.assertTrue(m.enviar(evento="pickup_created"))
        self.assertEqual(sum("sin clave evento=pickup_created" in a for a in m.avisos), 1)       # un aviso por evento y proceso
        self.assertEqual(len(m.salieron), 2)

    def test_sin_clave_se_deduce_del_evento_cuando_hay_una_sola_lectura(self):
        """Así salen los seis estados de Transporte (`_tr_notificar_cliente` no se toca)."""
        m = Mundo(filas={"tr_entrega_fallida": 0})
        self.assertFalse(m.enviar(modulo="transporte", evento="tracking_problema"))
        self.assertFalse(m.enviar(modulo="transporte", evento="tracking_entrega_fallida"))
        self.assertTrue(m.enviar(modulo="transporte", evento="tracking_en_ruta"))
        self.assertEqual(m.estados_log(), [("tr_entrega_fallida", "bloqueado"), ("tr_entrega_fallida", "bloqueado"),
                                           ("tr_en_ruta", "enviado")])

    def test_un_correo_protegido_sale_aunque_su_fila_diga_apagado(self):
        m = Mundo(filas={"acceso_olvido_clave": 0, "acceso_invitacion": 0, "acceso_cambio_clave": 0})
        for k in ("acceso_olvido_clave", "acceso_invitacion", "acceso_cambio_clave"):
            self.assertTrue(m.enviar(modulo="comunicacion_interna", correo=k, evento="password_access"), k)
        self.assertEqual(len(m.salieron), 3)

    def test_un_correo_protegido_igual_respeta_el_interruptor_general_y_su_llave(self):
        """Hoy y mañana: la seguridad no se salta el general ni la llave (como hoy). La pantalla lo avisa en ámbar."""
        m = Mundo(global_on=False)
        self.assertFalse(m.enviar(modulo="comunicacion_interna", correo="acceso_olvido_clave"))
        m = Mundo(llaves_cerradas={"comunicacion_interna"})
        self.assertFalse(m.enviar(modulo="comunicacion_interna", correo="acceso_olvido_clave"))

    def test_la_clave_es_un_parametro_propio_y_nunca_viaja_en_los_kwargs(self):
        """`correo` es keyword-only de `_send_ilus_email` (si fuera a **kwargs se reenviaría a _send_ilus_email_real, que no lo
        conoce → TypeError) y la cola la guarda en su columna, no en kwargs_json."""
        fn = ast.parse(FC.funcion_top("app.py", "_send_ilus_email")).body[0]
        self.assertIn("correo", [a.arg for a in fn.args.kwonlyargs])
        self.assertEqual([a.arg for a in fn.args.args], ["to_addr", "subject", "html_body"])      # la firma posicional no cambió
        for n in ast.walk(fn):
            if isinstance(n, ast.Call) and getattr(n.func, "id", "") == "_send_ilus_email_real":
                self.assertNotIn("correo", [k.arg for k in n.keywords])
            if isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "kw_persistir" for t in n.targets):
                self.assertNotIn("correo", ast.unparse(n))              # la tupla de claves que sí viajan en kwargs_json


class TestBaseCaida(unittest.TestCase):
    def test_si_la_lectura_falla_se_usa_el_ultimo_valor_conocido(self):
        m = Mundo(filas={"ret_cita_confirmada": 0})
        self.assertFalse(m.enviar(correo="ret_cita_confirmada"))                 # lee y recuerda: apagado
        m.db_cae = True
        m.amb["_CORREO_CTRL_CACHE"]["ts"] -= 100                                 # vence el caché: tiene que volver a leer
        self.assertFalse(m.enviar(correo="ret_cita_confirmada"), "con la base caída debe seguir apagado")
        self.assertTrue(m.enviar(correo="ret_propuesta"))
        self.assertTrue(any("último valor conocido" in a for a in m.avisos))

    def test_si_nunca_se_leyo_y_la_base_esta_caida_el_correo_sale(self):
        m = Mundo(filas={"ret_cita_confirmada": 0})
        m.db_cae = True
        self.assertTrue(m.enviar(correo="ret_cita_confirmada"))                  # igual que las llaves: no se frena por un error técnico
        self.assertEqual(len(m.salieron), 1)

    def test_el_estado_se_guarda_30_segundos(self):
        m = Mundo(filas={"ret_cita_confirmada": 0})
        for _ in range(5):
            m.enviar(correo="ret_cita_confirmada")
        self.assertEqual(m.lecturas, 1)
        m.amb["_CORREO_CTRL_CACHE"]["ts"] -= 31
        m.enviar(correo="ret_cita_confirmada")
        self.assertEqual(m.lecturas, 2)
        m.amb["_correo_ctrl_invalidar"]()                                         # lo que hace el PUT de la pantalla
        m.filas["ret_cita_confirmada"] = 1
        self.assertTrue(m.enviar(correo="ret_cita_confirmada"))


class TestCola(unittest.TestCase):
    def test_se_encola_encendido_con_la_clave_en_su_columna_y_no_en_kwargs_json(self):
        m = Mundo()
        ok = m.enviar(correo="ot_tec_asignada", modulo="ot_tecnicos", asincrono=True, cc="a@b.cl", evento="ot_asignada")
        self.assertTrue(ok)
        self.assertEqual(m.salieron, [])                                          # no se envió síncrono
        (sql, params), = m.cola
        self.assertIn("correo_clave", sql)
        self.assertIn("ot_tec_asignada", params)
        kwargs_json = next(p for p in params if isinstance(p, str) and p.startswith("{"))
        self.assertNotIn("correo", json.loads(kwargs_json))                       # provocaría TypeError al drenar
        self.assertEqual(json.loads(kwargs_json), {"cc": "a@b.cl"})

    def test_apagado_antes_de_encolar_ni_se_encola(self):
        m = Mundo(filas={"ot_tec_asignada": 0})
        self.assertFalse(m.enviar(correo="ot_tec_asignada", modulo="ot_tecnicos", asincrono=True))
        self.assertEqual(m.cola, [])
        self.assertEqual(m.estados_log(), [("ot_tec_asignada", "bloqueado")])

    def test_se_apaga_mientras_esperaba_en_la_cola_y_al_sacarlo_queda_bloqueado(self):
        m = Mundo()
        self.assertTrue(m.enviar(correo="ot_tec_asignada", modulo="ot_tecnicos", asincrono=True))
        (sql, params), = m.cola
        m.filas["ot_tec_asignada"] = 0
        m.amb["_correo_ctrl_invalidar"]()
        job = {"id": 9, "destinatario": "cliente@ejemplo.cl", "asunto": "Asunto", "cuerpo": "<p>x</p>",
               "evento": "ot_asignada", "modulo": "ot_tecnicos", "intentos": 1, "max_intentos": 5,
               "kwargs_json": None, "correo_clave": "ot_tec_asignada"}
        m.amb["_email_queue_send_job"](job)
        self.assertEqual(m.salieron, [])
        self.assertTrue(any("failed" in q and "Control de correos" in str(p) for q, p in m.jobs_estado), m.jobs_estado)
        ult = m.logs[-1]
        self.assertEqual(ult[0][3], "bloqueado")
        self.assertEqual(ult[1]["metadata"]["bloqueo"], "correo")
        self.assertEqual(ult[1]["correo"], "ot_tec_asignada")
        self.assertEqual(ult[1]["job_id"], 9)

    def test_encendido_el_job_sale_y_queda_enviado_con_su_clave(self):
        m = Mundo()
        job = {"id": 3, "destinatario": "a@b.cl", "asunto": "A", "cuerpo": "x", "evento": "ot_asignada",
               "modulo": "ot_tecnicos", "intentos": 1, "max_intentos": 5, "kwargs_json": '{"cc": "z@b.cl"}',
               "correo_clave": "ot_tec_asignada"}
        m.amb["_email_queue_send_job"](job)
        self.assertEqual(len(m.salieron), 1)
        self.assertEqual(m.salieron[0]["cc"], "z@b.cl")
        self.assertNotIn("correo", m.salieron[0])
        self.assertEqual(m.estados_log(), [("ot_tec_asignada", "enviado")])

    def test_un_job_anterior_a_la_columna_se_deduce_del_evento_o_sale_como_siempre(self):
        m = Mundo(filas={"ot_tec_asignada": 0})
        job = {"id": 4, "destinatario": "a@b.cl", "asunto": "A", "cuerpo": "x", "evento": "ot_asignada",
               "modulo": "ot_tecnicos", "intentos": 1, "max_intentos": 5, "kwargs_json": None}      # sin correo_clave
        m.amb["_email_queue_send_job"](job)
        self.assertEqual(m.salieron, [])                                                            # evento ot_asignada → ot_tec_asignada
        m2 = Mundo(filas={"ot_tec_asignada": 0})
        job2 = dict(job, id=5, evento="evento_raro")
        m2.amb["_email_queue_send_job"](job2)
        self.assertEqual(len(m2.salieron), 1)                                                       # sin lectura única: como siempre

    def test_sin_la_columna_correo_clave_igual_se_encola(self):
        m = Mundo(columna_cola=False)
        self.assertTrue(m.enviar(correo="ot_tec_asignada", modulo="ot_tecnicos", asincrono=True))
        (sql, params), = m.cola
        self.assertNotIn("correo_clave", sql)
        self.assertEqual(m.salieron, [])

    def test_sin_ninguna_columna_opcional_igual_se_encola(self):
        m = Mundo(columna_cola=False, columna_ref=False)
        self.assertTrue(m.enviar(correo="ot_tec_asignada", modulo="ot_tecnicos", asincrono=True,
                                 ref={"tipo": "OT", "id": 3, "codigo": "OT-3"}))
        (sql, params), = m.cola
        self.assertNotIn("correo_clave", sql)
        self.assertNotIn("ref_json", sql)

    def test_con_referencia_y_clave_van_las_dos_columnas(self):
        m = Mundo()
        self.assertTrue(m.enviar(correo="ot_tec_asignada", modulo="ot_tecnicos", asincrono=True,
                                 ref={"tipo": "OT", "id": 3, "codigo": "OT-3"}))
        (sql, params), = m.cola
        self.assertIn("ref_json", sql)
        self.assertIn("correo_clave", sql)

    def test_sin_la_columna_correo_pero_con_ref_se_conserva_la_referencia(self):
        m = Mundo(columna_cola=False)
        self.assertTrue(m.enviar(correo="ot_tec_asignada", modulo="ot_tecnicos", asincrono=True,
                                 ref={"tipo": "OT", "id": 3, "codigo": "OT-3"}))
        (sql, params), = m.cola
        self.assertIn("ref_json", sql)
        self.assertNotIn("correo_clave", sql)


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  4) ESTADO INICIAL: el día que se sube, nada cambia
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestEstadoInicial(unittest.TestCase):
    def test_con_la_tabla_vacia_solo_no_salen_los_que_ya_estaban_apagados_por_configuracion_o_codigo(self):
        # 2026-10-21 (REGLA #37): «OT terminada: aviso al cliente» NACE apagado (OT_TERMINADA_CLIENTE_ACTIVA, REGLA #22).
        # 2026-10-21 (solicitud de repuesto, REGLA #41): los tres correos INTERNOS de Repuestos nacen apagados (REPUESTOS_CORREOS_ACTIVOS).
        hoy_apagados_por_config = {"ret_comprobante_firma", "mant_recordatorio_visita", "mant_factura_pendiente",
                                   "prospecto_oferta_auto_1", "prospecto_oferta_auto_2", "ot_terminada_cliente",
                                   "rep_nueva", "rep_vencida", "rep_lista_tecnico"}
        hoy_apagados_por_codigo = {"acceso_alerta_login"}
        for c in R.listar():
            e = R.estado_efectivo(c, environ={})
            if c["clave"] in hoy_apagados_por_config:
                self.assertEqual(e["efectivo"], "config", c["clave"])
                self.assertFalse(e["puede_encender"], c["clave"])
            elif c["clave"] in hoy_apagados_por_codigo:
                self.assertEqual(e["efectivo"], "codigo", c["clave"])
            else:
                self.assertEqual(e["efectivo"], "sale", c["clave"])

    def test_sin_RETIROS_FIRMA_CORREO_el_comprobante_sigue_sin_salir_y_con_ella_sale(self):
        c = R.por_clave("ret_comprobante_firma")
        self.assertEqual(R.estado_efectivo(c, environ={})["efectivo"], "config")
        self.assertEqual(R.estado_efectivo(c, environ={"RETIROS_FIRMA_CORREO": "true"})["efectivo"], "config")   # exacto «1»
        self.assertEqual(R.estado_efectivo(c, environ={"RETIROS_FIRMA_CORREO": "1"})["efectivo"], "sale")

    def test_lectura_de_cada_variable_igual_que_el_codigo_que_la_usa(self):
        def sale(clave, **env):
            return R.estado_efectivo(R.por_clave(clave), environ=env)["efectivo"] == "sale"
        # MANT_RECORDATORIOS_EMAIL_ENABLED: 1/true/yes/on, por defecto apagada
        self.assertFalse(sale("mant_recordatorio_visita"))
        for v in ("1", "true", "YES", " on "):
            self.assertTrue(sale("mant_recordatorio_visita", MANT_RECORDATORIOS_EMAIL_ENABLED=v), v)
        self.assertFalse(sale("mant_recordatorio_visita", MANT_RECORDATORIOS_EMAIL_ENABLED="0"))
        # PROSPECTOS_OFERTA_AUTO: 1/true/si/yes
        self.assertFalse(sale("prospecto_oferta_auto_1"))
        self.assertTrue(sale("prospecto_oferta_auto_1", PROSPECTOS_OFERTA_AUTO="si"))
        self.assertFalse(sale("prospecto_oferta_auto_1", PROSPECTOS_OFERTA_AUTO="on"))

    def test_las_condiciones_de_solo_el_disparo_automatico_no_apagan_el_correo(self):
        """RETIROS_PREP_AUTO / RETIROS_RETIRO_AUTO solo condicionan el envío automático: el manual sigue funcionando."""
        e = R.estado_efectivo(R.por_clave("ret_en_preparacion"), environ={"RETIROS_PREP_AUTO": "0"})
        self.assertEqual(e["efectivo"], "sale")
        self.assertTrue(e["puede_encender"])
        self.assertTrue(any(x["variable"] == "RETIROS_PREP_AUTO" and not x["activa"] for x in e["condiciones"]))
        e = R.estado_efectivo(R.por_clave("ret_retiro_completado"), environ={})
        self.assertEqual(e["efectivo"], "sale")
        self.assertTrue(any(x["variable"] == "RETIROS_RETIRO_AUTO" and not x["activa"] for x in e["condiciones"]))

    def test_los_pasos_se_evaluan_en_el_orden_de_la_puerta(self):
        c = R.por_clave("ret_cita_confirmada")
        e = R.estado_efectivo(c, general_activo=False, llave_bloqueada=True, switch_activo=False)
        self.assertEqual((e["efectivo"], e["frena_en"]), ("general", 1))
        self.assertEqual([p["estado"] for p in e["pasos"]], ["apagado", "saltado", "saltado", "saltado"])
        e = R.estado_efectivo(c, llave_bloqueada=True, switch_activo=False)
        self.assertEqual((e["efectivo"], e["frena_en"]), ("llave", 2))
        e = R.estado_efectivo(c, switch_activo=False)
        self.assertEqual((e["efectivo"], e["frena_en"]), ("correo", 3))
        self.assertEqual([p["estado"] for p in e["pasos"]], ["ok", "ok", "apagado", "saltado"])
        e = R.estado_efectivo(c)
        self.assertEqual((e["efectivo"], e["frena_en"]), ("sale", 0))

    def test_un_protegido_muestra_candado_y_no_se_puede_apagar_ni_aunque_diga_apagado(self):
        c = R.por_clave("acceso_olvido_clave")
        e = R.estado_efectivo(c, switch_activo=False)           # una fila «apagada» NO lo frena
        self.assertEqual(e["efectivo"], "sale")
        self.assertEqual(e["pasos"][2]["estado"], "protegido")
        self.assertFalse(e["puede_apagar"])
        self.assertFalse(e["puede_encender"])


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  5) REINTENTAR
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestReintentar(unittest.TestCase):
    def _mundo(self, fila_log, **kw):
        m = Mundo(**kw)
        _extraer(m.amb, ("comm_log_reintentar",))
        m.amb.update({"mysql_fetchone": lambda sql, p=(): dict(fila_log), "jsonify": lambda d: d,
                      "_email_log_cuerpo_leer": lambda lid: ("<p>guardado</p>", 0),
                      "_comm_render_email_document": lambda t, b, s="": b})
        return m

    FILA = {"id": 7, "destinatario": "cliente@ejemplo.cl", "asunto": "Cita confirmada", "evento": "pickup_created",
            "modulo": "retiros", "metadata": "{}", "correo_clave": "ret_cita_confirmada", "ref_tipo": None, "ref_codigo": None,
            "ref_id": None}

    def test_el_reintento_sale_con_la_clave_original_y_si_ese_correo_esta_apagado_no_sale(self):
        m = self._mundo(self.FILA, filas={"ret_cita_confirmada": 0})
        r = m.amb["comm_log_reintentar"](7)
        self.assertFalse(r["ok"])
        self.assertEqual(r["estado"], "bloqueado")
        self.assertEqual(r["bloqueo"], "correo")
        self.assertIn("Cita de retiro confirmada", r["motivo"])
        self.assertEqual(m.salieron, [])
        self.assertEqual(m.estados_log(), [("ret_cita_confirmada", "bloqueado")])

    def test_encendido_el_reintento_sale(self):
        m = self._mundo(self.FILA)
        r = m.amb["comm_log_reintentar"](7)
        self.assertTrue(r["ok"])
        self.assertEqual(r["estado"], "enviado")
        self.assertEqual(len(m.salieron), 1)
        self.assertEqual(m.estados_log(), [("ret_cita_confirmada", "enviado")])

    def test_una_fila_antigua_sin_clave_se_deduce_del_evento_retry(self):
        fila = dict(self.FILA, evento="ticket_resuelto", modulo="tickets", correo_clave=None)
        m = self._mundo(fila, filas={"tk_resuelto": 0})
        r = m.amb["comm_log_reintentar"](7)
        self.assertEqual(r["estado"], "bloqueado")
        self.assertEqual(r["bloqueo"], "correo")

    def test_la_llave_cerrada_sigue_bloqueando_el_reintento(self):
        m = self._mundo(self.FILA, llaves_cerradas={"retiros"})
        r = m.amb["comm_log_reintentar"](7)
        self.assertEqual((r["estado"], r["bloqueo"]), ("bloqueado", "modulo"))


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  6) MIGRACIÓN BARATA (REGLA #18)
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
class _BaseSim:
    """information_schema de mentira: lo que existe se actualiza cuando un DDL se ejecuta."""

    def __init__(self, tablas=("email_log", "comm_email_jobs"), columnas=(), indices=()):
        self.tablas, self.columnas, self.indices = set(tablas), set(columnas), set(indices)
        self.ddl = []

    def fetchone(self, sql, params=()):
        if "information_schema.TABLES" in sql and "comm_correo_control" in sql:
            return {"x": 1} if "comm_correo_control" in self.tablas else None
        if "information_schema.STATISTICS" in sql:
            return {"x": 1} if "idx_correo_fecha" in self.indices else None
        return None

    def fetchall(self, sql, params=()):
        if "information_schema.COLUMNS" in sql:
            return [{"t": t} for t in ("email_log", "comm_email_jobs") if (t, "correo_clave") in self.columnas]
        if "information_schema.TABLES" in sql:
            return [{"t": t} for t in ("email_log", "comm_email_jobs") if t in self.tablas]
        return []

    def ejecutar(self, sql):
        self.ddl.append(sql)
        if sql.lstrip().upper().startswith("CREATE TABLE IF NOT EXISTS COMM_CORREO_CONTROL"):
            self.tablas.add("comm_correo_control")
        m = re.match(r"\s*ALTER TABLE (\w+) ADD COLUMN correo_clave", sql, re.I)
        if m:
            self.columnas.add((m.group(1), "correo_clave"))
        if re.match(r"\s*ALTER TABLE email_log ADD INDEX idx_correo_fecha", sql, re.I):
            self.indices.add("idx_correo_fecha")


class TestMigracion(unittest.TestCase):
    def _correr(self, sim):
        class Cur:
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def execute(self, sql, p=None):
                sim.ejecutar(sql)

        class Conexion:
            def cursor(self):
                return Cur()

            def commit(self):
                pass

            def close(self):
                pass

        amb = {"print": lambda *a, **k: None, "mysql_fetchone": sim.fetchone, "mysql_fetchall": sim.fetchall,
               "get_mysql": lambda: Conexion(), "_correo_ctrl_invalidar": lambda: None,
               "_EMAIL_LOG_TIENE_CORREO": {"v": True, "ts": 0}}
        _extraer(amb, ("_ensure_comm_correo_control",))
        return amb["_ensure_comm_correo_control"]()

    def test_la_primera_corrida_crea_lo_que_falta_y_la_segunda_no_ejecuta_ningun_ddl(self):
        sim = _BaseSim()
        hechos1 = self._correr(sim)
        self.assertEqual(len(sim.ddl), 4, sim.ddl)
        self.assertTrue(hechos1)
        antes = len(sim.ddl)
        hechos2 = self._correr(sim)
        self.assertEqual(len(sim.ddl), antes, "la segunda corrida ejecutó DDL: " + str(sim.ddl[antes:]))
        self.assertEqual(hechos2, [])

    def test_cada_sentencia_lleva_una_sola_clausula_y_la_tabla_no_se_siembra(self):
        sim = _BaseSim()
        self._correr(sim)
        for sql in sim.ddl:
            self.assertNotRegex(sql, r",\s*(ADD|MODIFY|CHANGE|DROP|ALTER|RENAME)\s", sql)       # REGLA #18: una sentencia, una cláusula
            self.assertNotRegex(sql, r"INSERT|UPDATE|DELETE", sql)                             # no hay relleno ni siembra
        crear = next(s for s in sim.ddl if s.upper().startswith("CREATE TABLE"))
        for col in ("clave VARCHAR(60) NOT NULL PRIMARY KEY", "activo TINYINT(1) NOT NULL DEFAULT 1", "motivo VARCHAR(300) NULL",
                    "actualizado_por VARCHAR(190) NULL", "actualizado_en DATETIME NULL"):
            self.assertIn(col, crear)

    def test_con_todo_aplicado_a_medias_solo_agrega_lo_que_falta(self):
        sim = _BaseSim(tablas=("email_log", "comm_email_jobs", "comm_correo_control"),
                       columnas=(("email_log", "correo_clave"),))
        self._correr(sim)
        self.assertEqual(len(sim.ddl), 2)
        self.assertTrue(any("comm_email_jobs ADD COLUMN correo_clave" in s for s in sim.ddl))
        self.assertTrue(any("ADD INDEX idx_correo_fecha (correo_clave, created_at)" in s for s in sim.ddl))

    def test_la_guardia_de_ddl_del_proyecto_tambien_reconoce_cada_sentencia_como_ya_aplicada(self):
        """Doble red: aunque la comprobación propia fallara, la guardia de `mysql_execute`/`get_mysql` salta lo ya hecho."""
        amb = {"re": re}
        _extraer(amb, ("_ddl_ya_aplicado",),
                 ("_DDL_RE", "_DDLG_MULTI", "_DDLG_CREATE_TBL", "_DDLG_ADD_IDX", "_DDLG_CREATE_IDX", "_DDLG_ADD_COL",
                  "_DDLG_MOD_ENUM", "_DDLG_ENUM_VAL", "_DDLG_NO_COLUMNA", "_ddl_guard_saltados"))
        sim = _BaseSim()
        self._correr(sim)
        for sql in sim.ddl:
            self.assertTrue(amb["_ddl_ya_aplicado"](lambda q, p: {"x": 1}, sql), f"la guardia no entiende: {sql}")


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  7) LA PANTALLA «CONTROL DE CORREOS» (fase 2): los endpoints REALES de app.py en un mundo de mentira
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
K = "ret_cita_confirmada"         # un correo a clientes, automático, con plantilla en el editor y sin variables de entorno
RUTAS_PANTALLA = (("comm_correos_listar", ()), ("comm_correos_cambiar", (K,)), ("comm_correos_vista_previa", (K,)),
                  ("comm_correos_prueba", (K,)), ("comm_correos_cambios", (K,)))
VARIABLES_DEL_REGISTRO = sorted({cond["variable"] for c in R.REGISTRO for cond in c["condiciones"]})
APAGADOS_HOY = {"ret_comprobante_firma", "mant_recordatorio_visita", "mant_factura_pendiente", "prospecto_oferta_auto_1",
                "prospecto_oferta_auto_2", "acceso_alerta_login", "ot_terminada_cliente",     # 2026-10-21: el aviso al cliente de la OT terminada nace apagado
                "rep_nueva", "rep_vencida", "rep_lista_tecnico"}      # 2026-10-21: los tres correos internos de Repuestos nacen apagados (REGLA #41)


class _PantallaBase(unittest.TestCase):
    """Un mundo nuevo por prueba y SIN variables de entorno del registro (el estado inicial no depende de la máquina)."""

    def setUp(self):
        avisos = warnings.catch_warnings()                  # `datetime.utcnow()` es el estilo de todo el proyecto: su aviso de Python no es de aquí
        avisos.__enter__()
        self.addCleanup(avisos.__exit__, None, None, None)
        warnings.simplefilter("ignore", DeprecationWarning)
        patch = mock.patch.dict(os.environ)
        patch.start()
        self.addCleanup(patch.stop)
        for v in VARIABLES_DEL_REGISTRO:
            os.environ.pop(v, None)
        self.m = MC.MundoPantalla()

    # atajos
    def listar(self):
        st, d, resp = self.m.llamar("comm_correos_listar")
        self.assertEqual(st, 200, d)
        return d

    def item(self, clave, datos=None):
        return next(c for c in (datos or self.listar())["correos"] if c["clave"] == clave)

    def put(self, clave, **cuerpo):
        return self.m.llamar("comm_correos_cambiar", clave, json_body=cuerpo)

    def version(self, clave):
        return self.item(clave)["estado"]["version"]

    def nada_cambio(self, m=None):
        m = m or self.m
        self.assertEqual(m.bd.ctrl, {})
        self.assertEqual(m.bd.auditoria, [])
        self.assertEqual(m.bd.log, [])
        self.assertEqual(m.enviados_reales, [])


class TestPantallaPermisos(_PantallaBase):
    """Solo el superadministrador: ni un administrador ni ningún técnico ven ni tocan esto (REGLAS #19/#26)."""

    def test_sin_sesion_responde_401_en_json_y_no_toca_la_base(self):
        m = MC.MundoPantalla()
        m.g.user = None
        for nombre, args in RUTAS_PANTALLA:
            st, d, _ = m.llamar(nombre, *args, json_body={"activo": False})
            self.assertEqual((st, d["error_codigo"]), (401, "SIN_SESION"), nombre)
        self.assertEqual(m.bd.sql, [])
        self.nada_cambio(m)

    def test_todo_rol_que_no_es_superadmin_recibe_403_sin_tocar_nada(self):
        for rol in ("admin", "supervisor", "ejecutivo", "tecnico", "tecnico_ejecutivo", "tecnico_externo"):
            m = MC.MundoPantalla(superadmin=False, rol=rol)
            for nombre, args in RUTAS_PANTALLA:
                st, d, _ = m.llamar(nombre, *args, json_body={"activo": False, "to": "otro@cliente.cl"})
                self.assertEqual((st, d["error_codigo"]), (403, "SOLO_SUPERADMIN"), f"{rol} {nombre}")
                self.assertEqual(set(d), {"ok", "error", "error_codigo"}, "la respuesta no debe traer datos")
            self.assertEqual(m.bd.sql, [], f"{rol}: ni siquiera debe consultar la base")
            self.nada_cambio(m)

    def test_el_superadmin_entra_a_las_cinco_rutas(self):
        for nombre, args in RUTAS_PANTALLA:
            st, d, _ = self.m.llamar(nombre, *args, json_body={"activo": True})
            self.assertNotIn(st, (401, 403), nombre)

    def test_cada_ruta_de_correos_declara_el_decorador_de_superadmin_con_el_route_primero(self):
        lineas = _leer("app.py").split("\n")
        rutas = [i for i, l in enumerate(lineas) if re.match(r'@app\.route\("/api/comunicaciones/correos', l)]
        self.assertEqual(len(rutas), 5, "se esperaban las 5 rutas de la pantalla")
        metodos = []
        for i in rutas:
            decs, j = [lineas[i].strip()], i + 1
            while not lineas[j].startswith("def "):
                decs.append(lineas[j].strip())
                j += 1
            self.assertIn("@_require_superadmin_api", decs, lineas[i])
            self.assertLess(0, decs.index("@_require_superadmin_api"), "el @app.route debe ir ANTES del decorador de permiso")
            metodos.append(re.search(r'methods=\["(\w+)"\]', lineas[i]).group(1))
        self.assertEqual(sorted(metodos), ["GET", "GET", "POST", "POST", "PUT"])

    def test_el_decorador_responde_json_y_no_redirige(self):
        codigo = FC.funcion_top("app.py", "_require_superadmin_api")
        self.assertNotIn("redirect(", codigo)
        self.assertIn("401", codigo)
        self.assertIn("403", codigo)


class TestPantallaListar(_PantallaBase):
    def test_lista_los_correos_de_ilus_y_los_avisos_de_terceros_y_ninguna_herramienta(self):
        st, d, resp = self.m.llamar("comm_correos_listar")
        self.assertEqual(st, 200)
        self.assertTrue(d["ok"])
        self.assertEqual([c["clave"] for c in d["correos"]], [c["clave"] for c in R.REGISTRO if c["tipo"] in ("ilus", "externo")])
        self.assertEqual(resp.headers["Cache-Control"], "no-store, max-age=0")
        self.assertEqual({h["clave"] for h in d["herramientas"]}, set(R.HERRAMIENTAS))
        self.assertNotIn("comunicaciones", [x["clave"] for x in d["modulos"]])
        self.assertEqual(d["resumen"]["total"], len(d["correos"]))

    def test_con_la_tabla_vacia_solo_no_salen_los_que_ya_estaban_apagados_por_configuracion_o_codigo(self):
        d = self.listar()
        r = d["resumen"]
        self.assertEqual(r["apagados_usuario"], 0)
        self.assertEqual({c["clave"] for c in d["correos"] if c["estado"]["efectivo"] in ("config", "codigo")}, APAGADOS_HOY)
        self.assertEqual(r["apagados_config"], len(APAGADOS_HOY))
        # 26-oct: los avisos internos nuevos nacen apagados («por encender»); no cuentan como apagados por una persona.
        self.assertEqual(r["por_encender"], 7)   # 27-oct: + los 2 recordatorios del SLA de Tickets
        self.assertEqual(r["salen"] + r["apagados_config"] + r["por_encender"], r["total"])
        self.assertEqual(r["frenados_llave"], 0)
        self.assertEqual(self.item("acceso_alerta_login", d)["estado"]["efectivo"], "codigo")
        self.assertEqual(self.item("ret_comprobante_firma", d)["estado"]["efectivo"], "config")
        self.assertEqual(d["avisos"], [])

    def test_cada_tarjeta_trae_lo_que_el_front_necesita(self):
        d = self.listar()
        modulos, dests = {x["clave"] for x in d["modulos"]}, {x["clave"] for x in d["destinatarios"]}
        for c in d["correos"]:
            for k in ("clave", "nombre", "modulo_ui", "llave", "llave_label", "destinatario", "destinatario_label", "disparo", "modo",
                      "modo_label", "avisa_a", "adjunta", "nota", "tipo", "protegido", "sensible", "tambien_whatsapp", "plantilla",
                      "estado", "cifras"):
                self.assertIn(k, c, c["clave"])
            self.assertIn(c["modulo_ui"], modulos, c["clave"])
            self.assertIn(c["destinatario"], dests, c["clave"])
            self.assertTrue(c["nombre"].strip() and c["disparo"].strip() and c["llave_label"].strip(), c["clave"])
            e = c["estado"]
            self.assertEqual([p["n"] for p in e["pasos"]], [1, 2, 3, 4], c["clave"])
            for k in ("efectivo", "frena_en", "condiciones", "puede_encender", "puede_apagar", "apagado_en_codigo", "switch_activo",
                      "version", "motivo", "apagado_por", "apagado_en", "llave_motivo"):
                self.assertIn(k, e, c["clave"])
            self.assertTrue({"enviados", "frenados", "fallidos", "ultimo"} <= set(c["cifras"]), c["clave"])
            if c["plantilla"]:
                self.assertTrue({"modulo", "estado", "existe", "activo", "texto_apagada"} <= set(c["plantilla"]), c["clave"])

    def test_las_ordenes_de_trabajo_son_su_propio_grupo_y_van_primero(self):
        d = self.listar()
        self.assertEqual(d["modulos"][0]["clave"], "ot")
        ot = [c for c in d["correos"] if c["modulo_ui"] == "ot"]
        self.assertTrue({"cliente", "tecnico"} <= {c["destinatario"] for c in ot})
        tickets = [c for c in d["correos"] if c["modulo_ui"] == "tickets"]
        self.assertTrue(tickets and not ({c["clave"] for c in tickets} & {c["clave"] for c in ot}))

    def test_el_interruptor_de_una_persona_se_ve_con_quien_cuando_y_por_que(self):
        self.m.bd.ctrl[K] = {"clave": K, "activo": 0, "motivo": "Pruebas hasta el viernes", "actualizado_por": "Daniel Aguilar",
                             "actualizado_en": _dtm.datetime(2026, 10, 20, 18, 5, 30)}
        d = self.listar()
        e = self.item(K, d)["estado"]
        self.assertEqual((e["efectivo"], e["frena_en"], e["switch_activo"]), ("correo", 3, False))
        self.assertEqual([p["estado"] for p in e["pasos"]], ["ok", "ok", "apagado", "saltado"])
        self.assertEqual((e["apagado_por"], e["motivo"]), ("Daniel Aguilar", "Pruebas hasta el viernes"))
        self.assertEqual(e["apagado_en"], "20/10/2026 15:05")           # 18:05 UTC = 15:05 hora Chile (UTC-3 en octubre)
        self.assertEqual(e["version"], "2026-10-20 18:05:30")
        self.assertEqual(d["resumen"]["apagados_usuario"], 1)
        self.assertEqual(d["resumen"]["salen"], d["resumen"]["total"] - len(APAGADOS_HOY) - 1 - d["resumen"]["por_encender"])

    def test_la_llave_cerrada_y_el_interruptor_general_se_ven_en_el_semaforo(self):
        self.m.bd.llaves["retiros"] = {"email_bloqueado": 1, "motivo": "Revisión de Joaquín"}
        d = self.listar()
        for c in d["correos"]:
            e = c["estado"]
            if c["llave"] == "retiros" and c["clave"] not in APAGADOS_HOY and not c["protegido"]:
                self.assertEqual((e["efectivo"], e["frena_en"]), ("llave", 2), c["clave"])
                self.assertEqual([p["estado"] for p in e["pasos"]], ["ok", "apagado", "saltado", "saltado"])
                self.assertEqual(e["llave_motivo"], "Revisión de Joaquín")
        self.assertEqual(self.item("tr_en_ruta", d)["estado"]["efectivo"], "sale")
        self.assertTrue(any("Retiros" in a["texto"] for a in d["avisos"]))
        self.assertGreater(d["resumen"]["frenados_llave"], 0)
        # el interruptor general manda sobre todo
        self.m.bd.global_on = False
        d2 = self.listar()
        self.assertTrue(all(c["estado"]["efectivo"] == "general" for c in d2["correos"]))
        self.assertTrue(any("general" in a["texto"] for a in d2["avisos"]))
        self.assertTrue(any("contraseña" in a["texto"] for a in d2["avisos"]))      # los de acceso tampoco salen si el general está apagado

    def test_las_cifras_cuentan_30_dias_y_no_cuentan_pruebas_ni_envios_sin_clave(self):
        bd = self.m.bd
        for dias in (1, 4, 10):
            bd.poner_log(K, "enviado", hace_dias=dias)
        bd.poner_log(K, "bloqueado", hace_dias=2)
        bd.poner_log(K, "fallido", hace_dias=3)
        bd.poner_log(K, "enviado", hace_dias=45)                                         # fuera de la ventana
        bd.poner_log(K, "enviado", hace_dias=0.5, evento="prueba_" + K)                  # una prueba de la pantalla no cuenta
        bd.poner_log(None, "enviado", hace_dias=1)                                       # sin clave: no es de ningún correo
        d = self.listar()
        c = self.item(K, d)["cifras"]
        self.assertEqual((c["enviados"], c["frenados"], c["fallidos"]), (3, 1, 1))
        ultimo = max(f["created_at"] for f in bd.log if f["correo_clave"] == K and f["estado"] == "enviado"
                     and not f["evento"].startswith("prueba_") and f["created_at"] > _dtm.datetime.utcnow() - _dtm.timedelta(days=30))
        self.assertEqual(c["ultimo"], MC.chile_fmt(ultimo))
        r = d["resumen"]
        self.assertEqual((r["enviados_30d"], r["frenados_30d"], r["fallidos_30d"]), (3, 1, 1))
        primero = min(f["created_at"] for f in bd.log if f["correo_clave"] == K and f["created_at"] > _dtm.datetime.utcnow()
                      - _dtm.timedelta(days=30) and not f["evento"].startswith("prueba_"))
        self.assertEqual(d["contando_desde"], MC.chile_fmt(primero, "%d/%m/%Y"))          # el conteo empezó hace poco: la pantalla lo dice
        self.assertTrue(d["cifras_disponibles"])

    def test_si_hay_datos_de_los_30_dias_completos_no_dice_contando_desde(self):
        self.m.bd.poner_log(K, "enviado", hace_dias=29.5)
        d = self.listar()
        self.assertIsNone(d["contando_desde"])
        self.assertEqual(self.item(K, d)["cifras"]["enviados"], 1)

    def test_sin_la_columna_correo_clave_no_hay_cifras_pero_la_pantalla_y_los_interruptores_funcionan(self):
        self.m.bd.columna_log = False
        d = self.listar()
        self.assertFalse(d["cifras_disponibles"])
        self.assertEqual(d["resumen"]["enviados_30d"], 0)
        st, r, _ = self.put(K, activo=False, motivo="x", estado_visto="")
        self.assertEqual(st, 200, r)

    def test_si_la_base_no_se_puede_leer_la_pantalla_lo_dice_y_no_inventa_que_todo_esta_encendido(self):
        self.m.bd.fallar_lectura = True
        st, d, _ = self.m.llamar("comm_correos_listar")
        self.assertEqual(st, 500)
        self.assertFalse(d["ok"])
        self.assertNotIn("la base no responde", d["error"])             # REGLA #4: ningún detalle interno al navegador
        self.assertNotIn("RuntimeError", d["error"])

    def test_la_puerta_de_envio_sigue_abierta_aunque_la_base_caiga_y_la_pantalla_sea_estricta(self):
        self.m.bd.fallar_lectura = True
        self.assertTrue(self.m.enviar(K))
        self.assertEqual(len(self.m.enviados_reales), 1)

    def test_el_listado_no_cambia_ningun_dato(self):
        self.listar()
        self.assertEqual(self.m.bd.ctrl, {})
        self.assertEqual(self.m.bd.auditoria, [])
        self.assertTrue(all(s.lstrip().upper().startswith("SELECT") for s, _p in self.m.bd.sql), "el listado solo puede leer")


class TestPantallaCambiar(_PantallaBase):
    def test_una_clave_que_no_existe_o_es_una_herramienta_es_404(self):
        for clave in ("no_existe_este", "comm_prueba_manual", "comm_reenviar_historial", ""):
            st, d, _ = self.put(clave or "  ", activo=False)
            self.assertEqual((st, d["error_codigo"]), (404, "NO_EXISTE"), clave)
        self.nada_cambio()

    def test_la_clave_se_normaliza(self):
        st, d, _ = self.put("  RET_CITA_CONFIRMADA ", activo=False, estado_visto="")
        self.assertEqual(st, 200, d)
        self.assertIn(K, self.m.bd.ctrl)

    def test_hay_que_decir_si_se_enciende_o_se_apaga(self):
        st, d, _ = self.put(K)
        self.assertEqual((st, d["error_codigo"]), (400, "FALTA_ACTIVO"))
        self.nada_cambio()

    def test_los_correos_de_seguridad_no_tienen_interruptor(self):
        for clave in ("acceso_olvido_clave", "acceso_invitacion", "acceso_cambio_clave"):
            for activo in (False, True):
                st, d, _ = self.put(clave, activo=activo)
                self.assertEqual((st, d["error_codigo"]), (409, "PROTEGIDO"), f"{clave} {activo}")
        self.nada_cambio()

    def test_la_pantalla_solo_apaga_un_correo_que_la_configuracion_o_el_codigo_mantienen_apagado(self):
        st, d, _ = self.put("ret_comprobante_firma", activo=True)
        self.assertEqual((st, d["error_codigo"]), (409, "CON_CANDADO"))
        st, d, _ = self.put("acceso_alerta_login", activo=True)                       # apagado en el código
        self.assertEqual((st, d["error_codigo"]), (409, "CON_CANDADO"))
        self.nada_cambio()
        # apagarlo sí se puede (queda apagado aunque alguien después encienda la variable)
        st, d, _ = self.put("ret_comprobante_firma", activo=False, motivo="Por si acaso", estado_visto="")
        self.assertEqual(st, 200, d)
        ver = d["correo"]["estado"]["version"]
        # y encenderlo sigue sin poder mientras la variable esté apagada
        st, d, _ = self.put("ret_comprobante_firma", activo=True, estado_visto=ver)
        self.assertEqual((st, d["error_codigo"]), (409, "CON_CANDADO"))
        self.assertEqual(self.m.bd.ctrl["ret_comprobante_firma"]["activo"], 0)
        # con la variable encendida (decisión de Daniel, REGLA #23) la pantalla sí puede encenderlo
        with mock.patch.dict(os.environ, {"RETIROS_FIRMA_CORREO": "1"}):
            st, d, _ = self.put("ret_comprobante_firma", activo=True, estado_visto=ver)
        self.assertEqual(st, 200, d)
        self.assertEqual(self.m.bd.ctrl["ret_comprobante_firma"]["activo"], 1)

    def test_apagar_guarda_la_fila_el_motivo_y_la_auditoria(self):
        st, d, _ = self.put(K, activo=False, motivo="  Pruebas hasta el viernes  ", estado_visto="")
        self.assertEqual(st, 200, d)
        self.assertTrue(d["ok"])
        fila = self.m.bd.ctrl[K]
        self.assertEqual((fila["activo"], fila["motivo"], fila["actualizado_por"]), (0, "Pruebas hasta el viernes", "Daniel Aguilar"))
        a, = self.m.bd.auditoria
        self.assertEqual((a["action"], a["target_type"], a["target_id"]), ("comm_correo_set", "comm_correo_control", K))
        self.assertEqual(json.loads(a["details"]), {"antes": True, "despues": False, "motivo": "Pruebas hasta el viernes",
                                                    "nombre": R.nombre_de(K), "por": "Daniel Aguilar"})
        e = d["correo"]["estado"]
        self.assertEqual((e["efectivo"], e["switch_activo"], e["apagado_por"]), ("correo", False, "Daniel Aguilar"))
        self.assertEqual(e["version"], fila["actualizado_en"].strftime("%Y-%m-%d %H:%M:%S"))
        self.assertRegex(e["apagado_en"], r"^\d{2}/\d{2}/\d{4} \d{2}:\d{2}$")             # día/mes/año, hora Chile (REGLA #6)

    def test_encender_limpia_el_motivo_y_deja_constancia(self):
        st, d, _ = self.put(K, activo=False, motivo="Pruebas", estado_visto="")
        st, d, _ = self.put(K, activo=True, motivo="este motivo se ignora al encender", estado_visto=d["correo"]["estado"]["version"])
        self.assertEqual(st, 200, d)
        fila = self.m.bd.ctrl[K]
        self.assertEqual((fila["activo"], fila["motivo"]), (1, None))
        self.assertEqual(d["correo"]["estado"]["efectivo"], "sale")
        self.assertEqual(len(self.m.bd.auditoria), 2)
        det = json.loads(self.m.bd.auditoria[1]["details"])
        self.assertEqual((det["antes"], det["despues"]), (False, True))

    def test_el_motivo_se_recorta_a_300_caracteres(self):
        st, d, _ = self.put(K, activo=False, motivo="x" * 500, estado_visto="")
        self.assertEqual(len(self.m.bd.ctrl[K]["motivo"]), 300)

    def test_si_la_pantalla_vio_una_version_vieja_es_409_y_no_cambia_nada(self):
        self.put(K, activo=False, motivo="primero", estado_visto="")
        self.m.bd.ctrl[K]["actualizado_por"] = "María Soto"
        antes = dict(self.m.bd.ctrl[K])
        st, d, _ = self.put(K, activo=True, estado_visto="")                  # esta pantalla leyó cuando aún no había fila
        self.assertEqual((st, d["error_codigo"]), (409, "CAMBIO_AJENO"))
        self.assertIn("María Soto ya cambió este correo", d["error"])
        self.assertRegex(d["error"], r"el \d{2}/\d{2}/\d{4} \d{2}:\d{2}")
        self.assertEqual(self.m.bd.ctrl[K], antes)
        self.assertEqual(len(self.m.bd.auditoria), 1)

    def test_si_alguien_cambia_la_fila_entre_la_lectura_y_la_escritura_no_se_pisa(self):
        self.put(K, activo=False, motivo="primero", estado_visto="")
        leida = dict(self.m.bd.ctrl[K])                                         # lo que este request alcanzó a leer
        self.m.bd.tick()
        self.m.bd.ctrl[K].update(motivo="de otra persona", actualizado_en=self.m.bd.ahora)      # alguien lo cambió justo después
        self.m.amb["_correo_ctrl_estado"] = lambda forzar=False, estricto=False: {K: leida}
        ver = leida["actualizado_en"].strftime("%Y-%m-%d %H:%M:%S")
        st, d, _ = self.put(K, activo=True, estado_visto=ver)
        self.assertEqual((st, d["error_codigo"]), (409, "CAMBIO_AJENO"))
        self.assertIn("justo ahora", d["error"])
        self.assertEqual(self.m.bd.ctrl[K]["motivo"], "de otra persona")
        self.assertEqual(self.m.amb["_CORREO_CTRL_CACHE"]["ts"], 0.0, "debe invalidar la caché para que se relea")
        self.assertEqual(len(self.m.bd.auditoria), 1)

    def test_sin_estado_visto_igual_funciona(self):
        st, d, _ = self.put(K, activo=False)
        self.assertEqual(st, 200, d)

    def test_si_la_base_falla_al_guardar_responde_500_amable_y_no_deja_rastro(self):
        def _boom():
            raise RuntimeError("secreto interno: password=abc123")
        self.m.amb["get_db"] = _boom
        st, d, _ = self.put(K, activo=False)
        self.assertEqual(st, 500)
        self.assertNotIn("abc123", json.dumps(d))
        self.assertNotIn("password", json.dumps(d))
        self.nada_cambio()

    def test_guardado_pero_sin_poder_releer_responde_ok_y_pide_recargar(self):
        original = self.m.amb["_correo_ctrl_estado"]

        def raro(forzar=False, estricto=False):
            if estricto:
                raise RuntimeError("se cayó justo ahora")
            return original(forzar=forzar, estricto=estricto)
        self.m.amb["_correo_ctrl_estado"] = raro
        st, d, _ = self.put(K, activo=False, estado_visto="")
        self.assertEqual(st, 200)
        self.assertTrue(d["ok"] and d["recargar"] and d["correo"] is None)
        self.assertEqual(self.m.bd.ctrl[K]["activo"], 0)
        self.assertEqual(len(self.m.bd.auditoria), 1)

    def test_apagar_y_encender_cambian_lo_que_decide_la_puerta_de_envio_al_instante(self):
        self.assertTrue(self.m.enviar(K))                                           # sin fila: sale como siempre
        self.m.enviados_reales.clear()
        st, d, _ = self.put(K, activo=False, motivo="x", estado_visto="")
        self.assertFalse(self.m.enviar(K))
        self.assertEqual(self.m.enviados_reales, [])
        ultimo = self.m.bd.log[-1]
        self.assertEqual((ultimo["estado"], ultimo["correo_clave"]), ("bloqueado", K))
        self.assertIn("Apagado en Control de correos", ultimo["error_msg"])
        self.assertEqual(ultimo["metadata"]["bloqueo"], "correo")
        self.assertTrue(self.m.enviar("ret_en_preparacion", evento="x"))            # los demás siguen saliendo
        st, d, _ = self.put(K, activo=True, estado_visto=d["correo"]["estado"]["version"])
        self.assertTrue(self.m.enviar(K))
        self.assertEqual(self.m.bd.log[-1]["estado"], "enviado")


class TestPantallaPrueba(_PantallaBase):
    def test_la_prueba_va_solo_al_correo_de_quien_esta_conectado_y_lleva_PRUEBA_en_el_asunto(self):
        self.m.sembrar_plantillas_de_ejemplo()
        st, d, _ = self.m.llamar("comm_correos_prueba", K, json_body={"to": "otra.persona@cliente.cl", "destinatario": "x@y.cl",
                                                                       "email": "z@w.cl", "cc": "q@r.cl"})
        self.assertEqual(st, 200, d)
        self.assertEqual(d, {"ok": True, "enviado_a": "daniel.aguilar@sphs.cl"})
        env, = self.m.enviados_reales
        self.assertEqual(env["to"], "daniel.aguilar@sphs.cl")
        self.assertTrue(env["asunto"].startswith("[PRUEBA] ILUS · "), env["asunto"])
        self.assertNotIn("cc", env)
        self.assertNotIn("{{", env["html"])
        log, = self.m.bd.log
        self.assertEqual((log["evento"], log["estado"], log["correo_clave"], log["destinatario"], log["actor"]),
                         ("prueba_" + K, "enviado", K, "daniel.aguilar@sphs.cl", "Daniel Aguilar"))
        a, = self.m.bd.auditoria
        self.assertEqual(a["action"], "comm_correo_prueba")
        self.assertEqual(json.loads(a["details"]), {"destino": "daniel.aguilar@sphs.cl", "enviado": True})

    def test_la_prueba_no_la_frena_el_interruptor_de_ese_correo_ni_su_llave(self):
        self.put(K, activo=False, motivo="x", estado_visto="")
        self.m.bd.llaves["retiros"] = {"email_bloqueado": 1}
        self.m.bd.log.clear()
        st, d, _ = self.m.llamar("comm_correos_prueba", K)
        self.assertEqual(st, 200, d)
        self.assertEqual(len(self.m.enviados_reales), 1)

    def test_la_prueba_respeta_el_interruptor_general(self):
        self.m.bd.global_on = False
        st, d, _ = self.m.llamar("comm_correos_prueba", K)
        self.assertEqual((st, d["error_codigo"]), (409, "CORREO_APAGADO_GENERAL"))
        self.nada_cambio()

    def test_sin_un_correo_valido_en_la_sesion_no_hay_a_donde_enviar(self):
        for usuario in ("daniel", "", "a" * 200 + "@x.cl"):
            self.m.g.user["username"] = usuario
            st, d, _ = self.m.llamar("comm_correos_prueba", K)
            self.assertEqual((st, d["error_codigo"]), (400, "SIN_CORREO"), usuario)
        self.nada_cambio()

    def test_si_el_proveedor_rechaza_queda_fallido_en_el_historial_y_no_se_filtra_el_detalle(self):
        def _boom(*a, **k):
            raise RuntimeError("smtp password=SECRETO")
        self.m.amb["_send_ilus_email_real"] = _boom
        st, d, _ = self.m.llamar("comm_correos_prueba", K)
        self.assertEqual((st, d["error_codigo"]), (502, "ENVIO_FALLIDO"))
        self.assertNotIn("SECRETO", json.dumps(d))                                   # REGLA #4
        log, = self.m.bd.log
        self.assertEqual((log["estado"], log["evento"]), ("fallido", "prueba_" + K))
        self.assertIn("SECRETO", log["error_msg"])                                   # el detalle queda en el Historial, no en el navegador
        self.m.amb["_send_ilus_email_real"] = lambda *a, **k: False
        st, d, _ = self.m.llamar("comm_correos_prueba", K)
        self.assertEqual(st, 502)

    def test_la_prueba_no_cuenta_en_las_cifras(self):
        self.m.llamar("comm_correos_prueba", K)
        c = self.item(K)["cifras"]
        self.assertEqual((c["enviados"], c["frenados"], c["fallidos"], c["ultimo"]), (0, 0, 0, ""))

    def test_un_correo_sin_plantilla_en_el_editor_manda_un_texto_de_muestra(self):
        clave = next(c["clave"] for c in R.listar("ilus") if not c["plantilla"])
        st, d, _ = self.m.llamar("comm_correos_prueba", clave)
        self.assertEqual(st, 200, d)
        html = self.m.enviados_reales[0]["html"]
        self.assertIn("prueba", html.lower())
        self.assertIn(_html.escape(R.nombre_de(clave)), html)

    def test_con_el_modo_prueba_de_la_empresa_la_redireccion_la_hace_el_envio_real(self):
        """La pantalla no decide destinatarios: llama a `_send_ilus_email_real`, que aplica ILUS_COMM_TEST_TO (REGLA #22)."""
        codigo = FC.funcion_top("app.py", "comm_correos_prueba")
        self.assertIn("_send_ilus_email_real(destino", codigo)
        self.assertNotRegex(codigo, r"\brequest\b", "el destinatario jamás sale de la petición")


class TestPantallaVistaPrevia(_PantallaBase):
    def test_la_vista_previa_no_envia_nada_ni_deja_variables_sin_resolver(self):
        self.m.sembrar_plantillas_de_ejemplo()
        con_plantilla = [c for c in R.listar("ilus") if c["plantilla"]]
        self.assertGreater(len(con_plantilla), 30)
        for c in con_plantilla:
            st, d, _ = self.m.llamar("comm_correos_vista_previa", c["clave"])
            self.assertEqual(st, 200, c["clave"])
            self.assertFalse(d["sin_plantilla"], c["clave"])
            self.assertNotIn("{{", d["html"], c["clave"])
            self.assertNotIn("{{", d["asunto"], c["clave"])
            self.assertTrue(d["asunto"].startswith("ILUS · "), d["asunto"])
        self.assertEqual(self.m.enviados_reales, [])
        self.assertEqual(self.m.bd.log, [])
        self.assertEqual(self.m.bd.auditoria, [])

    def test_los_datos_son_de_mentira_y_un_token_desconocido_se_ve_como_demo(self):
        self.m.sembrar_plantillas_de_ejemplo()
        st, d, _ = self.m.llamar("comm_correos_vista_previa", K)
        self.assertIn("Juan Daniel Aguilar", d["html"])            # nombre_cliente de mentira
        self.assertIn("RET-DEMO-7788", d["html"])
        self.assertIn("[no_existe demo]", d["html"])

    def test_solo_lee_plantillas_y_nunca_toca_retiros_ordenes_ni_clientes_reales(self):
        self.m.sembrar_plantillas_de_ejemplo()
        for c in R.listar("ilus"):
            self.m.llamar("comm_correos_vista_previa", c["clave"])
        tablas = set()
        for sql, _p in self.m.bd.sql:
            tablas |= set(re.findall(r"\bFROM\s+([A-Za-z_\.]+)", sql))
        self.assertLessEqual(tablas, {"comm_templates", "email_log", "information_schema.COLUMNS"}, tablas)

    def test_sin_plantilla_en_el_editor_lo_dice_y_apunta_al_ultimo_envio(self):
        clave = next(c["clave"] for c in R.listar("ilus") if not c["plantilla"])
        st, d, _ = self.m.llamar("comm_correos_vista_previa", clave)
        self.assertTrue(d["sin_plantilla"])
        self.assertIsNone(d["ultimo_log_id"])
        self.assertIn("Historial", d["mensaje"])
        primero = self.m.bd.poner_log(clave, "enviado")
        self.m.bd.poner_log(clave, "enviado", evento="prueba_" + clave)               # las pruebas no cuentan como «último enviado»
        st, d, _ = self.m.llamar("comm_correos_vista_previa", clave)
        self.assertEqual(d["ultimo_log_id"], primero)

    def test_con_la_plantilla_sin_guardar_en_el_editor_tambien_se_trata_como_sin_plantilla(self):
        st, d, _ = self.m.llamar("comm_correos_vista_previa", K)
        self.assertTrue(d["sin_plantilla"])
        self.assertEqual(self.item(K)["plantilla"]["existe"], False)

    def test_si_la_plantilla_esta_apagada_en_el_editor_avisa_lo_que_pasa_con_ese_correo(self):
        self.m.sembrar_plantillas_de_ejemplo()
        c = R.por_clave("ot_tec_asignada")
        self.assertEqual(c["plantilla_apagada"], "no_sale")
        self.m.poner_plantilla(*c["plantilla"], "Asunto", "<p>Cuerpo {{ot}}</p>", activo=0)
        st, d, _ = self.m.llamar("comm_correos_vista_previa", "ot_tec_asignada")
        self.assertFalse(d["plantilla_activa"])
        self.assertEqual(d["aviso_plantilla_apagada"], R.PLANTILLA_APAGADA["no_sale"])
        p = self.item("ot_tec_asignada")["plantilla"]
        self.assertEqual((p["existe"], p["activo"], p["texto_apagada"]), (True, False, R.PLANTILLA_APAGADA["no_sale"]))

    def test_vista_previa_de_una_clave_que_no_existe_es_404(self):
        st, d, _ = self.m.llamar("comm_correos_vista_previa", "no_existe_este")
        self.assertEqual((st, d["error_codigo"]), (404, "NO_EXISTE"))


class TestPantallaCambios(_PantallaBase):
    def test_lista_los_ultimos_5_cambios_del_mas_nuevo_al_mas_viejo_en_hora_chile(self):
        ver = ""
        for i in range(7):
            st, d, _ = self.put(K, activo=bool(i % 2), motivo=f"m{i}" if i % 2 == 0 else "", estado_visto=ver)
            self.assertEqual(st, 200, d)
            ver = d["correo"]["estado"]["version"]
        st, d, _ = self.m.llamar("comm_correos_cambios", K)
        self.assertEqual(st, 200)
        cambios = d["cambios"]
        self.assertEqual(len(cambios), 5)
        self.assertEqual([c["despues"] for c in cambios], [False, True, False, True, False])
        self.assertEqual([c["antes"] for c in cambios], [True, False, True, False, True])
        self.assertEqual(cambios[0]["motivo"], "m6")
        for c in cambios:
            self.assertRegex(c["fecha"], r"^\d{2}/\d{2}/\d{4} \d{2}:\d{2}$")
            self.assertEqual(c["usuario"], "Daniel Aguilar")                  # el nombre, no el correo de la sesión

    def test_un_correo_que_nadie_toco_no_tiene_cambios_y_una_clave_ajena_es_404(self):
        st, d, _ = self.m.llamar("comm_correos_cambios", K)
        self.assertEqual((st, d["cambios"]), (200, []))
        st, d, _ = self.m.llamar("comm_correos_cambios", "no_existe_este")
        self.assertEqual(st, 404)

    def test_solo_muestra_los_cambios_de_ese_correo(self):
        self.put(K, activo=False, estado_visto="")
        self.put("ret_en_preparacion", activo=False, estado_visto="")
        st, d, _ = self.m.llamar("comm_correos_cambios", K)
        self.assertEqual(len(d["cambios"]), 1)


# ────────────────────────────── Historial: la causa «Apagado: «Nombre»» y el filtro por clave ──────────────────────────────
class TestHistorialControlDeCorreos(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        amb = {"re": re, "json": json, "datetime": _dtm.datetime, "_ccr": R, "chile_fmt_filter": MC.chile_fmt,
               "_email_log_norm_ref": lambda ref, asunto: (None, None, None), "_EMAIL_LOG_MODULOS_PRIVADOS": (),
               "PICKUP_REQUESTS_TABLE": "pickup_requests", "mysql_fetchall": lambda *a, **k: []}
        _extraer(amb, ("_comm_hist_filtros", "_comm_hist_like", "_comm_hist_ramas", "_comm_hist_llave", "_comm_hist_serializar"),
                 ("_COMM_HIST_ESTADOS", "_COMM_HIST_PER_PAGE", "_COMM_HIST_MODULO_LABELS", "_COMM_HIST_PROVEEDOR_LABELS",
                  "_COMM_HIST_MARCA_DUP"))
        cls.amb = amb

    def fila(self, **kw):
        base = {"src": "email_log", "id": 5, "canal": "email", "destinatario": "cliente@ejemplo.cl", "asunto": "Asunto", "evento": "x",
                "estado": "bloqueado", "detalle": "", "enviado_por": "", "metadata": None, "created_at": _dtm.datetime(2026, 10, 20, 18, 0),
                "modulo": "retiros", "ref_tipo": None, "ref_id": None, "ref_codigo": None, "proveedor": None, "job_id": None,
                "cuerpo_redactado": None, "cuerpo_fila": 1, "cuerpo_html": 1, "adjuntos_json": None, "correo_clave": None}
        base.update(kw)
        return base

    def test_un_correo_apagado_se_explica_con_su_nombre(self):
        f = self.amb["_comm_hist_llave"]
        detalle = f"Apagado en Control de correos: «{R.nombre_de(K)}» ({K})"
        self.assertEqual(f(detalle, "retiros", {"bloqueo": "correo", "correo": K}), "Apagado: «" + R.nombre_de(K) + "»")
        self.assertEqual(f("", "retiros", {"bloqueo": "correo", "correo": K}), "Apagado: «" + K + "»")

    def test_las_otras_causas_se_explican_como_siempre(self):
        f = self.amb["_comm_hist_llave"]
        self.assertEqual(f("Email deshabilitado por superadmin (kill switch global ON)", "retiros", {}),
                         "Interruptor general de email (todos los módulos)")
        self.assertEqual(f('Llave de paso cerrada para módulo "retiros" — email no enviado', "retiros", {"killswitch": True}),
                         "Llave de paso de Retiros")

    def test_la_fila_bloqueada_por_un_correo_trae_su_nombre_su_clave_y_la_hora_chile(self):
        meta = json.dumps({"modulo": "retiros", "bloqueo": "correo", "correo": K})
        out, = self.amb["_comm_hist_serializar"]([self.fila(
            detalle=f"Apagado en Control de correos: «{R.nombre_de(K)}» ({K})", metadata=meta, correo_clave=K)])
        self.assertEqual(out["estado"], "bloqueado")
        self.assertEqual(out["motivo"], "Apagado: «" + R.nombre_de(K) + "»")
        self.assertEqual(out["llave"], out["motivo"])
        self.assertEqual((out["correo_clave"], out["correo_nombre"], out["bloqueo"]), (K, R.nombre_de(K), "correo"))
        self.assertEqual((out["fecha"], out["hora"]), ("20/10/2026", "15:00:00"))

    def test_una_fila_de_antes_de_la_columna_usa_la_clave_que_quedo_en_la_metadata(self):
        meta = json.dumps({"modulo": "retiros", "correo": K})
        out, = self.amb["_comm_hist_serializar"]([self.fila(estado="enviado", metadata=meta, correo_clave=None)])
        self.assertEqual(out["correo_clave"], K)
        self.assertIsNone(out["bloqueo"])

    def test_una_clave_que_ya_no_esta_en_el_registro_no_rompe_el_historial(self):
        out, = self.amb["_comm_hist_serializar"]([self.fila(estado="enviado", correo_clave="correo_que_ya_no_existe")])
        self.assertEqual(out["correo_clave"], "correo_que_ya_no_existe")
        self.assertIsNone(out["correo_nombre"])

    def test_el_filtro_por_clave_solo_acepta_claves_con_forma_de_clave(self):
        filtros = self.amb["_comm_hist_filtros"]
        self.assertEqual(filtros({"correo": "  RET_CITA_CONFIRMADA "})["correo"], K)
        for malo in ("", "a", "x'; DROP TABLE email_log;--", "ret cita", "ñandú", "r" * 61, "9abc", "ret-cita", "ret.cita"):
            self.assertEqual(filtros({"correo": malo})["correo"], "", malo)

    def test_el_filtro_por_clave_va_como_parametro_y_deja_solo_la_fuente_que_tiene_esa_columna(self):
        f = self.amb["_comm_hist_filtros"]({"correo": K})
        f["desde_utc"] = f["hasta_utc"] = None
        ramas = self.amb["_comm_hist_ramas"](f)
        self.assertEqual(set(ramas), {"email_log"})
        w, p = ramas["email_log"]
        self.assertIn("e.correo_clave=%s", w)
        self.assertEqual(p, [K])
        self.assertNotIn(K, " ".join(w))
        # sin el filtro siguen las tres fuentes de siempre
        g = self.amb["_comm_hist_filtros"]({})
        g["desde_utc"] = g["hasta_utc"] = None
        self.assertEqual(set(self.amb["_comm_hist_ramas"](g)), {"email_log", "comm_log", "cola"})


# ────────────────────────────── La pantalla en el navegador: HTML, JS, CSS y reglas del proyecto ──────────────────────────────
class TestInterfazEstatica(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.js = _leer("static/comunicaciones_correos.js")
        cls.css = _leer("static/comunicaciones_correos.css")
        cls.html = _leer("templates/comunicaciones/index.html")
        cls.hist_js = _leer("static/comunicaciones_index.js")

    def test_sin_popups_nativos_del_navegador(self):                                  # REGLA #1
        for nombre, texto in (("comunicaciones_correos.js", self.js),):
            self.assertNotRegex(texto, r"(?<![\w.$])(alert|confirm|prompt)\s*\(", nombre)
        self.assertIn("ilusPrompt(", self.js)
        self.assertIn("ilusConfirm(", self.js)

    def test_el_javascript_es_valido(self):
        node = shutil.which("node")
        if not node:
            self.skipTest("node no está instalado")
        for rel in ("static/comunicaciones_correos.js", "static/comunicaciones_index.js"):
            r = subprocess.run([node, "--check", os.path.join(RAIZ, rel.replace("/", os.sep))], capture_output=True, text=True)
            self.assertEqual(r.returncode, 0, rel + "\n" + r.stderr)

    def test_la_plantilla_es_jinja_valido_y_conserva_las_pestanas_de_antes(self):       # REGLA #4.2: no se quitó nada
        import jinja2
        jinja2.Environment().parse(self.html)
        for destino in ("#tabEmail", "#tabPlantillas", "#tabLog", "#tabCorreos"):
            self.assertIn(f'data-bs-target="{destino}"', self.html)
            self.assertIn(f'id="{destino[1:]}"', self.html)

    def test_todo_id_que_usa_el_js_existe_en_la_plantilla(self):
        ids = set(re.findall(r"\$\('([A-Za-z][\w-]*)'\)", self.js)) | set(re.findall(r"getElementById\('([A-Za-z][\w-]*)'\)", self.js))
        self.assertGreater(len(ids), 20)
        for i in sorted(ids):
            self.assertTrue(re.search(r'id="%s"' % re.escape(i), self.html), f"el JS usa #{i} y la plantilla no lo tiene")
        for i in re.findall(r"\$\('(hist[A-Za-z]*Correo)'\)|getElementById\('(histFiltroCorreo)'\)", self.hist_js):
            for nombre in i:
                if nombre:
                    self.assertIn(f'id="{nombre}"', self.html)

    def test_toda_accion_de_la_plantilla_y_del_js_tiene_su_caso(self):
        casos = set(re.findall(r"case '([\w-]+)':", self.js[self.js.index("switch (acc)"):]))
        usadas = set(re.findall(r'data-ctc="([\w-]+)"', self.html)) | set(re.findall(r'data-ctc="([\w-]+)"', self.js))
        usadas -= {"sel"}           # las casillas se atienden en el evento «change», no en el «click»
        self.assertGreater(len(usadas), 8)
        self.assertFalse(usadas - casos, f"acciones sin caso en el JS: {sorted(usadas - casos)}")

    def test_las_rutas_que_llama_el_js_existen_en_app_py(self):
        app = _leer("app.py")
        for ruta, metodo in (('/api/comunicaciones/correos"', "GET"), ('/api/comunicaciones/correos/<clave>"', "PUT"),
                             ('/api/comunicaciones/correos/<clave>/vista-previa"', "POST"),
                             ('/api/comunicaciones/correos/<clave>/prueba"', "POST"),
                             ('/api/comunicaciones/correos/<clave>/cambios"', "GET")):
            self.assertIn(f'@app.route("{ruta}, methods=["{metodo}"])', app)
        for fragmento in ("'/api/comunicaciones/correos'", "'/api/comunicaciones/correos/' + encodeURIComponent(c.clave)",
                          "'/vista-previa'", "'/prueba'", "'/cambios'"):
            self.assertIn(fragmento, self.js)

    def test_el_pie_de_tabla_es_el_de_etiquetas_y_el_tamano_de_pagina_lo_elige_el_usuario(self):     # REGLA #4.3
        for texto in ("Mostrando", 'id="ctcPerPage"', 'value="25"', 'value="50"', 'value="100"', "por página", "Anterior", "Siguiente",
                      "Página", 'id="ctcPrev"', 'id="ctcNext"'):
            self.assertIn(texto, self.html)
        self.assertIn("commCorreosPerPage", self.js)
        self.assertIn("cargar({ fresco: true });", self.js[self.js.index("function limpiarFiltros"):][:400])     # limpiar un filtro recarga

    def test_un_solo_boton_marca_y_desmarca_todo(self):                                # REGLA #14
        self.assertIn('id="ctcBtnTodos"', self.html)
        cuerpo = self.js[self.js.index("function alternarTodos"):][:500]
        self.assertIn("some(", cuerpo)
        self.assertIn("setSel(", cuerpo)
        self.assertIn("Deseleccionar todo", self.js)

    def test_los_nombres_no_se_cortan_y_el_celular_tiene_botones_de_44px_y_inputs_de_16px(self):    # REGLAS #15 y #3
        self.assertNotRegex(self.css, r"line-clamp|text-overflow\s*:\s*ellipsis")
        self.assertRegex(self.css, r"\.ctc-nombre\{[^}]*overflow-wrap:anywhere")
        for regla in (".ctc-btn,.ctc-switch{min-height:44px}", ".ctc-btn.chico{min-height:44px}",
                      "#tabCorreos .form-control,#tabCorreos .form-select{font-size:16px}"):
            i = self.css.find(regla)
            self.assertGreater(i, 0, regla)
            self.assertTrue(self.css[self.css.rfind("@media", 0, i):].startswith("@media (max-width:767.98px)"),
                            f"{regla} debe vivir dentro del bloque del celular")

    def test_lo_que_escribe_una_persona_siempre_se_escapa_antes_de_pintarse(self):
        # Funciones que arman el texto en una variable y lo pintan SOLO como esc(variable): se verifica que así sea.
        seguro_por_funcion = {"lineaEstado": "esc(t)"}
        for peligroso in ("e.motivo", "x.motivo", "e.apagado_por", "x.usuario", "c.nombre", "c.disparo", "c.avisa_a", "c.nota", "d.mensaje",
                          "d.asunto", "x.fecha"):
            for m in re.finditer(r"[+]\s*" + re.escape(peligroso) + r"\b", self.js):
                # Vale si el texto va dentro de esc(…) o de un ilusPrompt/ilusConfirm/ilusAlert (ilus_ui.js escapa título, mensaje y sub)
                seguro = False
                for llamada in ("esc(", "ilusPrompt(", "ilusConfirm(", "ilusAlert("):
                    ini = self.js.rfind(llamada, 0, m.start())
                    previo = self.js[ini:m.start()] if ini >= 0 else ""
                    seguro = seguro or (bool(previo) and previo.count("(") > previo.count(")"))     # sigue abierto su paréntesis
                if not seguro:
                    fn = re.findall(r"function (\w+)\(", self.js[:m.start()])[-1]
                    if fn in seguro_por_funcion:
                        cuerpo = self.js[self.js.rindex(f"function {fn}(", 0, m.start()):]
                        cuerpo = cuerpo[:cuerpo.index("\n  }\n")]
                        self.assertIn(seguro_por_funcion[fn], cuerpo, f"{fn} debe pintar su texto con {seguro_por_funcion[fn]}")
                        seguro = True
                self.assertTrue(seguro, f"se concatena {peligroso} sin esc(): …{self.js[max(0, m.start() - 40):m.end() + 20]}…")
        self.assertIn("function esc(", self.js)
        self.assertIn("esc(e.motivo)", self.js)
        ui = _leer("static/ilus_ui.js")                                        # el contrato en que se apoya lo de arriba
        self.assertIn("escapeHtml(title)", ui)
        self.assertIn("escapeHtml(message)", ui)

    def test_la_vista_previa_va_en_un_iframe_sin_permisos(self):
        self.assertRegex(self.js, r"<iframe[^>]*sandbox=\"\"")
        self.assertIn("srcdoc", self.js)

    def test_el_js_solo_apaga_lo_permitido_y_avisa_lo_demas(self):
        self.assertIn("puede_encender", self.js)
        self.assertIn("No se puede encender desde aquí", self.js)
        self.assertIn("c.protegido", self.js)

    def test_claude_md_tiene_la_regla_33_y_sigue_teniendo_la_32(self):
        md = _leer("CLAUDE.md")
        self.assertIn("REGLA #33", md)
        self.assertIn("REGLA #32", md)
        self.assertLess(md.index("REGLA #32"), md.index("REGLA #33"))
        self.assertIn("comm_correos_registro.py", md)

    def test_ningun_texto_del_registro_para_tecnicos_muestra_montos(self):               # REGLAS #19 y #26
        for c in R.REGISTRO:
            if c["destinatario"] == "tecnico":
                for campo in ("nombre", "disparo", "nota", "avisa_a", "adjunta"):
                    self.assertNotIn("$", c[campo], f"{c['clave']}.{campo}")


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  8) FASE 3: los avisos que envían FedEx y SimpliRoute (registro: tipo «externo»)
#     ILUS no los envía: decide si le ENTREGA al tercero el correo del destinatario. Obedecen al interruptor general y a su propio
#     interruptor (decisión de Daniel, 20-oct-2026); la llave de paso del módulo no los frena. Sin fila = como siempre.
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
EXTERNOS = ("ext_fedex_notif_envio", "ext_fedex_notif_retiro", "ext_fedex_export_email", "ext_simpliroute_contacto")
FUNCS_EXTERNOS = ("_correo_externo_bloqueo", "_correo_externo_permitido", "_correo_externo_frenado")


class TestExternosEnElRegistro(unittest.TestCase):
    def test_los_cuatro_avisos_de_terceros_son_externos_sin_plantilla_ni_proteccion(self):
        self.assertEqual(set(R.claves(tipo="externo")), set(EXTERNOS))
        for k in EXTERNOS:
            c = R.por_clave(k)
            self.assertEqual((c["tipo"], c["plantilla"], c["protegido"], c["sensible"]), ("externo", None, False, False), k)
            self.assertIn("interruptor general", c["nota"], k)

    def test_la_llave_del_modulo_no_frena_un_aviso_de_un_tercero(self):
        for k in EXTERNOS:
            e = R.estado_efectivo(R.por_clave(k), llave_bloqueada=True)
            self.assertEqual(e["efectivo"], "sale", k)
            self.assertEqual([p["estado"] for p in e["pasos"]], ["ok", "saltado", "ok", "ok"], k)
            self.assertFalse(e["aplica_llave"], k)

    def test_el_interruptor_general_y_el_propio_si_los_frenan(self):
        c = R.por_clave("ext_fedex_notif_envio")
        e = R.estado_efectivo(c, general_activo=False)
        self.assertEqual((e["efectivo"], e["frena_en"]), ("general", 1))
        self.assertEqual([p["estado"] for p in e["pasos"]], ["apagado", "saltado", "saltado", "saltado"])
        e = R.estado_efectivo(c, switch_activo=False)
        self.assertEqual((e["efectivo"], e["frena_en"]), ("correo", 3))
        self.assertEqual([p["estado"] for p in e["pasos"]], ["ok", "saltado", "apagado", "saltado"])
        self.assertTrue(e["puede_encender"])

    def test_un_correo_de_ilus_sigue_pasando_por_su_llave(self):
        e = R.estado_efectivo(R.por_clave("tr_en_ruta"), llave_bloqueada=True)
        self.assertEqual((e["efectivo"], e["frena_en"]), ("llave", 2))
        self.assertTrue(e["aplica_llave"])


class TestExternosLaPuerta(unittest.TestCase):
    """`_correo_externo_*` de app.py, funciones REALES: general > interruptor del aviso; la llave del módulo NO cuenta."""

    def mundo(self, **kw):
        m = Mundo(**kw)
        _extraer(m.amb, FUNCS_EXTERNOS)
        return m

    def test_sin_fila_se_le_entrega_el_correo_como_siempre(self):
        m = self.mundo()
        for k in EXTERNOS:
            self.assertIsNone(m.amb["_correo_externo_bloqueo"](k), k)
            self.assertTrue(m.amb["_correo_externo_permitido"](k), k)
        self.assertEqual(m.logs, [])

    def test_el_interruptor_general_lo_frena_aunque_el_del_aviso_este_encendido(self):
        m = self.mundo(global_on=False)
        self.assertEqual(m.amb["_correo_externo_bloqueo"]("ext_fedex_notif_envio"), "global")
        self.assertFalse(m.amb["_correo_externo_permitido"]("ext_fedex_notif_envio"))

    def test_el_interruptor_del_aviso_lo_frena_y_los_demas_siguen(self):
        m = self.mundo(filas={"ext_fedex_notif_envio": 0})
        self.assertEqual(m.amb["_correo_externo_bloqueo"]("ext_fedex_notif_envio"), "correo")
        for k in EXTERNOS[1:]:
            self.assertIsNone(m.amb["_correo_externo_bloqueo"](k), k)

    def test_la_llave_de_paso_del_modulo_no_frena_estos_avisos(self):
        m = self.mundo(llaves_cerradas={"transporte"})
        for k in EXTERNOS:
            self.assertTrue(m.amb["_correo_externo_permitido"](k), k)

    def test_base_caida_o_error_al_leer_nunca_impide_crear_una_guia(self):
        m = self.mundo()
        m.db_cae = True
        for k in EXTERNOS:
            self.assertTrue(m.amb["_correo_externo_permitido"](k), k)

        def _boom(canal):
            raise RuntimeError("no se pudo leer el interruptor general")
        m.amb["comm_is_enabled"] = _boom
        self.assertTrue(m.amb["_correo_externo_permitido"]("ext_fedex_notif_envio"))

    def test_lo_frenado_queda_en_el_historial_como_bloqueado_con_la_clave_del_aviso(self):
        m = self.mundo(filas={"ext_fedex_notif_envio": 0})
        m.amb["_correo_externo_frenado"]("ext_fedex_notif_envio", "cliente@ejemplo.cl", "correo", "guía FCV 123")
        (args, kw), = m.logs
        self.assertEqual(args[0], "cliente@ejemplo.cl")
        self.assertEqual(args[3], "bloqueado")
        self.assertEqual((kw["correo"], kw["modulo"]), ("ext_fedex_notif_envio", "transporte"))
        self.assertEqual(kw["metadata"]["bloqueo"], "correo")
        self.assertTrue(kw["metadata"]["externo"])
        self.assertIn("«" + R.nombre_de("ext_fedex_notif_envio") + "»", kw["error_msg"])
        self.assertIn("guía FCV 123", kw["error_msg"])
        self.assertIsNone(kw.get("html"))                                   # no hay cuerpo que guardar: ILUS no envió ningún correo

    def test_por_el_interruptor_general_el_motivo_lo_dice(self):
        m = self.mundo(global_on=False)
        m.amb["_correo_externo_frenado"]("ext_simpliroute_contacto", "cliente@ejemplo.cl", "global")
        (args, kw), = m.logs
        self.assertIn("global", kw["error_msg"].lower())
        self.assertNotIn("bloqueo", kw["metadata"])
        self.assertTrue(kw["metadata"]["killswitch"])

    def test_sin_decir_el_motivo_lo_calcula(self):
        m = self.mundo(filas={"ext_fedex_notif_retiro": 0})
        m.amb["_correo_externo_frenado"]("ext_fedex_notif_retiro", "x@y.cl")
        (args, kw), = m.logs
        self.assertEqual(kw["metadata"]["bloqueo"], "correo")

    def test_dejar_constancia_nunca_rompe_la_guia(self):
        m = self.mundo()

        def _boom(*a, **k):
            raise RuntimeError("email_log no responde")
        m.amb["_email_log"] = _boom
        m.amb["_correo_externo_frenado"]("ext_fedex_notif_envio", "x@y.cl", "correo")        # no lanza


class TestExternosEnElCodigo(unittest.TestCase):
    """Los cuatro sitios donde ILUS le entrega el correo a un tercero pasan por la puerta, ANTES de armar lo que se envía."""

    def test_cada_aviso_externo_se_consulta_en_el_codigo(self):
        app = _leer("app.py")
        for k in EXTERNOS:
            self.assertRegex(app, r'_correo_externo_(?:bloqueo|permitido)\("' + k + r'"\)', f"{k}: nadie lo consulta")

    def test_fedex_crea_la_guia_sin_correo_si_el_aviso_esta_apagado(self):
        s = FC.funcion_top("app.py", "_fedex_create_shipment")
        i_gate = s.index('_correo_externo_bloqueo("ext_fedex_notif_envio")')
        i_rec = s.index('"emailAddress": rec_email')
        i_notif = s.index('body["requestedShipment"]["emailNotificationDetail"]')
        self.assertLess(s.index('rec_email    = _fedex_clean_str(recipient.get("email"), 80)'), i_gate)
        self.assertLess(i_gate, i_rec)
        self.assertLess(i_gate, i_notif)
        self.assertIn('rec_email, notify_email = "", ""', s[i_gate:i_rec])
        self.assertIn("_correo_externo_frenado(", s[i_gate:i_rec])

    def test_fedex_programa_el_retiro_sin_correo_si_el_aviso_esta_apagado(self):
        s = FC.funcion_top("app.py", "_fedex_pickup_create")
        i_gate = s.index('_correo_externo_bloqueo("ext_fedex_notif_retiro")')
        i_notif = s.index('body["pickupNotificationDetail"]')
        self.assertLess(i_gate, i_notif)
        self.assertIn('notify_email = ""', s[i_gate:i_notif])

    def test_el_excel_de_fedex_y_el_de_simpliroute_dejan_vacia_la_celda_del_correo(self):
        s = FC.funcion_top("app.py", "_tr_manifiesto_export_impl")
        i_fx = s.index('_correo_externo_bloqueo("ext_fedex_export_email")')
        i_row_fx = s.index("recipient_email = fedex_test_notify_email if _fedex_test_mode")
        self.assertLess(i_fx, i_row_fx)
        self.assertIn('recipient_email = ""', s[i_row_fx:i_row_fx + 400])
        # SimpliRoute: la consulta va ANTES del título de la hoja (una prueba de «sin ceros» mira una ventana de texto desde ahí)
        i_sr = s.index('_correo_externo_bloqueo("ext_simpliroute_contacto")')
        self.assertLess(i_sr, s.index('ws.title = "Felca y Milling"'))
        self.assertIn("ws.cell(ri, 21).value = None", s[s.index('ws.title = "Felca y Milling"'):])
        # Clickex y el manifiesto genérico no tienen un aviso registrado: no se tocan
        self.assertNotIn("ext_clickex", s)

    def test_simpliroute_recibe_el_gancho_de_la_puerta(self):
        app = _leer("app.py")
        self.assertIn("_src_hook.CONTACTO_EMAIL_PERMITIDO = _sr_contacto_email_permitido", app)
        cuerpo = FC.funcion_top("app.py", "_sr_contacto_email_permitido")
        self.assertIn('_correo_externo_bloqueo("ext_simpliroute_contacto")', cuerpo)
        self.assertIn("return False", cuerpo)

    def test_lo_que_no_es_un_aviso_no_se_toca(self):
        """Las funciones INTOCABLES de Transporte (otras pruebas las comparan con main) no llevan nada de esto."""
        for nombre in ("_tr_notificar_cliente", "_simpliroute_poll_batch", "tr_item_simpliroute_cancelar", "_tr_bulk_sync_erp_mysql",
                       "_transporte_scheduler_loop"):
            self.assertNotIn("_correo_externo_", FC.funcion_top("app.py", nombre), nombre)


class TestExternosFedExReal(unittest.TestCase):
    """`_fedex_create_shipment` y `_fedex_pickup_create` REALES de app.py con la red de mentira: se captura el JSON que se le habría
    mandado a FedEx. Con el aviso apagado la guía y el retiro salen igual, sin el correo del destinatario."""

    DEST = {"nombre": "Gimnasio Demo", "empresa": "Demo SpA", "telefono": "+56912345678", "email": "cliente@ejemplo.cl",
            "direccion": "Av. Demo 123", "comuna": "Las Condes", "region": "RM", "cod_postal": "7550000"}
    PAQ = [{"peso_kg": 5, "largo": 10, "ancho": 10, "alto": 10}]
    CONTACTO = {"empresa": "ILUS Fitness", "nombre": "Despacho", "telefono": "56229465970"}
    DIRECCION = {"direccion": "Av. Industrial 4500", "comuna": "Quilicura", "cod_postal": "8730000"}

    def correr(self, funcion, *, filas=None, global_on=True, **kw):
        import sys
        m = Mundo(filas=filas, global_on=global_on)
        _extraer(m.amb, FUNCS_EXTERNOS)
        capturado = []

        class _Requests:                                          # el «módulo» requests: anota el JSON y corta (no hay red)
            @staticmethod
            def post(url, json=None, headers=None, timeout=None):
                capturado.append(json)
                raise RuntimeError("sin red: es una prueba")
        m.amb.update({
            "FEDEX_RATE_CLIENT_ID": "id", "FEDEX_RATE_CLIENT_SECRET": "secreto", "FEDEX_ACCOUNT": "123456789",
            "FEDEX_DEFAULT_SERVICE_CL": "FEDEX_PRIORITY", "FEDEX_DECLARED_CURRENCY": "", "FEDEX_ORIGIN_POSTAL": "8730000",
            "FEDEX_SHIP_URL": "https://fedex.invalid/ship", "FEDEX_PICKUP_CREATE_URL": "https://fedex.invalid/pickup",
            "ILUS_REMITENTE": {"nombre": "ILUS Fitness", "telefono": "56229465970", "bodega": "Av. Industrial 4500",
                               "direccion": "Av. Industrial 4500"},
            "_fedex_get_token": lambda: "token", "_fedex_clean_str": lambda v, n=0: (str(v or "").strip()[:n] if n else str(v or "").strip()),
            "_fedex_split_address": lambda a: [str(a or "")], "_fedex_postal_cl": lambda p: str(p or ""),
            "_resolve_comuna_erp": lambda c: c, "_cl_resolve_short_code": lambda c, region_hint=None: None,
            "_codigo_postal_chile": lambda c: "7550000", "_comuna_to_postal": lambda c: "7550000",
        })
        _extraer(m.amb, (funcion,))
        with mock.patch.dict(sys.modules, {"requests": _Requests}):
            res = m.amb[funcion](**kw)
        return capturado, res, m

    def guia(self, **kw):
        kw.setdefault("recipient", dict(self.DEST))
        kw.setdefault("packages", [dict(p) for p in self.PAQ])
        kw.setdefault("reference", "FCV 123")
        return self.correr("_fedex_create_shipment", **kw)

    @staticmethod
    def sin_correo(body):
        """La guía sin el correo del contacto ni la notificación: lo ÚNICO que debe cambiar cuando el aviso se apaga."""
        b = json.loads(json.dumps(body))
        b["requestedShipment"]["recipients"][0]["contact"].pop("emailAddress", None)
        b["requestedShipment"].pop("emailNotificationDetail", None)
        return b

    def test_con_el_aviso_encendido_la_guia_lleva_el_correo_y_la_notificacion(self):
        cap, res, m = self.guia(notify_email="cliente@ejemplo.cl", notify_name="Gimnasio Demo")
        body, = cap
        rs = body["requestedShipment"]
        self.assertEqual(rs["recipients"][0]["contact"]["emailAddress"], "cliente@ejemplo.cl")
        self.assertEqual(rs["emailNotificationDetail"]["emailNotificationRecipients"][0]["emailAddress"], "cliente@ejemplo.cl")
        self.assertEqual(m.logs, [])

    def test_con_el_aviso_apagado_la_guia_sale_igual_pero_sin_correo_y_con_una_constancia(self):
        con, _, _ = self.guia(notify_email="cliente@ejemplo.cl", notify_name="Gimnasio Demo")
        sin, res, m = self.guia(filas={"ext_fedex_notif_envio": 0}, notify_email="cliente@ejemplo.cl", notify_name="Gimnasio Demo")
        body, = sin
        rs = body["requestedShipment"]
        self.assertNotIn("emailAddress", rs["recipients"][0]["contact"])
        self.assertNotIn("emailNotificationDetail", rs)
        self.assertEqual(body, self.sin_correo(con[0]), "solo debe faltar el correo: el resto de la guía no cambia")
        self.assertIn("Red FedEx", res["error"])                    # llegó hasta el envío a FedEx: nada se cortó antes
        (args, kw), = m.logs
        self.assertEqual((args[0], args[3], kw["correo"]), ("cliente@ejemplo.cl", "bloqueado", "ext_fedex_notif_envio"))
        self.assertIn("guía FCV 123", kw["error_msg"])

    def test_el_interruptor_general_apagado_tambien_deja_la_guia_sin_correo(self):
        cap, _, m = self.guia(global_on=False, notify_email="cliente@ejemplo.cl")
        rs = cap[0]["requestedShipment"]
        self.assertNotIn("emailAddress", rs["recipients"][0]["contact"])
        self.assertNotIn("emailNotificationDetail", rs)
        (args, kw), = m.logs
        self.assertTrue(kw["metadata"]["killswitch"])

    def test_si_no_se_pidio_notificar_solo_se_quita_el_correo_del_contacto_y_no_hay_constancia(self):
        con, _, m_on = self.guia()
        sin, _, m_off = self.guia(filas={"ext_fedex_notif_envio": 0})
        self.assertEqual(con[0]["requestedShipment"]["recipients"][0]["contact"]["emailAddress"], "cliente@ejemplo.cl")
        self.assertNotIn("emailNotificationDetail", con[0]["requestedShipment"])
        self.assertNotIn("emailAddress", sin[0]["requestedShipment"]["recipients"][0]["contact"])
        self.assertEqual((m_on.logs, m_off.logs), ([], []))

    def test_un_destinatario_sin_correo_no_consulta_ni_deja_constancia(self):
        cap, _, m = self.guia(filas={"ext_fedex_notif_envio": 0}, recipient=dict(self.DEST, email=""))
        self.assertNotIn("emailAddress", cap[0]["requestedShipment"]["recipients"][0]["contact"])
        self.assertEqual(m.logs, [])

    def retiro(self, **kw):
        return self.correr("_fedex_pickup_create", contact=dict(self.CONTACTO), address=dict(self.DIRECCION),
                           ready_dt="2026-10-21T10:00:00", **kw)

    def test_el_retiro_con_el_aviso_encendido_lleva_la_notificacion(self):
        cap, _, m = self.retiro(notify_email="alguien@ilus.cl")
        self.assertEqual(cap[0]["pickupNotificationDetail"]["emailDetails"][0]["address"], "alguien@ilus.cl")
        self.assertEqual(m.logs, [])

    def test_el_retiro_con_el_aviso_apagado_sale_igual_sin_notificacion_y_con_constancia(self):
        con, _, _ = self.retiro(notify_email="alguien@ilus.cl")
        sin, res, m = self.retiro(filas={"ext_fedex_notif_retiro": 0}, notify_email="alguien@ilus.cl")
        self.assertNotIn("pickupNotificationDetail", sin[0])
        self.assertEqual(sin[0], {k: v for k, v in con[0].items() if k != "pickupNotificationDetail"})
        self.assertIn("FedEx Pickup", res["error"])
        (args, kw), = m.logs
        self.assertEqual((args[0], args[3], kw["correo"]), ("alguien@ilus.cl", "bloqueado", "ext_fedex_notif_retiro"))

    def test_el_retiro_sin_correo_no_consulta_nada(self):
        cap, _, m = self.retiro(filas={"ext_fedex_notif_retiro": 0})
        self.assertNotIn("pickupNotificationDetail", cap[0])
        self.assertEqual(m.logs, [])

    def test_cada_aviso_de_fedex_obedece_a_su_propio_interruptor(self):
        cap, _, _ = self.guia(filas={"ext_fedex_notif_retiro": 0}, notify_email="cliente@ejemplo.cl")      # apagar el del retiro no toca la guía
        self.assertIn("emailNotificationDetail", cap[0]["requestedShipment"])
        cap, _, _ = self.retiro(filas={"ext_fedex_notif_envio": 0}, notify_email="alguien@ilus.cl")
        self.assertIn("pickupNotificationDetail", cap[0])


class TestExternosElExcel(unittest.TestCase):
    """Los Excel de carga masiva (FedEx y SimpliRoute) generados por la función REAL `_tr_manifiesto_export_impl` de app.py, con
    openpyxl y en un mundo de mentira: con el aviso apagado la celda del correo queda VACÍA y el resto del archivo no cambia."""

    ITEMS = [
        {"nudo": "0000023102", "tido": "FCV", "cliente_nombre": "Gimnasio Demo SpA", "direccion": "Av. Demo 123", "comuna": "Las Condes",
         "telefono": "+56912345678", "email": "uno@cliente.cl", "n_bultos": 2, "peso": 40, "notas": "", "valor_bruto": 1000},
        {"nudo": "0000023103", "tido": "FCV", "cliente_nombre": "Club Deportivo Dos", "direccion": "Calle 2 45", "comuna": "Providencia",
         "telefono": "+56987654321", "email": "dos@cliente.cl", "n_bultos": 1, "peso": 10, "notas": "", "valor_bruto": 500},
        {"nudo": "0000023104", "tido": "BLV", "cliente_nombre": "Sin Correo Ltda", "direccion": "Av. Tres 3", "comuna": "Ñuñoa",
         "telefono": "+56911111111", "email": "", "n_bultos": 1, "peso": 5, "notas": "", "valor_bruto": 100},
    ]
    COURIER = {"fedex": "FedEx", "simplyroute": "Felca", "clickex": "Clickex", "generic": "Otro courier"}

    def correr(self, formato, *, filas=None, global_on=True, items=None):
        import io
        import openpyxl
        m = Mundo(filas=filas, global_on=global_on)
        _extraer(m.amb, FUNCS_EXTERNOS)
        man = {"id": 1, "courier": self.COURIER[formato], "correlativo": "MAN-0001"}
        lista = self.ITEMS if items is None else items
        m.amb.update({
            "os": os, "flash": lambda *a, **k: None, "redirect": lambda *a, **k: None, "url_for": lambda *a, **k: "#",
            "mysql_fetchone": lambda sql, p=(): {"formato_export": ""} if "transport_couriers" in sql else man,
            "_tr_manifiesto_items_export": lambda mid: [dict(i) for i in lista],
            "_peso_a_declarar": lambda it: float(it.get("peso") or 1),
            "_tr_sender_cfg": lambda: {"name": "ILUS Fitness", "line1": "Av. Industrial 4500", "line2": "", "company": "ILUS Fitness",
                                       "phone": "56229465970", "email": "despacho@ilus.cl", "postcode": "8730000", "state": "RM",
                                       "city": "Quilicura", "country": "CL"},
            "_codigo_postal_chile": lambda comuna: "7550000",
            "nudo_fmt_filter": lambda v: str(v or "").lstrip("0"),
            "_tr_log": lambda *a, **k: None,
        })
        _extraer(m.amb, ("_tr_manifiesto_export_impl",))
        datos, nombre = m.amb["_tr_manifiesto_export_impl"](1, devolver_bytes=True)
        return openpyxl.load_workbook(io.BytesIO(datos)).active, m

    def columna(self, ws, col):
        return [ws.cell(r, col).value for r in range(2, ws.max_row + 1)]

    def test_fedex_con_el_aviso_encendido_el_excel_lleva_el_correo_de_cada_cliente(self):
        ws, m = self.correr("fedex")
        self.assertEqual(ws.cell(1, 15).value, "recipientEmail")
        self.assertEqual(self.columna(ws, 15), ["uno@cliente.cl", "dos@cliente.cl", None])
        self.assertEqual(ws.cell(4, 15).fill.fill_type, "solid")                 # sin correo: se marca como faltante, como siempre
        self.assertEqual(m.logs, [])

    def test_fedex_con_el_aviso_apagado_la_columna_queda_vacia_y_se_deja_una_sola_constancia(self):
        ws, m = self.correr("fedex", filas={"ext_fedex_export_email": 0})
        self.assertEqual(self.columna(ws, 15), [None, None, None])
        for r in (2, 3, 4):
            self.assertIsNone(ws.cell(r, 15).fill.fill_type, "lo vacío a propósito no se marca como faltante")
        (args, kw), = m.logs                                                      # UNA por exportación, no una por fila
        self.assertEqual((args[0], args[3], kw["correo"]), ("2 correos de clientes", "bloqueado", "ext_fedex_export_email"))
        self.assertEqual(kw["metadata"]["bloqueo"], "correo")
        self.assertIn("MAN-0001", kw["error_msg"])

    def test_fedex_el_resto_del_archivo_no_cambia(self):
        con, _ = self.correr("fedex")
        sin, _ = self.correr("fedex", filas={"ext_fedex_export_email": 0})
        self.assertEqual((con.max_row, con.max_column), (sin.max_row, sin.max_column))
        for r in range(1, con.max_row + 1):
            for c in range(1, con.max_column + 1):
                if c != 15 or r == 1:                  # todo igual, salvo la celda del correo de cada fila (el encabezado sí se conserva)
                    self.assertEqual(con.cell(r, c).value, sin.cell(r, c).value, (r, c))

    def test_fedex_el_interruptor_general_apagado_tambien_vacia_la_columna(self):
        ws, m = self.correr("fedex", global_on=False)
        self.assertEqual(self.columna(ws, 15), [None, None, None])
        (args, kw), = m.logs
        self.assertTrue(kw["metadata"]["killswitch"])
        self.assertIn("global", kw["error_msg"].lower())

    def test_fedex_en_modo_prueba_el_correo_interno_tampoco_viaja_si_el_aviso_esta_apagado(self):
        with mock.patch.dict(os.environ, {"ILUS_FEDEX_TEST_MODE": "1"}):
            ws, _ = self.correr("fedex")
            self.assertEqual(set(self.columna(ws, 15)), {"alison.mellado@sphs.cl"})
            ws, m = self.correr("fedex", filas={"ext_fedex_export_email": 0})
            self.assertEqual(self.columna(ws, 15), [None, None, None])

    def test_si_ningun_cliente_trae_correo_no_hay_nada_que_frenar_ni_constancia(self):
        sin_correo = [dict(i, email="") for i in self.ITEMS]
        ws, m = self.correr("fedex", filas={"ext_fedex_export_email": 0}, items=sin_correo)
        self.assertEqual(self.columna(ws, 15), [None, None, None])
        self.assertEqual(m.logs, [])

    def test_simpliroute_con_el_aviso_encendido_el_excel_lleva_el_correo_de_contacto(self):
        ws, m = self.correr("simplyroute")
        self.assertEqual(ws.cell(1, 21).value, "Correo electrónico de contacto")
        self.assertEqual(self.columna(ws, 21), ["uno@cliente.cl", "dos@cliente.cl", None])
        self.assertEqual(ws.cell(4, 21).fill.fill_type, "solid")
        self.assertEqual(m.logs, [])

    def test_simpliroute_con_el_aviso_apagado_la_celda_queda_vacia_y_el_resto_igual(self):
        con, _ = self.correr("simplyroute")
        sin, m = self.correr("simplyroute", filas={"ext_simpliroute_contacto": 0})
        self.assertEqual(self.columna(sin, 21), [None, None, None])
        for r in (2, 3, 4):
            self.assertIsNone(sin.cell(r, 21).fill.fill_type)
        for r in range(1, con.max_row + 1):
            for c in range(1, con.max_column + 1):
                if c != 21 or r == 1:
                    self.assertEqual(con.cell(r, c).value, sin.cell(r, c).value, (r, c))
        (args, kw), = m.logs
        self.assertEqual((args[0], kw["correo"]), ("2 correos de clientes", "ext_simpliroute_contacto"))

    def test_cada_excel_obedece_a_su_propio_interruptor(self):
        ws, _ = self.correr("fedex", filas={"ext_simpliroute_contacto": 0})                   # apagar el de SimpliRoute no toca el de FedEx
        self.assertEqual(self.columna(ws, 15), ["uno@cliente.cl", "dos@cliente.cl", None])
        ws, _ = self.correr("simplyroute", filas={"ext_fedex_export_email": 0})
        self.assertEqual(self.columna(ws, 21), ["uno@cliente.cl", "dos@cliente.cl", None])

    def test_clickex_y_el_manifiesto_generico_no_dependen_de_estos_interruptores(self):
        for formato, col in (("clickex", 7), ("generic", 6)):
            ws, m = self.correr(formato, filas={k: 0 for k in EXTERNOS})
            self.assertEqual(self.columna(ws, col), ["uno@cliente.cl", "dos@cliente.cl", None], formato)
            self.assertEqual(m.logs, [], formato)


class TestExternosSimpliRoute(unittest.TestCase):
    """El gancho de `simpliroute_client.build_visit_payload`: por defecto el correo de contacto va como siempre."""

    ITEM = {"item_id": 1, "nudo": "123", "tido": "FCV", "cliente_nombre": "Gimnasio Demo SpA", "direccion": "Av. Demo 123",
            "comuna": "Las Condes", "region": "RM", "telefono": "+56912345678", "email": "cliente@ejemplo.cl", "n_bultos": 2,
            "peso_declarar": 40}

    def setUp(self):
        import simpliroute_client as src
        self.src = src
        self.addCleanup(setattr, src, "CONTACTO_EMAIL_PERMITIDO", src.CONTACTO_EMAIL_PERMITIDO)

    def test_por_defecto_el_correo_de_contacto_va_en_la_visita(self):
        p, errores = self.src.build_visit_payload(self.ITEM, planned_date="2026-10-21")
        self.assertEqual(errores, [])
        self.assertEqual(p["contact_email"], "cliente@ejemplo.cl")

    def test_con_el_aviso_apagado_la_visita_sale_igual_pero_sin_correo(self):
        con, _ = self.src.build_visit_payload(self.ITEM, planned_date="2026-10-21")
        vistos = []
        self.src.CONTACTO_EMAIL_PERMITIDO = lambda mail: vistos.append(mail) or False
        sin, errores = self.src.build_visit_payload(self.ITEM, planned_date="2026-10-21")
        self.assertEqual(errores, [])
        self.assertNotIn("contact_email", sin)
        self.assertEqual(vistos, ["cliente@ejemplo.cl"])
        self.assertEqual(sin, {k: v for k, v in con.items() if k != "contact_email"},
                         "solo debe faltar el correo: el resto de la visita no cambia")

    def test_sin_correo_en_el_item_ni_se_consulta_el_gancho(self):
        llamado = []
        self.src.CONTACTO_EMAIL_PERMITIDO = lambda mail: llamado.append(mail) or True
        p, _ = self.src.build_visit_payload(dict(self.ITEM, email=""), planned_date="2026-10-21")
        self.assertNotIn("contact_email", p)
        self.assertEqual(llamado, [])

    def test_si_el_gancho_falla_la_visita_sale_con_su_correo_como_siempre(self):
        def _boom(mail):
            raise RuntimeError("no se pudo leer el interruptor")
        self.src.CONTACTO_EMAIL_PERMITIDO = _boom
        p, errores = self.src.build_visit_payload(self.ITEM, planned_date="2026-10-21")
        self.assertEqual(errores, [])
        self.assertEqual(p["contact_email"], "cliente@ejemplo.cl")

    def test_con_la_funcion_real_de_app_py_cada_visita_con_correo_deja_su_constancia_o_sale_con_su_correo(self):
        items = [self.ITEM, dict(self.ITEM, item_id=2, nudo="124", email="otro@cliente.cl"), dict(self.ITEM, item_id=3, nudo="125", email="")]
        m = Mundo(filas={"ext_simpliroute_contacto": 0})
        _extraer(m.amb, FUNCS_EXTERNOS + ("_sr_contacto_email_permitido",))
        self.src.CONTACTO_EMAIL_PERMITIDO = m.amb["_sr_contacto_email_permitido"]
        payloads, saltados = self.src.build_visits_batch(items, planned_date="2026-10-21")
        self.assertEqual((len(payloads), saltados), (3, []))
        self.assertTrue(all("contact_email" not in p for _i, p in payloads))
        self.assertEqual([a[0] for a, _k in m.logs], ["cliente@ejemplo.cl", "otro@cliente.cl"])       # la visita sin correo no deja constancia
        self.assertTrue(all(a[3] == "bloqueado" and k["correo"] == "ext_simpliroute_contacto" for a, k in m.logs))
        # encendido: salen con su correo y sin constancias
        m2 = Mundo()
        _extraer(m2.amb, FUNCS_EXTERNOS + ("_sr_contacto_email_permitido",))
        self.src.CONTACTO_EMAIL_PERMITIDO = m2.amb["_sr_contacto_email_permitido"]
        payloads, _ = self.src.build_visits_batch(items, planned_date="2026-10-21")
        self.assertEqual([p.get("contact_email") for _i, p in payloads], ["cliente@ejemplo.cl", "otro@cliente.cl", None])
        self.assertEqual(m2.logs, [])

    def test_el_lote_completo_respeta_el_gancho(self):
        self.src.CONTACTO_EMAIL_PERMITIDO = lambda mail: False
        payloads, saltados = self.src.build_visits_batch([self.ITEM, dict(self.ITEM, item_id=2, nudo="124")],
                                                         planned_date="2026-10-21")
        self.assertEqual((len(payloads), saltados), (2, []))
        self.assertTrue(all("contact_email" not in p for _i, p in payloads))


class TestExternosLaPantalla(_PantallaBase):
    def test_aparecen_en_el_grupo_de_transporte_con_la_llave_como_no_aplica(self):
        d = self.listar()
        for k in EXTERNOS:
            c = self.item(k, d)
            self.assertEqual((c["tipo"], c["modulo_ui"]), ("externo", "transporte"), k)
            self.assertEqual(c["estado"]["pasos"][1]["estado"], "saltado", k)
        self.assertEqual(d["resumen"]["total"], len([x for x in R.REGISTRO if x["tipo"] in ("ilus", "externo")]))

    def test_la_llave_de_transporte_cerrada_frena_los_correos_de_ilus_pero_no_los_avisos_de_terceros(self):
        self.m.bd.llaves["transporte"] = {"email_bloqueado": 1, "motivo": "Revisión"}
        d = self.listar()
        self.assertEqual(self.item("tr_en_ruta", d)["estado"]["efectivo"], "llave")
        for k in EXTERNOS:
            self.assertEqual(self.item(k, d)["estado"]["efectivo"], "sale", k)

    def test_el_interruptor_general_los_frena_a_todos(self):
        self.m.bd.global_on = False
        d = self.listar()
        for k in EXTERNOS:
            self.assertEqual(self.item(k, d)["estado"]["efectivo"], "general", k)

    def test_se_apagan_y_se_encienden_como_cualquier_correo_y_la_puerta_lo_ve(self):
        k = "ext_fedex_notif_envio"
        self.assertTrue(self.m.amb["_correo_externo_permitido"](k))
        st, d, _ = self.put(k, activo=False, motivo="Pruebas de etiquetas", estado_visto="")
        self.assertEqual(st, 200, d)
        self.assertEqual(d["correo"]["estado"]["efectivo"], "correo")
        self.assertFalse(self.m.amb["_correo_externo_permitido"](k))
        st, d, _ = self.put(k, activo=True, estado_visto=d["correo"]["estado"]["version"])
        self.assertEqual(st, 200, d)
        self.assertTrue(self.m.amb["_correo_externo_permitido"](k))

    def test_no_tienen_plantilla_ni_prueba_porque_el_texto_es_del_tercero(self):
        st, d, _ = self.m.llamar("comm_correos_vista_previa", "ext_fedex_notif_envio")
        self.assertEqual(st, 200)
        self.assertTrue(d["sin_plantilla"])
        self.assertIn("FedEx", d["mensaje"])
        st, d, _ = self.m.llamar("comm_correos_prueba", "ext_fedex_notif_envio")
        self.assertEqual((st, d["error_codigo"]), (400, "NO_APLICA"))
        self.assertEqual(self.m.enviados_reales, [])
        self.assertEqual(self.m.bd.log, [])

    def test_las_cifras_cuentan_lo_frenado_y_no_inventan_enviados(self):
        self.m.bd.poner_log("ext_fedex_notif_envio", "bloqueado", hace_dias=1)
        self.m.bd.poner_log("ext_fedex_notif_envio", "bloqueado", hace_dias=2)
        c = self.item("ext_fedex_notif_envio")["cifras"]
        self.assertEqual((c["enviados"], c["frenados"]), (0, 2))


class TestExternosEnElHistorial(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        TestHistorialControlDeCorreos.setUpClass()
        cls.f = staticmethod(TestHistorialControlDeCorreos.amb["_comm_hist_llave"])      # la función real de app.py, sin «self»

    def test_un_aviso_de_tercero_frenado_por_su_interruptor_se_explica_con_su_nombre(self):
        nombre = R.nombre_de("ext_fedex_notif_envio")
        detalle = f"Apagado en Control de correos: «{nombre}» (ext_fedex_notif_envio). ILUS no le entregó el correo del destinatario."
        self.assertEqual(self.f(detalle, "transporte", {"bloqueo": "correo", "externo": True, "correo": "ext_fedex_notif_envio"}),
                         "Apagado: «" + nombre + "»")

    def test_frenado_por_el_interruptor_general_dice_general(self):
        detalle = "Email deshabilitado por superadmin (kill switch global ON). ILUS no le entregó el correo del destinatario."
        self.assertEqual(self.f(detalle, "transporte", {"killswitch": True, "externo": True}),
                         "Interruptor general de email (todos los módulos)")


class TestExternosEnLaInterfaz(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.js = _leer("static/comunicaciones_correos.js")

    def test_la_tarjeta_de_un_aviso_de_tercero_no_ofrece_plantilla_vista_previa_ni_prueba(self):
        i = self.js.index("function tarjeta(c)")
        cuerpo = self.js[i:self.js.index("function renderLista")]
        self.assertIn("c.tipo !== 'externo'", cuerpo)
        self.assertLess(cuerpo.index("c.tipo !== 'externo'"), cuerpo.index('data-ctc="plantilla"'))
        self.assertLess(cuerpo.index('data-ctc="prueba"'), cuerpo.index('data-ctc="historial"'))
        self.assertIn('data-ctc="cambios"', cuerpo)

    def test_el_paso_de_la_llave_dice_que_no_aplica(self):
        self.assertIn("No aplica a este aviso", self.js)

    def test_no_muestra_enviados_inventados_y_explica_quien_envia(self):
        self.assertIn("lo envía FedEx o SimpliRoute", self.js)
        self.assertIn("ILUS no los cuenta", self.js)


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  9) UN CORREO CUYO ENVÍO VIVE EN OTRA RAMA: «Falta de stock: aviso a bodega» (ot_que_paso_bodega)
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════
KB = "ot_que_paso_bodega"
# El envío TAL COMO LLEGARÁ desde la rama de «¿Qué pasó?» de la OT (sin `correo=`). Lo que importa: con el catálogo ya funciona.
ENVIO_DE_OTRA_RAMA = '''
def _ot_que_paso_avisar_bodega(vid, d_b, asunto, html, codigo):
    return _send_ilus_email(d_b, asunto, html, evento="ot_que_paso_bodega", modulo="comunicacion_interna",
                            asincrono=True, ref={"tipo": "OT", "id": vid, "codigo": codigo})
'''


def _llamadas_de_un_texto(texto):
    """Las llamadas de envío de un TEXTO de ejemplo, con el mismo escáner (`FC.llamadas`) y el mismo formato que
    `_llamadas_de_envio`: así se prueba la regla de cobertura con un envío que todavía no está en el código de esta rama."""
    rel = "__envio_de_otra_rama__.py"
    clave_cache = (rel, tuple(sorted(ENVOLTORIOS)))
    FC._TEXTOS[rel] = texto
    FC._LLAMADAS.pop(clave_cache, None)
    try:
        salida = []
        for ll in FC.llamadas(rel, set(ENVOLTORIOS)):
            kw = {n: ast.parse(v.strip(), mode="eval").body for n, v in ll["kw"].items() if n != "**"}
            salida.append((rel, ll["linea"], ll["nombre"], ll["funcion"], ll["params"], kw))
        return salida
    finally:
        FC._TEXTOS.pop(rel, None)
        FC._LLAMADAS.pop(clave_cache, None)


class TestCorreoDeOtraRama(unittest.TestCase):
    """«Falta de stock: aviso a bodega» (`ot_que_paso_bodega`): lo envía la función «¿Qué pasó?» de la OT, que vive en otra rama aún
    sin mezclar. El catálogo ya lo trae para que, el día de la mezcla, el interruptor exista y ninguna prueba falle."""

    def enviar_como_llegara(self, m):
        return m.amb["_send_ilus_email"]("bodega@ejemplo.cl", "Falta de stock", "<p>x</p>", evento="ot_que_paso_bodega",
                                         modulo="comunicacion_interna", asincrono=True,
                                         ref={"tipo": "OT", "id": 7, "codigo": "OT-2026-00007"})

    def test_esta_en_el_registro_como_aviso_interno_a_bodega_dentro_del_grupo_de_las_ot(self):
        c = R.por_clave(KB)
        self.assertIsNotNone(c)
        self.assertEqual((c["tipo"], c["modulo_ui"], c["llave"], c["destinatario"], c["modo"]),
                         ("ilus", "ot", "comunicacion_interna", "bodega", "automatico"))
        self.assertEqual(c["nombre"], "Falta de stock: aviso a bodega")
        self.assertIn("bodega", c["avisa_a"].lower())
        self.assertIn("nunca al cliente ni al técnico", c["avisa_a"])          # la leyenda de a quién le llega
        self.assertIsNone(c["plantilla"])                                       # su texto lo arma el código, no el editor
        self.assertFalse(c["protegido"] or c["sensible"])
        self.assertIn(KB, R.claves())
        self.assertIn("comunicacion_interna", _llaves_validas())
        self.assertEqual(R.validar(_llaves_validas()), [])

    def test_nace_encendido_y_sin_ninguna_condicion_de_configuracion(self):
        c = R.por_clave(KB)
        self.assertEqual(c["condiciones"], ())
        self.assertFalse(c["apagado_en_codigo"])
        e = R.estado_efectivo(c, environ={})
        self.assertEqual(e["efectivo"], "sale")
        self.assertTrue(e["puede_apagar"] and e["puede_encender"])

    def test_el_evento_se_resuelve_a_su_clave_con_o_sin_los_prefijos_del_historial_y_de_la_prueba(self):
        self.assertEqual(R.inferir("comunicacion_interna", "ot_que_paso_bodega"), KB)
        self.assertEqual(R.inferir("comunicacion_interna", "retry_ot_que_paso_bodega"), KB)        # «Reenviar» del Historial
        self.assertEqual(R.inferir("comunicacion_interna", "prueba_ot_que_paso_bodega"), KB)       # «Enviarme una prueba»

    def test_la_puerta_lo_deja_salir_encendido_y_lo_encola_con_su_clave_en_su_columna(self):
        m = Mundo()
        self.assertTrue(self.enviar_como_llegara(m))
        (sql, params), = m.cola
        self.assertIn("correo_clave", sql)
        self.assertIn(KB, params)
        self.assertIn("ref_json", sql)                                          # la referencia a la OT sigue viajando
        self.assertEqual(m.salieron, [])                                        # no se envió síncrono
        self.assertIsNone(m.amb["_email_ultimo_bloqueo"]())

    def test_apagado_no_se_encola_y_queda_bloqueado_con_su_clave_y_su_cuerpo(self):
        m = Mundo(filas={KB: 0})
        self.assertFalse(self.enviar_como_llegara(m))
        self.assertEqual(m.cola, [])
        self.assertEqual(m.salieron, [])
        self.assertEqual(m.estados_log(), [(KB, "bloqueado")])
        (args, kw), = m.logs
        self.assertEqual(kw["metadata"], {"modulo": "comunicacion_interna", "bloqueo": "correo", "correo": KB})
        self.assertIn("Falta de stock: aviso a bodega", kw["error_msg"])
        self.assertEqual(kw["html"], "<p>x</p>")
        self.assertEqual(kw["ref"], {"tipo": "OT", "id": 7, "codigo": "OT-2026-00007"})      # el Historial lo enlaza con su OT
        r = m.amb["_email_resultado"](False)
        self.assertEqual((r["causa"], r["correo"]), ("apagado", KB))

    def test_el_interruptor_general_y_la_llave_de_comunicacion_interna_van_antes(self):
        m = Mundo(global_on=False, filas={KB: 0})
        self.assertFalse(self.enviar_como_llegara(m))
        self.assertEqual(m.amb["_email_ultimo_bloqueo"]()["tipo"], "global")
        self.assertEqual(m.estados_log(), [(KB, "bloqueado")])
        m = Mundo(llaves_cerradas={"comunicacion_interna"}, filas={KB: 0})
        self.assertFalse(self.enviar_como_llegara(m))
        self.assertEqual(m.amb["_email_ultimo_bloqueo"]()["tipo"], "modulo")

    def test_si_se_apaga_mientras_espera_en_la_cola_no_sale(self):
        m = Mundo()
        self.assertTrue(self.enviar_como_llegara(m))
        m.filas[KB] = 0
        m.amb["_correo_ctrl_invalidar"]()
        job = {"id": 11, "destinatario": "bodega@ejemplo.cl", "asunto": "Falta de stock", "cuerpo": "<p>x</p>",
               "evento": "ot_que_paso_bodega", "modulo": "comunicacion_interna", "intentos": 1, "max_intentos": 5,
               "kwargs_json": None, "correo_clave": KB}
        m.amb["_email_queue_send_job"](job)
        self.assertEqual(m.salieron, [])
        ult = m.logs[-1]
        self.assertEqual((ult[0][3], ult[1]["correo"]), ("bloqueado", KB))
        # un job de antes de la columna (otra instancia en pleno deploy) igual se reconoce por su evento
        job_viejo = {k: v for k, v in job.items() if k != "correo_clave"}
        m.amb["_email_queue_send_job"](dict(job_viejo, id=12))
        self.assertEqual(m.salieron, [])

    def test_apagarlo_no_frena_los_otros_avisos_a_bodega(self):
        m = Mundo(filas={KB: 0})
        self.assertTrue(m.enviar(modulo="comunicacion_interna", correo="ret_aviso_bodega", evento="retiros_bodega"))
        self.assertEqual(len(m.salieron), 1)

    def test_la_regla_de_cobertura_acepta_el_envio_tal_como_llegara_y_sigue_vigilando(self):
        ll = _llamadas_de_un_texto(ENVIO_DE_OTRA_RAMA)
        self.assertEqual(len(ll), 1, "el escáner debe ver la llamada de ejemplo")
        self.assertEqual(ll[0][2:4], ("_send_ilus_email", "_ot_que_paso_avisar_bodega"))
        self.assertNotIn("correo", ll[0][5])                                    # sin correo=: lo resuelve su evento
        self.assertEqual(_problemas_de_claves(ll), [])
        self.assertEqual(_problemas_de_llave(ll), [])
        # un evento que el registro no conoce, sin correo=, SÍ se señala
        malo = _llamadas_de_un_texto(ENVIO_DE_OTRA_RAMA.replace("ot_que_paso_bodega", "ot_evento_que_nadie_registro"))
        self.assertEqual(len(_problemas_de_claves(malo)), 1)
        self.assertIn("falta correo=", _problemas_de_claves(malo)[0])
        # el evento bueno pero por otra llave de paso: la clave se deduce, la llave no coincide con el registro
        otra = _llamadas_de_un_texto(ENVIO_DE_OTRA_RAMA.replace('modulo="comunicacion_interna"', 'modulo="retiros"'))
        self.assertEqual(_problemas_de_claves(otra), [])
        self.assertEqual(len(_problemas_de_llave(otra)), 1)
        # con la clave explícita (lo recomendado al mezclar) también pasa
        explicita = _llamadas_de_un_texto(ENVIO_DE_OTRA_RAMA.replace("asincrono=True,", 'asincrono=True, correo="ot_que_paso_bodega",'))
        self.assertEqual((_problemas_de_claves(explicita), _problemas_de_llave(explicita)), ([], []))

    def test_la_lista_de_envios_en_otra_rama_solo_trae_correos_del_registro_con_su_razon(self):
        self.assertIn(KB, ENVIO_EN_OTRA_RAMA)
        for clave, razon in ENVIO_EN_OTRA_RAMA.items():
            c = R.por_clave(clave)
            self.assertIsNotNone(c, clave)
            self.assertEqual(c["tipo"], "ilus", clave)
            self.assertGreater(len(razon), 30, clave)


class TestCorreoDeOtraRamaEnLaPantalla(_PantallaBase):
    def test_aparece_en_el_grupo_de_las_ot_para_bodega_con_la_leyenda_y_nace_encendido(self):
        d = self.listar()
        c = self.item(KB, d)
        self.assertEqual((c["modulo_ui"], c["destinatario"], c["destinatario_label"], c["llave"], c["modo"]),
                         ("ot", "bodega", "Bodega", "comunicacion_interna", "automatico"))
        self.assertEqual(c["nombre"], "Falta de stock: aviso a bodega")
        self.assertIn("nunca al cliente ni al técnico", c["avisa_a"])
        self.assertEqual(c["llave_label"], "Comunicación interna")
        self.assertEqual(c["estado"]["efectivo"], "sale")
        self.assertTrue(c["estado"]["switch_activo"] and c["estado"]["puede_apagar"])
        self.assertIsNone(c["plantilla"])
        self.assertIn("bodega", {x["destinatario"] for x in d["correos"] if x["modulo_ui"] == "ot"})   # un subgrupo «Para: Bodega»
        self.assertEqual(d["resumen"]["apagados_usuario"], 0)

    def test_apagarlo_con_motivo_lo_frena_en_la_puerta_y_el_historial_lo_cuenta(self):
        st, d, _ = self.put(KB, activo=False, motivo="Bodega prefiere avisarse sola", estado_visto="")
        self.assertEqual(st, 200, d)
        self.assertEqual((d["correo"]["estado"]["efectivo"], d["correo"]["estado"]["apagado_por"]), ("correo", "Daniel Aguilar"))
        ok = self.m.enviar(None, modulo="comunicacion_interna", evento="ot_que_paso_bodega", to="bodega@ejemplo.cl")
        self.assertFalse(ok)
        self.assertEqual(self.m.enviados_reales, [])
        ultimo = self.m.bd.log[-1]
        self.assertEqual((ultimo["estado"], ultimo["correo_clave"]), ("bloqueado", KB))      # el Historial y las cifras lo cuentan por su clave
        self.assertEqual(ultimo["metadata"]["bloqueo"], "correo")

    def test_la_vista_previa_dice_que_el_texto_esta_en_el_codigo_y_no_envia_nada(self):
        st, d, _ = self.m.llamar("comm_correos_vista_previa", KB)
        self.assertEqual(st, 200)
        self.assertTrue(d["sin_plantilla"])
        self.assertIn("en el código", d["mensaje"])
        self.nada_cambio()

    def test_la_prueba_va_solo_a_quien_esta_conectado_nunca_a_bodega(self):
        st, d, _ = self.m.llamar("comm_correos_prueba", KB, json_body={"to": "bodega@ilusfitness.com"})
        self.assertEqual((st, d), (200, {"ok": True, "enviado_a": "daniel.aguilar@sphs.cl"}))
        env, = self.m.enviados_reales
        self.assertEqual(env["to"], "daniel.aguilar@sphs.cl")
        self.assertTrue(env["asunto"].startswith("[PRUEBA] "))
        self.assertEqual(self.m.bd.log[-1]["evento"], "prueba_" + KB)


if __name__ == "__main__":
    unittest.main()
