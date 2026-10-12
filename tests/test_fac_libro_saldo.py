"""📒 LIBRO DE SALDO DE FACTURAS + OT a pérdida con autorización (Daniel, 11-oct-2026, antes de la auditoría del martes).

Daniel: «Si ya utilicé una factura y se ocupó el saldo en otra orden de trabajo, ese saldo tiene que quedar bloqueado, a
menos que utilicemos un porcentaje nada más… si tengo 100 mil y uso 10 por un lado y 80 por otro, me quedan 10 mil».
«Solo consume el saldo la OT, que es la realización del servicio; [el ticket y la cotización] heredan». «Te va a indicar
si la factura va a pérdida y eso iría autorizado por mí».

Qué se vigila:
  · La lógica pura (ot_saldo_servicio.py): el ejemplo de los $100.000, los semáforos, los movimientos del libro y la
    pérdida autorizada HASTA un monto.
  · El libro en app.py con un mundo en memoria: consumo, sin cambios (la huella no escribe), liberación al cancelar,
    «carga inicial» para las OT que ya existían, y que el libro NUNCA edita ni borra un movimiento.
  · El candado de pérdida en el cierre con firma (después del costo del proveedor), su autorización nueva en la MISMA
    tabla y con la MISMA columna (REGLA #18: una sola sentencia define el ENUM), y el paso 6 del motor que la ofrece.
  · La herencia: la OT que nace de un ticket trae sus documentos (asistente), y la pantalla muestra dónde viaja la factura.
  · Confidencialidad (REGLA #26): toda ruta nueva con plata es @_no_tecnico.
Sin BD, sin Flask, sin ERP (REGLA #4.1).
Correr con:  py -m unittest tests.test_fac_libro_saldo
"""
import ast
import datetime as _dt
import hashlib
import json
import os
import re
import unittest

import ot_saldo_servicio as S
from tests.test_ot_saldo_servicio import _arbol, _fuente, _decoradores, _leer


def _exec_clase(nombre, amb):
    _codigo, arbol = _arbol()
    for n in arbol.body:
        if isinstance(n, ast.ClassDef) and n.name == nombre:
            exec(compile(ast.Module(body=[ast.parse(ast.unparse(n)).body[0]], type_ignores=[]), "<app>", "exec"), amb)
            return amb[nombre]
    raise AssertionError("no está " + nombre)


def _exec(nombre, amb):
    _codigo, arbol = _arbol()
    for n in arbol.body:
        if isinstance(n, ast.FunctionDef) and n.name == nombre:
            n2 = ast.parse(ast.unparse(n)).body[0]
            n2.decorator_list = []
            exec(compile(ast.Module(body=[n2], type_ignores=[]), "<app>", "exec"), amb)
            return amb[nombre]
    raise AssertionError("no está " + nombre)


# ───────────────────────── lógica pura ─────────────────────────
class TestEjemploDeDaniel(unittest.TestCase):
    def test_100_mil_usa_10_y_80_quedan_10(self):
        e = S.estado_factura({"total": 100000, "usado": 10000 + 80000}, {"total": 0, "usado": 0})
        self.assertEqual(e["estado"], "con_saldo")
        self.assertEqual(e["saldo"], 10000)
        self.assertEqual(e["pct_usado"], 90.0)
        v = S.puede_usar({"total": 100000, "usado": 90000}, {"total": 0, "usado": 0})
        self.assertTrue(v["ok"])
        self.assertIn("$10.000", v["texto"])

    def test_usada_completa_queda_bloqueada(self):
        e = S.estado_factura({"total": 100000, "usado": 100000}, {"total": 30000, "usado": 30000})
        self.assertEqual(e["estado"], "agotada")
        v = S.puede_usar({"total": 100000, "usado": 100000}, {"total": 30000, "usado": 30000})
        self.assertFalse(v["ok"])
        self.assertIn("autorización de Daniel", v["texto"])

    def test_pedido_mayor_al_saldo_no_alcanza(self):
        v = S.puede_usar({"total": 100000, "usado": 90000}, {"total": 0, "usado": 0}, {"servicio": 20000})
        self.assertFalse(v["ok"])
        self.assertEqual(v["titulo"], "No alcanza el saldo")

    def test_doble_cobro_se_ve_en_rojo(self):
        e = S.estado_factura({"total": 100000, "usado": 150000}, {"total": 0, "usado": 0})
        self.assertEqual(e["estado"], "excedida")
        self.assertEqual(e["exceso"], 50000)
        self.assertEqual(e["clase"], "rojo")

    def test_otros_estados(self):
        self.assertEqual(S.estado_factura({"total": 100, "usado": 0}, {})["estado"], "libre")
        self.assertEqual(S.estado_factura({"total": 0, "usado": 0}, {"total": 0, "usado": 0})["estado"], "sin_lineas")
        self.assertEqual(S.estado_factura({"total": 100, "usado": 0}, {}, leido=False)["estado"], "sin_leer")
        for k in S.ORDEN_ESTADOS:
            self.assertIn(k, S.ESTADOS_FACTURA)


class TestMovimientosDelLibro(unittest.TestCase):
    def test_consumo_liberacion_y_sin_cambio(self):
        a = {("FCV", "1", "servicio"): 50000, ("FCV", "1", "despacho"): 10000, ("FCV", "2", "servicio"): 5000}
        b = {("FCV", "1", "servicio"): 80000, ("FCV", "1", "despacho"): 10000}
        movs = S.movimientos_libro(a, b)
        self.assertEqual(len(movs), 2)
        m1 = next(m for m in movs if m["nudo"] == "1")
        self.assertEqual((m1["movimiento"], m1["monto"], m1["acumulado"]), ("consumo", 30000, 80000))
        m2 = next(m for m in movs if m["nudo"] == "2")
        self.assertEqual((m2["movimiento"], m2["monto"], m2["acumulado"]), ("liberacion", -5000, 0))
        self.assertEqual(S.movimientos_libro(b, dict(b)), [])

    def test_aportes_de_una_ot_solo_suyos_y_sin_canceladas(self):
        filas = [{"vid": 7, "estado": "programada", "clave": ("FCV", "1"), "servicio": 50000, "despacho": 0},
                 {"vid": 7, "estado": "programada", "clave": ("FCV", "1"), "servicio": 40000, "despacho": 9000},
                 {"vid": 8, "estado": "programada", "clave": ("FCV", "1"), "servicio": 999, "despacho": 0},
                 {"vid": 9, "estado": "cancelada", "clave": ("FCV", "1"), "servicio": 1, "despacho": 0}]
        self.assertEqual(S.aportes_de_ot(filas, 7), {("FCV", "1", "servicio"): 50000, ("FCV", "1", "despacho"): 9000})
        self.assertEqual(S.aportes_de_ot(filas, 9), {})


class TestAportesReales(unittest.TestCase):
    OT = {"estado": "cerrada", "zz_monto": 100000, "zz_envio_monto": 30000, "valor_origen": "zz"}

    def test_lo_antiguo_se_reparte_sin_duplicarse(self):
        # caso del revisor: FCV A solo servicio y FCV B solo despacho; antes se anotaban $260.000 de una OT de $130.000
        docs = [{"tido": "FCV", "nudo": "1", "es_cobro": 1, "serv": None, "env": None},
                {"tido": "FCV", "nudo": "2", "es_cobro": 1, "serv": None, "env": None}]
        caps = {("FCV", "1"): {"servicio": 100000, "despacho": 0}, ("FCV", "2"): {"servicio": 0, "despacho": 30000}}
        a = S.aportes_reales_ot(self.OT, docs, caps)
        self.assertEqual(a, {("FCV", "1", "servicio"): 100000, ("FCV", "2", "despacho"): 30000})
        self.assertEqual(sum(a.values()), 130000)

    def test_lo_anotado_manda_y_lo_reemplazado_no_cuenta(self):
        docs = [{"tido": "FCV", "nudo": "1", "es_cobro": 1, "serv": 60000, "env": None},
                {"tido": "NVV", "nudo": "VD00010667", "es_cobro": 1, "reemplazado": True, "serv": 60000, "env": None}]
        a = S.aportes_reales_ot(dict(self.OT, zz_monto=60000, zz_envio_monto=0), docs, {})
        self.assertEqual(a, {("FCV", "1", "servicio"): 60000})

    def test_no_consume_en_garantia_cancelada_o_con_cero_autorizado(self):
        docs = [{"tido": "FCV", "nudo": "1", "es_cobro": 1, "serv": 50000, "env": None}]
        for cambio in ({"estado": "cancelada"}, {"modalidad_cobro": "garantia"}, {"cobro_cero_motivo": "regalia"}):
            self.assertEqual(S.aportes_reales_ot(dict(self.OT, **cambio), docs, {}), {}, cambio)

    def test_valor_manual_no_es_plata_de_una_linea(self):
        docs = [{"tido": "FCV", "nudo": "1", "es_cobro": 1, "serv": None, "env": None}]
        self.assertEqual(S.aportes_reales_ot(dict(self.OT, valor_origen="manual"), docs, {}), {})

    def test_sin_tabla_puente_sale_de_factura_hasta_su_capacidad(self):
        ot = dict(self.OT, zz_monto=150000, zz_envio_monto=0, factura_tido="FCV", factura_nudo="0000011439")
        a = S.aportes_reales_ot(ot, [], {("FCV", "11439"): {"servicio": 100000, "despacho": 0}})
        self.assertEqual(a, {("FCV", "11439", "servicio"): 100000})

    def test_random_caido_es_lectura_incompleta(self):
        docs = [{"tido": "FCV", "nudo": "1", "es_cobro": 1, "serv": None, "env": None}]
        with self.assertRaises(S.LecturaIncompleta):
            S.aportes_reales_ot(self.OT, docs, {("FCV", "1"): None})


class TestPerdidaPura(unittest.TestCase):
    def test_sin_perdida_pasa(self):
        self.assertTrue(S.perdida_cierre(100000, 60000)["ok"])

    def test_perdida_sin_autorizar_no_pasa(self):
        r = S.perdida_cierre(100000, 125000)
        self.assertEqual((r["perdida"], r["ok"], r["falta_autorizar"]), (25000, False, 25000))

    def test_la_perdida_que_se_autoriza_es_la_del_servicio(self):
        # repuestos fuera (la factura los cobra como producto); despacho sin costo = $0 (mínimo)
        self.assertEqual(S.perdida_de_cuenta(50000, 200000, None)["perdida"], 150000)
        self.assertTrue(S.perdida_de_cuenta(100000, 60000, 20000)["ok"])

    def test_autorizada_hasta_un_monto(self):
        self.assertTrue(S.perdida_cierre(100000, 125000, autorizada=25000)["ok"])
        r = S.perdida_cierre(100000, 140000, autorizada=25000)   # creció después de autorizar: se vuelve a pedir
        self.assertFalse(r["ok"])
        self.assertEqual(r["falta_autorizar"], 15000)


# ───────────────────────── el libro en app.py (mundo en memoria) ─────────────────────────
class MundoLibro:
    def __init__(self):
        self.ots = {}        # vid → fila de mant_visitas (para la huella)
        self.aportes = {}    # vid → {(t, n, c): monto}
        self.libro = []      # filas de ot_fac_libro
        self.huellas = {}
        self.sql = []
        self.lock_ok = 1
        self.frescas = 0
        amb = {"print": lambda *a, **k: None, "_saldo": S, "datetime": _dt.datetime, "get_mysql": self._conn,
               "_fac_db_fresca": self._fresca, "get_db": self._db,
               "mysql_fetchone": self.fetchone, "mysql_fetchall": self.fetchall, "mysql_execute": self.execute,
               "_fac_aportes_ot": lambda vid: dict(self.aportes.get(int(vid), {})),
               "_ot_aut_usuario": lambda: (5, "Ejecutivo de SSTT"), "current_username": lambda: "ejecutivo",
               "chile_fmt_filter": lambda v, f="%d/%m/%Y %H:%M": v.strftime(f), "_FAC_LIBRO_DESDE": _dt.datetime(2026, 10, 11)}
        _exec("_fac_huella_ot", amb)
        _exec("_fac_libro_ultimo", amb)
        self.conciliar = _exec("_fac_libro_conciliar_ot", amb)

    def _fresca(self):
        self.frescas += 1

    def _db(self):
        mundo = self

        class D:
            def cursor(s):
                class K:
                    def __enter__(k): return k
                    def __exit__(k, *a): return False
                    def execute(k, sql, params=()): mundo.execute(sql, params)
                return K()
            def commit(s): mundo.commits = getattr(mundo, "commits", 0) + 1
            def rollback(s): pass
        return D()

    def _conn(self):
        mundo = self

        class C:
            def cursor(s):
                class K:
                    def __enter__(k): return k
                    def __exit__(k, *a): return False
                    def execute(k, *a): return None
                    def fetchone(k): return {"l": mundo.lock_ok}
                return K()
            def close(s): pass
        return C()

    def fetchone(self, sql, params=()):
        self.sql.append(sql)
        if "FROM mant_visitas v" in sql:
            v = self.ots.get(int(params[0]))
            if not v:
                return None
            return dict(v, h_prev=self.huellas.get(int(params[0])))
        if "FROM ot_fac_huella" in sql:
            h = self.huellas.get(int(params[0]))
            return {"huella": h} if h else None
        return None

    def fetchall(self, sql, params=()):
        self.sql.append(sql)
        if "FROM ot_fac_libro" in sql:
            vid = int(params[0])
            ult = {}
            for r in self.libro:
                if r["visita_id"] == vid:
                    ult[(r["doc_tido"], r["doc_nudo"], r["categoria"])] = r
            return list(ult.values())
        return []

    def execute(self, sql, params=()):
        self.sql.append(sql)
        if sql.startswith("INSERT INTO ot_fac_libro"):
            cols = ("doc_tido", "doc_nudo", "categoria", "movimiento", "monto", "acumulado_ot", "visita_id", "numero_ot",
                    "cliente", "origen", "usuario", "user_id", "detalle")
            self.libro.append(dict(zip(cols, params), id=len(self.libro) + 1))
        elif sql.startswith("INSERT INTO ot_fac_huella"):
            self.huellas[int(params[0])] = params[1]
        elif sql.startswith("DELETE FROM ot_fac_huella"):
            self.huellas.pop(int(params[0]), None)

    def ot(self, vid, creada=_dt.datetime(2026, 10, 12), **k):
        base = {"id": vid, "numero_ot": f"OT-2026-{vid:05d}", "created_at": creada, "estado": "programada",
                "modalidad_cobro": None, "cobro_cero_motivo": None, "valor_origen": "zz", "zz_monto": 0,
                "zz_envio_monto": 0, "factura_tido": "FCV", "factura_nudo": "0000011439", "cliente": "Gimnasio de ejemplo SpA",
                "docs": ""}
        base.update(k)
        self.ots[vid] = base


class TestLibroEnApp(unittest.TestCase):
    def test_consumo_y_sin_cambio_no_escribe(self):
        m = MundoLibro()
        m.ot(1, zz_monto=10000, docs="FCV:0000011439:1:-:10000:-")
        m.aportes[1] = {("FCV", "11439", "servicio"): 10000}
        r = m.conciliar(1, "crear")
        self.assertEqual(r["movimientos"], 1)
        self.assertEqual(m.libro[0]["movimiento"], "consumo")
        self.assertEqual(m.libro[0]["monto"], 10000)
        self.assertEqual(m.libro[0]["origen"], "crear")
        self.assertEqual(m.libro[0]["usuario"], "Ejecutivo de SSTT")
        # misma plata: la huella no cambió → ni una escritura más
        n = len(m.libro)
        self.assertEqual(m.conciliar(1, "edicion")["movimientos"], 0)
        self.assertEqual(len(m.libro), n)

    def test_dos_ot_10_y_80_y_al_cancelar_se_libera(self):
        m = MundoLibro()
        m.ot(1, zz_monto=10000, docs="a")
        m.aportes[1] = {("FCV", "11439", "servicio"): 10000}
        m.ot(2, zz_monto=80000, docs="b")
        m.aportes[2] = {("FCV", "11439", "servicio"): 80000}
        m.conciliar(1, "crear")
        m.conciliar(2, "crear")
        usado = sum(r["monto"] for r in m.libro)
        self.assertEqual(usado, 90000)       # de $100.000 quedan $10.000
        m.ots[2]["estado"] = "cancelada"     # la cancelada deja de consumir: se libera
        m.aportes[2] = {}
        r = m.conciliar(2, "estado")
        self.assertEqual(r["movimientos"], 1)
        self.assertEqual(m.libro[-1]["movimiento"], "liberacion")
        self.assertEqual(m.libro[-1]["monto"], -80000)
        self.assertEqual(sum(x["monto"] for x in m.libro), 10000)

    def test_ot_anterior_al_libro_es_carga_inicial(self):
        m = MundoLibro()
        m.ot(3, creada=_dt.datetime(2026, 9, 1), zz_monto=50000, docs="c")
        m.aportes[3] = {("FCV", "11439", "servicio"): 50000}
        m.conciliar(3, "revision")
        self.assertEqual(m.libro[0]["origen"], "inicial")
        self.assertEqual(m.libro[0]["usuario"], "Sistema")
        self.assertIn("11/10/2026", m.libro[0]["detalle"])

    def test_ot_eliminada_libera_todo(self):
        m = MundoLibro()
        m.ot(4, zz_monto=30000, docs="d")
        m.aportes[4] = {("FCV", "11439", "servicio"): 30000}
        m.conciliar(4, "crear")
        del m.ots[4]
        m.aportes[4] = {}
        r = m.conciliar(4, "eliminada")
        self.assertEqual(r["movimientos"], 1)
        self.assertEqual(m.libro[-1]["origen"], "eliminada")
        self.assertNotIn(4, m.huellas)

    def test_candado_ocupado_no_escribe_y_se_relee_con_foto_fresca(self):
        m = MundoLibro()
        m.ot(5, zz_monto=10000, docs="e")
        m.aportes[5] = {("FCV", "11439", "servicio"): 10000}
        m.lock_ok = 0
        self.assertEqual(m.conciliar(5, "crear")["movimientos"], 0)
        self.assertEqual(m.libro, [])
        self.assertNotIn(5, m.huellas)          # el barrido lo reintenta
        m.lock_ok = 1
        self.assertEqual(m.conciliar(5, "crear")["movimientos"], 1)
        self.assertGreaterEqual(m.frescas, 2)   # antes de leer y otra vez adentro del candado

    def test_random_caido_no_anota_ni_guarda_huella(self):
        m = MundoLibro()
        m.ot(6, zz_monto=10000, docs="f")

        def incompleta(vid):
            raise S.LecturaIncompleta("FCV 11439")
        m.conciliar.__globals__["_fac_aportes_ot"] = incompleta
        r = m.conciliar(6, "crear")
        self.assertTrue(r.get("incompleto"))
        self.assertEqual(m.libro, [])
        self.assertNotIn(6, m.huellas)

    def test_libro_y_huella_en_una_transaccion_y_huella_versionada(self):
        m = MundoLibro()
        m.ot(7, zz_monto=10000, docs="g")
        m.aportes[7] = {("FCV", "11439", "servicio"): 10000}
        m.conciliar(7, "crear")
        self.assertEqual(getattr(m, "commits", 0), 1)          # un solo commit para movimientos + huella
        self.assertIn('base = "v2|"', _fuente("_fac_huella_ot"))

    def test_el_libro_nunca_se_edita_ni_se_borra(self):
        app = _leer("app.py")
        self.assertNotRegex(app, r"UPDATE\s+ot_fac_libro")
        self.assertNotRegex(app, r"DELETE\s+FROM\s+ot_fac_libro")


class TestLecturaDeRandom(unittest.TestCase):
    """Revisión 2: «sin líneas» ≠ «no se leyó» ≠ «no existe o está anulado»."""

    def _f(self, cabecera, lineas, sql):
        olvidados = []
        amb = {"_mant_erp_doc_cached": lambda t, n: (cabecera, lineas),
               "_erp_zz_lineas": lambda t, n: (cabecera, [l for l in lineas if str(l.get("sku", "")).startswith("ZZ")], 0),
               "_saldo": S, "_ot_saldo_real": lambda k: (k[0], k[1].zfill(10)),
               "_fac_esdo": lambda t, n: (None if isinstance(sql, Exception) else sql),
               "_fac_olvidar_cache": lambda t, n: olvidados.append((t, n))}
        return _exec("_fac_leer_capacidad", amb)("FCV", "11439"), olvidados

    def test_leido_completo(self):
        (est, cap), _o = self._f({"x": 1}, [{"sku": "ZZINSTALACION", "monto": 100000}, {"sku": "BANCO1"}], [])
        self.assertEqual((est, cap["servicio"]), ("ok", 100000))

    def test_cabecera_sin_lineas_es_no_leido_y_sale_de_la_cache(self):
        (est, cap), olv = self._f({"x": 1}, [], [])
        self.assertEqual((est, cap), ("no_leido", None))
        self.assertEqual(olv, [("FCV", "0000011439")])   # se olvida en la forma real (la que usa la caché)

    def test_random_caido_es_no_leido(self):
        self.assertEqual(self._f(None, [], None)[0], ("no_leido", None))
        self.assertEqual(self._f(None, [], RuntimeError("x"))[0], ("no_leido", None))

    def test_anulada_o_inexistente_no_respalda_cobro(self):
        self.assertEqual(self._f(None, [], [])[0][0], "sin_doc")
        self.assertEqual(self._f(None, [], [{"ESDO": "NULO"}])[0], ("sin_doc", {"servicio": 0, "despacho": 0}))
        self.assertEqual(self._f(None, [], [{"ESDO": ""}])[0][0], "no_leido")   # existe, pero la lectura falló

    def test_solo_cache_nunca_sale_a_la_red(self):
        import threading
        import time as _t
        llamadas = []
        cache = {}
        amb = {"_mant_erp_doc_cached": lambda t, n: llamadas.append((t, n)) or (None, []),
               "_erp_zz_lineas": lambda t, n: (None, [], 0), "_saldo": S, "_ot_saldo_real": lambda k: k,
               "_fac_esdo": lambda t, n: llamadas.append("esdo") or [], "_fac_olvidar_cache": lambda t, n: None,
               "_MANT_ERP_DOC_CACHE": cache, "_MANT_ERP_DOC_CACHE_LOCK": threading.Lock(), "_MANT_ERP_DOC_CACHE_TTL": 300}
        f = _exec("_fac_leer_capacidad", amb)
        self.assertEqual(f("FCV", "11439", solo_cache=True), ("no_leido", None))
        self.assertEqual(llamadas, [])                                   # ni el ERP ni la consulta de existencia
        cache["FCV|11439"] = (_t.time(), {"x": 1}, [{"sku": "ZZINSTALACION", "monto": 5}])
        amb["_erp_zz_lineas"] = lambda t, n: ({"x": 1}, [{"sku": "ZZINSTALACION", "monto": 5}], 0)
        amb["_mant_erp_doc_cached"] = lambda t, n: (cache["FCV|11439"][1], cache["FCV|11439"][2])
        self.assertEqual(f("FCV", "11439", solo_cache=True)[0], "ok")    # con lectura completa en caché, sí responde

    def test_la_consulta_de_existencia_tiene_cache(self):
        import threading
        n = []
        amb = {"_random_sql_query": lambda *a, **k: n.append(1) or [{"ESDO": ""}], "_FAC_ESDO_CACHE": {},
               "_FAC_ESDO_CACHE_LOCK": threading.Lock()}
        f = _exec("_fac_esdo", amb)
        f("FCV", "0000011439")
        f("FCV", "0000011439")
        self.assertEqual(len(n), 1)

    def test_capacidades_perezosas_no_leen_si_la_ot_no_consume(self):
        leidas = []
        amb = {"_fac_leer_capacidad": lambda t, n, solo_cache=False: leidas.append((t, n)) or ("ok", {"servicio": 1, "despacho": 0})}
        Caps = _exec_clase("_FacCapsPerezosas", amb)
        caps = Caps({("FCV", "1"): ("FCV", "0000000001")})
        ot = {"estado": "cancelada", "zz_monto": 100000, "valor_origen": "zz"}
        S.aportes_reales_ot(ot, [{"tido": "FCV", "nudo": "1", "es_cobro": 1, "serv": None, "env": None}], caps)
        self.assertEqual(leidas, [])
        S.aportes_reales_ot(dict(ot, estado="cerrada"), [{"tido": "FCV", "nudo": "1", "es_cobro": 1, "serv": None, "env": None}], caps)
        self.assertEqual(leidas, [("FCV", "0000000001")])

    def test_fuente_2_solo_con_residuo_y_sin_red(self):
        u = _fuente("_ot_saldo_usos")
        self.assertIn("if not (s or e_):", u)
        self.assertIn("_fac_aportes_ot(_vr, solo_cache=True)", u)
        self.assertIn("_memo_ap", u)

    def test_pausa_tras_lectura_incompleta_y_fotos_que_rotan(self):
        self.assertIn("_FAC_INCOMPLETO_HASTA[vid] = time.time() + 300", _fuente("_fac_libro_tras_escritura"))
        self.assertIn("intentado_el", _fuente("_fac_libro_barrido"))
        self.assertIn("ADD COLUMN intentado_el", _fuente("_ensure_fac_libro"))

    def test_el_motor_no_toma_cabecera_sin_lineas_por_factura_sin_lineas(self):
        f = _fuente("_ot_saldo_doc")
        self.assertIn("if _hdr0 and not _lin0:", f)
        self.assertIn('res["usos_incompletos"] = bool(_err_usos)', f)
        u = _fuente("_ot_saldo_usos")
        self.assertEqual(u.count("errores.append("), 3)
        self.assertIn('if s.get("usos_incompletos"):', _fuente("_fac_doc_refrescar"))

    def test_cierre_falla_cerrado_si_no_puede_revisar_la_perdida(self):
        f = _fuente("mant_ot_aprobar_cierre")
        i = f.index("_fp_c = _fac_perdida_de_ot(vid)")
        self.assertIn("REVISION_NO_DISPONIBLE", f[i:i + 700])

    def test_el_barrido_sigue_a_las_ot_que_pasaron_por_el_libro(self):
        app = _leer("app.py")
        i = app.index("_FAC_SQL_OT_CON_DOC = ")
        self.assertIn("FROM ot_fac_huella h", app[i:i + 900])
        self.assertIn("FROM ot_fac_libro l", app[i:i + 900])
        self.assertIn("l.acumulado_ot>0", _fuente("_fac_libro_barrido"))


class TestCableado(unittest.TestCase):
    def test_tablas_al_arrancar_y_siempre(self):
        app = _leer("app.py")
        self.assertIn("_ensure_fac_libro()", app)
        f = _fuente("_ensure_fac_libro")
        for t in ("ot_fac_libro", "ot_fac_doc", "ot_fac_huella", "ot_fac_cursor"):
            self.assertIn("CREATE TABLE IF NOT EXISTS " + t, f)

    def test_registro_tras_cada_escritura(self):
        f = _fuente("_fac_libro_tras_escritura")
        self.assertIn('"/ot/"', f)
        self.assertIn('"/mantenciones/"', f)
        self.assertIn('"/tickets/"', f)
        self.assertIn("_fac_libro_conciliar_ot(", f)
        self.assertIn("_fac_doc_refrescar(", f)
        self.assertIn("app.after_request", _decoradores("_fac_libro_tras_escritura"))
        self.assertIn("_fac_libro_marcar(vid)", _fuente("_ot_saldo_reanotar_aportes"))

    def test_mismo_motor_que_el_candado(self):
        f = _fuente("_fac_aportes_ot")
        self.assertIn("_saldo.aportes_reales_ot(v, docs, _FacCapsPerezosas(formas, solo_cache=solo_cache))", f)
        self.assertIn("_ot_saldo_doc(", _fuente("_fac_doc_refrescar"))
        u = _fuente("_ot_saldo_usos")
        self.assertEqual(u.count("if solo:"), 3)

    def test_barrido_programado_cada_10_minutos(self):
        self.assertIn("_fac_libro_barrido(", _fuente("tr_cron_simpliroute_poll"))

    def test_rutas_nuevas_no_son_para_tecnicos(self):
        for r in ("ot_saldo_facturas_pagina", "ot_api_saldo_facturas", "ot_api_saldo_facturas_doc",
                  "ot_api_saldo_facturas_revisar", "ot_saldo_facturas_excel"):
            self.assertIn("_no_tecnico", _decoradores(r), r)
            self.assertIn("_ot_solo_superadmin", _decoradores(r), r)

    def test_los_tres_solo_superadmin(self):
        # Daniel 2026-10-11: «los tres solo superadmin» — Regularizar, Saldo de facturas y la bandeja de Autorizaciones.
        for r in ("ot_regularizar_pagina", "ot_api_regularizar", "ot_aut_pagina"):
            self.assertIn("_ot_solo_superadmin", _decoradores(r), r)
        # el detalle de UNA solicitud sigue abierto a quien la pidió (la API ya filtra: solo las propias)
        self.assertNotIn("_ot_solo_superadmin", _decoradores("ot_aut_pagina_detalle"))
        self.assertNotIn("_ot_solo_superadmin", _decoradores("ot_aut_api_detalle"))
        self.assertNotIn("url_for('ot_saldo_facturas_pagina')", _leer("templates/base.html"))

    def test_decorador_solo_superadmin(self):
        import types
        G = types.SimpleNamespace(user={"role": "ejecutivo"}, permissions={})
        req = types.SimpleNamespace(path="/ot/regularizar")
        amb = {"g": G, "request": req, "wraps": __import__("functools").wraps, "jsonify": lambda d: d,
               "flash": lambda *a, **k: None, "redirect": lambda u: ("redirect", u), "url_for": lambda n, **k: "/" + n}
        _exec("_ot_aut_es_superadmin", amb)
        _exec("_ot_fin_es_superadmin", amb)
        vista = _exec("_ot_solo_superadmin", amb)(lambda: "OK")
        self.assertEqual(vista(), ("redirect", "/ot2_panel"))
        req.path = "/ot/api/saldo-facturas"
        r = vista()
        self.assertEqual(r[1], 403)
        self.assertFalse(r[0]["ok"])
        G.user, G.permissions = {"role": "superadmin"}, {"superadmin": True}
        self.assertEqual(vista(), "OK")

    def test_esperando_lleva_a_su_solicitud(self):
        f = _exec("_ot_cierre_accion", {"_OT_CIERRE_ACCIONES": {"AUTORIZACION_PENDIENTE": ("esperar_autorizacion", "x")}})
        self.assertEqual(f("AUTORIZACION_PENDIENTE", 5, aid=9)["url"], "/ot/autorizaciones/9")
        self.assertEqual(f("AUTORIZACION_PENDIENTE", 5)["url"], "/ot/5")


class TestPerdidaEnElCierre(unittest.TestCase):
    def test_tipo_nuevo_en_la_misma_tabla_y_una_sola_sentencia(self):
        app = _leer("app.py")
        self.assertIn('"cerrar_con_perdida")', app.split("_OT_AUT_TIPOS = ", 1)[1].split("\n", 1)[0])
        self.assertIn("'exceder_saldo','cerrar_con_perdida') NOT NULL", app)
        self.assertNotIn("MODIFY tipo ENUM('crear_sin_documento','cobro_cero','cerrar_sin_documento','exceder_saldo') NOT NULL", app)

    def test_candado_despues_del_costo_del_proveedor_y_antes_del_anexo(self):
        f = _fuente("mant_ot_aprobar_cierre")
        i_prov = f.index('"error_codigo": "SIN_COSTO_PROVEEDOR"')
        i_per = f.index('"CIERRE_CON_PERDIDA"')
        self.assertLess(i_prov, i_per)
        self.assertIn("_fac_perdida_de_ot(vid)", f)
        self.assertIn('_cob_cierre == "cobra"', f[i_prov:i_per + 200])
        self.assertIn('"accion": _ot_cierre_accion("CIERRE_CON_PERDIDA", vid)', f)

    def test_perdida_info_lee_la_autorizacion_mayor(self):
        filas = [{"id": 1, "estado": "aprobada", "payload_json": json.dumps({"perdida": 10000}), "resuelto_por_nombre": "Daniel"},
                 {"id": 2, "estado": "aprobada", "payload_json": json.dumps({"perdida": 25000}), "resuelto_por_nombre": "Daniel Aguilar"},
                 {"id": 3, "estado": "pendiente", "payload_json": "{}", "resuelto_por_nombre": None}]
        amb = {"mysql_fetchall": lambda *a, **k: filas, "json": json, "_saldo": S, "print": lambda *a, **k: None}
        f = _exec("_fac_perdida_info", amb)
        fin = {"cobra": True, "cobre": {"total": 100000, "hay": True},
               "me_cobraron": {"total": 175000, "tecnico": 125000, "despacho": None, "repuestos": 50000, "falta_tecnico": False},
               "queda": {"mostrar": False}}   # despacho sin costo: igual se revisa (antes esto esquivaba el candado)
        r = f(9, fin)
        self.assertTrue(r["completa"])
        self.assertEqual(r["perdida"], 25000)          # repuestos ($50.000) aparte
        self.assertEqual(r["repuestos"], 50000)
        self.assertTrue(r["ok"])
        self.assertEqual(r["autorizada"], 25000)
        self.assertEqual(r["autorizada_por"], "Daniel Aguilar")
        self.assertEqual(r["pendiente_id"], 3)
        fin["me_cobraron"]["tecnico"] = 150000
        self.assertFalse(f(9, fin)["ok"])
        fin["me_cobraron"]["falta_tecnico"] = True
        self.assertFalse(f(9, fin)["completa"])

    def test_cerrar_sin_documento_no_exime_una_perdida_con_cobro(self):
        f = _exec("_fac_perdida_aplica", {})
        self.assertTrue(f("documento", {"cobre": 50000}))
        self.assertTrue(f("autorizacion_cerrar", {"cobre": 50000}))
        self.assertFalse(f("autorizacion_cerrar", {"cobre": 0}))
        self.assertIn("_fac_perdida_aplica(_via_cierre, _fp_c)", _fuente("mant_ot_aprobar_cierre"))
        self.assertIn('_fac_perdida_aplica(p_cerrar.get("via"), perdida)', _fuente("ot_api_recorrido"))

    def test_el_levantamiento_tambien_respeta_la_perdida_sin_mostrar_montos_al_tecnico(self):
        f = _fuente("mant_lev_cerrar")
        self.assertIn("_fac_perdida_de_ot(lev[\"visita_id\"])", f)
        self.assertIn("CIERRE_CON_PERDIDA", f)
        self.assertIn("if _es_rol_tecnico() else", f)

    def test_ver_la_solicitud_la_ve_quien_puede_aprobar_la_ot(self):
        f = _fuente("ot_aut_api_detalle")
        self.assertIn('_puede_ot_accion(a["visita_id"], "aprobar")', f)

    def test_solicitud_guarda_la_cuenta_y_aprobar_deja_constancia(self):
        c = _fuente("ot_aut_api_crear")
        self.assertIn('if tipo == "cerrar_con_perdida":', c)
        self.assertIn('"perdida": _fp["perdida"]', c)
        a = _fuente("ot_aut_api_aprobar")
        self.assertIn('elif tipo == "cerrar_con_perdida":', a)
        self.assertIn("cierre_con_perdida_autorizado", a)

    def test_paso_6_y_motor_ofrecen_pedir_la_autorizacion(self):
        r = _fuente("ot_api_recorrido")
        self.assertIn('"accion": "pedir_aut_perdida"', r)
        self.assertIn('"perdida": perdida', r)
        js = _leer("static/ot_fin_motor.js")
        self.assertIn("p.accion === 'pedir_aut_perdida'", js)
        self.assertIn("pedir_aut_perdida: 'pedirPerdida'", js)
        self.assertIn("tipo: 'cerrar_con_perdida'", js)
        self.assertIn("a.tipo === 'cerrar_con_perdida'", _leer("static/ot_autorizaciones.js"))


class TestExcelAuditoria(unittest.TestCase):
    def test_el_excel_se_arma_y_se_abre(self):
        import io as _io
        from openpyxl import load_workbook
        hoy = _dt.datetime(2026, 10, 11, 21, 0)
        docs = [{"doc_tido": "FCV", "doc_nudo": "11610", "titulo": "FCV 11610", "cliente": "Centro de ejemplo", "cliente_rut": "76.555.444-3",
                 "fecha_doc": "07/10/2026", "total_doc": 1850000, "cap_servicio": 100000, "cap_despacho": 30000,
                 "usado_servicio": 90000, "usado_despacho": 0, "n_ot": 2, "ots_txt": "OT-2026-00320, OT-2026-00322",
                 "estado": "con_saldo", "familia_de": None, "erp_leido": 1, "actualizado_el": hoy},
                {"doc_tido": "FCV", "doc_nudo": "11502", "titulo": "FCV 11502", "cliente": "Club de ejemplo", "cliente_rut": "",
                 "fecha_doc": "02/10/2026", "total_doc": 0, "cap_servicio": 300000, "cap_despacho": 0,
                 "usado_servicio": 360000, "usado_despacho": 0, "n_ot": 2, "ots_txt": "OT-2026-00311, OT-2026-00317",
                 "estado": "excedida", "familia_de": None, "erp_leido": 1, "actualizado_el": hoy},
                {"doc_tido": "VD", "doc_nudo": "10701", "titulo": "NVV VD 10701", "cliente": "Centro de ejemplo", "cliente_rut": "",
                 "fecha_doc": "05/10/2026", "total_doc": 0, "cap_servicio": 100000, "cap_despacho": 30000,
                 "usado_servicio": 90000, "usado_despacho": 0, "n_ot": 2, "ots_txt": "OT-2026-00320, OT-2026-00322",
                 "estado": "con_saldo", "familia_de": "FCV 11610", "erp_leido": 1, "actualizado_el": hoy}]
        libro = [{"id": 1, "creado_el": hoy, "doc_tido": "FCV", "doc_nudo": "11610", "categoria": "servicio", "movimiento": "consumo",
                  "monto": 80000, "acumulado_ot": 80000, "numero_ot": "OT-2026-00322", "cliente": "Centro de ejemplo",
                  "origen": "crear", "usuario": "Ejecutivo de SSTT", "detalle": ""}]
        salida = {}

        def fetchall(sql, params=()):
            if "FROM ot_fac_doc" in sql:
                return docs
            if "FROM ot_fac_libro ORDER BY id" in sql:
                return libro
            if "SELECT DISTINCT visita_id FROM ot_fac_libro" in sql:
                return [{"visita_id": 322}]
            return []

        def fetchone(sql, params=()):
            if "FROM ot_fac_huella" in sql:
                return {"r": hoy, "n": 860}
            if "numero_ot FROM mant_visitas" in sql:
                return {"numero_ot": "OT-2026-00322"}
            return {}

        def send_file(buf, **k):
            salida["buf"], salida["k"] = buf, k
            return "ok"
        amb = {"mysql_fetchall": fetchall, "mysql_fetchone": fetchone, "send_file": send_file, "io": _io, "datetime": _dt.datetime,
               "jsonify": lambda d: d, "print": lambda *a, **k: None, "_saldo": S, "current_username": lambda: "daniel",
               "chile_fmt_filter": lambda v, f="%d/%m/%Y %H:%M": v.strftime(f),
               "_ot_doc_real_a_usuario": lambda t, n: (t, n.lstrip("0")), "_ot_saldo_real": lambda k: k,
               "_FAC_ORIGEN_TXT": {"crear": "Creación de la OT"},
               "_FAC_LIBRO_DESDE": _dt.datetime(2026, 10, 11), "json": json,
               "_fac_finanzas_de_ots": lambda vids, fallos=None: {322: {"cobra": True, "cobre": 80000, "cobre_hay": True, "me_cobraron": 145000,
                                                          "me_tecnico": 95000, "me_despacho": 0, "repuestos": 50000,
                                                          "falta_tecnico": False, "queda": -65000, "numero_ot": "OT-2026-00322",
                                                          "cerrada_at": None, "frase": "x", "estado": "pendiente_aprobacion",
                                                          "estado_txt": "Pendiente de aprobación"}}}
        _exec("_fac_fila_lista", amb)
        f = _exec("ot_saldo_facturas_excel", amb)
        self.assertEqual(f(), "ok")
        self.assertTrue(salida["k"]["download_name"].startswith("Saldo_de_facturas_"))
        wb = load_workbook(salida["buf"])
        self.assertEqual(wb.sheetnames, ["Cómo funciona", "Resumen", "Facturas", "Alertas", "Movimientos"])
        alertas = [[c.value for c in r] for r in wb["Alertas"].iter_rows(min_row=2)]
        self.assertTrue(any(a[0] == "Doble cobro" and a[1] == "FCV 11502" for a in alertas))
        perd = [a for a in alertas if a[0] == "OT a pérdida"]
        self.assertEqual(len(perd), 1)
        self.assertEqual(perd[0][3], 15000)                 # $95.000 − $80.000 (los repuestos aparte)
        self.assertIn("repuestos aparte", perd[0][4])
        self.assertIn("Sin autorización", perd[0][5])
        self.assertEqual(sum(1 for a in alertas if a[0] == "Doble cobro"), 1)
        resumen = {r[0].value: r[1].value for r in wb["Resumen"].iter_rows(min_row=2)}
        self.assertEqual(resumen["Total de facturas"], 2)   # la nota de venta reemplazada no cuenta aparte (su factura está)
        self.assertEqual(wb["Facturas"].max_row, 4)          # pero sí aparece en la hoja Facturas
        self.assertEqual(wb["Movimientos"].cell(row=2, column=6).value, 80000)
        texto = " ".join(str(c.value or "") for c in wb["Cómo funciona"]["A"])
        self.assertIn("Solo la OT consume el saldo", texto)
        self.assertIn("ILUS Fitness", texto)


class TestExcelLecturaIncompleta(unittest.TestCase):
    def test_un_error_de_lectura_se_avisa_en_el_archivo(self):
        import io as _io
        from openpyxl import load_workbook
        hoy = _dt.datetime(2026, 10, 11, 21, 0)
        salida = {}

        def fetchall(sql, params=()):
            if "FROM ot_fac_doc" in sql or "FROM ot_fac_libro ORDER BY id" in sql:
                return []
            if "FROM mant_ot_autorizaciones" in sql or "SELECT id FROM mant_visitas" in sql:
                raise RuntimeError("timeout")
            return []
        amb = {"mysql_fetchall": fetchall, "mysql_fetchone": lambda *a, **k: {"r": hoy, "n": 1}, "io": _io,
               "send_file": lambda buf, **k: salida.update(buf=buf) or "ok", "datetime": _dt.datetime, "json": json,
               "jsonify": lambda d: d, "print": lambda *a, **k: None, "_saldo": S, "current_username": lambda: "daniel",
               "chile_fmt_filter": lambda v, f="%d/%m/%Y %H:%M": v.strftime(f), "_FAC_ORIGEN_TXT": {},
               "_ot_doc_real_a_usuario": lambda t, n: (t, n), "_ot_saldo_real": lambda k: k,
               "_FAC_LIBRO_DESDE": _dt.datetime(2026, 10, 11), "_fac_finanzas_de_ots": lambda vids, fallos=None: {}}
        _exec("_fac_fila_lista", amb)
        self.assertEqual(_exec("ot_saldo_facturas_excel", amb)(), "ok")
        wb = load_workbook(salida["buf"])
        self.assertIn("LECTURA INCOMPLETA", wb["Cómo funciona"]["A1"].value)
        self.assertEqual(wb["Alertas"].cell(row=2, column=1).value, "LECTURA INCOMPLETA")
        self.assertNotIn("Sin alertas", [c.value for c in wb["Alertas"]["A"]])


class TestHerenciaYPantalla(unittest.TestCase):
    def test_la_ot_del_ticket_hereda_sus_documentos(self):
        h = _leer("templates/ot2/_modal_crear.html")
        self.assertIn("function _o2mHeredarDocsTicket(docs, numeroTicket)", h)
        self.assertIn("_o2mHeredarDocsTicket(d.documentos, t.numero_ticket);", h)
        self.assertIn("_o2mHeredarDocsTicket(j && j.documentos, t.numero_ticket);", h)
        # el documento heredado entra por el MISMO camino que si la persona lo buscara
        i = h.index("function _o2mHeredarDocsTicket")
        self.assertIn("_o2mFinDocumentoSeleccionado([], {", h[i:i + 2500])

    def test_consulta_de_notas_de_venta_y_respuestas_viejas(self):
        html = _leer("templates/ot2/saldo_facturas.html")
        for v in ('value="VD"', 'value="WEB"', 'value="NVI"', 'value="BLE"'):
            self.assertIn(v, html)
        js = _leer("static/ot_saldo_facturas.js")
        self.assertIn("tipo = m[1] === 'NVV' ? 'VD' : m[1];", js)
        self.assertIn("if (seq !== seqLista) return;", js)
        self.assertIn("if (seq !== seqDoc) return;", js)
        f = _fuente("ot_api_saldo_facturas_doc")
        self.assertIn('if tido == "NVV" and nudo.isdigit():', f)
        self.assertIn("crear=False", f)

    def test_herencia_con_rut_real_y_se_suelta_al_cambiar(self):
        h = _leer("templates/ot2/_modal_crear.html")
        self.assertIn("function _o2mSoltarHerenciaTicket()", h)
        self.assertIn("if (v !== 'ticket'){ S.ticket_id = null; _o2mSoltarHerenciaTicket(); }", h)
        self.assertIn("dd.hdr.cliente_rut", h)

    def test_borrar_y_barrido_liberan(self):
        n = _fuente("_mant_visita_eliminar_core")
        self.assertIn("_fac_libro_marcar(vid)", n)
        b = _fuente("_fac_libro_barrido")
        self.assertIn("LEFT JOIN mant_visitas v ON v.id=l.visita_id", b)
        self.assertIn("erp_leido=0", b)

    def test_pantalla_sin_alert_nativo_y_paginada(self):
        js = _leer("static/ot_saldo_facturas.js")
        self.assertNotRegex(js, r"\b(alert|confirm|prompt)\(")
        self.assertIn("Mostrando ", js)
        self.assertIn("Anterior", js)
        self.assertIn("Siguiente", js)
        html = _leer("templates/ot2/saldo_facturas.html")
        self.assertIn("ot_saldo_facturas_excel", html)

    def test_marca_ilus_fitness(self):
        self.assertIn("ILUS Fitness", _leer("templates/ot2/saldo_facturas.html"))
        self.assertNotIn("Sport & Health", _leer("templates/ot2/saldo_facturas.html"))


if __name__ == "__main__":
    unittest.main()
