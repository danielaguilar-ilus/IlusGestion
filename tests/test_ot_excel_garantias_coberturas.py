# -*- coding: utf-8 -*-
"""2026-10-25 · Excel de OT (`/ot/reporte.xlsx`): hoja «Garantías y coberturas».

Daniel (con la OT-2026-00354): «garantías o coberturas asociadas a centro de costo y poder reportar esas situaciones». Toda OT que NO se cobra
(garantía, regalía, arriendo o leasing, cortesía, trabajo interno, mantención de contrato, desinstalación cubierta) se carga a un centro de costo; la
hoja suma cuántas son y cuánto nos costaron por cobertura × centro de costo y trae el detalle OT por OT. Solo el superadministrador ve además cómo
nació, quién autorizó y el argumento (REGLA #26: es información interna).

Sin BD ni Flask: `ot2_reporte_xlsx` y la cuenta única de la OT (`_ot_finanzas`) se extraen de app.py con ast y se ejecutan con lecturas de mentira; el
Excel que devuelve se abre con openpyxl. Correr con:  python -m pytest tests/test_ot_excel_garantias_coberturas.py
"""
import ast
import datetime
import io
import re
import types
import unittest

from tests.test_incidencias_bajas import _codigo_y_arbol

FUNCS = ("ot2_reporte_xlsx", "_ot_cobertura", "_ot_finanzas", "_ot_fin_num", "_ot_fin_clp", "_ot_fin_lote", "_ot_fin_agregado", "_ot_es_interna",
         "_vida_margen_clase")
CONSTS = ("_OT_FIN_COBERTURA_TXT", "_OT_FIN_COBERTURA_CORTA", "_OT_FIN_AGG_APARTE", "_OT_FIN_UMBRAL_BAJO", "_OT_FIN_ZZ_NO_SERVICIO",
          "_OT_FIN_ORIGENES_NO_COBRO", "_OT_FIN_ORIGENES_COBRO_RESPALDO", "_OT_FIN_FUENTE_COBRO", "_OT2_EXPORT_LIMIT")

_ARBOL = None
HOJA = "Garantías y coberturas"


def _arbol():
    global _ARBOL
    if _ARBOL is None:
        _ARBOL = _codigo_y_arbol()[1]
    return _ARBOL


def _fila(i, **kw):
    """Una fila de mant_visitas como la deja el SELECT del reporte (una OT que SE COBRA por defecto, con documento y su costo)."""
    f = {"id": i, "numero_ot": f"OT-2026-{i:05d}", "titulo": f"Trabajo {i}", "estado": "cerrada", "prioridad": "media",
         "fecha_programada": datetime.date(2026, 10, 20), "fecha_realizada": datetime.date(2026, 10, 20), "hora_inicio": None,
         "direccion_visita": "", "centro_costo": "sstt", "estado_facturacion": "facturada", "factura_tido": "FCV", "factura_nudo": str(1000 + i),
         "modalidad_cobro": None, "cubierto_por": None, "tipo": "instalacion", "cliente_id": 7, "contrato_id": None, "costo": None, "zz_monto": 100000,
         "cobro_cero_motivo": None, "zz_codigo": "ZZINSTALACION", "zz_envio_monto": None, "valor_origen": "zz", "costo_proveedor": 50000,
         "costo_despacho": 0, "proveedor_tipo": "externo", "valorizado_clp": None, "valorizado_fuente": None, "zz_motivo_manual": None,
         "ot_madre_id": None, "ot_relacion": None, "contrato_real": 0, "madre_valida": 0, "created_by": "aaron", "created_at": None, "cerrada_at": None,
         "responsable_nombre": "", "razon_social": f"Cliente {i}", "cliente_rut": "76.000.000-0", "cliente_direccion": "", "cliente_comuna": "",
         "tecnico_nombre": "Técnico Demo", "tecnico_role": "tecnico", "tecnico_proveedor_id": None, "numero_ticket": None, "desinstalacion_prevista": None,
         "madre_numero": None, "hijas_numeros": None, "pendientes_numeros": None, "n_tareas": 3, "n_completas": 3}
    f.update(kw)
    return f


def _garantia(i, **kw):
    base = dict(modalidad_cobro="garantia", cobro_cero_motivo="garantia", zz_monto=None, valor_origen=None, factura_tido=None, factura_nudo=None)
    base.update(kw)
    return _fila(i, **base)


def _ejecutar(filas, superadmin=True):
    """Corre ot2_reporte_xlsx de verdad con `filas` y devuelve el libro de Excel (openpyxl)."""
    import openpyxl
    amb = {"print": lambda *a, **k: None, "re": re, "datetime": datetime, "__builtins__": __builtins__}
    for nodo in _arbol().body:
        if isinstance(nodo, ast.FunctionDef) and nodo.name in FUNCS:
            nodo.decorator_list = []
            exec(compile(ast.Module(body=[nodo], type_ignores=[]), "<app>", "exec"), amb)
        elif isinstance(nodo, ast.Assign) and len(nodo.targets) == 1 and isinstance(nodo.targets[0], ast.Name) and nodo.targets[0].id in CONSTS:
            exec(compile(ast.Module(body=[nodo], type_ignores=[]), "<app>", "exec"), amb)
    faltan = [n for n in FUNCS + CONSTS if n not in amb]
    assert not faltan, faltan
    ids = [f["id"] for f in filas]
    amb.update({
        "request": types.SimpleNamespace(args={}), "flash": lambda *a, **k: None, "redirect": lambda *a, **k: None, "url_for": lambda *a, **k: "/",
        "_ot2_filtros_export": lambda args: ("", [], [], "reales"), "_ot_fin_cols_sql": lambda alias="v": "v.modalidad_cobro", "_ot_resp_ok": lambda: True,
        "_OT2_JOIN_TAREAS": "", "mysql_fetchall": lambda sql, params=None: [dict(f) for f in filas],
        "_now_chile": lambda: datetime.datetime(2026, 10, 25, 12, 0),
        "chile_fmt_filter": lambda d, *a, **k: d.strftime("%d/%m/%Y %H:%M") if hasattr(d, "strftime") else str(d),
        "_OT2_CENTROS_COSTO": (("sstt", "Servicio Técnico"), ("logistica", "Logística"), ("comercial", "Comercial"), ("marketing", "Marketing")),
        "_OT2_ESTADO_META": {"cerrada": ("Cerrada", "", "", "cerr"), "en_curso": ("En curso", "", "", "ejec")},
        "_OT2_FASE_LABELS": {"pend": "Pendiente", "ejec": "En ejecución", "comp": "Completada", "cerr": "Cerrada"},
        "_TIPO_OT_LABEL": {"garantia": "Garantía", "instalacion": "Instalación", "correctiva": "Correctiva", "preventiva": "Preventiva"},
        "_es_rol_tecnico": lambda: False, "_ot_aut_es_superadmin": lambda: superadmin,
        "_ot_docs_texto_lote": lambda i: {x: f"FCV {1000 + x}" for x in i}, "_ot_repuestos_desglose": lambda i: {},
        "_ot_fin_clasif_lote": lambda i: {x: {"texto": "Garantía autorizada por Daniel", "autorizado_por": "Daniel Aguilar", "argumento": "Falla de fábrica a los dos meses",
                                              "derivada": False, "otros": []} for x in i},
        "_qp": types.SimpleNamespace(es_sin_fecha=lambda f: False, REL_PENDIENTE_STOCK="pendiente_stock", chip_pendiente=lambda n: f"Pendiente de {n}",
                                     chip_madre=lambda n: f"Pendientes: {n}"),
        "_ot_tv_hhmm": lambda x: str(x or ""), "_mant_log": lambda *a, **k: None, "send_file": lambda buf, **k: buf,
        "_ot_fin_contrato_real_de": lambda vids, con_madre=False: {int(v): ((False, False) if con_madre else False) for v in vids},
    })
    buf = amb["ot2_reporte_xlsx"]()
    return openpyxl.load_workbook(io.BytesIO(buf.getvalue() if hasattr(buf, "getvalue") else buf.read()))


def _filas(ws):
    return [[c.value for c in fila] for fila in ws.iter_rows()]


FILAS = [
    _garantia(1, centro_costo="sstt", costo_proveedor=370000, costo_despacho=0),
    _garantia(2, centro_costo="logistica", costo_proveedor=100000, costo_despacho=20000),
    _fila(3, modalidad_cobro="sin_costo", cobro_cero_motivo="regalia", centro_costo="comercial", costo_proveedor=50000, zz_monto=None, valor_origen=None,
          factura_tido=None, factura_nudo=None),
    _garantia(4, centro_costo=None, costo_proveedor=None, costo_despacho=None),            # sin costo declarado y sin centro de costo
    _fila(5, centro_costo="sstt"),                                                          # se COBRA: no entra a la hoja
    _fila(6, cliente_id=None, centro_costo="sstt", costo_proveedor=0, zz_monto=None, valor_origen=None, factura_tido=None, factura_nudo=None),   # trabajo interno
]


try:
    import openpyxl  # noqa: F401
    _HAY_OPENPYXL = True
except ImportError:
    _HAY_OPENPYXL = False


@unittest.skipIf(not _HAY_OPENPYXL, "openpyxl no está instalado")
class TestHojaGarantiasYCoberturas(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = _ejecutar(FILAS, superadmin=True)
        cls.ws = cls.wb[HOJA]
        cls.f = _filas(cls.ws)

    def test_la_hoja_va_despues_de_por_tipo_y_antes_del_resumen(self):
        n = self.wb.sheetnames
        self.assertEqual(n[:6], ["OT", "Por técnico", "Por centro de costo", "Por tipo", HOJA, "Resumen"])

    def test_arriba_cuenta_cuantas_son_y_cuanto_costaron_por_cobertura_y_centro_de_costo(self):
        self.assertEqual(self.f[0][:6], ["Cobertura", "Centro de costo", "Cantidad de OT", "Nos costó", "Valorizado (referencia)", "OT sin costo declarado"])
        piv = {(r[0], r[1]): r for r in self.f[1:] if r[0] and r[1] and isinstance(r[2], (int, float))}
        self.assertEqual(piv[("Garantía", "Servicio Técnico")][2:4], [1, 370000.0])
        self.assertEqual(piv[("Garantía", "Logística")][2:4], [1, 120000.0])                      # técnico + despacho
        self.assertEqual(piv[("Regalía", "Comercial")][2:4], [1, 50000.0])
        self.assertEqual(piv[("Trabajo interno", "Servicio Técnico")][2:4], [1, 0.0])
        sin = piv[("Garantía", "(sin centro de costo)")]
        self.assertEqual((sin[2], sin[3], sin[5]), (1, 0.0, 1))                                  # sin costo declarado: se cuenta aparte, no se inventa un $0
        self.assertNotIn(("Se cobra", "Servicio Técnico"), piv)                                   # la OT que SE COBRA no entra

    def test_el_total_suma_todo_y_va_en_rojo_de_la_marca(self):
        total = next(r for r in self.f if r[0] == "TOTAL")
        self.assertEqual(total[2:4], [5, 540000.0])
        self.assertEqual(total[5], 1)
        celda = next(c for fila in self.ws.iter_rows() for c in fila if c.value == "TOTAL")
        self.assertEqual(celda.fill.fgColor.rgb[-6:], "DC2626")

    def test_el_detalle_lista_cada_ot_sin_cobro_con_su_cobertura_y_centro(self):
        i = next(n for n, r in enumerate(self.f) if r[0] == "Detalle OT por OT")
        cab = self.f[i + 1]
        self.assertEqual(cab[:9], ["Cobertura", "Centro de costo", "N° OT", "Nos costó", "Valorizado (referencia)", "Estado del costo", "Cliente", "Estado de la OT", "Documentos"])
        det = [r for r in self.f[i + 2:] if r[2] and str(r[2]).startswith("OT-2026-")]
        # agrupadas por cobertura (garantía, regalía, … interno: orden fijo) y, dentro de cada una, por centro de costo
        self.assertEqual([r[2] for r in det], ["OT-2026-00004", "OT-2026-00002", "OT-2026-00001", "OT-2026-00003", "OT-2026-00006"])
        self.assertNotIn("OT-2026-00005", [r[2] for r in det])
        falta = next(r for r in det if r[2] == "OT-2026-00004")
        self.assertEqual((falta[3], falta[5]), (None, "Falta lo que cobró el técnico"))           # sin costo: celda vacía y avisado
        ok = next(r for r in det if r[2] == "OT-2026-00001")
        self.assertEqual((ok[0], ok[1], ok[3], ok[5], ok[8]), ("Garantía", "Servicio Técnico", 370000.0, "Completo", "FCV 1001"))

    def test_el_superadmin_ve_como_nacio_quien_autorizo_y_el_argumento(self):
        i = next(n for n, r in enumerate(self.f) if r[0] == "Detalle OT por OT")
        cab = self.f[i + 1]
        self.assertEqual(cab[9:12], ["Cómo nació", "Autorizó", "Argumento"])
        fila = next(r for r in self.f[i + 2:] if r[2] == "OT-2026-00001")
        self.assertEqual(fila[9:12], ["Garantía autorizada por Daniel", "Daniel Aguilar", "Falla de fábrica a los dos meses"])

    def test_sin_superadmin_no_hay_columnas_de_autorizacion(self):
        ws = _ejecutar(FILAS, superadmin=False)[HOJA]
        filas = _filas(ws)
        i = next(n for n, r in enumerate(filas) if r[0] == "Detalle OT por OT")
        self.assertEqual([c for c in filas[i + 1] if c], ["Cobertura", "Centro de costo", "N° OT", "Nos costó", "Valorizado (referencia)", "Estado del costo",
                                                          "Cliente", "Estado de la OT", "Documentos"])
        texto = " ".join(str(c) for r in filas for c in r if c)
        for privado in ("Daniel Aguilar", "Falla de fábrica", "Autorizó", "Cómo nació"):
            self.assertNotIn(privado, texto)

    def test_con_dinero_en_pesos_chilenos_y_filtros_en_el_detalle(self):
        enc = next(c for fila in self.ws.iter_rows() for c in fila if c.value == "Cobertura" and c.row > 3)          # encabezado del detalle
        self.assertEqual(self.ws.auto_filter.ref.split(":")[0], f"A{enc.row}")
        celda = next(c for fila in self.ws.iter_rows() for c in fila if c.value == 370000.0 and c.column == 4)
        self.assertEqual(celda.number_format, '"$"#,##0')

    def test_un_texto_que_parte_con_igual_queda_como_texto_y_nunca_como_formula(self):
        """El cliente y el argumento los escribe una persona: «=HYPERLINK(…)» no puede ejecutarse al abrir el Excel."""
        peligro = '=HYPERLINK("http://ejemplo.invalid","clic")'
        ws = _ejecutar([_garantia(1, razon_social=peligro)])[HOJA]
        celdas = [c for fila in ws.iter_rows() for c in fila if c.value == peligro]
        self.assertTrue(celdas, "el cliente tiene que aparecer en el detalle")
        for c in celdas:
            self.assertEqual(c.data_type, "s", f"{c.coordinate} se guardó como {c.data_type}")
        self.assertFalse([c for fila in ws.iter_rows() for c in fila if c.data_type == "f"], "ninguna celda de la hoja es una fórmula")

    def test_sin_ninguna_ot_sin_cobro_la_hoja_no_existe(self):
        wb = _ejecutar([_fila(5), _fila(7)])
        self.assertNotIn(HOJA, wb.sheetnames)
        self.assertEqual(wb.sheetnames[:4], ["OT", "Por técnico", "Por centro de costo", "Por tipo"])


if __name__ == "__main__":
    unittest.main()
