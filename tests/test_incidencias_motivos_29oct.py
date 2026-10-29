"""Gestión de motivos de la bodega de incidencias — 2026-10-29 (Daniel: «sí, quiero la gestión de motivos»).

El motivo sigue siendo el TEXTO detallado de siempre; lo que se suma es el TIPO de motivo (catálogo administrable, clasificación por
incidencia con bitácora, «Por clasificar» con sugerencia y resumen por tipo). Se prueba con datos de mentira:

  · funciones puras: sugerencia (con su porqué), código/nombre de tipos nuevos, resumen por tipo;
  · guardado del tipo de una incidencia y migración, sobre una base de mentira;
  · permisos y reglas: el catálogo solo lo cambia un admin, un tipo NUNCA se borra, el tipo es opcional («alertar, no trabar») y no
    puede romper el alta ni la edición de una incidencia;
  · la pantalla: filtro, botón, modal, formulario y que lo nuevo no usa alert/confirm/prompt nativos (REGLA #1).
"""
import os
import re
import unittest
from datetime import date

from tests import _fuente_app as FA

TEMPLATE = os.path.join(FA.RAIZ, "templates", "mantenciones", "incidencias.html")

_PURAS = ["_inc_motivo_norm", "_inc_motivo_sugerir", "_inc_motivo_codigo_de", "_inc_motivo_nombre_validar", "_inc_motivo_resumen"]
_CONST = ["INC_MOTIVO_TIPOS_BASE", "INC_MOTIVO_NOMBRE_MIN", "INC_MOTIVO_NOMBRE_MAX", "_INC_MOTIVO_PISTAS"]


def _puras():
    return FA.cargar(funciones=_PURAS, constantes=_CONST)


class TestSugerencia(unittest.TestCase):
    def setUp(self):
        self.amb = _puras()
        self.sug = self.amb["_inc_motivo_sugerir"]

    def test_lee_el_texto_y_dice_por_que(self):
        cod, pistas = self.sug("Falta acrílico derecho del lado selector de pesos")
        self.assertEqual(cod, "pieza_faltante")
        self.assertEqual(pistas, ["falta"])

    def test_ignora_tildes_y_mayusculas_pero_devuelve_la_ortografia_natural(self):
        cod, pistas = self.sug("LLEGO DAÑADO por el TRANSPORTE, golpeada")
        self.assertEqual(cod, "dano_traslado")
        self.assertIn("transporte", pistas)
        self.assertIn("golpeada", pistas)
        cod, pistas = self.sug("Defecto de fabrica, aplica garantia")
        self.assertEqual(cod, "defecto_fabrica")
        self.assertEqual(pistas, ["defecto", "fábrica", "garantía"])        # se muestra con tildes

    def test_una_palabra_cuenta_solo_si_esta_entera(self):
        self.assertEqual(self.sug("el usuario reportó que la pieza trae todo")[0], None)   # «uso» no está en «usuario»
        self.assertEqual(self.sug("desgaste por uso intensivo")[0], "dano_uso")

    def test_gana_el_tipo_con_mas_palabras_y_en_empate_el_primero(self):
        cod, _ = self.sug("no enciende, hace ruido y la consola muestra error")
        self.assertEqual(cod, "falla_funcionamiento")
        cod, _ = self.sug("falta la pieza y llegó golpeada")        # 1 y 1: gana el primero de la lista
        self.assertEqual(cod, "pieza_faltante")

    def test_solo_sugiere_tipos_activos(self):
        cod, _ = self.sug("falta el tornillo", {"dano_traslado", "otro"})
        self.assertIsNone(cod)
        cod, _ = self.sug("falta el tornillo", {"pieza_faltante"})
        self.assertEqual(cod, "pieza_faltante")

    def test_sin_texto_o_sin_pistas_no_sugiere(self):
        for t in ("", None, "   ", "todo normal, sin novedades"):
            self.assertEqual(self.sug(t), (None, []))


class TestCatalogoNuevo(unittest.TestCase):
    def setUp(self):
        self.amb = _puras()

    def test_codigo_estable_sin_tildes_y_con_guiones_bajos(self):
        f = self.amb["_inc_motivo_codigo_de"]
        self.assertEqual(f("Daño por humedad"), "dano_por_humedad")
        self.assertEqual(f("  ¡Pieza   rota!  "), "pieza_rota")
        self.assertEqual(f("???"), "tipo")
        self.assertLessEqual(len(f("x" * 200)), 40)

    def test_si_el_codigo_ya_existe_se_numera_y_nunca_pisa_otro(self):
        f = self.amb["_inc_motivo_codigo_de"]
        self.assertEqual(f("Otro motivo", ["otro_motivo"]), "otro_motivo_2")
        self.assertEqual(f("Otro motivo", ["otro_motivo", "otro_motivo_2"]), "otro_motivo_3")
        self.assertLessEqual(len(f("y" * 60, ["y" * 40])), 40)

    def test_nombre_entre_3_y_80_y_sin_repetir(self):
        v = self.amb["_inc_motivo_nombre_validar"]
        self.assertEqual(v("  Daño   por humedad ", ["Otro motivo"]), ("Daño por humedad", None))
        self.assertIn("mínimo", v("ab")[1])
        self.assertIn("largo", v("x" * 81)[1])
        self.assertIn("Ya existe", v("DANO POR USO", ["Daño por uso"])[1])      # sin importar mayúsculas ni tildes


class TestResumen(unittest.TestCase):
    def setUp(self):
        self.res = _puras()["_inc_motivo_resumen"]
        self.tipos = [
            {"codigo": "a", "nombre": "Daño en traslado", "activo": True, "orden": 10},
            {"codigo": "b", "nombre": "Pieza faltante", "activo": True, "orden": 20},
            {"codigo": "c", "nombre": "Tipo viejo sin uso", "activo": False, "orden": 30},
            {"codigo": "d", "nombre": "Tipo viejo con uso", "activo": False, "orden": 40},
        ]
        self.hoy = date(2026, 10, 29)

    def test_cuenta_por_tipo_con_unidades_abiertas_y_dias(self):
        filas = [
            {"motivo_tipo": "a", "cantidad": 2, "estado": "abierta", "fecha_ingreso": "2026-10-19", "created_at": None},
            {"motivo_tipo": "a", "cantidad": 1, "estado": "abierta", "fecha_ingreso": None, "created_at": "2026-10-27 08:00:00"},
            {"motivo_tipo": "a", "cantidad": 5, "estado": "resuelta", "fecha_ingreso": "2026-01-01", "created_at": None},
            {"motivo_tipo": "b", "cantidad": 1, "estado": "abierta", "fecha_ingreso": date(2026, 10, 29), "created_at": None},
        ]
        r = self.res(filas, self.tipos, self.hoy)
        a = next(t for t in r["tipos"] if t["codigo"] == "a")
        self.assertEqual((a["n"], a["unidades"], a["abiertas"], a["resueltas"]), (3, 8, 2, 1))
        self.assertEqual((a["dias_prom"], a["dias_max"]), (6, 10))                 # solo las abiertas: 10 y 2 días
        b = next(t for t in r["tipos"] if t["codigo"] == "b")
        self.assertEqual((b["dias_prom"], b["dias_max"]), (0, 0))
        self.assertEqual(r["total"], 4)
        self.assertEqual(r["clasificadas"], 4)

    def test_lo_que_no_tiene_tipo_o_tiene_uno_desconocido_va_a_sin_clasificar(self):
        filas = [
            {"motivo_tipo": None, "cantidad": 3, "estado": "abierta"},
            {"motivo_tipo": "", "cantidad": 1, "estado": "resuelta"},
            {"motivo_tipo": "zzz", "cantidad": 2, "estado": "abierta"},
            {"motivo_tipo": "a", "cantidad": 1, "estado": "abierta"},
        ]
        r = self.res(filas, self.tipos, self.hoy)
        self.assertEqual(r["sin_clasificar"], {"n": 3, "unidades": 6, "abiertas": 2})
        self.assertEqual((r["total"], r["clasificadas"]), (4, 1))

    def test_un_tipo_desactivado_solo_aparece_si_tiene_incidencias(self):
        r = self.res([{"motivo_tipo": "d", "cantidad": 1, "estado": "abierta"}], self.tipos, self.hoy)
        codigos = [t["codigo"] for t in r["tipos"]]
        self.assertIn("d", codigos)                  # tiene incidencias: se conserva a la vista
        self.assertNotIn("c", codigos)               # desactivado y sin uso: no estorba
        self.assertEqual(codigos, ["a", "b", "d"])   # en el orden del catálogo

    def test_sin_filas_no_falla(self):
        r = self.res([], self.tipos, self.hoy)
        self.assertEqual((r["total"], r["clasificadas"], r["sin_clasificar"]["n"]), (0, 0, 0))


class _Mundo:
    """Base de mentira: catálogo + la columna `motivo_tipo` de cada incidencia."""
    def __init__(self, col_ok=True, tipos=("a", "b"), actual=None):
        self.col_ok, self.logs, self.updates, self.ejecutados = col_ok, [], [], []
        self.tipos = [{"codigo": c, "nombre": "Tipo " + c, "orden": 10, "activo": True} for c in tipos]
        self.actual = actual

    def amb(self):
        m = self
        extra = {
            "_inc_motivo_col_ok": lambda: m.col_ok,
            "_inc_motivo_tipos": lambda solo_activos=False: list(m.tipos),
            "mysql_fetchone": lambda sql, params=None: {"motivo_tipo": m.actual},
            "mysql_execute": lambda sql, params=None: m.updates.append((sql, params)),
            "_inc_log": lambda iid, accion, campo=None, antes=None, despues=None: m.logs.append((iid, accion, campo, antes, despues)),
        }
        return FA.cargar(funciones=["_inc_motivo_guardar_tipo"], extra=extra)


class TestGuardarTipo(unittest.TestCase):
    def test_guarda_y_deja_en_la_bitacora_quien_de_que_a_que(self):
        m = _Mundo(actual=None)
        m.amb()["_inc_motivo_guardar_tipo"](7, "a")
        self.assertEqual(len(m.updates), 1)
        self.assertEqual(m.updates[0][1], ("a", 7))
        self.assertEqual(m.logs, [(7, "clasificada", "motivo_tipo", None, "a")])

    def test_cambiar_de_tipo_registra_el_antes_y_el_despues(self):
        m = _Mundo(actual="a")
        m.amb()["_inc_motivo_guardar_tipo"](7, "b")
        self.assertEqual(m.logs, [(7, "clasificada", "motivo_tipo", "a", "b")])

    def test_vacio_deja_sin_clasificar(self):
        m = _Mundo(actual="a")
        m.amb()["_inc_motivo_guardar_tipo"](7, "")
        self.assertEqual(m.updates[0][1], (None, 7))
        self.assertEqual(m.logs, [(7, "clasificada", "motivo_tipo", "a", None)])

    def test_el_mismo_tipo_no_escribe_ni_registra_nada(self):
        m = _Mundo(actual="a")
        m.amb()["_inc_motivo_guardar_tipo"](7, "a")
        self.assertEqual((m.updates, m.logs), ([], []))

    def test_un_codigo_que_no_esta_en_el_catalogo_se_ignora(self):
        m = _Mundo(actual=None)
        m.amb()["_inc_motivo_guardar_tipo"](7, "inventado")
        self.assertEqual((m.updates, m.logs), ([], []))

    def test_si_la_columna_aun_no_existe_no_hace_nada_y_no_falla(self):
        m = _Mundo(col_ok=False)
        m.amb()["_inc_motivo_guardar_tipo"](7, "a")
        self.assertEqual((m.updates, m.logs), ([], []))

    def test_un_error_de_la_base_nunca_rompe_el_guardado_de_la_incidencia(self):
        m = _Mundo(actual=None)
        amb = m.amb()
        amb["mysql_execute"] = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("base caída"))
        amb["_inc_motivo_guardar_tipo"](7, "a")          # no levanta


class TestMigracion(unittest.TestCase):
    def amb(self, hay_tipos):
        self.sql, self.inserts = [], []

        def ejecutar(sql, params=None):
            (self.inserts if sql.startswith("INSERT") else self.sql).append((sql, params))

        extra = {"mysql_execute": ejecutar,
                 "mysql_fetchone": lambda sql, params=None: ({"x": 1} if hay_tipos else None)}
        return FA.cargar(funciones=["_ensure_incidencia_motivos"], constantes=["INC_MOTIVO_TIPOS_BASE"], extra=extra)

    def test_una_clausula_por_sentencia_y_barata_si_ya_existe(self):
        self.amb(hay_tipos=True)["_ensure_incidencia_motivos"]()
        self.assertEqual(len(self.sql), 3)                                     # tabla + columna + índice, cada una sola (REGLA #18)
        self.assertTrue(self.sql[0][0].startswith("CREATE TABLE IF NOT EXISTS mant_incidencia_motivo_tipos"))
        self.assertIn("ADD COLUMN motivo_tipo", self.sql[1][0])
        self.assertIn("ADD INDEX", self.sql[2][0])
        for s, _ in self.sql[1:]:
            self.assertEqual(s.count("ADD "), 1)                               # nunca dos cláusulas en un mismo ALTER
        self.assertEqual(self.inserts, [])                                     # el catálogo ya tenía filas: no se siembra ni se pisa nada

    def test_siembra_la_lista_base_solo_si_el_catalogo_esta_vacio(self):
        self.amb(hay_tipos=False)["_ensure_incidencia_motivos"]()
        codigos = [p[0] for _, p in self.inserts]
        self.assertEqual(codigos, ["dano_traslado", "pieza_faltante", "falla_funcionamiento", "defecto_fabrica", "dano_uso",
                                   "diferencia_inventario", "otro"])
        for s, _ in self.inserts:
            self.assertIn("INSERT IGNORE", s)                                  # dos procesos arrancando a la vez no duplican

    def test_el_arranque_la_llama_dentro_de_app_context(self):
        """mysql_execute y mysql_fetchone dependen de Flask `g`: sin `app.app_context()` fallan en silencio al arrancar y el catálogo
        nunca se crea ni se siembra (le pasó a las dos columnas de las bajas, corregido el 29-oct)."""
        t = FA.texto()
        for llamada in ("_ensure_incidencia_motivos()\nexcept", "_ensure_incidencia_bajas_posicion()\nexcept"):
            i = t.index(llamada)
            self.assertIn("with app.app_context():", t[i - 120:i], llamada)

    def test_se_llama_siempre_al_arrancar_incluso_sin_migraciones(self):
        t = FA.texto()
        i = t.index("_ensure_incidencia_motivos()\nexcept")
        self.assertIn("SIEMPRE, incluso skip-migrations", t[i - 400:i])
        self.assertNotIn("ILUS_SKIP_MIGRATIONS", t[i - 400:i])


class TestReglasDeLasRutas(unittest.TestCase):
    def test_todas_las_rutas_exigen_el_modulo_y_excluyen_a_los_tecnicos_de_afuera(self):
        for ruta in ("mant_api_incidencias_motivos", "mant_api_incidencias_motivos_por_clasificar", "mant_api_incidencias_motivos_sugerir",
                     "mant_api_incidencias_clasificar", "mant_api_incidencias_motivos_crear", "mant_api_incidencias_motivos_editar"):
            d = " ".join(FA.decoradores(ruta))
            self.assertIn("_mant_required", d, ruta)
            self.assertIn("_no_tecnico_salvo_taller", d, ruta)

    def test_el_catalogo_solo_lo_cambia_un_admin(self):
        for ruta in ("mant_api_incidencias_motivos_crear", "mant_api_incidencias_motivos_editar"):
            src = FA.fuente_funcion(ruta)
            self.assertIn("_inc_motivo_puede_editar_catalogo()", src, ruta)
            self.assertIn("403", src, ruta)
        amb = FA.cargar(funciones=["_inc_motivo_puede_editar_catalogo"], extra={"g": type("G", (), {"get": lambda s, k, d=None: {"admin": True}})()})
        self.assertTrue(amb["_inc_motivo_puede_editar_catalogo"]())
        amb = FA.cargar(funciones=["_inc_motivo_puede_editar_catalogo"], extra={"g": type("G", (), {"get": lambda s, k, d=None: {"mantenciones": True}})()})
        self.assertFalse(amb["_inc_motivo_puede_editar_catalogo"]())

    def test_un_tipo_nunca_se_borra(self):
        t = FA.texto()
        self.assertNotIn("DELETE FROM mant_incidencia_motivo_tipos", t)
        self.assertFalse(re.search(r'@app\.route\("/mantenciones/api/incidencias/motivos[^"]*",\s*methods=\["DELETE"\]', t))

    def test_clasificar_solo_acepta_un_tipo_activo_y_deja_cada_cambio_en_la_bitacora(self):
        src = FA.fuente_funcion("mant_api_incidencias_clasificar")
        self.assertIn("solo_activos=True", src)
        self.assertIn("INC_MOTIVO_CLASIFICAR_MAX", src)
        self.assertIn("_inc_motivo_guardar_tipo(", src)                        # ahí vive el registro en la bitácora

    def test_el_tipo_es_opcional_y_no_puede_romper_el_alta_ni_la_edicion(self):
        for ruta in ("mant_api_incidencias_crear", "mant_api_incidencias_editar"):
            src = FA.fuente_funcion(ruta)
            self.assertIn('if "motivo_tipo" in data:', src, ruta)              # solo si el formulario lo manda
            self.assertIn("_inc_motivo_guardar_tipo(", src, ruta)
        # el INSERT/UPDATE principal NO se toca: si la columna aún no existe, el guardado de siempre sigue funcionando
        self.assertNotIn("motivo_tipo", FA.fuente_funcion("mant_api_incidencias_crear").split("INSERT INTO mant_incidencias")[1].split("VALUES")[0])

    def test_la_lista_filtra_por_tipo_solo_si_la_columna_existe(self):
        src = FA.fuente_funcion("mant_api_incidencias_list")
        self.assertIn("f_motivo and _inc_motivo_col_ok()", src)
        self.assertIn('"sin"', src)
        self.assertIn("motivo_tipo_nombre", src)                               # el nombre completo, no el código

    def test_la_bitacora_y_el_texto_del_motivo_no_cambian(self):
        # el motivo sigue siendo el texto detallado (mínimo 30 caracteres en el formulario) y se guarda como siempre
        src = FA.fuente_funcion("mant_api_incidencias_editar")
        self.assertIn('"motivo": _sentence_case((data.get("motivo") or "").strip()) or None', src)


class TestExcelConTipoDeMotivo(unittest.TestCase):
    """El Excel de la bodega (solo admin) trae el tipo de motivo de cada incidencia y un resumen «Por tipo de motivo»."""

    def libro(self, filas):
        import datetime as dt
        import io
        from openpyxl import load_workbook
        amb = FA.cargar(funciones=["_inc_rango_dias", "_inc_excel_dt_chile", "_inc_excel_fecha", "_inc_excel_repuesto", "_inc_excel_bytes"],
                        constantes=["_INC_EXCEL_CHK_TEXTO"], extra={"to_chile_filter": lambda d: d - dt.timedelta(hours=3)})
        return load_workbook(io.BytesIO(amb["_inc_excel_bytes"](filas, dt.datetime(2026, 10, 29, 12, 0), {})))

    def filas(self):
        base = {"cantidad": 1, "estado": "abierta", "recomendacion": "UA1", "sku": "S", "descripcion": "Banco", "chk_estado": "coincide"}
        return [dict(base, id=1, motivo="Llegó golpeada por el transporte", motivo_tipo_nombre="Daño en traslado o transporte"),
                dict(base, id=2, motivo="Falta un tope en la caja", motivo_tipo_nombre="Daño en traslado o transporte"),
                dict(base, id=3, motivo="Sin novedades en la revisión de rutina", motivo_tipo_nombre="")]

    def test_la_columna_va_junto_al_motivo_y_el_enlace_de_la_foto_sigue_al_final(self):
        wb = self.libro(self.filas())
        ws = wb["Incidencias"]
        enc = [c.value for c in ws[1]]
        self.assertEqual(enc[enc.index("Motivo") + 1], "Tipo de motivo")
        self.assertEqual(enc[-1], "Enlace a la foto")                       # la columna del enlace se calcula como la última
        self.assertEqual(ws.cell(2, enc.index("Tipo de motivo") + 1).value, "Daño en traslado o transporte")
        self.assertIn(ws.cell(4, enc.index("Tipo de motivo") + 1).value, (None, ""))

    def test_el_resumen_cuenta_por_tipo_y_lo_que_falta_clasificar(self):
        wb = self.libro(self.filas())
        filas = [[c.value for c in f] for f in wb["Resumen"].iter_rows()]
        i = next(n for n, f in enumerate(filas) if f[0] == "Por tipo de motivo")
        bloque = {f[0]: f[1] for f in filas[i + 1:i + 4]}
        self.assertEqual(bloque.get("Daño en traslado o transporte"), 2)
        self.assertEqual(bloque.get("Sin clasificar"), 1)

    def test_si_el_catalogo_falla_el_excel_sale_igual(self):
        src = FA.fuente_funcion("mant_api_incidencias_excel")
        i = src.index("_inc_motivo_tipos()")
        self.assertIn("try:", src[max(0, i - 200):i])                      # el nombre del tipo es un extra: nunca tumba la descarga
        self.assertIn("except Exception", src[i:i + 600])


class TestPantalla(unittest.TestCase):
    def setUp(self):
        with open(TEMPLATE, encoding="utf-8") as fh:
            self.t = fh.read()
        i = self.t.index("Gestión de motivos (2026-10-29")
        self.js = self.t[i:self.t.index("document.getElementById('incFQ').addEventListener('input'", i)]

    def test_filtro_boton_y_modal_estan(self):
        for frag in ('id="incFMotivo"', 'id="incBtnMotivos"', 'id="incMotModal"', 'id="incMotivoTipo"', 'id="incMotivoTipoSug"',
                     'id="incMotTabResumen"', 'id="incMotTabPend"', 'id="incMotTabCat"'):
            self.assertIn(frag, self.t, frag)
        self.assertIn("modal-xl inc-modal-ancho", self.t[self.t.index('id="incMotModal"'):][:300])   # «me gustan los modales grandes»

    def test_el_tipo_se_manda_solo_si_el_catalogo_cargo(self):
        self.assertIn("if(incMotCatListo) payload.motivo_tipo = $('incMotivoTipo').value;", self.t)

    def test_filtrar_por_tipo_muestra_todo_y_limpiar_filtros_lo_recarga(self):
        self.assertIn("&f_motivo='+encodeURIComponent(incFMotivo)+'&ver_todo=1", self.t)
        self.assertIn("incFMotivo = ''; var selMot", self.t)
        self.assertIn("incQ || incEstado || incFChk || incFMotivo;", self.t)    # «hayFiltros» lo cuenta

    def test_lo_nuevo_no_usa_alert_confirm_prompt_nativos(self):
        self.assertFalse(re.search(r"(?<![\w.])(alert|confirm|prompt)\(", self.js), "REGLA #1: se usa ilusPrompt/ilusConfirm/ilusToast")
        self.assertIn("ilusPrompt(", self.js)
        self.assertIn("ilusConfirm(", self.js)

    def test_seleccionar_todo_es_un_toggle_y_la_lista_se_pagina(self):
        self.assertIn("marcar = Array.prototype.some.call(cks, function(c){ return !c.checked; })", self.js)    # REGLA #14
        for frag in ("Mostrando <b>", "por página", "incMotPrev", "incMotNext"):                                # REGLA #4.3
            self.assertIn(frag, self.js, frag)

    def test_nada_se_trunca_y_la_sugerencia_dice_por_que(self):
        self.assertNotIn("text-overflow", self.t[self.t.index(".inc-mot-badge"):self.t.index('id="incMotModal"')])
        self.assertIn("porque dice «", self.js)
        self.assertIn("Usar este", self.js)

    def test_el_catalogo_nunca_ofrece_borrar(self):
        self.assertNotIn("Eliminar", self.js)
        self.assertNotIn("Borrar", self.js)
        self.assertIn("Desactivar", self.js)


if __name__ == "__main__":
    unittest.main()
