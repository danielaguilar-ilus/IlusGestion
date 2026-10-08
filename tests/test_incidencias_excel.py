"""Excel de la bodega de Incidencias, solo admin (Daniel, 2026-10-08).

Sin Flask ni BD reales: las funciones de app.py se extraen con ast (mismo criterio que
tests/test_incidencias_bajas.py) y la ruta se ejecuta con una "base" falsa.

Correr con:  py -m unittest tests.test_incidencias_excel
"""
import ast
import datetime as dt
import io
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tests.test_incidencias_repuesto_tercera_fuente import _codigo_y_arbol  # noqa: E402

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PLANTILLA = os.path.join(RAIZ, "templates", "mantenciones", "incidencias.html")
FUNCIONES = ("_inc_rango_dias", "_inc_excel_dt_chile", "_inc_excel_fecha", "_inc_excel_repuesto",
             "_inc_excel_bytes", "mant_api_incidencias_excel")
RUTA = "/mantenciones/api/incidencias/excel"


class _Obj(dict):
    pass


def _cargar(request_args, permisos, filas, bajas=None, fotos=None):
    _, arbol = _codigo_y_arbol()
    consultas = []

    def mysql_fetchall(sql, params=None):
        consultas.append(sql)
        if "FROM mant_incidencia_bajas" in sql:
            return list(bajas or [])
        if "FROM mant_incidencia_fotos" in sql:
            return list(fotos or [])
        return [dict(f) for f in filas]

    class _Req:
        args = request_args
        host_url = "https://ilus.test/"

    class _Resp:
        def __init__(self, data, mimetype=None, headers=None):
            self.data, self.mimetype, self.headers = data, mimetype, headers or {}

    class _G:
        permissions = permisos

        def get(self, k, d=None):
            return getattr(self, k, d)

    ambito = {
        "g": _G(), "request": _Req(), "Response": _Resp,
        "jsonify": lambda d: _Obj(d),
        "mysql_fetchall": mysql_fetchall,
        "_inc_ids_en_baja": lambda: set(), "_inc_ids_no_cuadran": lambda: set(),
        "_mant_incidencia_row": lambda r: r,
        "_inc_anotar_ingreso": lambda rows: [
            dict(r, ingreso_fecha="2024-09-26", ingreso_origen="gri", ingreso_gri="1860") for r in rows],
        "_checkwms_stock_rows": lambda: [{"ua": "UA1", "ubicacion": "A-01"}],
        "_inc_comparativa_ua": lambda rows, w: [
            dict(r, chk_estado="coincide", chk_ua="UA1", chk_ubicacion="A-01") for r in rows],
        "_INC_CHK_ESTADOS": ("coincide", "distinta"),
        "INC_BAJA_MOTIVOS": {"duplicado_error": "Duplicado o error de registro"},
        "_now_chile": lambda: dt.datetime(2026, 10, 8, 12, 0),
        "_now_chile_str": lambda fmt: dt.datetime(2026, 10, 8, 12, 0).strftime(fmt),
        "to_chile_filter": lambda d: d - dt.timedelta(hours=3),
    }
    for nodo in arbol.body:
        if isinstance(nodo, ast.FunctionDef) and nodo.name in FUNCIONES:
            nodo.decorator_list = []
            exec(compile(ast.Module(body=[nodo], type_ignores=[]), "<app>", "exec"), ambito)
        if (isinstance(nodo, ast.Assign) and len(nodo.targets) == 1
                and isinstance(nodo.targets[0], ast.Name) and nodo.targets[0].id == "_INC_EXCEL_CHK_TEXTO"):
            exec(compile(ast.Module(body=[nodo], type_ignores=[]), "<app>", "exec"), ambito)
    return ambito, consultas


FILA = {
    "id": 7, "sku": "1118100137", "descripcion": "Banco", "cantidad": 2, "motivo": "Golpe en traslado",
    "req_repuesto": "si", "stock_repuesto": "no_hay", "descripcion_repuesto": "Tapiz",
    "observacion": "nota", "sugerencia": "", "fecha_resolucion": None, "recomendacion": "UA1",
    "estado": "abierta", "ubicacion": "A-01", "created_by": "ana", "created_at": "2026-10-01 15:00:00",
    "updated_by": "luis", "updated_at": "2026-10-02 15:00:00",
}


class TestIncidenciasExcel(unittest.TestCase):
    def test_la_ruta_existe_y_es_get(self):
        codigo, _ = _codigo_y_arbol()
        self.assertIn(f'@app.route("{RUTA}", methods=["GET"])', codigo)

    def test_tecnico_o_no_admin_recibe_403(self):
        amb, consultas = _cargar({}, {"admin": False, "superadmin": False}, [FILA])
        r = amb["mant_api_incidencias_excel"]()
        self.assertEqual(r[1], 403)
        self.assertEqual(consultas, [], "sin permiso no debe consultar nada")

    def test_admin_recibe_xlsx_con_las_hojas(self):
        from openpyxl import load_workbook
        amb, _ = _cargar({}, {"admin": True}, [FILA, dict(FILA, id=8, estado="resuelta", req_repuesto="no")],
                         bajas=[{"incidencia_id": 8, "motivo_codigo": "duplicado_error", "motivo_texto": "x",
                                 "baja_by": "jefe", "baja_at": dt.datetime(2026, 10, 3, 15, 0)}],
                         fotos=[{"incidencia_id": 7, "gcs_key": "inc/7/a.jpg"}])
        resp = amb["mant_api_incidencias_excel"]()
        self.assertIn("incidencias_bodega_08-10-2026.xlsx", resp.headers["Content-Disposition"])
        wb = load_workbook(io.BytesIO(resp.data))
        self.assertEqual(wb.sheetnames, ["Incidencias", "Resumen"])
        ws = wb["Incidencias"]
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(ws.auto_filter.ref.startswith("A1:"))
        enc = [c.value for c in ws[1]]
        for titulo in ("UA", "SKU", "Días en bodega", "N° GRI", "Dada de baja", "Enlace a la foto"):
            self.assertIn(titulo, enc)
        fila7 = {enc[i]: c.value for i, c in enumerate(ws[2])}
        self.assertEqual(fila7["Stock del repuesto"], "Sin stock")
        self.assertEqual(fila7["Días en bodega"], (dt.date(2026, 10, 8) - dt.date(2024, 9, 26)).days)
        self.assertEqual(fila7["Rango de días"], "Más de 365 días")
        self.assertEqual(fila7["Fecha de creación"], dt.datetime(2026, 10, 1, 12, 0))  # hora Chile
        self.assertEqual(ws.cell(2, enc.index("Enlace a la foto") + 1).hyperlink.target,
                         "https://ilus.test/f/inc/7/a.jpg")
        fila8 = {enc[i]: c.value for i, c in enumerate(ws[3])}
        self.assertEqual(fila8["Dada de baja"], "Sí")
        self.assertIn("Duplicado o error de registro", fila8["Motivo de la baja"])
        self.assertEqual(ws["A2"].font.name, "Arial")
        resumen = [[c.value for c in f] for f in wb["Resumen"].iter_rows()]
        textos = [str(f[0]) for f in resumen]
        for t in ("Por estado", "Por días en bodega", "Por repuesto", "Por ubicación (nuestra)", "Más de 365 días"):
            self.assertIn(t, textos)

    def test_todo_incluye_bajas_y_sin_todo_no(self):
        amb, _ = _cargar({"todo": "1"}, {"superadmin": True}, [])
        amb["_inc_ids_en_baja"] = lambda: self.fail("con todo=1 no se excluyen las bajas")
        amb["_inc_ids_no_cuadran"] = lambda: self.fail("con todo=1 no se excluyen las que no cuadran")
        amb["mant_api_incidencias_excel"]()
        amb2, _ = _cargar({}, {"admin": True}, [])
        llamado = []
        amb2["_inc_ids_en_baja"] = lambda: llamado.append(1) or set()
        amb2["mant_api_incidencias_excel"]()
        self.assertEqual(llamado, [1])

    def test_rangos_de_dias(self):
        f = _cargar({}, {}, [])[0]["_inc_rango_dias"]
        self.assertEqual([f(0), f(30), f(31), f(90), f(91), f(365), f(366), f(None)],
                         ["0-30 días", "0-30 días", "31-90 días", "31-90 días", "91-365 días",
                          "91-365 días", "Más de 365 días", "Sin fecha"])

    def test_boton_solo_admin_en_la_plantilla(self):
        html = open(PLANTILLA, encoding="utf-8").read()
        i = html.index("incBtnExcel")
        self.assertIn("permissions.admin or permissions.superadmin", html[max(0, i - 600):i])
        self.assertIn("Descargar Excel", html)


if __name__ == "__main__":
    unittest.main()
