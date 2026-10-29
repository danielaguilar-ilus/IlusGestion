# -*- coding: utf-8 -*-
"""2026-10-29 · Los Excel de OT y costos traen FECHAS REALES (no texto), en hora Chile, con formato español, filtro automático,
encabezado fijo y columnas de apoyo (Año · Mes · N° mes · Semana · Día · Año-Mes) para filtrar por mes sin fórmulas.

Daniel: «Cuando bajas las fechas en el reporte de la OT, la fecha viene con formato gringo y necesito que en el Excel, en español,
pueda filtrar por meses; que me salgan los meses, los días, y tener un filtro más amigable en Excel, en el reporte más que nada.»

Sin BD ni Flask: cada función se recorta de app.py con tests/_fuente_app.py (rápido: no se parsea app.py entero) y se ejecuta
con datos de mentira; el libro que devuelve se abre con openpyxl. Excel de cubiertos:
  · /ot/reporte.xlsx (ot2_reporte_xlsx) · Costeo de OT (ot_gerencial.excel_bytes) · Saldo de facturas (ot_saldo_facturas_excel)
  · Lote de Facturación de proveedor (_mfp_excel_lote) · OT por facturar (mant_facturas_proveedor_por_facturar_xlsx)
  · Control de cobros a externos (mant_facturas_proveedor_xlsx) · Facturación de proveedores (mant_facturacion_proveedores_xlsx)
  · Solicitudes de repuestos de OT (repstock_solicitudes_ot_export) · Acta de recepción de la revisión (_acta_xlsx_bytes)

Correr con:  python -m unittest tests.test_ot_excel_fechas
"""
import datetime as dt
import io
import json
import re
import sys
import types
import unittest

from tests import _fuente_app as FA

try:
    import openpyxl
    _HAY_OPENPYXL = True
except ImportError:                                   # pragma: no cover
    _HAY_OPENPYXL = False

import ot_intenciones
import ot_saldo_servicio
import xlsx_fechas as XF

F_D, F_DT = "dd/mm/yyyy", "dd/mm/yyyy hh:mm"
APOYO = ["Año", "Mes", "N° mes", "Semana", "Día", "Año-Mes"]


def _abrir(buf):
    datos = buf.getvalue() if hasattr(buf, "getvalue") else (buf.read() if hasattr(buf, "read") else buf)
    return openpyxl.load_workbook(io.BytesIO(datos))


def _cab(ws, fila=1):
    return [c.value for c in ws[fila]]


def _col(ws, nombre, fila=1):
    return _cab(ws, fila).index(nombre) + 1


def _apoyo_cab(ws, fila, de):
    """Las 6 columnas de apoyo con el «(de …)» de la fecha que manda: devuelve la columna de «Año (…)»."""
    cab = _cab(ws, fila)
    esperado = [f"{h} ({de})" for h in APOYO]
    i = cab.index(esperado[0])
    assert cab[i:i + 6] == esperado, cab[i:i + 6]
    return i + 1


def _apoyo_val(ws, fila, col):
    return [ws.cell(row=fila, column=col + k).value for k in range(6)]


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  /ot/reporte.xlsx
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════════
_FUNCS_OT = ("ot2_reporte_xlsx", "_ot_cobertura", "_ot_finanzas", "_ot_fin_num", "_ot_fin_clp", "_ot_fin_lote", "_ot_fin_agregado",
             "_ot_es_interna", "_vida_margen_clase")
_CONSTS_OT = ("_OT_FIN_COBERTURA_TXT", "_OT_FIN_COBERTURA_CORTA", "_OT_FIN_AGG_APARTE", "_OT_FIN_UMBRAL_BAJO", "_OT_FIN_ZZ_NO_SERVICIO",
              "_OT_FIN_ORIGENES_NO_COBRO", "_OT_FIN_ORIGENES_COBRO_RESPALDO", "_OT_FIN_FUENTE_COBRO", "_OT2_EXPORT_LIMIT")
CENTINELA = dt.date(2099, 12, 31)


def _fila_ot(i, **kw):
    f = {"id": i, "numero_ot": f"OT-2026-{i:05d}", "titulo": f"Trabajo {i}", "estado": "cerrada", "prioridad": "media",
         "fecha_programada": dt.date(2026, 10, 20), "fecha_realizada": dt.date(2026, 10, 21), "hora_inicio": None,
         "direccion_visita": "", "centro_costo": "sstt", "estado_facturacion": "facturada", "factura_tido": "FCV", "factura_nudo": str(1000 + i),
         "modalidad_cobro": None, "cubierto_por": None, "tipo": "instalacion", "cliente_id": 7, "contrato_id": None, "costo": None, "zz_monto": 100000,
         "cobro_cero_motivo": None, "zz_codigo": "ZZINSTALACION", "zz_envio_monto": None, "valor_origen": "zz", "costo_proveedor": 50000,
         "costo_despacho": 0, "proveedor_tipo": "externo", "valorizado_clp": None, "valorizado_fuente": None, "zz_motivo_manual": None,
         "ot_madre_id": None, "ot_relacion": None, "contrato_real": 0, "madre_valida": 0, "created_by": "aaron",
         "created_at": dt.datetime(2026, 10, 20, 2, 0),            # UTC → 19/10/2026 23:00 en Chile
         "cerrada_at": dt.datetime(2026, 4, 5, 3, 30),             # UTC, fin del horario de verano → 04/04/2026 23:30 en Chile
         "responsable_nombre": "", "razon_social": f"Cliente {i}", "cliente_rut": "76.000.000-0", "cliente_direccion": "", "cliente_comuna": "",
         "tecnico_nombre": "Técnico Demo", "tecnico_role": "tecnico", "tecnico_proveedor_id": None, "numero_ticket": "TK-1", "desinstalacion_prevista": None,
         "madre_numero": None, "hijas_numeros": None, "pendientes_numeros": None, "n_tareas": 3, "n_completas": 3}
    f.update(kw)
    return f


def _reporte_ot(filas, superadmin=True):
    amb = {
        "print": lambda *a, **k: None, "re": re, "datetime": dt.datetime,
        "request": types.SimpleNamespace(args={}), "flash": lambda *a, **k: None, "redirect": lambda *a, **k: None, "url_for": lambda *a, **k: "/",
        "_ot2_filtros_export": lambda args: ("", [], [], "reales"), "_ot_fin_cols_sql": lambda alias="v": "v.modalidad_cobro", "_ot_resp_ok": lambda: True,
        "_OT2_JOIN_TAREAS": "", "mysql_fetchall": lambda sql, params=None: [dict(f) for f in filas],
        "_now_chile": lambda: dt.datetime(2026, 10, 25, 12, 0),
        "chile_fmt_filter": lambda d, *a, **k: d.strftime("%d/%m/%Y %H:%M") if hasattr(d, "strftime") else str(d),
        "_OT2_CENTROS_COSTO": (("sstt", "Servicio Técnico"), ("logistica", "Logística")),
        "_OT2_ESTADO_META": {"cerrada": ("Cerrada", "", "", "cerr"), "en_curso": ("En curso", "", "", "ejec")},
        "_OT2_FASE_LABELS": {"pend": "Pendiente", "ejec": "En ejecución", "comp": "Completada", "cerr": "Cerrada"},
        "_TIPO_OT_LABEL": {"garantia": "Garantía", "instalacion": "Instalación"},
        "_es_rol_tecnico": lambda: False, "_ot_aut_es_superadmin": lambda: superadmin,
        "_ot_docs_texto_lote": lambda i: {x: f"FCV {1000 + x}" for x in i}, "_ot_repuestos_desglose": lambda i: {},
        "_ot_fin_clasif_lote": lambda i: {x: {"texto": "Garantía autorizada por Daniel", "autorizado_por": "Daniel Aguilar",
                                              "autorizado_at": "15/10/2026 10:20", "argumento": "Falla de fábrica", "derivada": False,
                                              "otros": []} for x in i},
        "_qp": types.SimpleNamespace(es_sin_fecha=lambda f: f == CENTINELA, REL_PENDIENTE_STOCK="pendiente_stock",
                                     chip_pendiente=lambda n: f"Pendiente de {n}", chip_madre=lambda n: f"Pendientes: {n}"),
        "_ot_tv_hhmm": lambda x: str(x or ""), "_mant_log": lambda *a, **k: None, "send_file": lambda buf, **k: buf,
        "_oint": ot_intenciones, "_ot_intenciones_catalogo": ot_intenciones.catalogo_semilla,
        "_ot_intenciones_de_ots": lambda ids: {},
        "_ot_fin_contrato_real_de": lambda vids, con_madre=False: {int(v): ((False, False) if con_madre else False) for v in vids},
    }
    amb = FA.cargar(funciones=_FUNCS_OT, constantes=_CONSTS_OT, extra=amb)
    return _abrir(amb["ot2_reporte_xlsx"]())


@unittest.skipIf(not _HAY_OPENPYXL, "openpyxl no está instalado")
class TestReporteOT(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        garantia = _fila_ot(2, modalidad_cobro="garantia", cobro_cero_motivo="garantia", zz_monto=None, valor_origen=None,
                            factura_tido=None, factura_nudo=None, fecha_programada=dt.date(2027, 1, 1), cerrada_at=None,
                            fecha_realizada=None, estado="en_curso")
        pendiente = _fila_ot(3, fecha_programada=CENTINELA, fecha_realizada=None, cerrada_at=None, created_at=None, estado="en_curso")
        cls.wb = _reporte_ot([_fila_ot(1), garantia, pendiente])
        cls.ws = cls.wb["OT"]
        cls.filas = {cls.ws.cell(row=r, column=1).value: r for r in range(3, cls.ws.max_row + 1)}

    def _celda(self, numero, nombre):
        return self.ws.cell(row=self.filas[numero], column=_col(self.ws, nombre, 2))

    def test_la_fecha_programada_es_una_fecha_de_excel_no_un_texto(self):
        c = self._celda("OT-2026-00001", "Fecha programada")
        self.assertIsInstance(c.value, dt.datetime)        # openpyxl lee toda fecha como datetime
        self.assertEqual(c.value.date(), dt.date(2026, 10, 20))
        self.assertEqual(c.number_format, F_D)
        self.assertTrue(c.is_date)
        c = self._celda("OT-2026-00001", "Fecha realizada")
        self.assertEqual((c.value.date(), c.number_format), (dt.date(2026, 10, 21), F_D))

    def test_creada_y_cerrada_en_hora_chile_con_hora(self):
        c = self._celda("OT-2026-00001", "Creada")
        self.assertEqual(c.value, dt.datetime(2026, 10, 19, 23, 0))           # 02:00 UTC del 20 = 23:00 del 19 en Chile
        self.assertEqual(c.number_format, F_DT)
        c = self._celda("OT-2026-00001", "Cerrada")
        self.assertEqual(c.value, dt.datetime(2026, 4, 4, 23, 30))            # día del cambio de horario
        self.assertEqual(c.number_format, F_DT)

    def test_vacios_son_celdas_vacias_y_la_ot_sin_fecha_lo_dice(self):
        self.assertIsNone(self._celda("OT-2026-00002", "Cerrada").value)
        self.assertIsNone(self._celda("OT-2026-00002", "Fecha realizada").value)
        self.assertEqual(self._celda("OT-2026-00003", "Fecha programada").value, "Sin fecha")
        self.assertIsNone(self._celda("OT-2026-00003", "Creada").value)
        textos = [c.value for fila in self.ws.iter_rows(min_row=3) for c in fila if isinstance(c.value, str)]
        self.assertNotIn("None", textos)

    def test_columnas_de_apoyo_de_la_fecha_programada_al_final(self):
        col = _apoyo_cab(self.ws, 2, "de la fecha programada")
        self.assertEqual(col + 5, self.ws.max_column, "van al final: no corren ninguna columna que ya existía")
        self.assertLess(_col(self.ws, "Responsable de la OT", 2), col)
        self.assertEqual(_apoyo_val(self.ws, self.filas["OT-2026-00001"], col), [2026, "Octubre", 10, 43, "Martes", "2026-10"])
        # 01/01/2027: año 2027, semana ISO 53 (de 2026)
        self.assertEqual(_apoyo_val(self.ws, self.filas["OT-2026-00002"], col), [2027, "Enero", 1, 53, "Viernes", "2027-01"])
        # la OT pendiente por falta de stock no inventa un mes
        self.assertEqual(_apoyo_val(self.ws, self.filas["OT-2026-00003"], col), [None] * 6)

    def test_filtro_y_encabezado_fijo_sobre_toda_la_tabla(self):
        self.assertEqual(self.ws.freeze_panes, "C3")
        ref = self.ws.auto_filter.ref
        self.assertTrue(ref.startswith("A2:"))
        self.assertEqual(ref.split(":")[1], openpyxl.utils.get_column_letter(self.ws.max_column) + str(self.ws.max_row))
        for hoja in ("Por técnico", "Por centro de costo", "Por tipo", "Clasificación"):
            self.assertTrue(self.wb[hoja].auto_filter.ref, hoja)
            self.assertEqual(self.wb[hoja].freeze_panes, "A2", hoja)
        # la fila TOTAL no entra al filtro (si no, al ordenar se movería)
        ws = self.wb["Por centro de costo"]
        fila_total = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(row=r, column=1).value or "").startswith("TOTAL"))
        self.assertLess(int(re.findall(r"\d+$", ws.auto_filter.ref)[0]), fila_total)

    def test_los_montos_siguen_igual(self):
        c = self._celda("OT-2026-00001", "Cobré")
        self.assertEqual((c.value, c.number_format), (100000, '"$"#,##0'))

    def test_cuando_autorizo_es_fecha_y_hora_real_en_garantias(self):
        ws = self.wb["Garantías y coberturas"]
        enc = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "Cobertura" and r > 3)
        col = _col(ws, "Cuándo autorizó", enc)
        c = ws.cell(row=enc + 1, column=col)
        self.assertEqual((c.value, c.number_format), (dt.datetime(2026, 10, 15, 10, 20), F_DT))

    def test_el_excel_abre_y_vuelve_a_guardarse_sin_errores(self):
        buf = io.BytesIO()
        self.wb.save(buf)
        self.assertGreater(len(buf.getvalue()), 1000)


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Costeo de OT (ot_gerencial.excel_bytes)
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════════
@unittest.skipIf(not _HAY_OPENPYXL, "openpyxl no está instalado")
class TestCosteoGerencial(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import ot_gerencial as G
        from tests import _ger_demo as D
        cls.d = D.tablero_demo()
        cls.wb = openpyxl.load_workbook(io.BytesIO(G.excel_bytes(cls.d)))

    def test_a_nombre_de_quien_trae_la_fecha_de_cierre_real_y_sus_columnas_de_apoyo(self):
        ws = self.wb["A nombre de quién"]
        col_f = _col(ws, "Fecha de cierre", 4)
        col_a = _apoyo_cab(ws, 4, "de la fecha de cierre")
        self.assertEqual(col_a, col_f + 1)
        casos = self.d["atribucion"]["casos"]
        self.assertTrue(casos)
        for i, caso in enumerate(casos):
            c = ws.cell(row=5 + i, column=col_f)
            esperado = XF.a_fecha_excel(caso["cerrada_at"])
            self.assertEqual(c.value, esperado)
            self.assertEqual(c.number_format, F_DT)
            self.assertEqual(_apoyo_val(ws, 5 + i, col_a), XF.columnas_apoyo(esperado))
            # el año y el mes NO llevan separador de miles (2026, no 2.026)
            self.assertNotIn("#,##0", ws.cell(row=5 + i, column=col_a).number_format)

    def test_cuando_autorizo_es_fecha_real_o_vacia(self):
        ws = self.wb["A nombre de quién"]
        col = _col(ws, "Cuándo autorizó", 4)
        for fila in ws.iter_rows(min_row=5, min_col=col, max_col=col):
            v = fila[0].value
            self.assertTrue(v is None or isinstance(v, dt.datetime), repr(v))
            if v is not None:
                self.assertEqual(fila[0].number_format, F_DT)

    def test_tendencia_con_el_lunes_y_el_domingo_como_fechas(self):
        ws = self.wb["Tendencia"]
        c1, c2 = _col(ws, "Desde (lunes)", 4), _col(ws, "Hasta (domingo)", 4)
        col_a = _apoyo_cab(ws, 4, "del lunes de la semana")
        lunes = ws.cell(row=5, column=c1).value
        self.assertEqual(lunes.weekday(), 0)
        self.assertEqual(ws.cell(row=5, column=c1).number_format, F_D)
        self.assertEqual((ws.cell(row=5, column=c2).value - lunes).days, 6)
        self.assertEqual(ws.cell(row=5, column=col_a + 4).value, "Lunes")

    def test_cada_hoja_tiene_filtro_y_encabezado_fijo(self):
        for ws in self.wb.worksheets:
            self.assertTrue(ws.auto_filter.ref and ws.auto_filter.ref.startswith("A4:"), ws.title)
            self.assertEqual(ws.freeze_panes, "A5", ws.title)
        # en «Garantías y coberturas» el filtro cubre solo la primera tabla (por familia), no «Por cliente» ni «Por mes»
        ws = self.wb["Garantías y coberturas"]
        fin = int(re.findall(r"\d+$", ws.auto_filter.ref)[0])
        self.assertEqual(fin, 4 + len(self.d["garantias"]["por_familia"]))


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Saldo de facturas (auditoría)
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════════
@unittest.skipIf(not _HAY_OPENPYXL, "openpyxl no está instalado")
class TestSaldoFacturas(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        hoy = dt.datetime(2026, 10, 11, 21, 0)
        docs = [{"doc_tido": "FCV", "doc_nudo": "11610", "titulo": "FCV 11610", "cliente": "Centro de ejemplo", "cliente_rut": "76.555.444-3",
                 "fecha_doc": "07/10/2026", "total_doc": 1850000, "cap_servicio": 100000, "cap_despacho": 30000,
                 "usado_servicio": 90000, "usado_despacho": 0, "n_ot": 1, "ots_txt": "OT-2026-00322",
                 "estado": "con_saldo", "familia_de": None, "erp_leido": 1, "actualizado_el": hoy},
                {"doc_tido": "FCV", "doc_nudo": "11502", "titulo": "FCV 11502", "cliente": "=HYPERLINK(1)", "cliente_rut": "",
                 "fecha_doc": "fecha rara", "total_doc": 0, "cap_servicio": 300000, "cap_despacho": 0,
                 "usado_servicio": 0, "usado_despacho": 0, "n_ot": 0, "ots_txt": "", "estado": "libre", "familia_de": None,
                 "erp_leido": 1, "actualizado_el": None}]
        libro = [{"id": 1, "creado_el": hoy, "doc_tido": "FCV", "doc_nudo": "11610", "categoria": "servicio", "movimiento": "consumo",
                  "monto": 80000, "acumulado_ot": 80000, "numero_ot": "OT-2026-00322", "cliente": "Centro de ejemplo",
                  "origen": "crear", "usuario": "Ejecutivo de SSTT", "detalle": ""}]
        salida = {}

        def fetchall(sql, params=()):
            if "FROM ot_fac_doc" in sql:
                return docs
            if "FROM ot_fac_libro ORDER BY id" in sql:
                return libro
            return []
        amb = {"mysql_fetchall": fetchall, "mysql_fetchone": lambda *a, **k: {"r": hoy, "n": 1}, "io": io,
               "send_file": lambda buf, **k: salida.update(buf=buf) or "ok", "datetime": dt.datetime, "json": json,
               "jsonify": lambda d: d, "print": lambda *a, **k: None, "_saldo": ot_saldo_servicio, "current_username": lambda: "daniel",
               "chile_fmt_filter": lambda v, f="%d/%m/%Y %H:%M": v.strftime(f), "_FAC_ORIGEN_TXT": {"crear": "Creación de la OT"},
               "_ot_doc_real_a_usuario": lambda t, n: (t, n), "_ot_saldo_real": lambda k: k,
               "_FAC_LIBRO_DESDE": dt.datetime(2026, 10, 11), "_fac_finanzas_de_ots": lambda vids, fallos=None: {},
               "_ot_hijas_vigentes_lote": lambda ids: {}}
        amb = FA.cargar(funciones=("_fac_fila_lista", "ot_saldo_facturas_excel"), extra=amb)
        assert amb["ot_saldo_facturas_excel"]() == "ok"
        cls.wb = _abrir(salida["buf"])

    def test_la_fecha_del_documento_es_una_fecha_real_con_su_mes(self):
        ws = self.wb["Facturas"]
        c = ws.cell(row=2, column=_col(ws, "Fecha"))
        self.assertEqual((c.value, c.number_format), (dt.datetime(2026, 10, 7), F_D))
        col_a = _apoyo_cab(ws, 1, "de la fecha del documento")
        self.assertEqual(_apoyo_val(ws, 2, col_a), [2026, "Octubre", 10, 41, "Miércoles", "2026-10"])
        r = ws.cell(row=2, column=_col(ws, "Revisada (hora Chile)"))
        self.assertEqual((r.value, r.number_format), (dt.datetime(2026, 10, 11, 21, 0), F_DT))

    def test_una_fecha_rara_queda_como_texto_y_no_rompe(self):
        ws = self.wb["Facturas"]
        self.assertEqual(ws.cell(row=3, column=_col(ws, "Fecha")).value, "fecha rara")
        self.assertIsNone(ws.cell(row=3, column=_col(ws, "Revisada (hora Chile)")).value)
        self.assertEqual(_apoyo_val(ws, 3, _apoyo_cab(ws, 1, "de la fecha del documento")), [None] * 6)
        self.assertEqual(ws.cell(row=3, column=_col(ws, "Cliente")).data_type, "s", "un texto con «=» nunca es fórmula")

    def test_movimientos_en_hora_chile(self):
        ws = self.wb["Movimientos"]
        c = ws.cell(row=2, column=_col(ws, "Fecha (hora Chile)"))
        self.assertEqual((c.value, c.number_format), (dt.datetime(2026, 10, 11, 18, 0), F_DT))     # 21:00 UTC = 18:00 Chile
        self.assertEqual(_apoyo_val(ws, 2, _apoyo_cab(ws, 1, "de la fecha del movimiento"))[:2], [2026, "Octubre"])
        self.assertEqual(ws.cell(row=2, column=6).value, 80000)
        for hoja in ("Facturas", "Movimientos", "Alertas"):
            self.assertEqual(self.wb[hoja].freeze_panes, "A2")
            self.assertTrue(self.wb[hoja].auto_filter.ref)


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Facturación de proveedor: lote, OT por facturar, control de cobros, facturación por proveedor
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def _item_mfp(i, **kw):
    it = {"id": i, "numero_ot": f"OT-2026-{i:05d}", "fecha": "08/10/2026", "fecha_programada": "01/10/2026", "fecha_realizada": "02/10/2026",
          "fecha_cierre": "08/10/2026", "creada": "30/09/2026 09:15", "hora": "10:00", "tecnico_nombre": "Marcelo Pérez",
          "proveedor": "Transportes felcarm SPA", "proveedor_rut": "77.230.690-4", "prov_key": "1", "fecha_raw": dt.date(2026, 10, 1),
          "cliente": f"Cliente {i}", "cliente_rut": "76.000.000-0", "comuna": "Ñuñoa", "direccion": "Av. Siempre Viva 123", "titulo": "Instalación",
          "tipo_label": "Instalación", "cobertura_txt": "", "cobertura_corta": "", "es_garantia": False, "estado": "cerrada", "no_cerrada": False,
          "anexo_numero": 221, "anexo_estado": "firmado", "anexo_firmado": True, "servicio": 50000.0, "despacho": 0.0, "sugerido": 50000.0,
          "fac_monto": 50000.0, "cobrado_cliente_serv": 100000.0, "cobrado_cliente_envio": 0.0, "cobrado_cliente": 100000.0,
          "margen_fila": 50000.0, "margen_fila_pct": 50.0, "cobrado_respaldado": 100000.0, "observacion": "", "usuario": "daniel",
          "asignada": "08/10/2026 21:24", "eq": None, "sin_costo": False, "ticket_numero": "", "prioridad": "media", "centro_costo": "Servicio Técnico",
          "sin_cobro_declarado": False, "fuente_cobro": "zz", "cubierto_por": "", "zz_codigo": "ZZINSTALACION", "documento_cliente": "FCV 1",
          "estado_facturacion": "facturada", "margen": 50000.0, "margen_pct": 50.0, "created_by": "aaron", "fin_avisos": []}
    it.update(kw)
    return it


@unittest.skipIf(not _HAY_OPENPYXL, "openpyxl no está instalado")
class TestLoteProveedor(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        amb = FA.cargar(funciones=("_mfp_excel_lote", "_mfp_eq_texto"), extra={
            "_mfp_totales_lote": lambda monto, items: {"monto_total": monto, "total_cobrado_declarado": 0, "n_no_cobran": 0, "n_sin_cerrar": 0},
            "_mfp_tiene_factura_real": lambda f: False, "_ot_fin_clp": lambda v: "$" + str(v)})
        items = [_item_mfp(1), _item_mfp(2, fecha="", asignada="")]
        eventos = [{"fecha": "08/10/2026 20:43", "usuario": "daniel", "accion_txt": "Solicitud de OC creada", "detalle": "-2+3"}]
        meta = {"fecha_doc": "07/10/2026", "creado_txt": "08/10/2026 20:43 · por daniel", "pagada_txt": "", "generado_txt": "x"}
        cls.wb = _abrir(amb["_mfp_excel_lote"]({"id": 17, "monto_total": 100000.0, "estado_pago": "pendiente"}, items, eventos,
                                               "Transportes felcarm SPA", "77.230.690-4", meta))

    def test_fecha_y_asignada_el_son_fechas_reales(self):
        ws = self.wb["OT del lote"]
        c = ws.cell(row=2, column=_col(ws, "Fecha"))
        self.assertEqual((c.value, c.number_format), (dt.datetime(2026, 10, 8), F_D))
        c = ws.cell(row=2, column=_col(ws, "Asignada el"))
        self.assertEqual((c.value, c.number_format), (dt.datetime(2026, 10, 8, 21, 24), F_DT))
        self.assertIsNone(ws.cell(row=3, column=_col(ws, "Fecha")).value)
        self.assertIsNone(ws.cell(row=3, column=_col(ws, "Asignada el")).value)

    def test_columnas_de_apoyo_al_final_y_el_filtro_las_cubre(self):
        ws = self.wb["OT del lote"]
        col = _apoyo_cab(ws, 1, "de la fecha de la OT")
        self.assertEqual(col, 24, "después de «Punto de equilibrio» (W): las fórmulas del Resumen siguen apuntando a M y V")
        self.assertEqual(_apoyo_val(ws, 2, col), [2026, "Octubre", 10, 41, "Jueves", "2026-10"])
        self.assertEqual(ws.auto_filter.ref, "A1:AC3")
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws["A4"].value, "TOTALES")
        self.assertEqual(ws["M4"].value, "=SUM(M2:M3)")

    def test_historial_y_resumen(self):
        ws = self.wb["Historial"]
        self.assertEqual((ws["A2"].value, ws["A2"].number_format), (dt.datetime(2026, 10, 8, 20, 43), F_DT))
        self.assertEqual(ws["D2"].data_type, "s")
        r = self.wb["Resumen"]
        fila = next(i for i in range(1, r.max_row + 1) if r.cell(row=i, column=1).value == "Fecha del documento")
        self.assertEqual((r.cell(row=fila, column=2).value, r.cell(row=fila, column=2).number_format), (dt.datetime(2026, 10, 7), F_D))


class _Args(dict):
    def get(self, k, d=None):
        return dict.get(self, k, d)


@unittest.skipIf(not _HAY_OPENPYXL, "openpyxl no está instalado")
class TestOtPorFacturar(unittest.TestCase):
    def _libro(self, version):
        items = [_item_mfp(1), _item_mfp(2, fecha_realizada="", fecha_cierre="", creada="", no_cerrada=True, estado="en_curso")]
        resumen = {"pagado_serv": 100000.0, "pagado_desp": 0.0, "pagado": 100000.0, "cobrado": 200000.0, "n_sin_anexo": 0, "margen": 100000.0,
                   "margen_pct": 50.0, "pagado_garantia": 0.0, "pagado_sin_cobro": 0.0, "cobrado_serv": 200000.0, "cobrado_env": 0.0}
        amb = FA.cargar(funciones=("mant_facturas_proveedor_por_facturar_xlsx",), extra={
            "_facprov_puede": lambda: True, "_mfp_403": lambda: ("no", 403), "request": types.SimpleNamespace(args=_Args(version=version)),
            "datetime": dt.datetime, "_mfp_por_facturar": lambda **k: [dict(x) for x in items], "_mfp_resumen_filas": lambda lst: resumen,
            "_now_chile": lambda: dt.datetime(2026, 10, 29, 10, 0), "current_username": lambda: "daniel", "ILUS_BRAND": "ILUS Fitness",
            "_formato_rut_chile": lambda r: r, "_OT2_ESTADO_META": {"cerrada": ("Cerrada", "", "", "cerr")},
            "send_file": lambda buf, **k: buf})
        return _abrir(amb["mant_facturas_proveedor_por_facturar_xlsx"]())

    def test_version_interna_fechas_reales_y_apoyo(self):
        wb = self._libro("interno")
        ws = wb[wb.sheetnames[1]]
        c = ws.cell(row=2, column=_col(ws, "Fecha programada"))
        self.assertEqual((c.value, c.number_format), (dt.datetime(2026, 10, 1), F_D))
        self.assertEqual(ws.cell(row=2, column=_col(ws, "Fecha cierre")).value, dt.datetime(2026, 10, 8))
        c = ws.cell(row=2, column=_col(ws, "Creada"))
        self.assertEqual((c.value, c.number_format), (dt.datetime(2026, 9, 30, 9, 15), F_DT))
        self.assertIsNone(ws.cell(row=3, column=_col(ws, "Fecha cierre")).value)
        col = _apoyo_cab(ws, 1, "de la fecha programada")
        self.assertEqual(col + 5, len([x for x in _cab(ws) if x]))
        self.assertEqual(_apoyo_val(ws, 2, col)[:3], [2026, "Octubre", 10])
        self.assertTrue(ws.auto_filter.ref.endswith(openpyxl.utils.get_column_letter(col + 5) + "3"))
        # «Revisar» sigue marcando en rojo su propia columna (no la última de apoyo)
        self.assertEqual(ws.cell(row=3, column=_col(ws, "Revisar")).value, "OT no cerrada")

    def test_version_proveedor_tambien(self):
        wb = self._libro("proveedor")
        ws = wb[wb.sheetnames[1]]
        self.assertEqual(ws.cell(row=2, column=_col(ws, "Fecha programada")).value, dt.datetime(2026, 10, 1))
        _apoyo_cab(ws, 1, "de la fecha programada")
        self.assertNotIn("Cobrado total", _cab(ws))


@unittest.skipIf(not _HAY_OPENPYXL, "openpyxl no está instalado")
class TestControlCobrosExternos(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fila = {"id": 5, "numero_ot": "OT-2026-00005", "tipo": "instalacion", "estado": "cerrada", "fecha_programada": dt.date(2026, 10, 14),
                "hora_inicio": dt.timedelta(hours=9, minutes=30), "cerrada_at": dt.datetime(2026, 10, 15, 2, 0), "cliente_id": 7,
                "razon_social": "Cliente 7", "comuna": "Ñuñoa", "factura_tido": "FCV", "factura_nudo": "1", "documentos_extra": None,
                "fac_id": 17, "fac_tipo": "factura", "fac_numero": "5400", "fac_estado": "pagada", "fac_monto": 50000,
                "fac_created_at": dt.datetime(2026, 10, 20, 2, 0), "fac_created_by": "daniel", "fac_pagada_at": dt.datetime(2026, 10, 25, 15, 0),
                "fac_pagada_por": "daniel", "fac_ref": "", "fac_archivo": None, "fac_obs": "", "n_obl": 0, "n_obl_ok": 0,
                "firma_tecnico_at": 1, "firma_cliente_at": 1, "anexo_numero": None, "diagnostico": "", "zz_motivo_manual": None,
                "garantia_motivo": None, "prov_rut": "1-9"}
        lote = {"id": 17, "proveedor_nombre": "Felca", "proveedor_rut": "1-9", "tipo_documento": "factura", "numero_documento": "5400",
                "fecha": dt.date(2026, 10, 16), "monto_total": 50000, "n_ot": 1, "suma": 50000, "ots": "OT-2026-00005", "estado_pago": "pagada",
                "created_at": dt.datetime(2026, 10, 20, 2, 0), "created_by": "daniel", "pagada_at": None, "pagada_por": None,
                "pago_referencia": None, "archivo_url": None, "notas": None}

        def fetchall(sql, params=None):
            if "FROM mant_facturas_proveedor f" in sql:
                return [lote]
            if sql.startswith("SELECT id, cliente_id, fecha_programada"):
                return []
            return [fila]
        cobro = {"fin": {"cobra": True, "me_cobraron": {"tecnico": 50000, "despacho": 0}, "a_pagar_proveedor": 50000, "cobertura": "cobra",
                         "cobertura_txt": "Se cobra"}, "servicio": 100000, "envio": 0, "fuente": "zz", "sin_cobro": False, "avisos": []}
        amb = FA.cargar(funciones=("mant_facturas_proveedor_xlsx", "_mfp_xl_dt"), extra={
            "g": types.SimpleNamespace(permissions={"superadmin": True}), "jsonify": lambda d: d, "datetime": dt.datetime, "json": json,
            "request": types.SimpleNamespace(args=_Args(desde="2026-10-01", hasta="2026-10-31")), "_now_chile": lambda: dt.datetime(2026, 10, 29, 10),
            "_MFP_SQL_OT_EXTERNA": "1=1", "_ot_fin_sql_contrato_real": lambda a: "0", "_mfp_sql_ficha_de_tecnico": lambda a: "0",
            "mysql_fetchall": fetchall, "_mfp_nombre_proveedor_ot": lambda r: "Felca", "_ot_cobro_facprov": lambda r: cobro,
            "_OT2_ESTADO_META": {"cerrada": ("Cerrada", "", "", "cerr")}, "_TIPO_OT_LABEL": {"instalacion": "Instalación"},
            "_mant_log": lambda *a, **k: None, "current_username": lambda: "daniel", "Response": lambda body, **k: body})
        cls.wb = openpyxl.load_workbook(io.BytesIO(amb["mant_facturas_proveedor_xlsx"]()))

    def test_detalle_con_formato_espanol_y_apoyo_de_la_fecha_de_instalacion(self):
        ws = self.wb["Detalle"]
        c = ws.cell(row=2, column=1)
        self.assertEqual((c.value, c.number_format), (dt.datetime(2026, 10, 14), F_D))
        self.assertEqual(ws.cell(row=2, column=2).value, 42, "«Sem» se conserva")
        ag = ws.cell(row=2, column=_col(ws, "Agendada"))
        self.assertEqual((ag.value, ag.number_format), (dt.datetime(2026, 10, 14, 9, 30), F_DT))      # hora local: no se convierte
        reg = ws.cell(row=2, column=_col(ws, "Registrada"))
        self.assertEqual((reg.value, reg.number_format), (dt.datetime(2026, 10, 19, 23, 0), F_DT))    # UTC → Chile
        col = _apoyo_cab(ws, 1, "de la fecha de instalación")
        self.assertEqual(_apoyo_val(ws, 2, col), [2026, "Octubre", 10, 42, "Miércoles", "2026-10"])
        self.assertTrue(ws.auto_filter.ref.endswith(openpyxl.utils.get_column_letter(col + 5) + "2"))

    def test_facturas_y_filtros_en_formato_espanol(self):
        ws = self.wb["Facturas"]
        self.assertEqual(ws.cell(row=2, column=5).number_format, F_D)
        self.assertEqual(ws.cell(row=2, column=11).number_format, F_DT)
        self.assertEqual(self.wb["Filtros"]["B3"].number_format, F_D)


@unittest.skipIf(not _HAY_OPENPYXL, "openpyxl no está instalado")
class TestFacturacionProveedores(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        provs = [{"nombre": "Felca", "rut": "1-9", "n_ot": 1, "instalacion": 50000.0, "mantencion": 0.0, "otros": 0.0, "despacho": 0.0,
                  "total": 50000.0, "cobrado": 100000.0, "margen": 50000.0, "margen_pct": 50.0, "n_sin_cobro": 0}]
        tot = dict(provs[0], nombre="TOTAL")
        detalle = [{"numero_ot": "OT-2026-00005", "cerrada": "14/10/2026", "proveedor": "Felca", "cliente": "Cliente 7", "tipo_label": "Instalación",
                    "documento": "FCV 1", "servicio": 50000.0, "despacho": 0.0, "total": 50000.0, "cobrado_servicio": 100000.0,
                    "cobrado_envio": 0.0, "cobrado": 100000.0, "margen": 50000.0, "garantia": False, "fuente_cobro": "zz", "cobertura_txt": ""},
                   dict(numero_ot="OT-2026-00006", cerrada="", proveedor="Felca", cliente="X", tipo_label="", documento="", servicio=0.0,
                        despacho=0.0, total=0.0, cobrado_servicio=0.0, cobrado_envio=0.0, cobrado=0.0, margen=0.0, garantia=True,
                        fuente_cobro="", cobertura_txt="Garantía")]
        amb = FA.cargar(funciones=("mant_facturacion_proveedores_xlsx",), extra={
            "_facprov_puede": lambda: True, "jsonify": lambda d: d, "io": io,
            "_facprov_rango": lambda: (dt.datetime(2026, 10, 1), dt.datetime(2026, 11, 1), "2026-10", "2026-10"),
            "_facprov_datos": lambda d, h: (provs, detalle, tot), "send_file": lambda buf, **k: buf})
        cls.wb = _abrir(amb["mant_facturacion_proveedores_xlsx"]())

    def test_cerrada_es_fecha_real_con_apoyo_de_la_fecha_de_cierre(self):
        ws = self.wb["Detalle por OT"]
        c = ws.cell(row=2, column=_col(ws, "Cerrada"))
        self.assertEqual((c.value, c.number_format), (dt.datetime(2026, 10, 14), F_D))
        self.assertIsNone(ws.cell(row=3, column=_col(ws, "Cerrada")).value)
        col = _apoyo_cab(ws, 1, "de la fecha de cierre")
        self.assertEqual(_apoyo_val(ws, 2, col)[:5], [2026, "Octubre", 10, 42, "Miércoles"])
        self.assertEqual(ws.cell(row=2, column=_col(ws, "Margen")).number_format, '"$"#,##0;[Red]-"$"#,##0')
        self.assertTrue(ws.auto_filter.ref)
        self.assertEqual(ws.freeze_panes, "A2")

    def test_resumen_con_montos_con_miles_y_porcentaje_sin_cambiar_la_cifra(self):
        ws = self.wb["Resumen"]
        self.assertEqual(ws["H2"].number_format, '"$"#,##0;[Red]-"$"#,##0')
        self.assertEqual((ws["K2"].value, ws["K2"].number_format), (50.0, '0.0"%"'))
        self.assertEqual(ws.auto_filter.ref, "A1:L2", "la fila TOTAL queda fuera del filtro")


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Solicitudes de repuestos de OT · Acta de recepción
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════════
@unittest.skipIf(not _HAY_OPENPYXL, "openpyxl no está instalado")
class TestSolicitudesRepuestos(unittest.TestCase):
    def _libro(self, oculta=False):
        sols = [{"id": 1, "estado_label": "Solicitado", "contexto": "ot", "cliente_nombre": "Cliente", "maquina_nombre": "Trotadora",
                 "numero_ot": "OT-2026-00001", "numero_ticket": "", "repuesto_nombre": "Correa", "repuesto_sku": "SKU1", "cantidad": 1.0,
                 "proveedor_nombre": "Prov", "oc_numero": "OC1", "motivo": "", "solicitado_por": "lenin",
                 "created_at": dt.datetime(2026, 10, 20, 2, 0), "dias_transcurridos": 9.0},
                {"id": 2, "estado_label": "Solicitado", "contexto": "manual", "created_at": None, "dias_transcurridos": None}]
        amb = FA.cargar(funciones=("repstock_solicitudes_ot_export",), extra={
            "_otrep_filtros_query": lambda: ("1=1", []), "_OTREP_SQL_SOL": "SELECT", "mysql_fetchall": lambda sql, p=None: sols,
            "_otrep_fila": lambda r: r, "_oculta_proveedores": lambda: oculta, "jsonify": lambda d: d, "print": lambda *a, **k: None,
            "send_file": lambda buf, **k: buf, "_now_chile": lambda: dt.datetime(2026, 10, 29, 10)})
        return _abrir(amb["repstock_solicitudes_ot_export"]()).active

    def test_creado_era_utc_crudo_y_ahora_es_hora_chile_con_formato_espanol(self):
        ws = self._libro()
        c = ws.cell(row=2, column=_col(ws, "Creado (hora Chile)"))
        self.assertEqual((c.value, c.number_format), (dt.datetime(2026, 10, 19, 23, 0), F_DT))
        self.assertIsNone(ws.cell(row=3, column=_col(ws, "Creado (hora Chile)")).value)
        col = _apoyo_cab(ws, 1, "de la fecha de creación")
        self.assertEqual(_apoyo_val(ws, 2, col), [2026, "Octubre", 10, 43, "Lunes", "2026-10"])
        self.assertEqual(ws.auto_filter.ref, "A1:" + openpyxl.utils.get_column_letter(col + 5) + "3")
        self.assertEqual(ws.freeze_panes, "A2")

    def test_sin_proveedor_ni_oc_las_columnas_siguen_alineadas(self):
        ws = self._libro(oculta=True)
        self.assertNotIn("Proveedor", _cab(ws))
        self.assertEqual(ws.cell(row=2, column=_col(ws, "Creado (hora Chile)")).value, dt.datetime(2026, 10, 19, 23, 0))
        self.assertEqual(ws.cell(row=2, column=_apoyo_cab(ws, 1, "de la fecha de creación")).value, 2026)


@unittest.skipIf(not _HAY_OPENPYXL, "openpyxl no está instalado")
class TestActaRecepcion(unittest.TestCase):
    def test_fechas_del_acta_reales_y_tabla_con_filtro(self):
        amb = FA.cargar(funciones=("_acta_xl_celda", "_acta_xlsx_bytes"), extra={"ILUS_BRAND": "ILUS Fitness", "io": io})
        datos = {"numero_ot": "OT-2026-00400", "tipo_nombre": "Revisión de equipos", "documentos": ["NVI 687"], "revisores": ["Lenin"],
                 "ultima_revision": dt.datetime(2026, 10, 20, 2, 0), "generado_en": dt.datetime(2026, 10, 29, 15, 0), "n": 2,
                 "totales": [{"resultado": "Apto", "n": 2}],
                 "productos": [{"sku": "A", "nombre": "Banca", "cantidad": 1, "resultado": "Apto", "revisado_por": "Lenin",
                                "revisado_en": dt.datetime(2026, 10, 20, 2, 0)},
                               {"sku": "B", "nombre": "=cmd", "cantidad": 1, "resultado": "Apto", "revisado_en": None}]}
        ws = openpyxl.load_workbook(io.BytesIO(amb["_acta_xlsx_bytes"](datos))).active
        info = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2) for r in range(3, 9)}
        self.assertEqual((info["Última revisión"].value, info["Última revisión"].number_format), (dt.datetime(2026, 10, 19, 23, 0), F_DT))
        self.assertEqual(info["Generada"].value, dt.datetime(2026, 10, 29, 12, 0))
        enc = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "N°")
        c = ws.cell(row=enc + 1, column=9)
        self.assertEqual((c.value, c.number_format), (dt.datetime(2026, 10, 19, 23, 0), F_DT))
        self.assertIsNone(ws.cell(row=enc + 2, column=9).value)
        self.assertEqual(ws.cell(row=enc + 2, column=3).data_type, "s")
        self.assertEqual(ws.auto_filter.ref, f"A{enc}:I{enc + 2}")


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Un solo helper: nadie arma su propia lógica de fechas en estos Excel
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestUnSoloHelper(unittest.TestCase):
    FUNCS = ("ot2_reporte_xlsx", "ot_saldo_facturas_excel", "_mfp_excel_lote", "mant_facturas_proveedor_por_facturar_xlsx",
             "mant_facturas_proveedor_xlsx", "mant_facturacion_proveedores_xlsx", "repstock_solicitudes_ot_export", "_acta_xlsx_bytes")

    def test_cada_excel_usa_xlsx_fechas(self):
        for n in self.FUNCS:
            self.assertIn("import xlsx_fechas as _xf", FA.fuente_funcion(n), n)
        self.assertIn("import xlsx_fechas as XF", FA.fuente_funcion("excel_bytes", "ot_gerencial.py"))
        self.assertIn("_xf.a_fecha_excel(dt, utc=True)", FA.fuente_funcion("_mfp_xl_dt"))

    def test_ninguno_escribe_una_fecha_como_texto(self):
        # Antes: chile_fmt_filter(...) / strftime en una celda. Ahora esas columnas pasan por el helper.
        self.assertNotIn('chile_fmt_filter(r["creado_el"])', FA.fuente_funcion("ot_saldo_facturas_excel"))
        self.assertNotIn('s.get("created_at") or ""', FA.fuente_funcion("repstock_solicitudes_ot_export"))
        self.assertNotIn('_acta_fecha(pr.get("revisado_en"))', FA.fuente_funcion("_acta_xlsx_bytes"))
        self.assertNotIn('"DD-MM-YYYY"', FA.fuente_funcion("mant_facturas_proveedor_xlsx"))
        self.assertNotIn("fecha_hora_chile_txt(c[\"autorizo\"][\"at\"])", FA.fuente_funcion("excel_bytes", "ot_gerencial.py"))


if __name__ == "__main__":
    unittest.main()
