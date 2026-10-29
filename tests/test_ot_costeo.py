# -*- coding: utf-8 -*-
"""🔎 Costeo de OT · Explorar (FASE 1, Daniel 2026-10-29) — el lado Python: política de margen, dataset por OT, documento de
HOY, suma única, filtros del Excel, la ruta (solo superadmin, solo lectura) y la plantilla con sus dos pestañas.

Datos de MENTIRA (REGLA #22). La cuenta de cada OT es la REAL de app.py (`_ot_finanzas`, sacada por texto) y la suma se
compara contra `_ot_fin_agregado` REAL: las cifras del tablero tienen que ser las mismas del resto del proyecto.

Correr con:  python -m unittest tests.test_ot_costeo
"""
import functools
import io
import json
import os
import re
import sys
import time
import types
import unittest
from datetime import date, datetime, timedelta, timezone

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import ot_costeo as C  # noqa: E402
import ot_gerencial as G  # noqa: E402
import ot_margen_politica as POL  # noqa: E402
from tests import _ger_app as A  # noqa: E402
from tests import _ger_demo as D  # noqa: E402

HOY = D.HOY


def _utc(d, h=15):
    return datetime(d.year, d.month, d.day, h + 3, 0)


def extras():
    """OT de mentira que el demo del Resumen no trae: contrato real, desinstalación cubierta, cortesía, cobro declarado en $0
    con costo, despacho cobrado sin su costo, dos personas de la MISMA empresa externa y una OT de otra empresa."""
    base = dict(cliente_id=104, cliente="Hotel Épsilon Spa & Gym")
    ext_mili = dict(tecnico_user_id=41, tecnico_nombre="Persona Uno Mili", tecnico_role="tecnico_externo", te_id=9,
                    te_empresa="Mili Servicios SpA", proveedor_tipo="externo", proveedor_nombre="Mili Servicios SpA")
    ext_mili2 = dict(ext_mili, tecnico_user_id=42, tecnico_nombre="Persona Dos Mili")
    ext_dap = dict(tecnico_user_id=43, tecnico_nombre="Daniel Pulgar", tecnico_role="tecnico_externo", te_id=10,
                   te_empresa="DAP Servicios", proveedor_tipo="externo", proveedor_nombre="DAP Servicios")
    vs = [
        A.visita(501, cerrada_at=_utc(date(2026, 10, 20)), tipo="preventiva", contrato_real=1, zz_monto=None, zz_codigo=None,
                 costo_proveedor=80000, centro_costo="sstt", **dict(base, **ext_mili)),
        A.visita(502, cerrada_at=_utc(date(2026, 10, 21)), tipo="desinstalacion", ot_relacion="desinstalacion", ot_madre_id=9,
                 madre_valida=1, zz_monto=None, zz_codigo=None, costo_proveedor=60000, **dict(base, **ext_mili2)),
        A.visita(503, cerrada_at=_utc(date(2026, 10, 22)), modalidad_cobro="sin_costo", zz_monto=None, zz_codigo=None,
                 costo_proveedor=45000, valorizado_clp=70000, centro_costo="comercial", **dict(base, **ext_dap)),
        A.visita(504, cerrada_at=_utc(date(2026, 10, 22)), zz_monto=0, costo_proveedor=50000, **dict(base, **ext_dap)),
        A.visita(505, cerrada_at=_utc(date(2026, 10, 19)), zz_monto=300000, zz_envio_monto=20000, costo_proveedor=150000,
                 costo_despacho=None, **dict(base, **ext_mili)),
        A.visita(506, cerrada_at=_utc(date(2026, 10, 19)), zz_monto=500000, costo_proveedor=300000, costo_despacho=10000,
                 centro_costo="marketing", **dict(base, **ext_mili2)),
    ]
    return vs


def construir(p=None, docs=None, auts=None):
    """[(v, fin, fila_del_tablero)] del demo + extras, con la cuenta REAL."""
    p = p or C.leer_parametros({"desde": "2026-09-01", "hasta": "2026-10-25"}, HOY)
    vs, ctxs = D.dataset()
    vs = vs + extras()
    out = []
    for v in vs:
        ctx = dict(ctxs.get(v["id"], {}))
        ctx.setdefault("nombres", D.USUARIOS)
        ctx["tipo_txt"] = {"instalacion": "Instalación", "correctiva": "Mantención correctiva", "preventiva": "Mantención preventiva",
                           "desinstalacion": "Desinstalación", "revision_interna": "Trabajo de bodega"}.get(v["tipo"], v["tipo"])
        ctx["centro_txt"] = {"sstt": "Servicio Técnico", "comercial": "Comercial", "marketing": "Marketing"}.get(v["centro_costo"], "Sin centro de costo")
        fin = A.fin_de(v)
        f = G.armar_fila(v, fin, ctx)
        out.append((v, fin, C.fila(v, f, (docs or {}).get(v["id"], []), (auts or {}).get(v["id"], []), p)))
    return out, p


def agregado():
    return A.cuenta_real()["_ot_fin_agregado"]


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestPolitica(unittest.TestCase):
    """REGLA #57: 30 % objetivo; tolerancia desde 20 %; bajo 20 % amarillo; bajo 0 % rojo. Solo pinta."""

    def test_bandas_y_bordes(self):
        casos = [(45, "objetivo"), (30, "objetivo"), (29.9, "tolerancia"), (20, "tolerancia"), (19.99, "amarillo"),
                 (0, "amarillo"), (-0.1, "rojo"), (-80, "rojo")]
        for pct, esperado in casos:
            self.assertEqual(POL.banda(pct), esperado, pct)
        self.assertEqual((POL.OBJETIVO_PCT, POL.TOLERANCIA_PCT), (30.0, 20.0))

    def test_sin_porcentaje_no_se_inventa_salvo_una_perdida_real(self):
        self.assertIsNone(POL.banda(None))
        self.assertIsNone(POL.banda(None, 0))
        self.assertEqual(POL.banda(None, -50000), "rojo", "cobro declarado en $0 con costo: se perdió plata")
        self.assertIsNone(POL.banda("x"))

    def test_el_navegador_recibe_los_cortes_desde_aqui(self):
        j = POL.como_json()
        self.assertEqual((j["objetivo"], j["tolerancia"], j["perdida"]), (30.0, 20.0, 0.0))
        self.assertEqual(j["orden"], ["objetivo", "tolerancia", "amarillo", "rojo"])
        js = A.leer("static", "ot_costeo_explorar.js")
        self.assertIn("meta().politica", js, "el JS usa los cortes que manda el servidor")

    def test_no_toca_la_cuenta_ni_los_candados(self):
        src = A.leer("ot_margen_politica.py")
        for prohibido in ("import app", "_ot_finanzas", "_OT_FIN_UMBRAL_BAJO =", "mysql", "flask"):
            self.assertNotIn(prohibido, src.split('"""', 2)[-1])
        # 2026-10-29: el umbral de la ficha se alineó con la política (20 %): lo vigila tests/test_ot_margen_alineado.py
        self.assertIn("_OT_FIN_UMBRAL_BAJO = 20", A.app_src(), "el umbral de la ficha es la tolerancia de la política")


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestParametros(unittest.TestCase):
    def test_defecto_los_ultimos_12_meses_por_cierre_solo_cerradas(self):
        p = C.leer_parametros({}, HOY)
        self.assertEqual((p["desde"], p["hasta"], p["fecha"], p["alcance"]), (date(2025, 11, 1), HOY, "cierre", "cerradas"))
        self.assertEqual(p["carga_desde"], p["desde"] - timedelta(days=p["dias"]), "carga también el período anterior")
        self.assertEqual(p["previo_hasta"], p["desde"] - timedelta(days=1))

    def test_nada_se_confia_del_navegador(self):
        p = C.leer_parametros({"desde": "2026-10-20", "hasta": "2026-10-01", "fecha": "DROP", "alcance": "x"}, HOY)
        self.assertEqual((p["desde"], p["hasta"]), (date(2026, 10, 1), date(2026, 10, 20)), "al revés se da vuelta")
        self.assertEqual((p["fecha"], p["alcance"]), ("cierre", "cerradas"))
        p = C.leer_parametros({"desde": "2015-01-01", "hasta": "2026-10-20"}, HOY)
        self.assertEqual(p["dias"], C.MAX_RANGO_DIAS, "más de 3 años se acorta")
        p = C.leer_parametros({"hasta": "2030-01-01", "fecha": "programada", "alcance": "todas"}, HOY)
        self.assertEqual(p["hasta"], HOY + timedelta(days=C.MAX_FUTURO_DIAS))
        self.assertEqual((p["fecha"], p["alcance"]), ("programada", "todas"))
        p = C.leer_parametros({"desde": "basura", "hasta": "2026-13-45"}, HOY)
        self.assertEqual(p["hasta"], HOY)


class TestFechas(unittest.TestCase):
    def test_el_cierre_se_mira_en_hora_de_chile(self):
        # 02:00 UTC del 21 = 23:00 del 20 en Chile (verano, -3): cuenta para el 20
        self.assertEqual(C.fecha_local(datetime(2026, 10, 21, 2, 0), True), date(2026, 10, 20))
        self.assertEqual(C.fecha_local(date(2026, 10, 21), False), date(2026, 10, 21), "una fecha DATE no se mueve")
        self.assertIsNone(C.fecha_local(None, True))

    def test_dia_semana_numero_de_semana_mes_y_anio(self):
        d = C.derivar(date(2026, 10, 21))
        self.assertEqual(d, {"dia": "2026-10-21", "sem": "2026-10-19", "nsem": "43", "mes": "2026-10", "anio": "2026"})
        self.assertEqual(C.derivar(date(2027, 1, 1))["nsem"], "53", "N° de semana ISO (el 1-ene-2027 cae en la 53 de 2026)")
        self.assertEqual(C.derivar(None)["dia"], "")

    def test_etiquetas_en_espanol_nunca_iso(self):
        self.assertEqual(C.etiqueta_tiempo("mes", "2026-10"), "Octubre 2026")
        self.assertEqual(C.etiqueta_tiempo("nsem", "07"), "Semana 7")
        self.assertEqual(C.etiqueta_tiempo("dia", "2026-10-21"), "miércoles 21/10/2026")
        self.assertEqual(C.etiqueta_tiempo("sem", "2026-09-28"), "Semana 40 · 28 de sep al 4 de oct de 2026")
        self.assertEqual(C.etiqueta_tiempo("anio", "2026"), "2026")


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestDocumentoHoy(unittest.TestCase):
    """Con qué documento se respalda HOY lo cobrado (no cómo nació la OT)."""

    def test_nacio_con_nota_de_venta_y_ya_tiene_factura_cuenta_como_factura(self):
        docs = [{"origen": "erp", "tido": "NVV", "nudo": "VD00010667", "es_cobro": True},
                {"origen": "erp", "tido": "FCV", "nudo": "11348", "es_cobro": True}]
        self.assertEqual(C.documento_hoy(True, docs), "factura")
        self.assertEqual(C.documento_hoy(True, docs[:1]), "nota_venta")

    def test_orden_factura_boleta_nota_cotizacion_otro(self):
        self.assertEqual(C.documento_hoy(True, [{"tido": "BLV"}, {"tido": "NVI"}]), "boleta")
        self.assertEqual(C.documento_hoy(True, [{"tido": "COV"}]), "cotizacion")
        self.assertEqual(C.documento_hoy(True, [{"origen": "cotizacion", "tido": None}]), "cotizacion")
        self.assertEqual(C.documento_hoy(True, [{"tido": "GDV"}]), "otro_doc")
        self.assertEqual(C.documento_hoy(True, [{"tido": "WEB"}]), "nota_venta")

    def test_manda_el_documento_que_lleva_el_cobro(self):
        docs = [{"tido": "FCV", "es_cobro": False}, {"tido": "NVV", "es_cobro": True}]
        self.assertEqual(C.documento_hoy(True, docs), "nota_venta", "una factura de referencia no respalda el cobro")
        self.assertEqual(C.documento_hoy(True, [{"tido": "FCV", "es_cobro": False}]), "factura",
                         "si ninguno lleva el cobro, se mira el que hay")

    def test_sin_documento_autorizado_o_no(self):
        self.assertEqual(C.documento_hoy(True, []), "sin_documento")
        self.assertEqual(C.documento_hoy(True, [], [{"tipo": "crear_sin_documento", "estado": "aprobada"}]), "autorizado_sin_doc")
        self.assertEqual(C.documento_hoy(True, [], [{"tipo": "cerrar_con_perdida", "estado": "aprobada"}]), "sin_documento")
        self.assertEqual(C.documento_hoy(True, [], [{"tipo": "crear_sin_documento", "estado": "rechazada"}]), "sin_documento")

    def test_una_ot_que_no_se_cobra_no_aplica(self):
        self.assertEqual(C.documento_hoy(False, [{"tido": "NVV"}]), "no_aplica")

    def test_el_principal_cuenta_solo_si_se_valido_contra_random(self):
        v = {"factura_tido": "FCV", "factura_nudo": "00011348", "factura_asociada_por": "aaron"}
        self.assertEqual(C.docs_con_principal([], v), [{"origen": "erp", "tido": "FCV", "nudo": "00011348", "es_cobro": True}])
        self.assertEqual(C.docs_con_principal([{"tido": "FCV", "nudo": "11348"}], v), [{"tido": "FCV", "nudo": "11348"}],
                         "el mismo documento no se repite")
        self.assertEqual(C.docs_con_principal([], {"factura_tido": "FCV", "factura_nudo": "1"}), [], "escrito a mano: no cuenta")
        self.assertEqual(len(C.docs_con_principal([], {"factura_tido": "FCV", "factura_nudo": "1", "factura_rut_ok": 0})), 1)


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestDatasetYSumaUnica(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.trios, cls.p = construir()
        cls.filas = [t[2] for t in cls.trios]

    def comparar(self, filas, fins, que):
        k = C.kpis(filas)
        ag = agregado()(fins)
        for clave in ("cobre", "me_cobraron", "queda", "pct", "n", "n_fuera", "cobre_fuera", "n_no_cobra", "costo_no_cobra"):
            self.assertEqual(k[clave], ag[clave], f"{que}: {clave}")
        self.assertEqual({c: (a["n"], a["costo"]) for c, a in k["aparte"].items()},
                         {c: (a["n"], a["costo"]) for c, a in ag["aparte"].items()}, que + ": aparte")

    def test_las_cifras_son_iguales_a_ot_fin_agregado(self):
        self.comparar(self.filas, [t[1] for t in self.trios], "todas")
        act = C.actuales(self.filas)
        ids = {r["id"] for r in act}
        self.comparar(act, [t[1] for t in self.trios if t[2]["id"] in ids], "rango")

    def test_y_tambien_en_cada_grupo_filtrado(self):
        for dim in ("tec", "emp", "cc", "cob", "doc", "clase", "mes", "ttipo"):
            for valor in {r[dim] for r in self.filas}:
                sub = C.filtrar(self.filas, {dim: [valor]})
                ids = {r["id"] for r in sub}
                self.comparar(sub, [t[1] for t in self.trios if t[2]["id"] in ids], f"{dim}={valor}")

    def test_hay_de_todo_en_la_cuenta(self):
        clases = {r["clase"] for r in self.filas}
        self.assertTrue({"ganada", "perdida", "cobertura", "cubierta", "aparte", "sin_medir"} <= clases, clases)
        self.assertEqual({r["cob"] for r in self.filas if r["clase"] == "aparte"}, {"interno", "contrato"})

    def test_garantia_e_interno_nunca_con_margen_inventado(self):
        for r in self.filas:
            if r["clase"] in ("cobertura", "cubierta", "aparte"):
                self.assertIsNone(r["pct"], r["numero"])
                self.assertEqual(r["banda"], "no_cobra", r["numero"])
            if r["clase"] == "sin_medir":
                self.assertEqual(r["banda"], "sin_medir")
                self.assertIsNone(r["pct"], "sin el costo del técnico, _ot_finanzas daría 100 %: aquí no se inventa")
                self.assertIsNone(r["queda"])
            if r["clase"] == "aparte":
                self.assertIsNone(r["queda"])
            if r["cob"] == "interno":
                self.assertEqual(r["doc"], "no_aplica")
        cero = next(r for r in self.filas if r["id"] == 504)
        self.assertEqual((cero["pct"], cero["banda"]), (None, "rojo"), "cobro en $0 con costo: pérdida real, sin % inventado")

    def test_la_banda_de_cada_ot_es_la_de_la_politica(self):
        for r in self.filas:
            if r["clase"] in ("ganada", "perdida") and r["pct"] is not None:
                self.assertEqual(r["banda"], POL.banda(r["pct"]))

    def test_proveedores_agrupados_por_empresa(self):
        mili = [r for r in self.filas if r["id"] in (501, 502, 505, 506)]
        self.assertEqual({r["emp"] for r in mili}, {"emp:mili servicios spa"}, "dos personas, UNA empresa")
        self.assertEqual({r["tec"] for r in mili}, {"u:41", "u:42"}, "y cada persona por su lado")
        self.assertEqual({r["_lab"]["emp"] for r in mili}, {"Mili Servicios SpA"})
        dap = next(r for r in self.filas if r["id"] == 503)
        self.assertEqual((dap["emp"], dap["_lab"]["tec"], dap["ttipo"]), ("emp:dap servicios", "Daniel Pulgar", "externo"))
        propios = [r for r in self.filas if r["ttipo"] in ("interno", "elevado")]
        self.assertTrue(propios)
        self.assertEqual({r["emp"] for r in propios}, {"propios"})
        self.assertEqual({r["_lab"]["emp"] for r in propios}, {"Técnicos propios"})

    def test_el_periodo_anterior_viene_marcado(self):
        for r in self.filas:
            self.assertEqual(r["prev"], 1 if r["dia"] < "2026-09-01" else 0, r["numero"])
        self.assertTrue(any(r["prev"] for r in self.filas) and any(not r["prev"] for r in self.filas))

    def test_el_dataset_va_por_columnas_y_es_json(self):
        d = C.armar_dataset(self.filas, self.p, {"generado": "23/10/2026 18:00"})
        self.assertEqual(d["cols"], list(C.COLS))
        self.assertTrue(all(len(f) == len(C.COLS) for f in d["rows"]))
        self.assertEqual(len(d["rows"]), len(self.filas))
        texto = json.dumps(d, ensure_ascii=False)          # sin date ni Decimal sueltos
        self.assertNotIn("_lab", texto)
        self.assertEqual(d["meta"]["desde"], "2026-09-01")
        self.assertEqual(d["meta"]["politica"]["objetivo"], 30.0)
        self.assertEqual(d["meta"]["clasificacion"], "Confidencial · solo gerencia")
        for dim in C.DIMS:
            self.assertTrue(d["dims"][dim], dim)
            for clave, nombre in d["dims"][dim].items():
                self.assertTrue(nombre and isinstance(nombre, str), (dim, clave))
        self.assertEqual(d["tec_info"]["u:41"], {"ttipo": "externo", "emp": "emp:mili servicios spa"})

    def test_el_documento_de_hoy_sale_de_los_documentos_vivos(self):
        docs = {1: [{"origen": "erp", "tido": "NVV", "nudo": "1", "es_cobro": True}, {"origen": "erp", "tido": "FCV", "nudo": "2", "es_cobro": True}],
                2: [{"origen": "erp", "tido": "NVV", "nudo": "3", "es_cobro": True}]}
        trios, _ = construir(docs=docs, auts={3: [{"tipo": "crear_sin_documento", "estado": "aprobada"}]})
        por = {t[2]["id"]: t[2] for t in trios}
        self.assertEqual(por[1]["doc"], "factura")
        self.assertEqual(por[2]["doc"], "nota_venta")
        self.assertEqual(por[3]["doc"], "autorizado_sin_doc")
        self.assertEqual(por[4]["doc"], "sin_documento")
        self.assertEqual(por[1]["nac"], "factura", "cómo nació va aparte: sale de la clasificación guardada")


class TestFiltros(unittest.TestCase):
    def test_validar_filtros_no_confia_en_nada(self):
        f = C.validar_filtros(json.dumps({"tec": ["u:1", "u:1", 5], "malo": ["x"], "cc": "sstt", "mes": [True, "2026-10"],
                                          "doc": [], "reg": [{"x": 1}]}))
        self.assertEqual(f, {"tec": ["u:1", "5"], "mes": ["2026-10"]})
        self.assertEqual(C.validar_filtros("{no es json"), {})
        self.assertEqual(C.validar_filtros(None), {})
        self.assertEqual(len(C.validar_filtros({"cli": [str(i) for i in range(900)]})["cli"]), 500)

    def test_filtrar_y_excepto(self):
        filas = [{"id": 1, "tec": "a", "cc": "x"}, {"id": 2, "tec": "b", "cc": "x"}, {"id": 3, "tec": "a", "cc": "y"},
                 {"id": 4, "tec": None, "cc": "y"}]
        self.assertEqual([r["id"] for r in C.filtrar(filas, {"tec": ["a"], "cc": ["x"]})], [1])
        self.assertEqual([r["id"] for r in C.filtrar(filas, {"tec": ["a", "b"]})], [1, 2, 3])
        self.assertEqual([r["id"] for r in C.filtrar(filas, {"tec": ["a"], "cc": ["x"]}, ["tec"])], [1, 2])
        self.assertEqual([r["id"] for r in C.filtrar(filas, {"tec": [""]})], [4], "null se compara como ''")
        self.assertEqual(len(C.filtrar(filas, {})), 4, "sin filtros vuelven todas (limpiar recalcula)")

    def test_leer_ids(self):
        self.assertEqual(C.leer_ids("1, 2,x,3,"), {1, 2, 3})
        self.assertEqual(C.leer_ids(None), set())


class TestExcel(unittest.TestCase):
    def test_fechas_reales_columnas_de_tiempo_y_confidencial(self):
        import openpyxl
        trios, p = construir()
        filas = C.filtrar(C.actuales([t[2] for t in trios]), {"emp": ["emp:mili servicios spa"]})
        dims = C.armar_dataset([t[2] for t in trios], p)["dims"]
        contenido = C.excel_bytes(filas, p, {"emp": ["emp:mili servicios spa"]}, dims, generado="29/10/2026 10:00")
        wb = openpyxl.load_workbook(io.BytesIO(contenido))
        self.assertEqual(wb.sheetnames, ["OT filtradas", "Resumen"])
        ws = wb["OT filtradas"]
        self.assertIn("Confidencial · solo gerencia", ws["A1"].value)
        self.assertIn("Empresa o proveedor: Mili Servicios SpA", ws["A2"].value)
        enc = [c.value for c in ws[5]]
        for col in ("Fecha de cierre", "Año", "Mes", "N° semana", "Día", "Día de la semana", "Cobré ($)", "Me cobraron ($)",
                    "Margen ($)", "Margen %", "Margen según la política", "Documento hoy", "Cómo nació la OT", "Empresa o proveedor"):
            self.assertIn(col, enc)
        self.assertEqual(ws.max_row - 5, len(filas))
        i_f, i_mes, i_sem, i_dia = (enc.index("Fecha de cierre") + 1, enc.index("Mes") + 1, enc.index("N° semana") + 1,
                                    enc.index("Día") + 1)
        c = ws.cell(row=6, column=i_f)
        self.assertIsInstance(c.value, datetime, "fecha REAL de Excel")
        self.assertEqual(c.number_format, "dd/mm/yyyy")
        self.assertIn(ws.cell(row=6, column=i_mes).value, [m.capitalize() for m in G.MESES_ES])
        self.assertIsInstance(ws.cell(row=6, column=i_sem).value, int)
        self.assertIsInstance(ws.cell(row=6, column=i_dia).value, int)
        rs = wb["Resumen"]
        self.assertEqual(rs["A5"].value, "Cobré ($)")
        self.assertEqual(rs["B5"].value, C.kpis(filas)["cobre"])

    def test_una_ot_aparte_o_sin_medir_no_lleva_margen_inventado(self):
        import openpyxl
        trios, p = construir()
        filas = [t[2] for t in trios if t[2]["clase"] in ("aparte", "sin_medir")]
        wb = openpyxl.load_workbook(io.BytesIO(C.excel_bytes(filas, p)))
        ws = wb["OT filtradas"]
        enc = [c.value for c in ws[5]]
        for fila in range(6, ws.max_row + 1):
            self.assertIsNone(ws.cell(row=fila, column=enc.index("Margen ($)") + 1).value)
            self.assertIsNone(ws.cell(row=fila, column=enc.index("Margen %") + 1).value)

    def test_un_excel_vacio_no_revienta(self):
        _, p = construir()
        self.assertTrue(C.excel_bytes([], p, {}, {}).startswith(b"PK"))


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def bloque_costeo():
    """El bloque de Explorar en app.py: desde su encabezado hasta el encabezado del costeo gerencial que lo sigue."""
    src = A.app_src()
    ini = src.index("#  🔎 COSTEO DE OT · EXPLORAR")
    ini = src.rfind("# ═", 0, ini)
    fin = src.index("#  📊 COSTEO DE ÓRDENES DE TRABAJO — el tablero gerencial")
    fin = src.rfind("# ═", 0, fin)
    return src[ini:fin]


class TestAppBloque(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.src = A.app_src()
        cls.bloque = bloque_costeo()

    def test_va_antes_y_fuera_del_bloque_gerencial(self):
        self.assertEqual(self.src.count("#  🔎 COSTEO DE OT · EXPLORAR"), 1)
        self.assertNotIn("costos-dataset", A.bloque_gerencial(), "el bloque del Resumen sigue con sus dos rutas")
        self.assertIn("def _costeo_cargar", self.bloque)

    def test_rutas_solo_superadmin_y_con_costo_en_el_nombre(self):
        for ruta, fn in (('@app.route("/ot/api/costeo/costos-dataset", methods=["GET"])', "ot_costeo_costos_dataset"),
                         ('@app.route("/ot/api/costeo/costos-excel", methods=["GET"])', "ot_costeo_costos_excel")):
            i = self.bloque.index(ruta)
            cab = self.bloque[i:self.bloque.index("def " + fn, i)]
            self.assertIn("@_mant_required\n@_no_tecnico\n@_costeo_solo_superadmin\n", cab, fn)
            self.assertIn("costo", fn, "la prueba de confidencialidad la atrapa por el nombre")
        self.assertEqual(self.bloque.count("@app.route("), 2)
        self.assertIn("return _ot_solo_superadmin(view)(*a, **kw)", self.bloque)
        conf = A.leer("tests", "test_confidencialidad_financiera_tecnicos.py")
        self.assertIn("/ot/gerencial|/ot/api/costeo", conf)

    def test_solo_lectura_sin_random_ni_check_ni_ddl(self):
        sin_comentarios = "\n".join(l for l in self.bloque.split("\n") if not l.strip().startswith("#"))
        for verbo in ("INSERT", "UPDATE", "DELETE", "ALTER", "CREATE", "DROP", "TRUNCATE", "REPLACE", "GRANT", "MERGE"):
            self.assertIsNone(re.search(r"\b%s\b" % verbo, sin_comentarios), verbo)
        for llamada in ("mysql_execute", "conn.commit", ".commit(", "requests.", "_random_sql_query", "_checkwms_get",
                        "erp_engine", "fetch_document", "_send_ilus_email", "_mant_log", "_ot_finanzas_guardar"):
            self.assertNotIn(llamada, self.bloque, llamada)

    def test_cada_columna_que_lee_existe(self):
        usadas = set(re.findall(r"\bv\.(\w+)", self.bloque)) - {"id"}
        creada = self.src[self.src.index("CREATE TABLE IF NOT EXISTS mant_visitas ("):][:1800]
        declaradas = set(re.findall(r"\n\s+(\w+)\s+[A-Z]", creada))
        declaradas |= set(re.findall(r"ALTER TABLE mant_visitas ADD COLUMN (\w+)", self.src))
        i_cc = self.src.index('("centro_costo",  "VARCHAR(20) NULL')
        declaradas |= set(re.findall(r'\("(\w+)",\s+"[A-Z]', self.src[i_cc - 1500:i_cc + 1500]))
        fac = self.src[self.src.index("def _ensure_mant_visitas_factura_cols"):][:1500]
        declaradas |= set(re.findall(r'"(\w+)":\s+"[A-Z]', fac))
        for col in sorted(usadas):
            self.assertIn(col, declaradas, "mant_visitas no declara " + col)
        self.assertIn("ALTER TABLE mant_clientes ADD COLUMN region", self.src)
        vd = self.src[self.src.index("CREATE TABLE IF NOT EXISTS mant_visita_documentos"):][:1200]
        for col in ("visita_id", "origen", "erp_tido", "erp_nudo", "es_cobro"):
            self.assertRegex(vd, r"\n\s+%s\s" % col)
        self.assertIn("ALTER TABLE mant_visita_documentos ADD COLUMN heredado_de_visita_id", self.src)
        for col in ("fecha_programada", "fecha_realizada", "created_at", "cerrada_at"):
            self.assertIn(col, declaradas)
        self.assertEqual(set(C.FECHA_COLUMNA.values()), {"cerrada_at", "fecha_realizada", "created_at", "fecha_programada"})

    def test_reutiliza_la_cuenta_y_no_la_reimplementa(self):
        self.assertIn("_ot_fin_lote(filas_sql)", self.bloque)
        self.assertIn("_ger.armar_fila(r, fin, ctx)", self.bloque)
        self.assertNotIn("costo_proveedor", self.bloque.split("def _costeo_where")[1], "la plata no se recalcula aquí")


class TestAppEjecutado(unittest.TestCase):
    """Ejecuta el bloque REAL (y el del Resumen, del que toma las lecturas) con una base de mentira."""

    def montar(self, es_superadmin=True, docs=None, args=None):
        vs, _ = D.dataset()
        vs = vs + extras()
        cuenta = A.cuenta_real()
        docs = docs or {}

        def responder(sql, params):
            if "information_schema.COLUMNS" in sql:
                return []
            if "c.region AS cliente_region" in sql:
                ini, fin, limite = params
                return [dict(v) for v in vs if v.get("cerrada_at") and ini <= v["cerrada_at"] < fin][:limite]
            if "FROM mant_visita_documentos" in sql:
                return [dict(d, visita_id=vid) for vid, ds in docs.items() if vid in params for d in ds]
            if "FROM mant_ot_autorizaciones" in sql:
                return []
            raise AssertionError("SQL no previsto: " + sql[:160])

        db = A.BaseFalsa(responder)
        avisos = []

        class Resp(dict):
            def __init__(self, d):
                super().__init__(d)
                self.headers = {}

        amb = {
            "__name__": "app_bloque", "app": types.SimpleNamespace(route=lambda *a, **k: (lambda f: f)), "wraps": functools.wraps,
            "time": time, "io": io, "datetime": datetime, "timedelta": timedelta, "timezone": timezone, "date": date,
            "_mant_required": lambda f: f, "_no_tecnico": lambda f: f,
            "mysql_fetchall": db.fetchall, "mysql_fetchone": db.fetchone, "_ot_resp_ok": lambda: True,
            "_ot_fin_cols_sql": lambda alias="v": ", ".join(f"{alias}.{c}" for c in (
                "modalidad_cobro", "cubierto_por", "tipo", "cliente_id", "contrato_id", "costo", "zz_monto", "cobro_cero_motivo", "zz_codigo",
                "zz_envio_monto", "valor_origen", "costo_proveedor", "costo_despacho", "proveedor_tipo", "valorizado_clp", "valorizado_fuente",
                "zz_motivo_manual", "ot_madre_id", "ot_relacion")) + ", 0 AS contrato_real, 0 AS madre_valida",
            "_ot_fin_lote": lambda rows: {int(r["id"]): cuenta["_ot_finanzas"](r) for r in rows},
            "_ot_finanzas": cuenta["_ot_finanzas"], "_ot_fin_agregado": cuenta["_ot_fin_agregado"],
            "_ot_fin_clasif_lote": lambda ids: {}, "_ot_intenciones_catalogo": lambda: [],
            "_ot_intencion_info": lambda *a, **k: None,
            "_now_chile": lambda: datetime(HOY.year, HOY.month, HOY.day, 15, 0, tzinfo=timezone(timedelta(hours=-3))),
            "_TIPO_OT_LABEL": {"instalacion": "Instalación"}, "_OT2_CENTROS_COSTO": (("sstt", "Servicio Técnico"), ("comercial", "Comercial")),
            "_ot_aut_es_superadmin": lambda: es_superadmin, "_friendly_error_page": lambda t, m, c: ("error", t, c),
            "render_template": lambda plantilla, **kw: ("render", plantilla, kw),
            "request": types.SimpleNamespace(args=args or {}, path="/ot/api/costeo/costos-dataset"),
            "flash": lambda *a, **k: avisos.append(a), "redirect": lambda u: ("redirect", u), "url_for": lambda e, **k: "/" + e,
            "jsonify": lambda d: Resp(d), "send_file": lambda f, **k: ("archivo", f.getvalue(), k),
        }
        exec(compile(A.bloque(A.app_src(), "_ot_solo_superadmin"), "<app:_ot_solo_superadmin>", "exec"), amb)
        exec(compile(A.bloque_gerencial(), "<app:bloque_gerencial>", "exec"), amb)
        exec(compile(bloque_costeo(), "<app:bloque_costeo>", "exec"), amb)
        return amb, db

    def test_arma_el_dataset_leyendo_solo_con_select_y_por_lotes(self):
        amb, db = self.montar()
        r = amb["ot_costeo_costos_dataset"]()
        self.assertTrue(r["ok"])
        self.assertEqual(r.headers.get("Cache-Control"), "no-store")
        self.assertEqual(r["cols"], list(C.COLS))
        self.assertGreater(len(r["rows"]), 20)
        self.assertTrue(all(s.split(" ", 1)[0].upper() == "SELECT" for s in db.sqls))
        self.assertLessEqual(len(db.sqls), 5, "nada de una consulta por OT")
        for s in db.sqls:
            self.assertNotIn("2026-", s, "las fechas viajan como parámetros")

    def test_los_limites_viajan_en_utc_e_incluyen_el_periodo_anterior(self):
        amb, db = self.montar()
        capt = []
        base = db.responder

        def espia(sql, params):
            if "c.region AS cliente_region" in sql:
                capt.append(params)
            return base(sql, params)

        db.responder = espia
        p = C.leer_parametros({"desde": "2026-10-01", "hasta": "2026-10-31"}, HOY)
        amb["_costeo_cargar"](p)
        ini, fin, limite = capt[0]
        self.assertEqual(ini, datetime(2026, 8, 31, 4, 0), "31 días antes del 1 de octubre, a medianoche de Chile en UTC")
        self.assertEqual(fin, datetime(2026, 11, 1, 3, 0), "hasta el fin del 31 de octubre en Chile (ya en horario de verano)")
        self.assertEqual(limite, C.LIMITE_FILAS)

    def test_la_fecha_que_manda_cambia_la_columna_y_las_fechas_date_no_se_mueven(self):
        amb, _ = self.montar()
        w, params, col = amb["_costeo_where"](C.leer_parametros({"fecha": "programada", "alcance": "todas", "desde": "2026-10-01",
                                                                 "hasta": "2026-10-31"}, HOY))
        self.assertEqual(col, "fecha_programada")
        self.assertIn("v.fecha_programada >= %s AND v.fecha_programada <= %s", w)
        self.assertIn("NOT IN ('cancelada','anulada')", w)
        self.assertEqual(params, (date(2026, 8, 31), date(2026, 10, 31)))
        w, _, _ = amb["_costeo_where"](C.leer_parametros({}, HOY))
        self.assertIn("v.estado = 'cerrada'", w)

    def test_la_cache_de_90_segundos(self):
        amb, db = self.montar()
        p = C.leer_parametros({}, HOY)
        amb["_costeo_datos"](p)
        n = len(db.sqls)
        amb["_costeo_datos"](p)
        self.assertEqual(len(db.sqls), n)
        amb["_costeo_datos"](p, forzar=True)
        self.assertGreater(len(db.sqls), n)
        self.assertEqual(amb["_COSTEO_CACHE_SEG"], 90)

    def test_un_tecnico_no_recibe_nada_y_ni_se_lee_la_base(self):
        amb, db = self.montar(es_superadmin=False)
        for vista in ("ot_costeo_costos_dataset", "ot_costeo_costos_excel"):
            r = amb[vista]()
            self.assertEqual(r[1], 403, vista)
            self.assertFalse(r[0]["ok"])
        self.assertEqual(db.sqls, [], "ni siquiera lee la base")

    def test_el_documento_de_hoy_se_lee_por_lote(self):
        docs = {1: [{"origen": "erp", "erp_tido": "NVV", "erp_nudo": "VD1", "es_cobro": 1},
                    {"origen": "erp", "erp_tido": "FCV", "erp_nudo": "99", "es_cobro": 1}]}
        amb, _ = self.montar(docs=docs)
        d = amb["_costeo_cargar"](C.leer_parametros({"desde": "2026-08-01", "hasta": "2026-10-25"}, HOY))
        por = {r["id"]: r for r in d["filas"]}
        self.assertEqual(por[1]["doc"], "factura")

    def test_el_excel_reaplica_los_filtros_y_la_seleccion(self):
        import openpyxl
        args = {"desde": "2026-08-01", "hasta": "2026-10-25", "f": json.dumps({"emp": ["emp:mili servicios spa"]})}
        amb, _ = self.montar(args=args)
        r = amb["ot_costeo_costos_excel"]()
        self.assertEqual(r[0], "archivo")
        self.assertTrue(r[2]["download_name"].startswith("Costeo_OT_Explorar_01-08-2026_al_25-10-2026"))
        ws = openpyxl.load_workbook(io.BytesIO(r[1]))["OT filtradas"]
        self.assertEqual(ws.max_row - 5, 4, "las 4 OT de Mili")
        amb["request"].args = dict(args, ids="501,506")
        r = amb["ot_costeo_costos_excel"]()
        ws = openpyxl.load_workbook(io.BytesIO(r[1]))["OT filtradas"]
        self.assertEqual(ws.max_row - 5, 2, "solo las elegidas")
        self.assertIn("Solo las 2 OT seleccionadas", ws["A2"].value)


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestPlantilla(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import jinja2
        stub = ("<!doctype html><html><head>{% block head_extra %}{% endblock %}</head><body>"
                "{% block content %}{% endblock %}{% block scripts %}{% endblock %}</body></html>")
        env = jinja2.Environment(loader=jinja2.ChoiceLoader([jinja2.DictLoader({"base.html": stub}),
                                                             jinja2.FileSystemLoader(os.path.join(A.RAIZ, "templates"))]),
                                 autoescape=True)

        def url_for(endpoint, **kw):
            if endpoint == "static":
                return "/static/" + kw["filename"]
            if endpoint == "ot2_detalle":
                return "/ot/%s" % kw["vid"]
            return "/" + endpoint
        env.globals["url_for"] = url_for
        cls.html = env.get_template("ot2/gerencial.html").render(d=D.tablero_demo(), G=G, p={"modo": "semana", "dias": 60})

    def test_dos_pestanas_resumen_primero_y_explorar_escondida(self):
        h = self.html
        self.assertLess(h.index('data-gc-pestana="resumen"'), h.index('data-gc-pestana="explorar"'))
        self.assertIn('id="gcApp" role="tabpanel"', h)
        self.assertRegex(h, r'<section class="cx" id="cxApp" role="tabpanel" aria-labelledby="gcPestExplorar" hidden')
        self.assertIn('data-url="/ot_costeo_costos_dataset"', h)
        self.assertIn('data-excel="/ot_costeo_costos_excel"', h)
        self.assertIn('data-ot="/ot/0"', h)
        self.assertIn("/static/ot_costeo_explorar.js", h)
        self.assertIn("/static/ot_costeo_explorar.css", h)

    def test_el_resumen_no_cambia(self):
        h = self.html
        self.assertEqual(h.count("¿Cómo se calcula?"), 5, "el Resumen conserva sus cinco tarjetas (REGLA #4.2)")
        for ancla in ('id="gcT-cobre"', 'id="gcGarantias"', 'id="gcTendencia"', 'id="gcReinc"', 'id="gcCamino"'):
            self.assertIn(ancla, h)
        self.assertEqual(h.count('data-gc-solo="resumen"'), 2, "Excel del período y Actualizar se esconden en Explorar")


if __name__ == "__main__":
    unittest.main()
