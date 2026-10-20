# -*- coding: utf-8 -*-
"""🧰 2026-10-18 — Revisión de equipos, ETAPA 3: acciones arriba de cada producto, repuestos sugeridos y acta.

Daniel: «arriba debe tener las opciones para dar de baja, para poner fuera de servicio, para decir que el producto no
está, para solicitar repuestos, solicitar piolas si es selectora de peso, solicitar cintas si es trotadora…
inteligentemente por las plantillas» y «generar un reporte para enviar a logística de lo que se está recibiendo y lo que
no, de la nota de venta, para que ellos puedan hacer la recepción y la gestión documental».

Se prueba:
  1. La barra de cada producto: Dar de baja / Fuera de servicio / No llegó ESCRIBEN el Resultado (mismo guardado), y
     Solicitar repuesto / Informar incidencia con los repuestos sugeridos de SUS bloques (node).
  2. Repuestos sugeridos: semilla de la propuesta (solo nombres), normalización, siembra UNA sola vez sin pisar lo
     editado, editor de plantillas y checklist de la ficha.
  3. El flujo de repuestos de la OT para un PRODUCTO de trabajo interno (rutas /producto/<id>, candados, sin tocar las
     de cliente) y la incidencia desde el producto (mismo registro de Incidencias).
  4. Técnicos: no ven proveedores ni costos (mismos recortes del servidor), ni el acta.
  5. El acta: armado puro, PDF (plantilla) y Excel sin datos internos, envío SOLO a dominios internos con registro y
     con el interruptor de correo respetado.

Sin BD, sin Flask, sin ERP (REGLA #4.1), sin correos reales (REGLA #22).
Correr:  python -m pytest tests/test_ot_revision_acta.py -q -p no:warnings
"""
import ast
import datetime as _dt
import io
import json
import re
import unicodedata
import unittest

import jinja2

from tests.test_fac_libro_saldo import _exec
from tests.test_ot_revision_v3 import HAY_NODE, _js_funcion, _node
from tests.test_ot_saldo_servicio import _arbol, _decoradores, _fuente, _leer

D = _leer("templates/ot2/detalle.html")
P = _leer("templates/mantenciones/plantillas.html")
CONF = _leer("templates/mantenciones/configuracion.html")
ACTA_TPL = _leer("templates/ot2/acta_recepcion_pdf.html")


def _cargar(funcs=(), consts=(), amb=None):
    """Ejecuta en un mismo espacio las funciones (sin decoradores) y constantes de app.py indicadas."""
    amb = {} if amb is None else amb
    amb.setdefault("json", json)
    amb.setdefault("unicodedata", unicodedata)
    amb.setdefault("re", re)
    amb.setdefault("io", io)
    _codigo, arbol = _arbol()
    for n in arbol.body:
        if isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and t.id in consts for t in n.targets):
            exec(compile(ast.Module(body=[n], type_ignores=[]), "<app>", "exec"), amb)
    for n in arbol.body:
        if isinstance(n, ast.FunctionDef) and n.name in funcs:
            n2 = ast.parse(ast.unparse(n)).body[0]
            n2.decorator_list = []
            exec(compile(ast.Module(body=[n2], type_ignores=[]), "<app>", "exec"), amb)
    faltan = [x for x in tuple(funcs) + tuple(consts) if x not in amb]
    assert not faltan, f"no están en app.py: {faltan}"
    return amb


def _js_desde(marca, src=D):
    """Fuente desde `marca` hasta la llave que cierra su primer bloque (para `window.x = function(...){...}`)."""
    i = src.index(marca)
    j = src.index("{", i)
    prof, k, comillas = 0, j, None
    while k < len(src):
        ch = src[k]
        if comillas:
            if ch == "\\":
                k += 2
                continue
            if ch == comillas:
                comillas = None
        elif src.startswith("//", k):
            k = src.index("\n", k)
            continue
        elif src.startswith("/*", k):
            k = src.index("*/", k) + 2
            continue
        elif ch in "'\"`":
            comillas = ch
        elif ch == "{":
            prof += 1
        elif ch == "}":
            prof -= 1
            if prof == 0:
                return src[i:k + 1] + ";"
        k += 1
    raise AssertionError("sin cierre: " + marca)


_PLATA = re.compile(r"proveedor_|[\"']proveedor|costo|precio|monto|margen|zz_|sin_saldo|factura_monto", re.I)
_NATIVO = re.compile(r"(?<![\w.])(alert|confirm|prompt)\(")


# ═══════════════════════════════ 1) la barra de cada producto ═══════════════════════════════
class TestBarraDelProducto(unittest.TestCase):
    def test_dar_de_baja_fuera_de_servicio_y_no_llego_escriben_el_resultado(self):
        barra = _js_funcion("otdChkBarraResultado", D)
        for o in ("'dar de baja'", "'fuera de servicio'", "'no llego'"):
            self.assertIn(o, barra)
        rap = _js_funcion("otdChkResultadoRapido", D)
        self.assertIn("const d = await otdChkGuardar(tid, opcion);", rap)   # MISMO guardado (regla dura incluida)
        self.assertIn("if (d) otdChkPedirArgumento(found.mid, n);", rap)   # y pide el argumento en la Observación
        g = _js_funcion("otdChkGuardar", D)
        self.assertIn("/tareas/${tid}/respuesta", g)
        self.assertIn("d.error_codigo === 'RESULTADO_CON_FALLA'", g)
        arg = _js_funcion("otdChkPedirArgumento", D)
        self.assertIn("['no llego', 'dar de baja', 'fuera de servicio']", arg)
        self.assertIn("startsWith('observacion')", arg)
        self.assertIn("'otdChkVal-' + obs.id", arg)

    def test_la_barra_del_producto_se_dibuja_bajo_el_resultado(self):
        r = _js_funcion("otdRenderChecklistEquipo", D)
        self.assertIn("otdChkBarraResultado(_resultado, _ev.no_llego.length > 0)", r)
        self.assertIn("+ otdChkBarraProducto(mid, grupos, _resultado, _ev.no_llego.length > 0) + grupos.map(g => {", r)
        b = _js_funcion("otdChkBarraProducto", D)
        self.assertIn("Solicitar repuesto", b)
        self.assertIn("Informar incidencia", b)
        self.assertIn("otdProdRepuesto(${imid}, this.dataset.n)", b)
        self.assertIn("!(res.item_manual_id && !res.maquina_id)", b)   # solo productos de trabajo interno
        self.assertIn("const OTD_PUEDE_PRODUCTO = {{ 'true' if ((puede_ejecutar or puede_configurar) and not "
                      "ot_solo_lectura) else 'false' }};", D)

    def test_sin_popups_nativos(self):   # REGLA #1
        for n in ("otdChkBarraProducto", "otdProdRepuesto", "otdProdIncidencia", "otdChkPedirArgumento",
                  "otdActaEnviar", "otincEnviar", "otincAgregarArchivos", "otincRefrescar", "otdChkResultadoRapido"):
            self.assertIsNone(_NATIVO.search(_js_funcion(n, D)), n)
        self.assertIsNone(_NATIVO.search(_js_desde("window.otrepUsarSugerido = function(")))
        self.assertIsNone(_NATIVO.search(_js_desde("window.otdIncAbrir = function(")))

    def test_toques_de_44px(self):   # REGLA #3
        for css in (".otd-prod-btn{min-height:44px;", ".otd-prod-chip{min-height:44px;",
                    ".otd-acta-btn{min-height:44px;", "#modalRepSol .otrep-sug-chip{min-height:44px;"):
            self.assertIn(css, D)
        self.assertIn(".pl-rsug-chip button{border:0;background:none;color:#dc2626;min-width:44px;min-height:44px;", P)


@unittest.skipUnless(HAY_NODE, "node no está instalado")
class TestBarraEnNode(unittest.TestCase):
    # otdEsc se reemplaza por uno equivalente: su regex /'/g confunde al extractor de funciones (_js_funcion).
    PREVIO = ("const OTD_VID = 77; let OTD_PUEDE_PRODUCTO = true; let OTD_PUEDE_INCIDENCIA = true;\n"
              "function otdEsc(s){ return String(s == null ? '' : s).replace(/&/g, '&amp;').replace(/</g, '&lt;')"
              ".replace(/>/g, '&gt;').replace(/\"/g, '&quot;'); }\n"
              "function otdChkEstaBloqueada(){ return false; } function otdBloqueadaPorCierre(){ return false; }\n")
    FUNCS = ("otdChkNorm", "otdChkVj", "otdChkOpciones", "otdChkEsResultado", "otdChkBaseMarcoObs",
             "otdChkSugeridos", "otdChkBarraProducto")

    def _js(self, cuerpo):
        fuente = self.PREVIO + "\n".join(_js_funcion(n, D) for n in self.FUNCS)
        return _node(fuente + "\n" + cuerpo)

    def test_chips_de_los_bloques_del_producto_sin_repetir(self):
        grupos = [{"repuestos_sugeridos": []},
                  {"repuestos_sugeridos": ["Piola (cable de acero) o correa de transmisión", "Pin selector con cordel"]},
                  {"repuestos_sugeridos": ["PIN SELECTOR CON CORDEL", "Kit de pernos y tuercas"]}]
        res = {"id": 5, "item_manual_id": 3, "maquina_id": None}
        out = self._js("const h = otdChkBarraProducto(3, " + json.dumps(grupos, ensure_ascii=False) + ", "
                       + json.dumps(res) + ");\n"
                       "console.log(JSON.stringify({h: h, n: (h.match(/otd-prod-chip/g) || []).length}));")
        self.assertIn("Solicitar repuesto", out["h"])
        self.assertIn("Informar incidencia", out["h"])
        self.assertIn("otdProdRepuesto(3)", out["h"])
        self.assertIn("otdProdIncidencia(3)", out["h"])
        self.assertEqual(out["n"], 3)                                     # 3 sugeridos distintos (sin duplicar)
        self.assertIn("Piola (cable de acero) o correa de transmisión", out["h"])   # nombre COMPLETO (REGLA #15)

    def test_condicional_cerrado_y_no_llego_sin_chips(self):
        """✂️ 2026-10-19 (revisión de la Etapa 3): solo los sugeridos de los bloques que aplican. El condicional
        («Estructura y tapiz») aporta los suyos solo abierto; un producto «No llegó» no muestra chips."""
        base = {"tareas": [{"id": 1, "tipo_respuesta": "lista", "titulo": "Resultado del control",
                            "opciones_lista_json": json.dumps(["Apto", "Apto con observaciones", "Reparar"]),
                            "valor_json": {"opcion": "Apto"}}], "repuestos_sugeridos": []}
        sist = {"tareas": [], "repuestos_sugeridos": ["Polea con rodamiento", "Pin selector con cordel"]}
        cond = {"tareas": [], "condicional": True, "completas": 0,
                "repuestos_sugeridos": ["Kit de pernos y tuercas", "Pintura de retoque"]}
        res = {"id": 1, "item_manual_id": 3, "maquina_id": None}
        reparar = json.loads(json.dumps(base))
        reparar["tareas"][0]["valor_json"] = {"opcion": "Reparar"}
        out = self._js(
            "const n = h => (h.match(/otd-prod-chip/g) || []).length;\n"
            "const G = " + json.dumps({"cerrado": [base, sist, cond], "abierto": [reparar, sist, cond]},
                                      ensure_ascii=False) + ";\n"
            "const R = " + json.dumps(res) + ";\n"
            "console.log(JSON.stringify({cerrado: n(otdChkBarraProducto(3, G.cerrado, R)),"
            " abierto: n(otdChkBarraProducto(3, G.abierto, R)),"
            " no_llego: n(otdChkBarraProducto(3, G.abierto, R, true)),"
            " pintura: otdChkBarraProducto(3, G.cerrado, R).includes('Pintura de retoque')}));")
        self.assertEqual(out, {"cerrado": 2, "abierto": 4, "no_llego": 0, "pintura": False})

    def test_informar_incidencia_solo_con_el_candado_del_modulo(self):
        """🔒 2026-10-19: sin OTD_PUEDE_INCIDENCIA (técnico de campo común o externo) no hay botón de incidencia."""
        res = {"id": 1, "item_manual_id": 3, "maquina_id": None}
        out = self._js("OTD_PUEDE_INCIDENCIA = false;\n"
                       "const h = otdChkBarraProducto(3, [], " + json.dumps(res) + ");\n"
                       "console.log(JSON.stringify([h.includes('Solicitar repuesto'), h.includes('otdProdIncidencia')]));")
        self.assertEqual(out, [True, False])

    def test_no_aparece_en_equipos_de_ficha_ni_sin_permiso_ni_sin_resultado(self):
        out = self._js(
            "const a = otdChkBarraProducto(9, [], {id: 1, maquina_id: 9, item_manual_id: null});\n"
            "const b = otdChkBarraProducto(3, [], null);\n"
            "OTD_PUEDE_PRODUCTO = false;\n"
            "const c = otdChkBarraProducto(3, [], {id: 1, item_manual_id: 3, maquina_id: null});\n"
            "console.log(JSON.stringify([a, b, c]));")
        self.assertEqual(out, ["", "", ""])

    def test_rutas_del_modal_para_un_producto(self):
        fuente = ("const OTD_VID = 77;\n" + _js_funcion("_otrepEsProducto", D) + "\n" + _js_funcion("_otrepBase", D))
        out = _node(fuente + "\nconsole.log(JSON.stringify([_otrepBase('i5'), _otrepBase(12), _otrepEsProducto('i5'),"
                             " _otrepEsProducto(5), _otrepEsProducto('5'), _otrepEsProducto('i5x')]));")
        self.assertEqual(out, ["/ot/api/77/producto/5", "/ot/api/77/equipo/12", True, False, False, False])

    def test_un_toque_en_el_chip_deja_la_linea_en_la_lista(self):
        js = ("var window = globalThis; var _otrep = {modo: 'repuesto', lista: [], sel: null};\n"
              "var _cant = {value: ''}; var document = {getElementById: function(){ return _cant; },"
              " querySelectorAll: function(){ return []; }};\n"
              "var toasts = []; function ilusToast(m){ toasts.push(m); }\n"
              "function otrepAgregarALaLista(){ _otrep.lista.push(Object.assign({cantidad: Number(_cant.value)}, _otrep.sel));"
              " _otrep.sel = null; }\n"
              + _js_desde("window.otrepUsarSugerido = function(") + "\n"
              "otrepUsarSugerido('Cinta (banda)'); otrepUsarSugerido('cinta (banda)');\n"
              "console.log(JSON.stringify({lista: _otrep.lista, toasts: toasts}));")
        out = _node(js)
        self.assertEqual(len(out["lista"]), 1)
        it = out["lista"][0]
        self.assertEqual((it["nombre"], it["origen"], it["stock_id"], it["cantidad"]), ("Cinta (banda)", "manual", None, 1))
        self.assertTrue(out["toasts"])                                  # el repetido avisa, no duplica


# ═══════════════════════════════ 2) repuestos sugeridos ═══════════════════════════════
class TestRepuestosSugeridos(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.k = _cargar(("_chk_txt_norm", "_rsug_normalizar", "_rsug_a_sembrar", "_rsug_unir"),
                        ("_RSUG_MAX", "_RSUG_LARGO", "_CHK_REPUESTOS_SUGERIDOS_SEED", "_CHK_BLOQUES_SEED",
                         "_CHK_RESULTADO_OPCIONES", "_CHK_BASE_CONTROL_CALIDAD"))

    def test_semilla_de_la_propuesta_en_los_bloques_de_sistema(self):
        seed = self.k["_CHK_REPUESTOS_SUGERIDOS_SEED"]
        bloques = {b[0] for b in self.k["_CHK_BLOQUES_SEED"]}
        self.assertEqual(set(seed), bloques - {self.k["_CHK_BASE_CONTROL_CALIDAD"]})   # 9 bloques; la base no
        self.assertIn("Piola (cable de acero) o correa de transmisión", seed["Revisión — Selector de peso"])
        self.assertIn("Pin selector con cordel", seed["Revisión — Selector de peso"])
        self.assertIn("Cinta (banda)", seed["Revisión — Trotadora (cinta y motor)"])
        self.assertIn("Tabla (deck)", seed["Revisión — Trotadora (cinta y motor)"])
        for nombre, lista in seed.items():
            self.assertTrue(4 <= len(lista) <= 6, nombre)
            self.assertEqual(self.k["_rsug_normalizar"](lista), lista, nombre)   # ya vienen limpias
            for n in lista:                                                      # solo NOMBRES de pieza
                self.assertIsNone(re.search(r"\$|\d{3,}|proveedor|costo|precio", n, re.I), n)

    def test_normalizar(self):
        f = self.k["_rsug_normalizar"]
        self.assertEqual(f(["  Pin   selector ", "pín SELECTOR", "", "x", "Pad de asiento, respaldo o rodillo"]),
                         ["Pin selector", "Pad de asiento, respaldo o rodillo"])   # la coma NO separa
        self.assertEqual(f("Cinta (banda); Tabla (deck)\nMotor"), ["Cinta (banda)", "Tabla (deck)", "Motor"])
        self.assertEqual(f('["Biela", "biela", "Consola"]'), ["Biela", "Consola"])
        self.assertEqual(f(["Cinta $45.000", "Cinta"]), ["Cinta"])                 # nunca un costo
        self.assertEqual(len(f([f"Pieza {i}" for i in range(40)])), self.k["_RSUG_MAX"])
        self.assertEqual(len(f(["a" * 200])[0]), self.k["_RSUG_LARGO"])
        self.assertEqual(f(None), [])
        self.assertEqual(f(42), [])
        self.assertEqual(self.k["_rsug_unir"]([["A1", "B1"], None, ["b1", "C1"]]), ["A1", "B1", "C1"])

    def test_sembrar_solo_lo_vacio_y_nunca_pisar_lo_editado(self):
        filas = [
            {"id": 1, "nombre": "Revisión — Selector de peso", "repuestos_sugeridos_json": None},
            {"id": 2, "nombre": "Revisión — Trotadora (cinta y motor)", "repuestos_sugeridos_json": '["Cinta"]'},
            {"id": 3, "nombre": "Revisión — Rack y soportes", "repuestos_sugeridos_json": "[]"},
            {"id": 4, "nombre": "Revisión — Registro por equipo", "repuestos_sugeridos_json": None},
            {"id": 5, "nombre": "Otra plantilla", "repuestos_sugeridos_json": None},
        ]
        plan = self.k["_rsug_a_sembrar"](filas)
        self.assertEqual([p["id"] for p in plan], [1, 3])
        self.assertIn("Pin selector con cordel", plan[0]["lista"])

    def test_siembra_una_sola_vez_y_barata(self):
        f = _fuente("_ensure_repuestos_sugeridos_seed")
        self.assertIn("_RSUG_MARCA", f)
        self.assertIn("SELECT id FROM mant_logs WHERE entidad=%s AND entidad_id=%s AND accion=%s LIMIT 1", f)
        self.assertIn("(repuestos_sugeridos_json IS NULL OR TRIM(repuestos_sugeridos_json) IN ('','[]'))", f)
        self.assertIn("return out   # sin marca: el próximo arranque reintenta", f)
        self.assertIn("return out   # los bloques todavía no existen", f)        # sin bloques no se marca
        col = _fuente("_ensure_repuestos_sugeridos_col")
        self.assertIn("information_schema.COLUMNS", col)                        # REGLA #18: barata ya aplicada
        self.assertIn("ADD COLUMN repuestos_sugeridos_json TEXT NULL", col)
        codigo, _arb = _arbol()
        self.assertIn("        _ensure_repuestos_sugeridos_col()\n        _rsug_seed = _ensure_repuestos_sugeridos_seed()",
                      codigo.replace("\r\n", "\n"))

    def test_editor_de_plantillas(self):
        self.assertIn('"repuestos_sugeridos": _rsug_normalizar(plant.get("repuestos_sugeridos_json") or ""),',
                      _fuente("mant_plantilla_detalle"))
        for nombre in ("mant_plantilla_crear", "mant_plantilla_actualizar"):
            f = _fuente(nombre)
            self.assertIn('if "repuestos_sugeridos" in d and _rsug_ok():', f, nombre)   # un editor viejo no borra
            self.assertIn("UPDATE mant_tarea_plantillas SET repuestos_sugeridos_json=%s WHERE id=%s", f, nombre)
        self.assertIn('id="ePRsugWrap"', P)
        self.assertIn("function rsugAgregar(){", P)
        self.assertIn("function rsugQuitar(i){", P)
        self.assertIn("repuestos_sugeridos: _rsug.slice(),", P)
        self.assertIn("_rsug = Array.isArray(p.repuestos_sugeridos) ? p.repuestos_sugeridos.slice() : [];", P)
        self.assertIn("Solo el nombre de la pieza: aquí no va el costo.", P)
        self.assertIn("if (_rsugPend && String(_rsugPend.value || '').trim()) rsugAgregar();", P)   # lo escrito no se pierde

    def test_la_ficha_recibe_los_sugeridos_de_cada_bloque(self):
        f = _fuente("mant_visita_checklist_get")
        self.assertIn('_rsug = _rsug_de_plantillas({grp["plantilla_id"] for grp in grupos if grp["plantilla_id"]})', f)
        self.assertIn('grp["repuestos_sugeridos"] = _rsug.get(grp["plantilla_id"]) or []', f)
        for fn in ("_rsug_de_plantillas", "_otrep_sugeridos_de_producto", "_rsug_normalizar"):
            self.assertIsNone(_PLATA.search(_fuente(fn).split('"""', 2)[-1]), fn)   # código sin plata ni proveedor


# ═══════════════════════════════ 3) repuestos e incidencias de un PRODUCTO ═══════════════════════════════
class _BD:
    def __init__(self, filas):
        self.filas = filas
        self.sql = []

    def __call__(self, sql, params=None):
        self.sql.append((sql, params))
        for marca, fila in self.filas:
            if marca in sql:
                return fila
        return None


def _amb_producto(bd):
    errs = []

    def _ot2_err(msg, cod, http=400, **kw):
        errs.append((cod, http))
        return ({"ok": False, "error": msg, "error_codigo": cod}, http)
    amb = _cargar(("_otrep_nombre_producto", "_otrep_cargar_ot_y_producto"), ("_OTPROD_NO_LLEGO_SUFIJO",),
                  amb={"mysql_fetchone": bd, "_ot2_err": _ot2_err})
    return amb, errs


class TestProductoEnElFlujoDeRepuestos(unittest.TestCase):
    def test_candados_del_producto(self):
        bd = _BD([("FROM mant_visitas", {"id": 9, "numero_ot": "OT-2026-00400", "cliente_id": 15, "tipo": "x",
                                         "estado": "en_curso"})])
        amb, errs = _amb_producto(bd)
        v, p, err = amb["_otrep_cargar_ot_y_producto"](9, 1)
        self.assertEqual((v, p), (None, None))
        self.assertEqual(errs[-1], ("OT_CON_CLIENTE", 403))                # OT de cliente: va por la ficha del equipo
        bd = _BD([("FROM mant_visitas", {"id": 9, "numero_ot": "OT-2026-00400", "cliente_id": None,
                                         "tipo": "control_calidad", "estado": "en_curso"})])
        amb, errs = _amb_producto(bd)
        self.assertIsNone(amb["_otrep_cargar_ot_y_producto"](9, 3)[1])
        self.assertEqual(errs[-1], ("PRODUCTO_NO_EN_OT", 403))             # producto de OTRA OT: no
        self.assertIn("maquina_id IS NULL AND item_manual_id=%s", bd.sql[-1][0])
        self.assertEqual(bd.sql[-1][1], (9, 3))
        bd = _BD([("FROM mant_visitas", {"id": 9, "numero_ot": "OT-2026-00400", "cliente_id": None,
                                         "tipo": "control_calidad", "estado": "en_curso"}),
                  ("FROM mant_visita_tareas", {"titulo": "Multi Hip Machine SD45HEP — NO LLEGÓ",
                                               "item_manual_sku": "SD45HEP", "cantidad": 2, "plantilla_id": None})])
        amb, errs = _amb_producto(bd)
        v, p, err = amb["_otrep_cargar_ot_y_producto"](9, 3)
        self.assertIsNone(err)
        self.assertEqual((p["id"], p["nombre"], p["sku"], p["cantidad"], p["interno"]),
                         (3, "Multi Hip Machine SD45HEP", "SD45HEP", 2, True))

    def test_rutas_y_permisos(self):
        for fn in ("ot2_api_producto_repuestos_opciones", "ot2_api_producto_solicitar_repuestos",
                   "ot2_api_producto_evidencia", "ot2_api_producto_incidencia"):
            decs = _decoradores(fn)
            self.assertIn("_mant_required", decs, fn)
            self.assertIn("_ot_can_configurar", decs, fn)                  # técnico asignado o gestión, como el de cliente
            self.assertIn("_otrep_cargar_ot_y_producto(vid, imid)", _fuente(fn), fn)
        for fn in ("ot2_producto_manual_ver", "ot2_producto_manual_legado_ver"):
            self.assertIn("_ot_can_view", _decoradores(fn))
            self.assertIn("interno=True", _fuente(fn))
        codigo = _arbol()[0]
        for ruta in ('"/ot/api/<int:vid>/producto/<int:imid>/repuestos-opciones", methods=["GET"]',
                     '"/ot/api/<int:vid>/producto/<int:imid>/solicitar-repuestos", methods=["POST"]',
                     '"/ot/api/<int:vid>/producto/<int:imid>/evidencia", methods=["POST"]',
                     '"/ot/api/<int:vid>/producto/<int:imid>/incidencia", methods=["POST"]'):
            self.assertIn("@app.route(" + ruta, codigo)

    def test_reutiliza_el_flujo_de_la_ot_sin_tocar_el_de_cliente(self):
        f = _fuente("ot2_api_producto_solicitar_repuestos")
        for pieza in ("_otrep_validar_lineas_lote(lineas_in, stock_por_id)", "_otrep_primer_archivo()",
                      "_otrep_crear_fila_solicitud(\n                    vid, None, None,",
                      "UPDATE mant_ot_repuesto_solicitudes SET item_manual_id=%s WHERE id IN",
                      "_otrep_subir_evidencia_generica(", "_otrep_ticket_para_ot(vid, v, user)",
                      "_otrep_ticket_anotar_producto(", "_otrep_nombres_alt_guardar(", "_otrep_evento(",
                      '_mant_log("visita", vid, "repuesto_solicitado"', "len(motivo) < 10", "EVIDENCIA_REQUERIDA",
                      "_OTPROD_ESTADOS_FIN", "_otrep_contexto_producto(v, p, docs)"):
            self.assertIn(pieza, f, pieza)
        # sin máquina: nada de fuera de servicio de una ficha ni de anotar un maquina_id en el ticket
        for prohibido in ("_ot_equipo_fuera_servicio_marcar", "_otrep_ticket_anotar_lote(", "mant_maquinas"):
            self.assertNotIn(prohibido, f, prohibido)
        self.assertIn("maquina_id IS NULL", _fuente("_otrep_ticket_anotar_producto"))
        # el de cliente sigue igual
        cli = _fuente("ot2_api_equipo_solicitar_repuestos")
        self.assertIn("v, m, err = _otrep_cargar_ot_y_equipo(vid, mid)", cli)
        self.assertIn("_ot_equipo_fuera_servicio_marcar(vid, mid, motivo, v, m, crear_urgencia=False)", cli)
        ev = _fuente("ot2_api_producto_evidencia")
        self.assertIn("AND visita_id=%s AND maquina_id IS NULL", ev)
        self.assertIn('" AND item_manual_id=%s" if _otrep_im_ok() else ""', ev)

    def test_columna_item_manual_barata(self):
        f = _fuente("_ensure_otrep_item_manual_col")
        self.assertIn("information_schema.COLUMNS", f)
        self.assertIn('"ALTER TABLE mant_ot_repuesto_solicitudes ADD COLUMN item_manual_id INT NULL "', f)
        self.assertIn('"ALTER TABLE mant_ot_repuesto_solicitudes ADD INDEX idx_otrep_item_manual (visita_id, item_manual_id)"', f)
        self.assertIn("_ensure_otrep_item_manual_col()", _arbol()[0])

    def test_manuales_con_url_propia_del_producto(self):
        amb = _cargar(("_otrep_manuales_de_modelos",), amb={})
        filas = {"cat_producto_manuales": [{"id": 3, "producto_id": 10, "gcs_key": "k", "nombre_archivo": "M.pdf",
                                            "size_kb": 1, "orden": 1}], "manual_pdf_key": []}

        def fetchall(sql, params=None):
            return filas["cat_producto_manuales"] if "cat_producto_manuales" in sql else []
        amb["_otrep_manuales_de_modelos"].__globals__["mysql_fetchall"] = fetchall
        fn = amb["_otrep_manuales_de_modelos"]
        mod = [{"id": 10, "sku": "SD45HEP", "nombre": "Multi Hip"}]
        self.assertEqual(fn(mod, 5, 7)[0]["url"], "/ot/5/equipo/7/manual/3")          # cliente: igual que antes
        self.assertEqual(fn(mod, 5, 2, base="/ot/5/producto/2")[0]["url"], "/ot/5/producto/2/manual/3")
        self.assertIn("(_otrep_cargar_ot_y_producto(vid, mid) if interno else _otrep_cargar_ot_y_equipo(vid, mid))",
                      _fuente("_otrep_servir_manual"))

    def test_el_modal_usa_las_rutas_del_producto(self):
        abrir = _js_desde("window.otdRepAbrir = async function(")
        self.assertIn("fetch(_otrepBase(mid) + '/repuestos-opciones'", abrir)
        self.assertIn("if (_st4) _st4.style.display = _esProd ? 'none' : '';", abrir)   # sin «fuera de servicio» aquí
        self.assertIn("_otrepBuscador.setModeloBase({ id: j.modelo.id, sku: j.modelo.sku, nombre: j.modelo.nombre }, false)",
                      abrir)
        self.assertIn("otrepPintarSugeridos(j.repuestos_sugeridos || []);", abrir)
        prep = _js_funcion("otrepBuscadorPreparar", D)
        i_prod = prep.index("if (_otrepEsProducto(mid)){")
        self.assertLess(i_prod, prep.index("_otrepBuscador.setEquipo(eq)"))   # nunca un maquina_id con el producto
        self.assertIn("fetch(_otrepBase(mid) + '/solicitar-repuestos'", _js_funcion("otrepEnviarLote", D))
        for n in ("_otrepSubirExtra", "_otrepSubirExtraLote"):
            self.assertIn("fetch(_otrepBase(mid) + '/evidencia'", _js_funcion(n, D), n)
        ref = _js_funcion("otdRepRefrescar", D)
        self.assertIn("(!s.maquina_id && Number(s.item_manual_id) === Number(_midDom))", ref)
        self.assertIn("{% for e in equipos if e._interno %}_OTREP_EQ_NOMBRE['i{{ e.id }}']", D)
        self.assertIn('id="otdRepLista-i{{ e.id }}"', D)

    def test_tecnico_no_ve_proveedores_ni_costos(self):   # REGLA #19/#26
        f = _fuente("ot2_api_producto_repuestos_opciones")
        self.assertIn("_otrep_fmt_stock(r, para_ot=True)", f)              # mismo recorte que el de cliente
        self.assertIn("para_ot=True)", f.split("solicitudes = _otrep_listar(", 1)[1])
        self.assertIn('"repuestos_sugeridos": _otrep_sugeridos_de_producto(vid, imid)', f)
        sol = _fuente("ot2_api_producto_solicitar_repuestos")
        self.assertIn('item.pop("repuesto_stock_id", None)', sol)          # externo: ni stock ni disponible
        self.assertIn('item.pop("stock_disponible", None)', sol)
        self.assertIsNone(re.search(r"proveedor_nombre|costo_unitario", _fuente("_otrep_ticket_anotar_producto")))
        self.assertIsNone(re.search(r"proveedor|costo", _fuente("_otrep_contexto_producto").split('"""', 2)[-1]))
        fila = _fuente("_otrep_fila")                                        # la fila sigue recortando a técnicos
        self.assertIn("if _oculta_proveedores():", fila)
        self.assertIn("if _es_rol_tecnico():", fila)


class TestIncidenciaDesdeElProducto(unittest.TestCase):
    def test_mismo_candado_que_el_modulo_de_incidencias(self):
        """🔒 2026-10-19 (revisión de la Etapa 3): técnico de campo común y externo → 403; técnico ELEVADO solo con el
        permiso «taller» (mant_taller); gestión siempre. Igual que @_no_tecnico_salvo_taller (Daniel 2026-09-23)."""
        f = _fuente("ot2_api_producto_incidencia")
        self.assertIn("if not _puede_informar_incidencia():", f)
        self.assertIn('"SIN_PERMISO_INCIDENCIA", http=403)', f)
        self.assertLess(f.index("_puede_informar_incidencia()"), f.index("INSERT INTO mant_incidencias"))
        self.assertIn("puede_incidencia=_puede_informar_incidencia(),", _fuente("ot2_detalle"))
        self.assertIn("const OTD_PUEDE_INCIDENCIA = {{ 'true' if (puede_incidencia and", D)

        class _G:
            permissions = {}
        g = _G()
        rol = {"r": "tecnico"}

        def familia(r):
            return "tecnico" if (r or "").startswith("tecnico") else (r or "")
        amb = {"g": g,
               "_es_rol_tecnico": lambda: familia(rol["r"]) == "tecnico",
               "_es_tecnico_externo": lambda: rol["r"].startswith("tecnico_externo"),
               "_es_tecnico_elevado": lambda: (familia(rol["r"]) == "tecnico" and rol["r"] != "tecnico"
                                               and not rol["r"].startswith("tecnico_externo"))}
        puede = _exec("_puede_informar_incidencia", amb)
        casos = []
        for r, taller in (("tecnico", True), ("tecnico_externo", True), ("tecnico_ejecutivo", False),
                          ("tecnico_ejecutivo", True), ("supervisor", False), ("admin", False)):
            rol["r"] = r
            g.permissions = {"mant_taller": taller}
            casos.append((r, taller, puede()))
        self.assertEqual(casos, [("tecnico", True, False), ("tecnico_externo", True, False),
                                 ("tecnico_ejecutivo", False, False), ("tecnico_ejecutivo", True, True),
                                 ("supervisor", False, True), ("admin", False, True)])

    def test_mismo_registro_de_incidencias(self):
        f = _fuente("ot2_api_producto_incidencia")
        self.assertIn("INSERT INTO mant_incidencias (sku, descripcion, cantidad, motivo, observacion, created_by)", f)
        self.assertIn("_inc_foto_guardar(iid, a, t_a, user)", f)
        self.assertIn('_inc_log(iid, "creada", "sku", None, p.get("sku") or "(sin SKU)")', f)
        self.assertIn("_otrep_clasificar_archivo(a)", f)
        self.assertIn("n_f > INC_MAX_FOTOS or n_v > INC_MAX_VIDEOS", f)
        self.assertIn('_mant_log("visita", vid, "incidencia_informada"', f)
        self.assertIn("cantidad > max(1, int(p.get(\"cantidad\") or 1))", f)   # tope = cantidad del documento
        self.assertIn("Informada desde la {numero_ot}", f)
        # las columnas existen en el CREATE TABLE de mant_incidencias (REGLA #5)
        codigo = _arbol()[0]
        crear = codigo[codigo.index("CREATE TABLE IF NOT EXISTS mant_incidencias ("):][:1500]
        for col in ("sku ", "descripcion ", "cantidad ", "motivo ", "observacion ", "created_by "):
            self.assertIn(col, crear, col)

    def test_las_fotos_se_guardan_igual_que_en_la_ficha_de_incidencias(self):
        vieja = _fuente("mant_api_incidencia_foto_subir")
        self.assertIn("_ok_f, _r_f = _inc_foto_guardar(iid, f, tipo, current_username())", vieja)
        self.assertIn("tope = INC_MAX_FOTOS if tipo == \"foto\" else INC_MAX_VIDEOS", vieja)   # el tope sigue ahí
        self.assertIn("_no_tecnico_salvo_taller", _decoradores("mant_api_incidencia_foto_subir"))
        llamadas = []
        intentos = {"n": 0}

        def ejecutar(sql, params=None):
            intentos["n"] += 1
            if intentos["n"] == 1:
                raise RuntimeError("Duplicate entry (carrera)")
            llamadas.append(params)
        amb = {"_uploader_upload": lambda f, folder=None, resource_type=None: {"public_id": "incidencias/abc"},
               "mysql_fetchone": lambda sql, p=None: {"prox": 3}, "mysql_execute": ejecutar,
               "_inc_log": lambda *a, **k: llamadas.append(("log",) + a)}
        fn = _exec("_inc_foto_guardar", amb)
        self.assertEqual(fn(7, object(), "foto", "dave"), (True, "incidencias/abc"))
        self.assertEqual(llamadas[0], (7, "incidencias/abc", 4, "foto", "dave"))   # reintento con el orden siguiente
        self.assertEqual(llamadas[1][:3], ("log", 7, "foto_agregada"))
        amb["_uploader_upload"] = lambda *a, **k: {}
        self.assertEqual(fn(7, object(), "foto", "dave"), (False, "Subida sin resultado válido."))

    def test_el_modal_de_la_ficha(self):
        self.assertIn('id="modalIncProd"', D)
        self.assertIn("modal-fullscreen-sm-down", D)
        abrir = _js_desde("window.otdIncAbrir = function(")
        self.assertIn("c.max = String(_otinc.max);", abrir)
        env = _js_funcion("otincEnviar", D)
        self.assertIn("'/ot/api/' + OTD_VID + '/producto/' + _otinc.imid + '/incidencia'", env)
        self.assertIn("fd.append('fotos', a.file", env)
        self.assertIn("{% for e in equipos if e._interno %}_OTINC_PROD[{{ e.id }}]", D)


# ═══════════════════════════════ 4) el acta ═══════════════════════════════
OPC = json.dumps(["Apto", "Apto con observaciones", "Reparar", "Dar de baja", "Fuera de servicio", "No llegó",
                  "Pendiente de diagnóstico"])
T0 = _dt.datetime(2026, 10, 18, 15, 0, 0)


def _tarea(i, im, tr, titulo, vj=None, pid=7, comp_por=None, comp_at=None, cant=1, sku=None, desc=None):
    return {"id": i, "maquina_id": None, "item_manual_id": im, "item_manual_sku": sku, "plantilla_id": pid,
            "titulo": titulo, "descripcion": desc, "cantidad": cant, "tipo_respuesta": tr,
            "opciones_lista_json": OPC if tr == "lista" else None,
            "valor_json": json.dumps(vj) if vj is not None else None,
            "completada_por": comp_por, "completada_at": comp_at}


def _revision():
    t = []
    # 1) Apto con foto, revisado por Lenin
    t += [_tarea(1, 1, "check", "Multi Hip Machine SD45HEP", pid=None, cant=2, sku="SD45HEP"),
          _tarea(2, 1, "foto", "Foto general del equipo", comp_por="lenin", comp_at=T0),
          _tarea(3, 1, "sino", "¿Faltan piezas o accesorios?", vj={"valor": "no"}, comp_por="lenin", comp_at=T0),
          _tarea(4, 1, "lista", "Resultado del control", vj={"opcion": "Apto"}, comp_por="lenin",
                 comp_at=T0 + _dt.timedelta(minutes=5)),
          _tarea(5, 1, "texto", "Observación del equipo", vj={"texto": "=cmd|' /C calc'!A0"}, comp_por="lenin",
                 comp_at=T0)]
    # 2) No llegó (Resultado)
    t += [_tarea(10, 2, "check", "Banco plano ajustable FID", pid=None, cant=1, sku="FID01"),
          _tarea(11, 2, "lista", "Resultado del control", vj={"opcion": "No llegó"}, comp_por="dave",
                 comp_at=T0 + _dt.timedelta(minutes=30)),
          _tarea(12, 2, "texto", "Observación del equipo", vj={"texto": "Se esperaba el 15/10; bodega no lo tiene"})]
    # 3) No llegó declarado al crear la OT (sin checklist)
    t += [_tarea(20, 3, "check", "Rack de mancuernas 10 pares — NO LLEGÓ", pid=None, cant=1, sku="RK10",
                 desc="Faltante declarado al crear la OT por jaizer el 18/10/2026 15:00.")]
    # 4) Reparar con faltantes y piezas contadas
    t += [_tarea(30, 4, "check", "Set mancuernas 2,5 a 25 kg", pid=None, cant=20, sku="MANC"),
          _tarea(31, 4, "sino", "¿Faltan piezas o accesorios?", vj={"valor": "si"}),
          _tarea(32, 4, "numero", "Cantidad contada (piezas)", vj={"numero": 18}),
          _tarea(33, 4, "lista", "Resultado del control", vj={"opcion": "Reparar"})]
    # 5) sin resultado todavía
    t += [_tarea(40, 5, "check", "Trotadora", pid=None, cant=1, sku="TR1"),
          _tarea(41, 5, "lista", "Resultado del control")]
    return t


class TestActaPura(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.k = _cargar(("_tarea_vj", "_chk_txt_norm", "_chk_opciones", "_chk_es_resultado", "_otrep_nombre_producto",
                         "_acta_resultado_canon", "_acta_armar", "_acta_correos_normalizar"),
                        ("_CHK_RES_APTO", "_CHK_RES_APTO_OBS", "_OTPROD_NO_LLEGO_SUFIJO", "_ACTA_RESULTADOS",
                         "_ACTA_SEMAFORO", "_ACTA_DOMINIOS_INTERNOS", "_ACTA_RX_CORREO"))

    def test_producto_por_producto(self):
        a = self.k["_acta_armar"](_revision(), {2: ["/f/fotos/general.jpg"]})
        self.assertEqual(a["n"], 5)
        pr = {p["item_manual_id"]: p for p in a["productos"]}
        self.assertEqual((pr[1]["nombre"], pr[1]["sku"], pr[1]["cantidad"], pr[1]["resultado"], pr[1]["semaforo"]),
                         ("Multi Hip Machine SD45HEP", "SD45HEP", 2, "Apto", "verde"))
        self.assertEqual(pr[1]["foto"], "/f/fotos/general.jpg")
        self.assertEqual(pr[1]["faltantes"], "Completo")
        self.assertEqual((pr[1]["revisado_por"], pr[1]["revisado_en"]), ("lenin", T0 + _dt.timedelta(minutes=5)))
        self.assertEqual((pr[2]["resultado"], pr[2]["semaforo"]), ("No llegó", "gris"))
        self.assertEqual(pr[2]["faltantes"], "No llegó (1 de 1)")
        self.assertEqual(pr[2]["observacion"], "Se esperaba el 15/10; bodega no lo tiene")
        self.assertEqual((pr[3]["nombre"], pr[3]["resultado"]), ("Rack de mancuernas 10 pares", "No llegó"))
        self.assertIn("Faltante declarado al crear la OT", pr[3]["observacion"])
        self.assertEqual(pr[4]["resultado"], "Reparar")
        self.assertEqual(pr[4]["faltantes"], "Faltan piezas o accesorios (detalle en la observación) · Contadas 18 pieza(s)")
        self.assertEqual((pr[5]["resultado"], pr[5]["foto"]), ("Sin resultado", None))
        self.assertEqual([(t["resultado"], t["n"]) for t in a["totales"]],
                         [("Apto", 1), ("Reparar", 1), ("No llegó", 2), ("Sin resultado", 1)])
        self.assertEqual(a["revisores"], ["lenin", "dave"])
        self.assertEqual(a["ultima_revision"], T0 + _dt.timedelta(minutes=30))

    def test_sin_plata_ni_proveedores(self):
        a = self.k["_acta_armar"](_revision(), {})
        for p in a["productos"]:
            self.assertEqual(set(p), {"item_manual_id", "sku", "nombre", "cantidad", "resultado", "semaforo",
                                      "observacion", "faltantes", "foto", "revisado_por", "revisado_en"})
        for fn in ("_acta_armar", "_ot_acta_datos", "_acta_xlsx_bytes", "_acta_html"):
            self.assertIsNone(_PLATA.search(_fuente(fn).split('"""', 2)[-1]), fn)
        self.assertIn("maquina_id IS NULL AND item_manual_id IS NOT NULL", _fuente("_ot_acta_datos"))

    def test_resultado_de_la_version_larga(self):
        c = self.k["_acta_resultado_canon"]
        self.assertEqual(c("No apto - requiere reparación"), "Reparar")
        self.assertEqual(c("  no LLEGO "), "No llegó")
        self.assertEqual(c("Apto con observaciones"), "Apto con observaciones")
        self.assertEqual(c(""), "")

    def test_correos_solo_internos(self):
        f = self.k["_acta_correos_normalizar"]
        ok, mal = f("Logistica@SPHS.cl, juan@ilusfitness.com; logistica@sphs.cl\ncliente@gmail.com "
                    "x@mail.sphs.cl x@sphs.cl.com x@evilsphs.cl <bodega@sphs.cl> no-es-correo x..y@sphs.cl")
        self.assertEqual(ok, ["logistica@sphs.cl", "juan@ilusfitness.com", "bodega@sphs.cl"])
        self.assertEqual(mal, ["cliente@gmail.com", "x@mail.sphs.cl", "x@sphs.cl.com", "x@evilsphs.cl",
                               "no-es-correo", "x..y@sphs.cl"])
        self.assertEqual(f(["a@sphs.cl", "b@ilusfitness.com"])[0], ["a@sphs.cl", "b@ilusfitness.com"])
        self.assertEqual(f(""), ([], []))


class TestActaArchivos(unittest.TestCase):
    def test_pdf_plantilla_sin_datos_internos(self):
        env = jinja2.Environment(autoescape=True)
        tpl = env.from_string(ACTA_TPL)
        datos = {"numero_ot": "OT-2026-00400", "tipo_nombre": "Revisión de equipos", "documentos": ["NVI 687"],
                 "revisores": ["lenin"], "n": 1, "ot": {"estado": "en_curso", "tecnico_nombre": "Lenin"},
                 "totales": [{"resultado": "Apto", "n": 1, "semaforo": "verde"}]}
        prods = [{"nombre": "Selectorized Pendulum Kick Back (Glute Builder con stack de 100 kg)", "sku": "SPKB",
                  "cantidad": 1, "resultado": "Apto", "semaforo": "verde", "observacion": "sin novedad",
                  "faltantes": "Completo", "foto": "data:image/jpeg;base64,AAAA", "revisado_por": "lenin",
                  "revisado_en_txt": "18/10/2026 12:00"}]
        html = tpl.render(d=datos, productos=prods, marca="ILUS Fitness", generado_txt="18/10/2026 12:30",
                          ultima_txt="18/10/2026 12:00")
        self.assertIn("Selectorized Pendulum Kick Back (Glute Builder con stack de 100 kg)", html)   # completo
        self.assertIn("NVI 687", html)
        self.assertIn("ILUS <span>Fitness</span>", html)                   # REGLA #0
        self.assertIn('class="chip verde"><b>1</b> Apto', html)
        self.assertIn("Recibe · Logística", html)
        self.assertNotIn("$", html)
        expr = str(re.findall(r"\{\{(.*?)\}\}|\{%(.*?)%\}", ACTA_TPL, flags=re.S))   # solo lo que se imprime
        self.assertIsNone(re.search(r"monto|costo|precio|proveedor|saldo|zz_|margen", expr, re.I))

    def test_excel_con_texto_peligroso_como_texto(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl no está instalado")
        k = _cargar(("_tarea_vj", "_chk_txt_norm", "_chk_opciones", "_chk_es_resultado", "_otrep_nombre_producto",
                     "_acta_resultado_canon", "_acta_armar", "_acta_xl_celda", "_acta_xlsx_bytes", "_acta_fecha"),
                    ("_CHK_RES_APTO", "_CHK_RES_APTO_OBS", "_OTPROD_NO_LLEGO_SUFIJO", "_ACTA_RESULTADOS",
                     "_ACTA_SEMAFORO"),
                    amb={"ILUS_BRAND": "ILUS Fitness", "chile_fmt_filter": lambda d, f="": d.strftime("%d/%m/%Y %H:%M")})
        datos = k["_acta_armar"](_revision(), {})
        datos.update({"numero_ot": "OT-2026-00400", "tipo_nombre": "Revisión de equipos", "documentos": ["NVI 687"],
                      "ot": {}, "generado_en": T0})
        wb = openpyxl.load_workbook(io.BytesIO(k["_acta_xlsx_bytes"](datos)))
        ws = wb.active
        valores = [[c.value for c in fila] for fila in ws.iter_rows()]
        planos = [v for fila in valores for v in fila if v is not None]
        self.assertIn("ILUS Fitness · Acta de recepción de la revisión", planos)
        self.assertIn("NVI 687", planos)
        cab = next(f for f in valores if f and f[0] == "N°")
        self.assertEqual(cab[:9], ["N°", "SKU", "Producto", "Cant. documento", "Resultado", "Observación", "Faltantes",
                                   "Revisó", "Fecha de revisión"])
        obs = next(c for fila in ws.iter_rows() for c in fila if c.value == "=cmd|' /C calc'!A0")
        self.assertEqual(obs.data_type, "s")                               # texto, nunca fórmula
        self.assertIsNone(re.search(r"monto|costo|proveedor", json.dumps(planos, default=str), re.I))

    def test_rutas_del_acta_solo_gestion(self):
        for fn in ("ot2_acta_recepcion_pdf", "ot2_acta_recepcion_xlsx", "ot2_api_acta_recepcion_enviar"):
            decs = _decoradores(fn)
            self.assertIn("_mant_required", decs, fn)
            self.assertIn("_no_tecnico", decs, fn)                         # nunca un técnico
            self.assertIn("_ot_can_view", decs, fn)
        self.assertIn("if not _es_rol_tecnico() and _ot_acta_aplica(v):", _fuente("ot2_detalle"))
        self.assertIn("{% if acta and not es_tecnico %}", D)
        self.assertIn('href="/ot/{{ v.id }}/acta-recepcion.pdf"', D)
        self.assertIn('href="/ot/{{ v.id }}/acta-recepcion.xlsx"', D)
        aplica = _fuente("_ot_acta_aplica")
        self.assertIn('if not v or v.get("cliente_id"):', aplica)            # nunca de una OT de cliente
        self.assertIn("_chk_tipos_herencia()", aplica)
        pdf = _fuente("_acta_pdf_bytes")
        self.assertIn("_pw_pdf(", pdf)                                       # el motor de PDF de siempre
        self.assertIn("except PDFEngineUnavailable:", pdf)


class _Req:
    def __init__(self, body):
        self._b = body

    def get_json(self, silent=False):
        return self._b


def _amb_envio(extra, config=("logistica@sphs.cl",), email_on=True, envio_ok=True):
    reg = {"enviados": [], "logs": [], "errs": []}

    def _ot2_err(msg, cod, http=400, **kw):
        reg["errs"].append(cod)
        return ({"ok": False, "error": msg, "error_codigo": cod}, http)

    def _send(to, subject, html, **kw):
        reg["enviados"].append((to, kw.get("modulo"), kw.get("ref"), [a[0] for a in kw.get("attachments") or []]))
        return envio_ok
    k = _cargar(("_acta_correos_normalizar",), ("_ACTA_DOMINIOS_INTERNOS", "_ACTA_RX_CORREO", "_ACTA_MAX_DEST"))
    amb = dict(k)
    amb.update({
        "request": _Req({"extra": extra}),
        "_ot_acta_datos": lambda vid: {"numero_ot": "OT-2026-00400", "tipo_nombre": "Revisión de equipos",
                                       "documentos": ["NVI 687"], "n": 2, "revisores": ["lenin"],
                                       "totales": [{"resultado": "Apto", "n": 2, "semaforo": "verde"}]},
        "_ot2_err": _ot2_err, "_acta_destinatarios_config": lambda: list(config),
        "comm_is_enabled": lambda canal: email_on, "current_username": lambda: "daniel",
        "_acta_pdf_bytes": lambda datos: (b"%PDF", "<html>"), "_acta_xlsx_bytes": lambda datos: b"PK",
        "_acta_nombre_archivo": lambda datos, ext: "Acta." + ext, "_ilus_email_master": lambda ctx: "<html/>",
        "_brand_subject": lambda t: "ILUS · " + t, "_ilus_email_rows": lambda pares: "",
        "_public_base_url": lambda: "https://x", "_send_ilus_email": _send,
        "_mant_log": lambda *a: reg["logs"].append(a), "_now_chile": lambda: T0,
        "jsonify": lambda d=None, **kw: (d if d is not None else kw),
    })
    fn = _exec("ot2_api_acta_recepcion_enviar", amb)
    return fn, reg


class TestEnvioALogistica(unittest.TestCase):
    def test_un_correo_externo_frena_todo_antes_de_enviar(self):
        fn, reg = _amb_envio("cliente@gimnasio.cl")
        r = fn(400)
        self.assertEqual(r[1], 400)
        self.assertEqual(reg["errs"], ["CORREO_NO_INTERNO"])
        self.assertEqual(reg["enviados"], [])                               # REGLA #22: nada sale
        self.assertEqual(reg["logs"], [])

    def test_envia_solo_a_internos_con_adjuntos_y_registro(self):
        fn, reg = _amb_envio("bodega@ilusfitness.com")
        r = fn(400)
        self.assertTrue(r["ok"])
        self.assertEqual([e[0] for e in reg["enviados"]], ["logistica@sphs.cl", "bodega@ilusfitness.com"])
        for to, modulo, ref, adj in reg["enviados"]:
            self.assertEqual(modulo, "mantenciones")                         # llave del módulo + historial
            self.assertEqual(ref, {"tipo": "OT", "id": 400, "codigo": "OT-2026-00400"})
            self.assertEqual(adj, ["Acta.pdf", "Acta.xlsx"])
        self.assertEqual(len(reg["logs"]), 1)
        ent, vid, accion, detalle = reg["logs"][0]
        self.assertEqual((ent, vid, accion), ("visita", 400, "acta_recepcion_enviada"))
        self.assertIn("daniel", detalle)                                    # quién
        self.assertIn("18/10/2026", detalle)                                # cuándo (hora Chile)
        self.assertIn("logistica@sphs.cl, bodega@ilusfitness.com", detalle)  # a quién

    def test_sin_destinatarios_o_con_el_correo_apagado_no_sale_nada(self):
        fn, reg = _amb_envio("", config=())
        self.assertEqual(fn(400)[1], 400)
        self.assertEqual(reg["errs"], ["SIN_DESTINATARIOS"])
        fn, reg = _amb_envio("", email_on=False)
        self.assertEqual(fn(400)[1], 409)
        self.assertEqual(reg["errs"], ["CORREO_APAGADO"])
        self.assertEqual(reg["enviados"], [])

    def test_si_no_sale_ninguno_lo_dice(self):
        fn, reg = _amb_envio("", envio_ok=False)
        r = fn(400)
        self.assertEqual(r[1], 502)
        self.assertEqual(r[0]["error_codigo"], "NO_SALIO")
        self.assertIn("no salió a: logistica@sphs.cl", reg["logs"][0][3])

    def test_configuracion_solo_acepta_correos_internos(self):
        g = _fuente("mant_reglas_guardar")
        self.assertIn("if str(_cl_acta).strip().lower() == _ACTA_EMAILS_REGLA:", g)
        self.assertIn('"error_codigo": "CORREO_NO_INTERNO"', g)
        self.assertIn('"error_codigo": "LISTA_DEMASIADO_LARGA"', g)          # valor es VARCHAR(500)
        self.assertLess(g.index("_ACTA_EMAILS_REGLA"), g.index("for clave, valor in cambios.items():"))   # antes de guardar
        regla = _fuente("_ensure_acta_recepcion_regla")
        self.assertIn("INSERT IGNORE INTO mant_reglas_negocio", regla)
        self.assertIn('(_ACTA_EMAILS_REGLA, "", "string", "acta", "Correos del acta de recepción",', regla)   # vacía
        self.assertIn("acta:       {icon:'📦', label:'Acta de recepción (logística)'},", CONF)
        self.assertIn("_ensure_acta_recepcion_regla()", _arbol()[0])

    def test_la_ficha_pregunta_con_ilus_prompt_y_valida_el_dominio(self):
        env = _js_funcion("otdActaEnviar", D)
        self.assertIn("await ilusPrompt({", env)
        self.assertIn("_OTD_ACTA_DOMINIO.test(otro)", env)
        self.assertIn("/ot/api/${OTD_VID}/acta-recepcion/enviar", env)
        self.assertIn(r"const _OTD_ACTA_DOMINIO = /^[^@\s]+@(sphs\.cl|ilusfitness\.com)$/i;", D)


if __name__ == "__main__":
    unittest.main()
