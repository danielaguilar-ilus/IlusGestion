# -*- coding: utf-8 -*-
"""2026-10-29 · `xlsx_fechas.py`: las fechas de los Excel de OT y costos son FECHAS reales, en hora Chile y filtrables.

Daniel: «Cuando bajas las fechas en el reporte de la OT, la fecha viene con formato gringo y necesito que en el Excel, en
español, pueda filtrar por meses; que me salgan los meses, los días…». Módulo puro: sin BD ni Flask.

Correr con:  python -m unittest tests.test_xlsx_fechas
"""
import io
import unittest
from datetime import date, datetime, timedelta, timezone

import xlsx_fechas as XF

try:
    import openpyxl
    _HAY_OPENPYXL = True
except ImportError:                                   # pragma: no cover
    _HAY_OPENPYXL = False


class TestUtcAChile(unittest.TestCase):
    """REGLA #6: MySQL guarda en UTC; el Excel muestra hora Chile, también en los días del cambio de horario."""

    def test_un_datetime_sin_zona_es_utc_y_pasa_a_chile(self):
        self.assertEqual(XF.a_fecha_excel(datetime(2026, 10, 20, 15, 0)), datetime(2026, 10, 20, 12, 0))   # verano: UTC-3
        self.assertEqual(XF.a_fecha_excel(datetime(2026, 7, 1, 15, 0)), datetime(2026, 7, 1, 11, 0))       # invierno: UTC-4

    def test_medianoche_utc_cae_el_dia_anterior_en_chile(self):
        v = XF.a_fecha_excel(datetime(2026, 10, 20, 2, 0))
        self.assertEqual(v, datetime(2026, 10, 19, 23, 0))
        # y las columnas de apoyo salen del día CHILENO (lunes 19), no del UTC (martes 20)
        self.assertEqual(XF.columnas_apoyo(v), [2026, "Octubre", 10, 43, "Lunes", "2026-10"])

    def test_fin_de_mes_y_de_ano_en_utc_es_el_mes_anterior_en_chile(self):
        self.assertEqual(XF.a_fecha_excel(datetime(2026, 1, 1, 2, 59)), datetime(2025, 12, 31, 23, 59))
        self.assertEqual(XF.columnas_apoyo(XF.a_fecha_excel(datetime(2026, 1, 1, 2, 59)))[5], "2025-12")
        self.assertEqual(XF.a_fecha_excel(datetime(2026, 1, 1, 3, 0)), datetime(2026, 1, 1, 0, 0))

    def test_fin_del_horario_de_verano_abril(self):
        # 05/04/2026: Chile vuelve de UTC-3 a UTC-4 a la medianoche → dos instantes UTC distintos dan la misma hora local
        self.assertEqual(XF.a_fecha_excel(datetime(2026, 4, 5, 2, 30)), datetime(2026, 4, 4, 23, 30))
        self.assertEqual(XF.a_fecha_excel(datetime(2026, 4, 5, 3, 30)), datetime(2026, 4, 4, 23, 30))

    def test_inicio_del_horario_de_verano_septiembre(self):
        # 06/09/2026: Chile salta de UTC-4 a UTC-3 (no existe la hora 00:xx local)
        self.assertEqual(XF.a_fecha_excel(datetime(2026, 9, 6, 3, 30)), datetime(2026, 9, 5, 23, 30))
        self.assertEqual(XF.a_fecha_excel(datetime(2026, 9, 6, 4, 30)), datetime(2026, 9, 6, 1, 30))

    def test_con_zona_propia_siempre_se_convierte(self):
        memphis = timezone(timedelta(hours=-5))
        self.assertEqual(XF.a_fecha_excel(datetime(2026, 10, 20, 14, 0, tzinfo=memphis)), datetime(2026, 10, 20, 16, 0))
        self.assertEqual(XF.a_fecha_excel(datetime(2026, 10, 20, 14, 0, tzinfo=memphis), utc=False), datetime(2026, 10, 20, 16, 0))

    def test_utc_false_deja_la_hora_local_tal_cual(self):
        self.assertEqual(XF.a_fecha_excel(datetime(2026, 10, 20, 9, 30), utc=False), datetime(2026, 10, 20, 9, 30))

    def test_el_resultado_nunca_lleva_zona_ni_microsegundos(self):
        v = XF.a_fecha_excel(datetime(2026, 10, 20, 15, 0, 7, 999999))
        self.assertIsNone(v.tzinfo)
        self.assertEqual(v.microsecond, 0)

    def test_una_fecha_pura_no_se_convierte(self):
        self.assertEqual(XF.a_fecha_excel(date(2026, 10, 20)), date(2026, 10, 20))
        self.assertNotIsInstance(XF.a_fecha_excel(date(2026, 10, 20)), datetime)


class TestTextos(unittest.TestCase):
    def test_texto_dia_mes_ano_se_lee_como_chileno(self):
        self.assertEqual(XF.a_fecha_excel("07/10/2026"), date(2026, 10, 7))       # 7 de octubre, nunca 10 de julio
        self.assertEqual(XF.a_fecha_excel("08/10/2026 21:24"), datetime(2026, 10, 8, 21, 24))   # ya está en hora Chile
        self.assertEqual(XF.a_fecha_excel("08-10-2026 21:24:05"), datetime(2026, 10, 8, 21, 24, 5))
        self.assertEqual(XF.a_fecha_excel(" 1/2/2026 "), date(2026, 2, 1))

    def test_texto_iso(self):
        self.assertEqual(XF.a_fecha_excel("2026-10-07"), date(2026, 10, 7))
        self.assertEqual(XF.a_fecha_excel("2026-10-07 10:15:00"), datetime(2026, 10, 7, 10, 15))      # sin zona: tal cual
        self.assertEqual(XF.a_fecha_excel("2026-07-29T21:00:00-04:00"), datetime(2026, 7, 29, 21, 0))
        self.assertEqual(XF.a_fecha_excel("2026-10-20T15:00:00Z"), datetime(2026, 10, 20, 12, 0))

    def test_none_y_vacio_son_celda_vacia(self):
        for v in (None, "", "   "):
            self.assertIsNone(XF.a_fecha_excel(v), repr(v))
        self.assertEqual(XF.columnas_apoyo(None), [None] * 6)

    def test_texto_raro_queda_como_el_texto_original(self):
        for raro in ("Sin fecha", "07/10/2026 → 15/10/2026", "31/02/2026", "2026-13-45", "0000-00-00", "mañana", "10/2026",
                     "=HYPERLINK(1)"):
            self.assertEqual(XF.a_fecha_excel(raro), raro)
            self.assertEqual(XF.columnas_apoyo(XF.a_fecha_excel(raro)), [None] * 6)

    def test_otros_tipos_no_revientan(self):
        self.assertEqual(XF.a_fecha_excel(12), 12)
        self.assertEqual(XF.a_fecha_excel(True), "True")
        self.assertEqual(XF.a_fecha_excel(date(1800, 1, 1)), "1800-01-01")      # Excel no tiene fechas antes de 1900
        self.assertEqual(XF.a_fecha_excel(object.__new__(type("X", (), {"__str__": lambda s: "x"}))), "x")


class TestColumnasDeApoyo(unittest.TestCase):
    def test_mes_y_dia_en_espanol(self):
        self.assertEqual(XF.columnas_apoyo(date(2026, 10, 29)), [2026, "Octubre", 10, 44, "Jueves", "2026-10"])
        self.assertEqual(XF.columnas_apoyo(datetime(2026, 3, 1, 8, 0)), [2026, "Marzo", 3, 9, "Domingo", "2026-03"])
        self.assertEqual([XF.MESES_ES[m - 1] for m in (1, 9, 12)], ["Enero", "Septiembre", "Diciembre"])
        self.assertEqual(XF.DIAS_ES, ("Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"))

    def test_semana_iso_en_el_cambio_de_ano(self):
        self.assertEqual(XF.columnas_apoyo(date(2026, 12, 31))[3], 53)
        self.assertEqual(XF.columnas_apoyo(date(2027, 1, 1))[:4], [2027, "Enero", 1, 53])    # año calendario 2027, semana ISO 53
        self.assertEqual(XF.columnas_apoyo(date(2025, 12, 29))[:4], [2025, "Diciembre", 12, 1])
        self.assertEqual(XF.columnas_apoyo(date(2027, 1, 4))[3], 1)

    def test_ano_mes_es_texto_ordenable(self):
        meses = [XF.columnas_apoyo(date(2026, m, 1))[5] for m in (11, 2, 10, 1)]
        self.assertEqual(sorted(meses), ["2026-01", "2026-02", "2026-10", "2026-11"])
        self.assertTrue(all(isinstance(m, str) for m in meses))

    def test_encabezados_dicen_de_que_fecha_salen(self):
        self.assertEqual(XF.encabezados_apoyo(), ["Año", "Mes", "N° mes", "Semana", "Día", "Año-Mes"])
        self.assertEqual(XF.encabezados_apoyo("de la fecha de cierre")[1], "Mes (de la fecha de cierre)")
        self.assertEqual(XF.encabezados_apoyo("del lunes de la semana")[3], "Semana (del lunes de la semana)")
        self.assertEqual(len(XF.ANCHOS_APOYO), len(XF.APOYO))

    def test_formato_espanol(self):
        self.assertEqual(XF.formato(datetime(2026, 1, 1, 10, 0)), "dd/mm/yyyy hh:mm")
        self.assertEqual(XF.formato(date(2026, 1, 1)), "dd/mm/yyyy")
        self.assertIsNone(XF.formato("01/01/2026"))
        self.assertIsNone(XF.formato(None))


@unittest.skipIf(not _HAY_OPENPYXL, "openpyxl no está instalado")
class TestEscribirEnExcel(unittest.TestCase):
    def _libro(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        XF.escribir_fecha(ws.cell(row=1, column=1), datetime(2026, 10, 20, 2, 0))       # UTC → 19/10 23:00 Chile
        XF.escribir_fecha(ws.cell(row=1, column=2), date(2026, 10, 20))
        XF.escribir_fecha(ws.cell(row=1, column=3), "Sin fecha")
        XF.escribir_fecha(ws.cell(row=1, column=4), None)
        XF.escribir_fecha(ws.cell(row=1, column=5), "=1+1")
        XF.escribir_apoyo(ws, 1, 6, XF.a_fecha_excel(datetime(2026, 10, 20, 2, 0)))
        ws.append([date(2026, 1, 5)])
        XF.formatear(ws.cell(row=2, column=1))
        buf = io.BytesIO()
        wb.save(buf)
        return openpyxl.load_workbook(io.BytesIO(buf.getvalue())).active

    def test_la_celda_es_fecha_de_excel_con_formato_espanol_y_vuelve_igual(self):
        ws = self._libro()
        self.assertIsInstance(ws["A1"].value, datetime)
        self.assertEqual(ws["A1"].value, datetime(2026, 10, 19, 23, 0))
        self.assertEqual(ws["A1"].number_format, "dd/mm/yyyy hh:mm")
        self.assertTrue(ws["A1"].is_date)
        self.assertEqual(ws["B1"].number_format, "dd/mm/yyyy")
        self.assertTrue(ws["B1"].is_date)
        self.assertEqual(ws["A2"].number_format, "dd/mm/yyyy")

    def test_texto_raro_y_vacio_no_rompen_y_el_igual_queda_como_texto(self):
        ws = self._libro()
        self.assertEqual(ws["C1"].value, "Sin fecha")
        self.assertIsNone(ws["D1"].value)
        self.assertEqual(ws["E1"].value, "=1+1")
        self.assertEqual(ws["E1"].data_type, "s")

    def test_columnas_de_apoyo_escritas(self):
        ws = self._libro()
        self.assertEqual([ws.cell(row=1, column=c).value for c in range(6, 12)], [2026, "Octubre", 10, 43, "Lunes", "2026-10"])


if __name__ == "__main__":
    unittest.main()
