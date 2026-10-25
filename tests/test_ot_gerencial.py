# -*- coding: utf-8 -*-
"""📊 Costeo de órdenes de trabajo (Daniel, 2026-10-25) — el tablero gerencial de Servicio Técnico.

Sin base de datos ni Flask. Usa la cuenta REAL de la OT (`_ot_finanzas`, `_ot_fin_agregado`) sacada de app.py por texto y
datos de MENTIRA (REGLA #22). Fija: los períodos en hora Chile (también en el cambio de hora), las cinco tarjetas
(Cobré · Me cobraron · Gané · Perdí · Neto), que Neto = Gané − Perdí − coberturas, la atribución SEPARADA (originó ·
autorizó · lleva · ejecutó), la reincidencia de externos, que todo sea solo lectura y solo superadmin, y la plantilla.

Correr con:  py -m pytest tests/test_ot_gerencial.py -q --tb=short
"""
import functools
import io
import os
import re
import sys
import time
import types
import unittest
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import ot_gerencial as G  # noqa: E402
from tests import _ger_app as A  # noqa: E402
from tests import _ger_demo as D  # noqa: E402

TZ = ZoneInfo("America/Santiago")
HOY = D.HOY


def utc(y, m, d, h=12, mi=0, s=0):
    return datetime(y, m, d, h, mi, s)


def agregar():
    return A.cuenta_real()["_ot_fin_agregado"]


def filas_de(visitas, ctxs=None):
    out = []
    for v in visitas:
        ctx = dict((ctxs or {}).get(v["id"], {}))
        ctx.setdefault("nombres", D.USUARIOS)
        out.append(G.armar_fila(v, A.fin_de(v), ctx))
    return out


def en_semana(dia=0, h=15, **kw):
    """Una OT cerrada la semana del 19 al 25 de octubre de 2026 (hora Chile), `dia` 0 = lunes."""
    f = date(2026, 10, 19) + timedelta(days=dia)
    return dict(cerrada_at=datetime(f.year, f.month, f.day, h + 3, 0), **kw)


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestPeriodos(unittest.TestCase):
    def test_la_semana_va_de_lunes_a_domingo_y_el_anterior_es_la_semana_previa(self):
        p = G.rango_periodo(date(2026, 10, 23))
        self.assertEqual((p["modo"], p["desde"], p["hasta"]), ("semana", date(2026, 10, 19), date(2026, 10, 25)))
        self.assertEqual((p["previo"]["desde"], p["previo"]["hasta"]), (date(2026, 10, 12), date(2026, 10, 18)))
        self.assertEqual(p["etiqueta"], "Semana del 19 al 25 de octubre de 2026")
        self.assertTrue(p["en_curso"])
        self.assertEqual(p["transcurridos"], 5)

    def test_una_semana_que_cruza_de_mes_se_nombra_con_los_dos_meses(self):
        p = G.rango_periodo(date(2026, 11, 3))
        self.assertEqual((p["desde"], p["hasta"]), (date(2026, 11, 2), date(2026, 11, 8)))
        p = G.rango_periodo(date(2026, 10, 30))
        self.assertIn("26 de oct al 1 de nov", p["etiqueta"])

    def test_los_bordes_de_la_semana_son_en_hora_de_chile_no_en_utc(self):
        p = G.rango_periodo(date(2026, 10, 23))
        # octubre: Chile en horario de verano (UTC-3) → el lunes empieza a las 03:00 UTC
        self.assertEqual(p["ini_utc"], datetime(2026, 10, 19, 3, 0))
        self.assertEqual(p["fin_utc"], datetime(2026, 10, 26, 3, 0))
        ultimo_instante = datetime(2026, 10, 26, 2, 59, 59)
        self.assertTrue(p["ini_utc"] <= ultimo_instante < p["fin_utc"], "las 23:59:59 del domingo (Chile) son de esta semana")
        self.assertFalse(p["ini_utc"] <= datetime(2026, 10, 26, 3, 0) < p["fin_utc"], "las 00:00 del lunes (Chile) ya son de la siguiente")
        self.assertFalse(p["ini_utc"] <= datetime(2026, 10, 19, 2, 59, 59) < p["fin_utc"])

    def test_el_cambio_de_hora_no_deja_huecos_ni_solapes(self):
        # Chile pasa a horario de verano en la madrugada del domingo 6 de septiembre de 2026
        a = TZ.utcoffset(datetime(2026, 8, 31, 12))
        b = TZ.utcoffset(datetime(2026, 9, 7, 12))
        self.assertNotEqual(a, b, "la prueba solo vale si el huso cambia esa semana")
        sem = G.rango_periodo(date(2026, 9, 2))
        sig = G.rango_periodo(date(2026, 9, 9))
        ant = G.rango_periodo(date(2026, 8, 26))
        self.assertEqual(sem["ini_utc"], datetime(2026, 8, 31, 4, 0), "agosto: UTC-4 → 04:00 UTC")
        self.assertEqual(sem["fin_utc"], datetime(2026, 9, 7, 3, 0), "septiembre: UTC-3 → 03:00 UTC")
        self.assertEqual(sem["fin_utc"], sig["ini_utc"])
        self.assertEqual(ant["fin_utc"], sem["ini_utc"])
        self.assertEqual(sem["fin_utc"] - sem["ini_utc"], timedelta(days=7) - timedelta(hours=1))

    def test_semana_y_mes_de_un_cierre_en_el_borde(self):
        self.assertEqual(G.semana_de(datetime(2026, 10, 26, 2, 59, 59)), date(2026, 10, 19))
        self.assertEqual(G.semana_de(datetime(2026, 10, 26, 3, 0, 0)), date(2026, 10, 26))
        self.assertEqual(G.mes_de(datetime(2026, 11, 1, 2, 59)), "2026-10", "31-oct 23:59 en Chile todavía es octubre")
        self.assertEqual(G.mes_de(datetime(2026, 11, 1, 3, 0)), "2026-11")
        self.assertIsNone(G.semana_de(None))

    def test_mes_y_trimestre_con_su_periodo_anterior(self):
        m = G.rango_periodo(date(2026, 10, 23), "mes")
        self.assertEqual((m["desde"], m["hasta"]), (date(2026, 10, 1), date(2026, 10, 31)))
        self.assertEqual((m["previo"]["desde"], m["previo"]["hasta"]), (date(2026, 9, 1), date(2026, 9, 30)))
        self.assertEqual(m["etiqueta"], "Octubre de 2026")
        ene = G.rango_periodo(date(2026, 1, 15), "mes")
        self.assertEqual((ene["previo"]["desde"], ene["previo"]["hasta"]), (date(2025, 12, 1), date(2025, 12, 31)))
        t = G.rango_periodo(date(2026, 10, 23), "trimestre")
        self.assertEqual((t["desde"], t["hasta"]), (date(2026, 10, 1), date(2026, 12, 31)))
        self.assertEqual((t["previo"]["desde"], t["previo"]["hasta"]), (date(2026, 7, 1), date(2026, 9, 30)))
        t1 = G.rango_periodo(date(2026, 2, 10), "trimestre")
        self.assertEqual((t1["previo"]["desde"], t1["previo"]["hasta"]), (date(2025, 10, 1), date(2025, 12, 31)))

    def test_un_rango_mal_dado_se_corrige_sin_reventar(self):
        r = G.rango_periodo(HOY, "rango", date(2026, 10, 10), date(2026, 10, 1))
        self.assertEqual((r["desde"], r["hasta"]), (date(2026, 10, 1), date(2026, 10, 10)), "al revés → se ordena")
        self.assertEqual((r["previo"]["desde"], r["previo"]["hasta"]), (date(2026, 9, 21), date(2026, 9, 30)))
        falta = G.rango_periodo(HOY, "rango", None, date(2026, 10, 10))
        self.assertEqual(falta["modo"], "semana", "sin las dos puntas vuelve a la semana")
        largo = G.rango_periodo(HOY, "rango", date(2020, 1, 1), date(2026, 10, 10))
        self.assertLessEqual(largo["dias"], G.MAX_RANGO_DIAS)
        self.assertEqual(G.rango_periodo(HOY, "cualquier cosa")["modo"], "semana")

    def test_la_navegacion_no_deja_ir_al_futuro(self):
        p = G.rango_periodo(date(2026, 10, 23))
        self.assertEqual(p["ancla_previa"], date(2026, 10, 12))
        self.assertIsNone(p["ancla_siguiente"], "la semana en curso no tiene siguiente")
        anterior = G.rango_periodo(date(2026, 10, 23), ancla=date(2026, 10, 14))
        self.assertEqual(anterior["ancla_siguiente"], date(2026, 10, 19))
        self.assertFalse(anterior["en_curso"])
        self.assertIsNone(G.rango_periodo(HOY, "rango", date(2026, 10, 1), date(2026, 10, 5))["ancla_previa"])

    def test_leer_parametros_no_confia_en_nada(self):
        p = G.leer_parametros({"modo": "MES", "ancla": "2026-10-01", "dias": "90"}, HOY)
        self.assertEqual((p["modo"], p["ancla"], p["dias"]), ("mes", date(2026, 10, 1), 90))
        raro = G.leer_parametros({"modo": "x'; drop", "ancla": "mañana", "desde": "1/2", "dias": "7"}, HOY)
        self.assertEqual((raro["modo"], raro["ancla"], raro["desde"], raro["dias"]), ("semana", None, None, 60))
        self.assertEqual(G.leer_parametros({"ancla": "2030-01-01"}, HOY)["ancla"], HOY, "una fecha futura se corrige a hoy")
        self.assertEqual(G.leer_parametros({"dias": "abc"}, HOY)["dias"], G.DIAS_REINCIDENCIA_DEFECTO)

    def test_la_ventana_de_carga_cubre_anterior_12_semanas_y_6_meses(self):
        p = G.rango_periodo(date(2026, 10, 23))
        ini, fin = G.ventana_carga(p)
        self.assertEqual(fin, date(2026, 10, 25))
        self.assertLessEqual(ini, p["previo"]["desde"])
        self.assertLessEqual(ini, date(2026, 5, 1), "6 meses calendario hacia atrás")
        self.assertLessEqual(ini, date(2026, 10, 19) - timedelta(days=77))
        sem = G.ultimas_semanas(date(2026, 10, 23), 12)
        self.assertEqual((len(sem), sem[0], sem[-1]), (12, date(2026, 8, 3), date(2026, 10, 19)))
        self.assertEqual(G.ultimos_meses(date(2026, 2, 10), 4), [(2025, 11), (2025, 12), (2026, 1), (2026, 2)])


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestFormatos(unittest.TestCase):
    def test_clp_y_porcentaje(self):
        self.assertEqual(G.clp(1234567), "$1.234.567")
        self.assertEqual(G.clp(-5000), "-$5.000")
        self.assertEqual(G.clp(0), "$0")
        self.assertEqual(G.clp(None), "—")
        self.assertEqual(G.clp(1500, True), "+$1.500")
        self.assertEqual(G.pct_txt(12.5), "12,5 %")
        self.assertEqual(G.pct_txt(40.0), "40 %")
        self.assertEqual(G.pct_txt(None), "—")

    def test_fechas_en_chile_nunca_iso_crudo(self):
        self.assertEqual(G.fecha_hora_chile_txt(datetime(2026, 10, 21, 18, 30)), "21/10/2026 15:30")
        self.assertEqual(G.fecha_hora_chile_txt(datetime(2026, 7, 21, 18, 30)), "21/07/2026 14:30", "en invierno es UTC-4")
        self.assertEqual(G.fecha_hora_chile_txt(None), "")
        self.assertEqual(G.fecha_txt(date(2026, 10, 5)), "05/10/2026")


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestResumen(unittest.TestCase):
    def test_una_semana_vacia_no_divide_por_cero_ni_inventa(self):
        r = G.resumen([], agregar())
        for k in ("cobre", "me_cobraron", "neto", "ganamos", "perdimos", "regalado_costo", "pago_tecnicos"):
            self.assertEqual(r[k], 0, k)
        self.assertIsNone(r["pct"])
        self.assertEqual(r["clase"], "gris")
        per = G.rango_periodo(HOY)
        t = G.tarjetas(r, r, False)
        self.assertEqual([x["titulo"] for x in t], ["Cobré", "Me cobraron", "Gané", "Perdí", "Neto"])
        self.assertTrue(all(x["cmp"]["flecha"] == "sin_base" for x in t))
        v = G.veredicto(r, r, False, [], per)
        self.assertEqual(v["clase"], "gris")
        self.assertIn("Todavía no hay OT medidas", v["titulo"])
        d = G.armar_tablero(per, [], agregar(), HOY)
        self.assertEqual(d["resumen"]["n_ot"], 0)
        self.assertEqual(len(d["tendencia"]), 12)
        self.assertEqual(d["atribucion"]["casos"], [])
        self.assertTrue(d["svg_tendencia"].startswith("<svg"))

    def _mezcla(self):
        vs = [
            A.ot_cobrada(1, 500000, 300000, **en_semana(0)),                                    # gana 200.000
            A.ot_cobrada(2, 400000, 100000, 20000, **en_semana(1)),                              # gana 280.000
            A.ot_cobrada(3, 180000, 240000, **en_semana(2)),                                     # PIERDE 60.000
            A.ot_garantia(4, 95000, 150000, **en_semana(2)),                                     # garantía: cuesta 95.000
            A.visita(5, modalidad_cobro="sin_costo", cobro_cero_motivo="regalia", costo_proveedor=40000, valorizado_clp=90000,
                     zz_monto=None, **en_semana(3)),                                             # regalía
            A.visita(6, tipo="desinstalacion", ot_relacion="desinstalacion", ot_madre_id=1, madre_valida=1, zz_monto=None,
                     costo_proveedor=30000, **en_semana(3)),                                     # desinstalación cubierta
            A.visita(7, cliente_id=None, cliente=None, tipo="revision_interna", modalidad_cobro="interno", zz_monto=None,
                     costo_proveedor=0, proveedor_tipo="interno", **en_semana(4)),               # interno: aparte
            A.visita(8, tipo="preventiva", contrato_real=1, zz_monto=None, costo_proveedor=50000, **en_semana(4)),  # contrato: aparte
            A.ot_cobrada(9, 300000, None, **en_semana(4)),                                       # sin medir (falta el costo)
        ]
        return filas_de(vs)

    def test_neto_es_gane_menos_perdi_menos_coberturas(self):
        filas = self._mezcla()
        clases = {f["id"]: f["clase_neto"] for f in filas}
        self.assertEqual(clases, {1: "ganada", 2: "ganada", 3: "perdida", 4: "cobertura", 5: "cobertura", 6: "cubierta",
                                  7: "aparte", 8: "aparte", 9: "sin_medir"})
        r = G.resumen(filas, agregar())
        self.assertEqual(r["ganamos"], 480000)
        self.assertEqual(r["perdimos"], 60000)
        self.assertEqual(r["regalado_costo"], 135000)
        self.assertEqual(r["cubiertas_costo"], 30000)
        self.assertEqual(r["neto"], 480000 - 60000 - 135000 - 30000)
        self.assertEqual(r["neto"], r["ag"]["queda"], "Neto sale de la suma única del proyecto")
        self.assertEqual(r["neto"], round(r["ganamos"] - r["perdimos"] - r["regalado_costo"] - r["cubiertas_costo"], 2))
        self.assertEqual(r["cobre"], 500000 + 400000 + 180000)
        self.assertEqual(r["me_cobraron"], r["ag"]["me_cobraron"])
        self.assertEqual((r["n_ganadas"], r["n_perdidas"], r["regalado_n"], r["cubiertas_n"], r["n_sin_medir"]), (2, 1, 2, 1, 1))
        self.assertEqual(set(r["aparte"]), {"interno", "contrato"})
        self.assertEqual(r["aparte_costo"], 50000)

    def test_lo_cobrado_se_reparte_por_el_documento_que_lo_respalda(self):
        vs = [A.ot_cobrada(1, 500000, 300000, **en_semana(0)), A.ot_cobrada(2, 210000, 100000, **en_semana(1)),
              A.ot_cobrada(3, 90000, 40000, **en_semana(2)), A.ot_cobrada(4, 70000, 40000, **en_semana(2))]
        ctx = {1: {"clasif": {"codigo": "factura", "codigos": ["factura"]}},
               2: {"clasif": {"codigo": "nota_venta", "codigos": ["nota_venta"]}},
               3: {"clasif": {"codigo": "sin_documento", "codigos": ["sin_documento"]}},
               4: {"clasif": {"codigo": "documento_antiguo", "codigos": ["documento_antiguo", "factura"]}}}
        r = G.resumen(filas_de(vs, ctx), agregar())
        self.assertEqual(r["respaldo"], {"factura": 570000, "nota_venta": 210000, "sin_documento": 90000, "otro": 0})
        self.assertEqual(sum(r["respaldo"].values()), r["cobre"], "las partes suman lo cobrado")

    def test_una_ot_a_perdida_se_cuenta_como_perdi_y_su_autorizacion_se_ve(self):
        f = filas_de([A.ot_cobrada(1, 180000, 240000, **en_semana(0))],
                     {1: {"autorizaciones": [{"id": 5, "tipo": "cerrar_con_perdida", "estado": "aprobada", "resuelto_por_nombre": "Gerente Demo",
                                              "resuelto_at": utc(2026, 10, 20), "argumento": "Cliente estratégico"}]}})
        r = G.resumen(f, agregar())
        self.assertEqual((r["perdimos"], r["ganamos"], r["neto"], r["n_perdidas"], r["perdidas_autorizadas"]), (60000, 0, -60000, 1, 1))
        self.assertEqual(r["clase"], "rojo")
        tarj = {t["clave"]: t for t in G.tarjetas(r, r, True)}
        self.assertEqual(tarj["perdimos"]["tono"], "perdida")
        v = G.veredicto(r, r, True, [], G.rango_periodo(HOY))
        self.assertEqual(v["clase"], "rojo")
        self.assertIn("perdió $60.000", v["titulo"])

    def test_una_garantia_que_cuesta_mas_de_lo_que_vale_se_marca(self):
        f = filas_de([A.ot_garantia(1, 130000, 100000, **en_semana(0)), A.ot_garantia(2, 40000, 100000, **en_semana(1))])
        g = G.garantias(f, f, date(2026, 10, 25))
        self.assertEqual(g["total_regalado"]["n_costo_mayor"], 1)
        self.assertEqual([x["numero"] for x in g["costo_mayor_valor"]], ["OT-2026-00001"])
        fam = g["por_familia"][0]
        self.assertEqual((fam["cobertura"], fam["n"], fam["costo"], fam["valor"], fam["pct_costo_valor"]), ("garantia", 2, 170000, 200000, 85.0))

    def test_un_cero_sin_costo_declarado_no_se_inventa(self):
        f = filas_de([A.visita(1, modalidad_cobro="garantia", cobro_cero_motivo="garantia", costo_proveedor=None, valorizado_clp=80000,
                               zz_monto=None, **en_semana(0))])
        self.assertTrue(f[0]["falta_tecnico"])
        self.assertIsNone(f[0]["costo_cobertura"])
        self.assertEqual(f[0]["clase_neto"], "sin_medir")
        g = G.garantias(f, f, date(2026, 10, 25))
        self.assertEqual((g["total_regalado"]["n"], g["total_regalado"]["costo"], g["total_regalado"]["n_sin_costo"]), (1, 0, 1))
        at = G.atribucion(f)
        self.assertEqual(len(at["casos"]), 1)
        self.assertIsNone(at["casos"][0]["monto"], "aparece en la trazabilidad pero sin monto inventado")
        self.assertEqual(G.resumen(f, agregar())["regalado_costo"], 0)

    def test_comparar_con_el_periodo_anterior_y_su_tono(self):
        s = G.comparar(150, 100, "sube")
        self.assertEqual((s["flecha"], s["delta"], s["delta_pct"], s["tono"]), ("sube", 50, 50.0, "bueno"))
        self.assertEqual(G.comparar(80, 100, "sube")["tono"], "malo")
        self.assertEqual(G.comparar(150, 100, "baja")["tono"], "malo", "que suba lo perdido es malo")
        self.assertEqual(G.comparar(50, 100, "baja")["tono"], "bueno")
        self.assertEqual(G.comparar(150, 100, "neutro")["tono"], "neutro", "lo que cobran los técnicos no es bueno ni malo por sí solo")
        self.assertEqual(G.comparar(100, 100)["flecha"], "igual")
        self.assertEqual(G.comparar(100, None)["flecha"], "sin_base")
        self.assertIsNone(G.comparar(10, 0)["delta_pct"], "sin porcentaje cuando el período anterior fue cero")

    def test_las_cinco_tarjetas_van_en_el_orden_de_daniel(self):
        r = G.resumen(self._mezcla(), agregar())
        t = G.tarjetas(r, r, True)
        self.assertEqual([x["titulo"] for x in t], ["Cobré", "Me cobraron", "Gané", "Perdí", "Neto"])
        self.assertEqual([x["clave"] for x in t], ["cobre", "me_cobraron", "ganamos", "perdimos", "neto"])
        neto = t[4]
        self.assertEqual(neto["valor"], r["neto"])
        etiquetas = [e[0] for e in neto["detalle"]]
        self.assertEqual(etiquetas[:3], ["Gané", "Perdí", "Garantías y coberturas"])
        self.assertIn("Desinstalaciones cubiertas", etiquetas)
        self.assertEqual(sum(e[1] for e in neto["detalle"]), r["neto"], "el detalle del Neto suma el Neto")
        for x in t:
            self.assertTrue(x["ayuda"], "cada cifra explica cómo se calcula")
            for _e, _v, tipo in x["detalle"]:
                self.assertIn(tipo, ("$", "n"))

    def test_el_veredicto_solo_es_rojo_si_se_pierde_plata(self):
        per = G.rango_periodo(HOY)
        # muchas garantías pero el neto sigue positivo → NO rojo
        f = filas_de([A.ot_cobrada(1, 900000, 300000, **en_semana(0)), A.ot_garantia(2, 250000, 150000, **en_semana(1))])
        r = G.resumen(f, agregar())
        self.assertGreater(r["neto"], 0)
        v = G.veredicto(r, r, True, [], per)
        self.assertEqual(v["clase"], "ok")
        self.assertIn("garantías y coberturas costaron", v["texto"])
        # margen entre 0 y 10 % → ámbar
        f2 = filas_de([A.ot_cobrada(1, 100000, 95000, **en_semana(0))])
        r2 = G.resumen(f2, agregar())
        v2 = G.veredicto(r2, r2, True, [], per)
        self.assertEqual(v2["clase"], "bajo")
        self.assertIn("justo", v2["titulo"])
        # pérdida real → rojo
        f3 = filas_de([A.ot_cobrada(1, 100000, 160000, **en_semana(0))])
        r3 = G.resumen(f3, agregar())
        self.assertEqual(G.veredicto(r3, r3, True, [], per)["clase"], "rojo")

    def test_el_veredicto_compara_con_el_periodo_anterior_en_puntos(self):
        per = G.rango_periodo(HOY)
        act = G.resumen(filas_de([A.ot_cobrada(1, 100000, 60000, **en_semana(0))]), agregar())      # 40 %
        pre = G.resumen(filas_de([A.ot_cobrada(2, 100000, 80000, cerrada_at=datetime(2026, 10, 15, 18, 0))]), agregar())  # 20 %
        v = G.veredicto(act, pre, True, [], per)
        self.assertIn("20 puntos sobre el del período anterior (20 %)", v["texto"])

    def test_la_tendencia_junta_por_semana_chilena_y_deja_las_vacias(self):
        vs = [A.ot_cobrada(1, 200000, 100000, cerrada_at=datetime(2026, 10, 26, 2, 59)),   # domingo 25 23:59 Chile → semana 19
              A.ot_cobrada(2, 200000, 100000, cerrada_at=datetime(2026, 10, 26, 3, 0))]    # lunes 26 00:00 Chile → semana 26
        t = G.tendencia(filas_de(vs), date(2026, 10, 30), agregar(), 12)
        por = {s["lunes"]: s for s in t}
        self.assertEqual(len(t), 12)
        self.assertEqual(por[date(2026, 10, 19)]["n_medidas"], 1)
        self.assertEqual(por[date(2026, 10, 26)]["n_medidas"], 1)
        self.assertEqual(por[date(2026, 10, 12)]["n_medidas"], 0)
        self.assertIsNone(por[date(2026, 10, 12)]["pct"])
        self.assertEqual(por[date(2026, 10, 19)]["neto"], 100000)

    def test_el_margen_por_semana_es_una_linea_con_huecos_no_con_ceros(self):
        vs = [A.ot_cobrada(1, 200000, 100000, **en_semana(0)),                                         # 50 %
              A.ot_cobrada(2, 200000, 160000, cerrada_at=datetime(2026, 10, 7, 18, 0))]                # 20 %, dos semanas antes
        t = G.tendencia(filas_de(vs), date(2026, 10, 25), agregar(), 12)
        tp = G.tendencia_pct(t)
        self.assertEqual(len(tp["serie"]), 12)
        self.assertEqual((tp["ultimo"], tp["min"], tp["max"], tp["n_semanas_con_dato"]), (50.0, 20.0, 50.0, 2))
        self.assertEqual(sum(1 for p in tp["serie"] if p is None), 10, "las semanas sin cobros quedan como hueco")
        self.assertIn("<polyline", tp["svg"])
        vacia = G.tendencia_pct(G.tendencia([], date(2026, 10, 25), agregar(), 12))
        self.assertEqual((vacia["ultimo"], vacia["svg"]), (None, ""))

    def test_donde_se_va_la_plata_suma_todo_lo_que_nos_cobraron(self):
        f = self._mezcla()
        d = G.donde_se_va(f)
        self.assertEqual(d["total"], round(sum(x["me_total"] for x in f), 2))
        for corte in ("por_tipo", "por_centro", "por_proveedor"):
            self.assertAlmostEqual(sum(e["costo"] for e in d[corte]), d["total"], places=2)
            costos = [e["costo"] for e in d[corte]]
            self.assertEqual(costos, sorted(costos, reverse=True), "Pareto: de mayor a menor")
            self.assertAlmostEqual(d[corte][-1]["acum"], 100.0, delta=0.2)
        self.assertEqual(G.donde_se_va([])["total"], 0)
        self.assertEqual(G.donde_se_va([])["por_tipo"], [])
        propio = dict(tecnico_role="tecnico", te_id=None, te_empresa=None, proveedor_tipo="interno", tecnico_nombre="Propio Demo")
        f2 = f + filas_de([A.ot_cobrada(30, 100000, 0, **en_semana(0), **propio),
                           A.ot_cobrada(31, 100000, 0, **en_semana(1), **dict(propio, tecnico_nombre="Otro Propio", tecnico_user_id=33))])
        d2 = G.donde_se_va(f2)
        propios = [e for e in d2["por_proveedor"] if e["clave"] == "propios"]
        self.assertEqual([(e["etiqueta"], e["n"]) for e in propios], [("Técnicos propios", 2)], "los técnicos propios van juntos")


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestAtribucion(unittest.TestCase):
    """Originó · Autorizó · Lleva · Ejecutó: cuatro papeles separados, cada uno con su fuente."""

    def _caso(self):
        v = A.ot_garantia(1, 80000, 150000, **en_semana(1), created_by="ejecutivo1", responsable_nombre="Aarón Urbina",
                          responsable_at=utc(2026, 10, 15, 17), tk_id=7, numero_ticket="TK-2026-0007", tk_origen="backoffice",
                          tk_creado_por="ejecutivo2", tk_creado_at=utc(2026, 10, 10, 14))
        ctx = {"autorizaciones": [{"id": 31, "tipo": "cobro_cero", "estado": "aprobada", "motivo": "garantia", "argumento": "Garantía de fábrica",
                                   "solicitado_por_nombre": "Ejecutiva Demo Uno", "solicitado_at": utc(2026, 10, 11, 12),
                                   "resuelto_por_nombre": "Gerente Demo", "resuelto_at": utc(2026, 10, 12, 18, 40)}]}
        return filas_de([v], {1: ctx})

    def test_cada_papel_sale_de_su_propia_fuente(self):
        f = self._caso()[0]
        self.assertEqual(f["origino"]["nombre"], "Ejecutivo Demo Dos", "el ticket lo creó ejecutivo2, no quien creó la OT")
        self.assertEqual(f["origino"]["fuente"], "Ticket TK-2026-0007")
        self.assertEqual(f["autorizo"]["nombre"], "Gerente Demo", "autoriza quien RESOLVIÓ, no quien lo pidió")
        self.assertEqual(f["autorizo"]["solicito"], "Ejecutiva Demo Uno")
        self.assertEqual(f["autorizo"]["argumento"], "Garantía de fábrica")
        self.assertEqual(f["autorizo"]["at_txt"], "12/10/2026 15:40")
        self.assertEqual(f["autorizo"]["fuente"], "Autorización #31 · Cobro en $0")
        self.assertEqual(f["lleva"]["nombre"], "Aarón Urbina")
        self.assertEqual(f["ejecutor"]["nombre"], "Servicios Norte Ltda.")
        self.assertEqual(f["ejecutor"]["persona"], "Técnico Externo Uno")

    def test_la_fila_de_la_trazabilidad_trae_los_cuatro_papeles_y_el_porque(self):
        at = G.atribucion(self._caso())
        c = at["casos"][0]
        self.assertEqual((c["tipo"], c["cobertura_txt"], c["monto"]), ("cobertura", "Garantía", 80000))
        self.assertEqual(c["origino"]["nombre"], "Ejecutivo Demo Dos")
        self.assertEqual(c["autorizo"]["nombre"], "Gerente Demo")
        self.assertEqual(c["lleva"]["nombre"], "Aarón Urbina")
        self.assertEqual(c["ejecuto"]["nombre"], "Servicios Norte Ltda.")
        self.assertEqual(c["argumento"], "Garantía de fábrica")
        self.assertIsNone(c["asume_area"])
        self.assertIsNone(c["intencion"])
        self.assertEqual(c["centro_costo"], "sstt")
        self.assertFalse(c["sin_autorizacion"])

    def test_cada_persona_suma_solo_por_el_papel_que_tuvo(self):
        at = G.atribucion(self._caso())
        por = {p["nombre"]: p for p in at["personas"]}
        solo = lambda p, rol: {r: por[p][r]["n"] for r in ("origino", "autorizo", "lleva", "ejecuto")}[rol]
        self.assertEqual({r: por["Gerente Demo"][r]["n"] for r in ("origino", "autorizo", "lleva", "ejecuto")},
                         {"origino": 0, "autorizo": 1, "lleva": 0, "ejecuto": 0}, "autorizar no es ejecutar")
        self.assertEqual({r: por["Ejecutivo Demo Dos"][r]["n"] for r in ("origino", "autorizo", "lleva", "ejecuto")},
                         {"origino": 1, "autorizo": 0, "lleva": 0, "ejecuto": 0})
        self.assertEqual({r: por["Aarón Urbina"][r]["n"] for r in ("origino", "autorizo", "lleva", "ejecuto")},
                         {"origino": 0, "autorizo": 0, "lleva": 1, "ejecuto": 0})
        self.assertEqual({r: por["Servicios Norte Ltda."][r]["n"] for r in ("origino", "autorizo", "lleva", "ejecuto")},
                         {"origino": 0, "autorizo": 0, "lleva": 0, "ejecuto": 1})
        self.assertNotIn("Ejecutiva Demo Uno", por, "quien solo PIDIÓ la autorización no aparece como quien autorizó ni originó")
        self.assertEqual(por["Servicios Norte Ltda."]["ejecuto"]["costo"], 80000)
        self.assertEqual(solo("Gerente Demo", "autorizo"), 1)

    def test_el_mismo_nombre_escrito_distinto_es_una_sola_persona(self):
        v1 = A.ot_garantia(1, 50000, 90000, **en_semana(0), responsable_nombre="Aarón Urbina")
        v2 = A.ot_garantia(2, 60000, 90000, **en_semana(1), responsable_nombre="aarón urbina ")
        at = G.atribucion(filas_de([v1, v2]))
        lleva = [p for p in at["personas"] if p["lleva"]["n"]]
        self.assertEqual(len(lleva), 1)
        self.assertEqual(lleva[0]["lleva"]["n"], 2)

    def test_sin_autorizacion_registrada_es_un_dato_visible_no_un_error(self):
        f = filas_de([A.ot_garantia(1, 80000, 150000, **en_semana(0))])[0]
        self.assertTrue(f["autorizo"]["sin_autorizacion"])
        self.assertIsNone(f["autorizo"]["nombre"])
        at = G.atribucion([f])
        self.assertEqual(at["n_sin_autorizacion"], 1)
        self.assertEqual(at["n_sin_argumento"], 1)
        self.assertEqual(at["n_sin_clasificar"], 1)
        self.assertFalse(any(p["autorizo"]["n"] for p in at["personas"]))

    def test_si_la_clasificacion_guardada_al_crear_trae_quien_autorizo_se_usa(self):
        ctx = {"clasif": {"codigo": "garantia", "codigos": ["garantia"], "texto": "$0 · garantía", "autorizado_por": "Gerente Demo",
                          "argumento": "Autorizado al crear la OT", "autorizado_at": "12/10/2026 15:40", "pedido_por": "Ejecutiva Demo Uno"}}
        f = filas_de([A.ot_garantia(1, 80000, 150000, **en_semana(0))], {1: ctx})[0]
        self.assertEqual(f["autorizo"]["nombre"], "Gerente Demo")
        self.assertEqual(f["autorizo"]["at_txt"], "12/10/2026 15:40")
        self.assertEqual(f["autorizo"]["fuente"], "Clasificación guardada al crear la OT")
        self.assertFalse(f["autorizo"]["sin_autorizacion"])

    def test_una_autorizacion_de_otro_asunto_no_se_le_acredita_a_un_cero(self):
        ctx = {"autorizaciones": [{"id": 9, "tipo": "documento_observado", "estado": "aprobada", "resuelto_por_nombre": "Gerente Demo",
                                   "resuelto_at": utc(2026, 10, 12), "argumento": "x"}]}
        f = filas_de([A.ot_garantia(1, 80000, 150000, **en_semana(0))], {1: ctx})[0]
        self.assertTrue(f["autorizo"]["sin_autorizacion"])
        # las rechazadas o pendientes tampoco
        ctx2 = {"autorizaciones": [{"id": 9, "tipo": "cobro_cero", "estado": "rechazada", "resuelto_por_nombre": "Gerente Demo", "resuelto_at": utc(2026, 10, 12)},
                                   {"id": 10, "tipo": "cobro_cero", "estado": "pendiente", "resuelto_por_nombre": None}]}
        self.assertTrue(filas_de([A.ot_garantia(2, 80000, 150000, **en_semana(0))], {2: ctx2})[0]["autorizo"]["sin_autorizacion"])

    def test_con_varias_autorizaciones_manda_la_mas_reciente_del_tipo_que_corresponde(self):
        auts = [{"id": 1, "tipo": "cobro_cero", "estado": "aprobada", "resuelto_por_nombre": "Primero", "resuelto_at": utc(2026, 10, 1), "argumento": "a"},
                {"id": 2, "tipo": "cobro_cero", "estado": "aprobada", "resuelto_por_nombre": "Segundo", "resuelto_at": utc(2026, 10, 5), "argumento": "b"},
                {"id": 3, "tipo": "cerrar_con_perdida", "estado": "aprobada", "resuelto_por_nombre": "Otro", "resuelto_at": utc(2026, 10, 9), "argumento": "c"}]
        a = G.armar_autorizo("garantia", auts)
        self.assertEqual((a["nombre"], a["argumento"]), ("Segundo", "b"))
        p = G.armar_autorizo("cobra", auts)
        self.assertEqual(p["nombre"], "Otro", "una OT cobrada a pérdida mira la autorización de pérdida")
        self.assertEqual(len(a["todas"]), 3)

    def test_el_origen_distingue_formulario_erp_ticket_y_ot_directa(self):
        base = dict(tk_id=1, numero_ticket="TK-1", tk_creado_at=utc(2026, 10, 1))
        self.assertEqual(G.armar_origen(dict(base, tk_origen="form", tk_creado_por=None))["nombre"], "Cliente (formulario web)")
        self.assertEqual(G.armar_origen(dict(base, tk_origen="erp"))["nombre"], "ERP (documento de venta)")
        self.assertEqual(G.armar_origen(dict(base, tk_origen="backoffice", tk_creado_por="Ejecutivo1"), {"ejecutivo1": "Ejecutiva Demo Uno"})["nombre"],
                         "Ejecutiva Demo Uno")
        directa = G.armar_origen({"created_by": "admin1", "created_at": utc(2026, 10, 2)}, {"admin1": "Administrador Demo"})
        self.assertEqual((directa["nombre"], directa["fuente"], directa["ticket"]), ("Administrador Demo", "OT creada directamente (sin ticket)", None))

    def test_una_ot_sin_responsable_lo_dice(self):
        f = filas_de([A.ot_garantia(1, 80000, 150000, **en_semana(0), responsable_nombre=None, responsable_at=None)])[0]
        self.assertIsNone(f["lleva"]["nombre"])
        self.assertEqual(f["lleva"]["fuente"], "Sin responsable declarado")

    def test_incluye_las_coberturas_y_las_perdidas_y_nada_mas(self):
        f = filas_de([A.ot_cobrada(1, 500000, 300000, **en_semana(0)), A.ot_cobrada(2, 100000, 160000, **en_semana(1)),
                      A.ot_garantia(3, 70000, 100000, **en_semana(2))])
        at = G.atribucion(f)
        self.assertEqual({c["numero"]: c["tipo"] for c in at["casos"]}, {"OT-2026-00002": "perdida", "OT-2026-00003": "cobertura"})
        perdida = next(c for c in at["casos"] if c["tipo"] == "perdida")
        self.assertEqual((perdida["monto"], perdida["cobertura_txt"]), (60000, "Cobrada a pérdida"))
        por_persona_perdida = [p for p in at["personas"] if p["ejecuto"]["perdida"]]
        self.assertTrue(por_persona_perdida)

    def test_los_puntos_de_extension_se_usan_cuando_existen(self):
        ctx = {"intencion": "Evento de marketing", "asume_area": "Marketing"}
        f = filas_de([A.ot_garantia(1, 80000, 150000, **en_semana(0))], {1: ctx})[0]
        self.assertEqual((f["intencion"], f["asume_area"]), ("Evento de marketing", "Marketing"))
        at = G.atribucion([f])
        self.assertEqual(at["n_sin_clasificar"], 0)
        self.assertEqual(at["casos"][0]["asume_area"], "Marketing")
        g = G.garantias([f], [f], date(2026, 10, 25))
        self.assertEqual([(e["etiqueta"], e["sin_clasificar"]) for e in g["por_area"]], [("Marketing", False)])
        self.assertEqual([(e["etiqueta"], e["sin_clasificar"]) for e in g["por_intencion"]], [("Evento de marketing", False)])


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestGarantias(unittest.TestCase):
    def test_por_familia_separa_lo_regalado_de_lo_que_no_es_un_regalo(self):
        vs = [A.ot_garantia(1, 80000, 150000, **en_semana(0)),
              A.visita(2, modalidad_cobro="sin_costo", cobro_cero_motivo="regalia", costo_proveedor=40000, valorizado_clp=90000, zz_monto=None, **en_semana(1)),
              A.visita(3, modalidad_cobro="sin_costo", cobro_cero_motivo="arriendo_leasing", costo_proveedor=30000, valorizado_clp=70000, zz_monto=None, **en_semana(1)),
              A.visita(4, tipo="desinstalacion", ot_relacion="desinstalacion", ot_madre_id=1, madre_valida=1, zz_monto=None, costo_proveedor=20000, **en_semana(2))]
        g = G.garantias(filas_de(vs), filas_de(vs), date(2026, 10, 25))
        self.assertEqual([x["cobertura"] for x in g["por_familia"]], ["garantia", "regalia", "arriendo_leasing", "cubierta"])
        self.assertEqual([x["regalada"] for x in g["por_familia"]], [True, True, True, False])
        self.assertEqual(g["total_regalado"]["n"], 3, "la desinstalación cubierta no es un regalo")
        self.assertEqual(g["total_regalado"]["costo"], 150000)
        self.assertEqual(g["total_regalado"]["valor"], 310000)
        self.assertEqual(g["n_cubiertas"], 1)

    def test_por_cliente_ordena_por_lo_que_costo_y_deja_las_ot_a_un_clic(self):
        vs = [A.ot_garantia(1, 90000, 150000, **en_semana(0), cliente_id=1, cliente="Gimnasio Uno"),
              A.ot_garantia(2, 30000, 100000, **en_semana(1), cliente_id=2, cliente="Gimnasio Dos"),
              A.ot_garantia(3, 40000, 100000, **en_semana(2), cliente_id=2, cliente="Gimnasio Dos")]
        g = G.garantias(filas_de(vs), filas_de(vs), date(2026, 10, 25))
        self.assertEqual([(e["cliente"], e["n"], e["costo"]) for e in g["por_cliente"]], [("Gimnasio Uno", 1, 90000), ("Gimnasio Dos", 2, 70000)])
        self.assertEqual([o["id"] for o in g["por_cliente"][1]["ots"]], [2, 3])

    def test_por_mes_cuenta_el_mes_de_chile_y_trae_seis_meses(self):
        vs = [A.ot_garantia(1, 50000, 90000, cerrada_at=datetime(2026, 11, 1, 2, 30)),     # 31-oct 23:30 Chile → octubre
              A.ot_garantia(2, 60000, 90000, cerrada_at=datetime(2026, 11, 1, 3, 30)),     # 1-nov 00:30 Chile → noviembre
              A.ot_garantia(3, 10000, 90000, cerrada_at=datetime(2026, 5, 20, 15, 0))]
        g = G.garantias([], filas_de(vs), date(2026, 11, 10))
        por = {m["clave"]: m for m in g["por_mes"]}
        self.assertEqual(len(g["por_mes"]), 6)
        self.assertEqual((por["2026-10"]["n"], por["2026-10"]["costo"]), (1, 50000))
        self.assertEqual((por["2026-11"]["n"], por["2026-11"]["costo"]), (1, 60000))
        self.assertNotIn("2026-05", por, "mayo queda fuera de los últimos 6 meses terminados en noviembre")
        self.assertEqual(max(m["pct"] for m in g["por_mes"]), 100.0)

    def test_sin_columnas_de_extension_todo_queda_sin_clasificar(self):
        f = filas_de([A.ot_garantia(1, 50000, 90000, **en_semana(0)), A.ot_garantia(2, 60000, 90000, **en_semana(1))])
        g = G.garantias(f, f, date(2026, 10, 25))
        self.assertEqual([(e["etiqueta"], e["n"], e["sin_clasificar"]) for e in g["por_area"]], [(G.SIN_CLASIFICAR, 2, True)])
        self.assertEqual([(e["etiqueta"], e["n"], e["sin_clasificar"]) for e in g["por_intencion"]], [(G.SIN_CLASIFICAR, 2, True)])
        self.assertEqual(g["clasificadas"], 0)

    def test_por_centro_de_costo_funciona_hoy(self):
        f = filas_de([A.ot_garantia(1, 50000, 90000, **en_semana(0), centro_costo="marketing"),
                      A.ot_garantia(2, 60000, 90000, **en_semana(1), centro_costo="comercial")],
                     {1: {"centro_txt": "Marketing"}, 2: {"centro_txt": "Comercial"}})
        g = G.garantias(f, f, date(2026, 10, 25))
        self.assertEqual([(e["etiqueta"], e["costo"]) for e in g["por_centro"]], [("Comercial", 60000), ("Marketing", 50000)])


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestTecnicos(unittest.TestCase):
    def _filas(self):
        ext = dict(tecnico_role="tecnico_externo", te_id=5, te_empresa="Servicios Norte Ltda.", proveedor_tipo="externo")
        return filas_de([
            A.ot_cobrada(1, 500000, 200000, 20000, **en_semana(0), tecnico_user_id=21, tecnico_nombre="Pedro Ejemplo", **ext),
            A.ot_cobrada(2, 400000, 100000, **en_semana(1), tecnico_user_id=23, tecnico_nombre="Ana Ejemplo", **ext),   # misma empresa, otra persona
            A.ot_cobrada(3, 300000, 150000, **en_semana(2), tecnico_user_id=22, tecnico_nombre="Marta Demo", tecnico_role="tecnico_externo",
                         te_id=6, te_empresa="Técnica Sur SpA", proveedor_tipo="externo"),
            A.ot_cobrada(4, 250000, 0, **en_semana(3), tecnico_user_id=31, tecnico_nombre="Jaizer Demo", tecnico_role="tecnico_ejecutivo",
                         te_id=None, te_empresa=None, proveedor_tipo="interno"),
            A.ot_cobrada(5, 250000, 0, **en_semana(3), tecnico_user_id=32, tecnico_nombre="Técnico Propio Demo", tecnico_role="tecnico",
                         te_id=None, te_empresa=None, proveedor_tipo="interno"),
            A.ot_cobrada(6, 100000, None, **en_semana(4), tecnico_user_id=22, tecnico_nombre="Marta Demo", tecnico_role="tecnico_externo",
                         te_id=6, te_empresa="Técnica Sur SpA", proveedor_tipo="externo"),
        ])

    def test_los_externos_se_agrupan_por_empresa_y_los_propios_por_persona(self):
        t = G.tecnicos(self._filas(), [], self._filas(), date(2026, 10, 25))
        por = {g["nombre"]: g for g in t["grupos"]}
        self.assertEqual(set(por), {"Servicios Norte Ltda.", "Técnica Sur SpA", "Jaizer Demo", "Técnico Propio Demo"})
        norte = por["Servicios Norte Ltda."]
        self.assertEqual((norte["tipo"], norte["n"], norte["cobran"]), ("externo", 2, 320000))
        self.assertEqual(sorted(norte["personas"]), ["Ana Ejemplo", "Pedro Ejemplo"], "una empresa, dos personas enviadas")
        self.assertEqual(por["Jaizer Demo"]["tipo"], "elevado")
        self.assertEqual(por["Técnico Propio Demo"]["tipo"], "interno")
        self.assertEqual(por["Jaizer Demo"]["cobran"], 0, "un técnico propio no cobra por OT")
        self.assertEqual([x["tipo"] for x in t["tipos"]], ["externo", "elevado", "interno"])

    def test_porcentaje_del_costo_promedio_y_orden(self):
        t = G.tecnicos(self._filas(), [], self._filas(), date(2026, 10, 25))
        self.assertEqual(t["total"], 470000)
        self.assertAlmostEqual(sum(g["pct_costo"] for g in t["grupos"]), 100.0, delta=0.2)
        costos = [g["cobran"] for g in t["grupos"]]
        self.assertEqual(costos, sorted(costos, reverse=True))
        sur = next(g for g in t["grupos"] if g["nombre"] == "Técnica Sur SpA")
        self.assertEqual((sur["n"], sur["n_sin_costo"], sur["promedio"]), (2, 1, 150000), "el promedio ignora la OT sin costo declarado")

    def test_evolucion_contra_el_periodo_anterior_y_serie_semanal(self):
        actual = self._filas()
        previo = filas_de([A.ot_cobrada(20, 300000, 100000, cerrada_at=datetime(2026, 10, 15, 18, 0), tecnico_user_id=21,
                                        tecnico_nombre="Pedro Ejemplo", tecnico_role="tecnico_externo", te_id=5,
                                        te_empresa="Servicios Norte Ltda.", proveedor_tipo="externo")])
        t = G.tecnicos(actual, previo, actual + previo, date(2026, 10, 25))
        norte = next(g for g in t["grupos"] if g["nombre"] == "Servicios Norte Ltda.")
        self.assertEqual(norte["previo"], 100000)
        self.assertEqual((norte["cmp"]["flecha"], norte["cmp"]["delta"], norte["cmp"]["delta_pct"]), ("sube", 220000, 220.0))
        self.assertEqual(len(norte["serie"]), 8)
        self.assertEqual(norte["serie"][-1], 320000)
        self.assertEqual(norte["serie"][-2], 100000)
        self.assertEqual({o["numero"] for o in norte["ots"]}, {"OT-2026-00001", "OT-2026-00002"})

    def test_cumplimiento_agendadas_hechas_y_vencidas(self):
        ej = G.armar_ejecutor({"tecnico_role": "tecnico_externo", "te_id": 5, "te_empresa": "Servicios Norte Ltda.", "tecnico_user_id": 21})
        per = G.rango_periodo(date(2026, 10, 23))
        ag = [{"id": 1, "numero": "A", "estado": "cerrada", "fecha_programada": date(2026, 10, 19), "ejecutor": ej},
              {"id": 2, "numero": "B", "estado": "programada", "fecha_programada": date(2026, 10, 20), "ejecutor": ej},     # venció
              {"id": 3, "numero": "C", "estado": "programada", "fecha_programada": date(2026, 10, 24), "ejecutor": ej},     # aún no
              {"id": 4, "numero": "D", "estado": "cancelada", "fecha_programada": date(2026, 10, 20), "ejecutor": ej},
              {"id": 5, "numero": "E", "estado": "reagendada", "fecha_programada": date(2026, 10, 20), "ejecutor": ej},
              {"id": 6, "numero": "F", "estado": "programada", "fecha_programada": date(2026, 10, 12), "ejecutor": ej},     # otra semana
              {"id": 7, "numero": "G", "estado": "completada", "fecha_programada": date(2026, 12, 31), "ejecutor": ej}]
        c = G.cumplimiento(ag, per, date(2026, 10, 23))["emp:servicios norte ltda."]
        self.assertEqual((c["agendadas"], c["hechas"], c["vencidas"], c["pct_hechas"]), (3, 1, 1, 33.3))
        self.assertEqual([x["numero"] for x in c["casos_vencidos"]], ["B"])

    def test_el_tipo_de_tecnico_sale_del_rol_y_de_la_ficha(self):
        self.assertEqual(G.tipo_ejecutor("tecnico_externo", None, None, 1), "externo")
        self.assertEqual(G.tipo_ejecutor("tecnico", "externo", None, 1), "externo", "el proveedor_tipo de la OT manda")
        self.assertEqual(G.tipo_ejecutor("tecnico", None, 9, 1), "externo", "tiene ficha de proveedor")
        self.assertEqual(G.tipo_ejecutor("tecnico_ejecutivo", "interno", None, 2), "elevado")
        self.assertEqual(G.tipo_ejecutor("tecnico_jr", None, None, 2), "interno")
        self.assertEqual(G.tipo_ejecutor(None, None, None, None), "sin_asignar")
        self.assertEqual(G.armar_ejecutor({})["nombre"], "Sin técnico asignado")


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestReincidencia(unittest.TestCase):
    AHORA = datetime(2026, 10, 23, 18, 0)
    NORTE = G.armar_ejecutor({"tecnico_role": "tecnico_externo", "te_id": 5, "te_empresa": "Servicios Norte Ltda.", "tecnico_user_id": 21})
    SUR = G.armar_ejecutor({"tecnico_role": "tecnico_externo", "te_id": 6, "te_empresa": "Técnica Sur SpA", "tecnico_user_id": 22})
    PROPIO = G.armar_ejecutor({"tecnico_role": "tecnico", "tecnico_user_id": 32, "tecnico_nombre": "Propio"})

    def cerrada(self, id=1, dias_atras=100, ej=None, cli=100, maq=(1,)):
        return {"id": id, "numero": f"OT-{id}", "cliente_id": cli, "cliente": "Gimnasio Alfa SpA",
                "cerrada_at": self.AHORA - timedelta(days=dias_atras), "ejecutor": ej or self.NORTE, "maquinas": set(maq)}

    def cand(self, id=50, despues_de=None, dias=20, tipo="correctiva", cli=100, maq=(1,), **kw):
        base = despues_de or self.cerrada()
        d = {"id": id, "numero": f"OT-{id}", "cliente_id": cli, "tipo": tipo, "es_garantia": False, "estado": "programada",
             "created_at": base["cerrada_at"] + timedelta(days=dias), "ot_madre_id": None, "ot_relacion": None, "maquinas": set(maq),
             "ejecutor": self.SUR}
        d.update(kw)
        return d

    def evaluar(self, cerradas, cands, dias=60):
        return G.detectar_reincidencias(cerradas, cands, self.AHORA, dias)

    def test_una_correctiva_sobre_el_mismo_equipo_dentro_de_la_ventana_es_reincidencia(self):
        r = self.evaluar([self.cerrada()], [self.cand()])
        e = r["por_empresa"][0]
        self.assertEqual((e["n_cerradas"], e["n_reincidentes"], e["n_evaluadas"], e["tasa"]), (1, 1, 1, 100.0))
        c = r["casos"][0]
        self.assertEqual((c["x_numero"], c["y_numero"], c["dias_despues"], c["nivel"], c["misma_empresa"]), ("OT-1", "OT-50", 20, "equipo", False))

    def test_los_bordes_de_la_ventana(self):
        x = self.cerrada()
        dentro = self.cand(id=2, despues_de=x, dias=60)
        fuera = self.cand(id=3, despues_de=x, dias=60, created_at=x["cerrada_at"] + timedelta(days=60, seconds=1))
        antes = self.cand(id=4, despues_de=x, created_at=x["cerrada_at"] - timedelta(seconds=1))
        mismo_instante = self.cand(id=5, despues_de=x, created_at=x["cerrada_at"])
        self.assertEqual(len(self.evaluar([x], [dentro])["casos"]), 1, "exactamente N días cuenta")
        self.assertEqual(self.evaluar([x], [fuera])["casos"], [], "un segundo más ya no")
        self.assertEqual(self.evaluar([x], [antes])["casos"], [], "una OT anterior al cierre no es consecuencia de él")
        self.assertEqual(self.evaluar([x], [mismo_instante])["casos"], [])

    def test_la_ventana_se_ajusta_a_30_60_o_90_dias(self):
        x = self.cerrada(dias_atras=150)
        y = self.cand(despues_de=x, dias=45)
        self.assertEqual(len(self.evaluar([x], [y], 30)["casos"]), 0)
        self.assertEqual(len(self.evaluar([x], [y], 60)["casos"]), 1)
        self.assertEqual(len(self.evaluar([x], [y], 90)["casos"]), 1)

    def test_no_cuenta_lo_que_es_continuacion_planificada_ni_lo_cancelado(self):
        x = self.cerrada()
        self.assertEqual(self.evaluar([x], [self.cand(ot_madre_id=1, ot_relacion="desinstalacion")])["casos"], [], "hija de la misma OT")
        self.assertEqual(self.evaluar([x], [self.cand(ot_madre_id=1, ot_relacion="pendiente_stock")])["casos"], [])
        self.assertEqual(self.evaluar([x], [self.cand(ot_madre_id=1, ot_relacion="seguimiento")])["casos"], [], "seguimiento de un hallazgo suyo")
        self.assertEqual(self.evaluar([x], [self.cand(estado="cancelada")])["casos"], [])
        self.assertEqual(self.evaluar([x], [self.cand(estado="anulada")])["casos"], [])
        self.assertEqual(self.evaluar([x], [self.cand(id=1)])["casos"], [], "la misma OT no se cuenta a sí misma")

    def test_solo_cuentan_correctivas_o_de_garantia(self):
        x = self.cerrada()
        self.assertEqual(self.evaluar([x], [self.cand(tipo="instalacion")])["casos"], [])
        self.assertEqual(self.evaluar([x], [self.cand(tipo="preventiva")])["casos"], [])
        self.assertEqual(len(self.evaluar([x], [self.cand(tipo="visita_correctiva")])["casos"]), 1)
        self.assertEqual(len(self.evaluar([x], [self.cand(tipo="garantia")])["casos"]), 1)
        self.assertEqual(len(self.evaluar([x], [self.cand(tipo="instalacion", es_garantia=True)])["casos"]), 1, "una OT en garantía por modalidad cuenta")

    def test_otro_cliente_u_otro_equipo_no_es_reincidencia(self):
        x = self.cerrada()
        self.assertEqual(self.evaluar([x], [self.cand(cli=999)])["casos"], [])
        self.assertEqual(self.evaluar([x], [self.cand(maq=(2,))])["casos"], [], "ambas OT tienen equipos y ninguno coincide")
        sin = self.evaluar([x], [self.cand(maq=())])["casos"]
        self.assertEqual([c["nivel"] for c in sin], ["cliente"], "sin equipos registrados se compara por cliente y se dice")

    def test_solo_se_evaluan_los_externos(self):
        r = self.evaluar([self.cerrada(ej=self.PROPIO)], [self.cand()])
        self.assertEqual(r["por_empresa"], [])
        self.assertEqual(r["n_evaluadas"], 0)
        self.assertIsNone(r["tasa_global"])

    def test_lo_que_aun_esta_en_observacion_no_entra_al_denominador(self):
        reciente = self.cerrada(id=1, dias_atras=10)
        vieja = self.cerrada(id=2, dias_atras=100, cli=101, maq=(2,))
        r = self.evaluar([reciente, vieja], [])
        e = r["por_empresa"][0]
        self.assertEqual((e["n_cerradas"], e["n_en_observacion"], e["n_evaluadas"], e["n_reincidentes"], e["tasa"]), (2, 1, 1, 0, 0.0))
        # una reciente que YA reincidió cuenta (el resultado ya se conoce)
        r2 = self.evaluar([reciente], [self.cand(despues_de=reciente, dias=3)])
        e2 = r2["por_empresa"][0]
        self.assertEqual((e2["n_en_observacion"], e2["n_evaluadas"], e2["tasa"]), (0, 1, 100.0))
        # y sin ninguna evaluada no hay tasa (no 0 %)
        sola = self.evaluar([reciente], [])["por_empresa"][0]
        self.assertIsNone(sola["tasa"])
        self.assertFalse(sola["revisar"])

    def test_conviene_revisar_exige_tasa_alta_y_suficientes_casos_y_es_un_aviso_no_una_sentencia(self):
        cerradas = [self.cerrada(id=i, cli=100 + i, maq=(i,)) for i in range(1, 5)]
        cands = [self.cand(id=50 + i, despues_de=cerradas[i - 1], cli=100 + i, maq=(i,)) for i in (1, 2)]
        r = self.evaluar(cerradas, cands)
        e = r["por_empresa"][0]
        self.assertEqual((e["n_evaluadas"], e["n_reincidentes"], e["tasa"], e["revisar"]), (4, 2, 50.0, True))
        self.assertEqual(r["n_revisar"], 1)
        poco = self.evaluar(cerradas[:2], cands[:1])["por_empresa"][0]
        self.assertEqual((poco["tasa"], poco["revisar"]), (50.0, False), "con 2 OT evaluadas no se saca conclusión")
        bajo = self.evaluar(cerradas, cands[:1])["por_empresa"][0]
        self.assertEqual((bajo["tasa"], bajo["revisar"]), (25.0, False))

    def test_volvio_la_misma_empresa_se_cuenta_aparte(self):
        x = self.cerrada()
        r = self.evaluar([x], [self.cand(ejecutor=self.NORTE)])
        self.assertTrue(r["casos"][0]["misma_empresa"])
        self.assertEqual((r["n_repite_visita"], r["por_empresa"][0]["n_repite_visita"]), (1, 1))
        r2 = self.evaluar([x], [self.cand(ejecutor=self.SUR)])
        self.assertFalse(r2["casos"][0]["misma_empresa"])
        self.assertEqual(r2["casos"][0]["y_ejecutor"], "Técnica Sur SpA")

    def test_una_ot_con_dos_posteriores_cuenta_una_reincidencia_y_lista_los_dos_casos(self):
        x = self.cerrada()
        r = self.evaluar([x], [self.cand(id=60, despues_de=x, dias=10), self.cand(id=61, despues_de=x, dias=30)])
        e = r["por_empresa"][0]
        self.assertEqual((e["n_reincidentes"], len(e["casos"])), (1, 2))
        self.assertEqual([c["y_numero"] for c in e["casos"]], ["OT-60", "OT-61"], "los más cercanos primero")


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestGraficosYExcel(unittest.TestCase):
    def test_las_barras_marcan_verde_rojo_y_vacio_y_son_accesibles(self):
        sem = [{"lunes": date(2026, 9, 7) + timedelta(days=7 * i), "etiqueta": f"Semana {i}", "corta": f"{i}/9", "n": 3, "n_medidas": 3,
                "neto": v, "cobre": 0, "me_cobraron": 0, "pct": None} for i, v in enumerate([100000, -50000, 0, 250000])]
        sem[2]["n_medidas"] = 0
        svg = G.svg_barras_neto(sem)
        self.assertTrue(svg.startswith("<svg") and svg.endswith("</svg>"))
        self.assertIn('role="img"', svg)
        self.assertIn('aria-label="Neto semanal de las últimas 4 semanas"', svg)
        self.assertEqual(svg.count("<rect"), 4)
        self.assertIn("gc-b-pos", svg)
        self.assertIn("gc-b-neg", svg)
        self.assertIn("gc-b-vacia", svg)
        self.assertIn("$250.000", svg, "la última semana lleva su cifra")
        self.assertEqual(G.svg_barras_neto([]), "")

    def test_el_sparkline_necesita_dos_puntos_y_tolera_huecos(self):
        self.assertEqual(G.svg_sparkline([1]), "")
        self.assertEqual(G.svg_sparkline([None, None]), "")
        s = G.svg_sparkline([0, None, 5, 3])
        self.assertIn("<polyline", s)
        self.assertIn("<circle", s)
        self.assertIn("aria-label", s)
        self.assertIn("<polygon", G.svg_sparkline([2, 2, 2]), "una línea plana no divide por cero")

    def test_el_excel_trae_una_hoja_por_seccion_y_la_etiqueta_de_confidencial(self):
        import openpyxl
        d = D.tablero_demo()
        wb = openpyxl.load_workbook(io.BytesIO(G.excel_bytes(d)))
        self.assertEqual(wb.sheetnames, ["Resumen", "Tendencia", "Dónde se va la plata", "Garantías y coberturas", "A nombre de quién",
                                         "Técnicos y proveedores", "Reincidencia"])
        for ws in wb.worksheets:
            self.assertIn("Confidencial · solo gerencia", str(ws["A1"].value), ws.title)
            self.assertIn("Costeo de órdenes de trabajo", str(ws["A1"].value))
        res = {r[0].value: r[1].value for r in wb["Resumen"].iter_rows(min_row=5) if r[0].value}
        self.assertEqual(res["Cobré"], d["resumen"]["cobre"])
        self.assertEqual(res["Neto"], d["resumen"]["neto"])
        cab = [c.value for c in wb["A nombre de quién"][4]]
        for h in ("Originó", "Autorizó", "Cuándo autorizó", "Argumento guardado", "Lleva", "Ejecutó", "Área que asume", "Intención"):
            self.assertIn(h, cab)
        filas = [[c.value for c in r] for r in wb["A nombre de quién"].iter_rows(min_row=5)]
        self.assertTrue(any("Sin clasificar" in str(x) for f in filas for x in f), "lo que falta clasificar se dice")

    def test_el_excel_de_una_semana_vacia_no_revienta(self):
        import openpyxl
        d = G.armar_tablero(G.rango_periodo(HOY), [], agregar(), HOY, reinc={"cerradas": [], "candidatas": [], "dias": 60, "hoy_utc": datetime(2026, 10, 23)})
        wb = openpyxl.load_workbook(io.BytesIO(G.excel_bytes(d)))
        self.assertIn("Resumen", wb.sheetnames)


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestTableroDemo(unittest.TestCase):
    """El tablero completo con la cuenta real y los datos de mentira."""

    @classmethod
    def setUpClass(cls):
        cls.d = D.tablero_demo()

    def test_las_cifras_cuadran_entre_si(self):
        r = self.d["resumen"]
        self.assertEqual(r["neto"], r["ag"]["queda"])
        self.assertEqual(r["neto"], round(r["ganamos"] - r["perdimos"] - r["regalado_costo"] - r["cubiertas_costo"], 2))
        self.assertEqual(sum(r["respaldo"].values()), r["cobre"])
        self.assertAlmostEqual(r["pct"], r["neto"] / r["cobre"] * 100, delta=0.1)
        self.assertEqual(self.d["tendencia"][-1]["neto"], r["neto"], "la última semana de la tendencia es la semana elegida")
        self.assertEqual(r["pago_tecnicos"] + r["repuestos"], r["me_cobraron"])

    def test_trae_todas_las_secciones(self):
        for k in ("tarjetas", "veredicto", "tendencia", "svg_tendencia", "donde_se_va", "garantias", "atribucion", "tecnicos",
                  "reincidencia", "abiertas", "cxp", "perdidas", "sin_medir", "periodo", "clasificacion"):
            self.assertIn(k, self.d)
        self.assertEqual(self.d["clasificacion"], "Confidencial · solo gerencia")
        self.assertEqual(self.d["nombre"], "Costeo de órdenes de trabajo")

    def test_el_resumen_de_un_mes_y_un_trimestre_tambien_arma(self):
        for modo in ("mes", "trimestre"):
            d = D.tablero_demo(modo)
            self.assertEqual(d["periodo"]["modo"], modo)
            self.assertGreater(d["resumen"]["n_ot"], self.d["resumen"]["n_ot"])

    def test_los_externos_traen_su_reincidencia_y_cumplimiento(self):
        norte = next(g for g in self.d["tecnicos"]["grupos"] if g["nombre"] == "Servicios Norte Ltda.")
        self.assertEqual(norte["reincidencia"]["n_reincidentes"], 2)
        self.assertTrue(norte["reincidencia"]["revisar"])
        self.assertEqual(norte["cumplimiento"]["agendadas"], 3)
        propio = next(g for g in self.d["tecnicos"]["grupos"] if g["tipo"] == "interno")
        self.assertIsNone(propio.get("reincidencia"))


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestPlantilla(unittest.TestCase):
    """La plantilla se renderiza de verdad (Jinja) con datos de mentira y un base.html de mentira."""

    @classmethod
    def setUpClass(cls):
        import jinja2
        stub = ("<!doctype html><html><head>{% block head_extra %}{% endblock %}</head><body>"
                "{% block content %}{% endblock %}{% block scripts %}{% endblock %}</body></html>")
        cls.env = jinja2.Environment(loader=jinja2.ChoiceLoader([jinja2.DictLoader({"base.html": stub}),
                                                                 jinja2.FileSystemLoader(os.path.join(A.RAIZ, "templates"))]),
                                     autoescape=True)

        def url_for(endpoint, **kw):
            if endpoint == "static":
                return "/static/" + kw["filename"]
            if endpoint == "ot2_detalle":
                return "/ot/%s" % kw["vid"]
            q = "&".join("%s=%s" % (k, v) for k, v in sorted(kw.items()) if v is not None)
            return "/" + endpoint + ("?" + q if q else "")

        cls.env.globals["url_for"] = url_for
        cls.d = D.tablero_demo()
        cls.html = cls.render(cls.d)

    @classmethod
    def render(cls, d, modo="semana"):
        return cls.env.get_template("ot2/gerencial.html").render(d=d, G=G, p={"modo": modo, "dias": 60})

    def texto(self, html=None):
        plano = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html or self.html))
        return re.sub(r" ([,.:;)])", r"\1", plano)

    def test_renderiza_y_se_llama_costeo_de_ordenes_de_trabajo(self):
        self.assertIn("<h4><i class=\"bi bi-bullseye me-2\"></i>Costeo de órdenes de trabajo</h4>", self.html)
        self.assertIn("Costeo de órdenes de trabajo", re.search(r"<title>|{% block title", self.env.loader.get_source(self.env, "ot2/gerencial.html")[0]) and
                      self.env.loader.get_source(self.env, "ot2/gerencial.html")[0])

    def test_abre_con_las_cinco_tarjetas_y_luego_garantias_y_recien_despues_el_detalle(self):
        h = self.html
        orden = [h.index('id="gcT-cobre"'), h.index('id="gcT-me_cobraron"'), h.index('id="gcT-ganamos"'), h.index('id="gcT-perdimos"'),
                 h.index('id="gcT-neto"'), h.index('class="gc-veredicto'), h.index('id="gcGarantias"'), h.index('id="gcTraza"'),
                 h.index('id="gcTendencia"'), h.index('id="gcDonde"'), h.index('id="gcTecnicos"'), h.index('id="gcReinc"'),
                 h.index('id="gcPersonas"'), h.index('id="gcCamino"')]
        self.assertEqual(orden, sorted(orden), "tarjetas → veredicto → garantías (con su trazabilidad) → detalle")
        for titulo in ("Cobré", "Me cobraron", "Gané", "Perdí", "Neto"):
            self.assertIn(f"<span>{titulo}</span>", h)
        self.assertEqual(h.count("¿Cómo se calcula?"), 5, "cada cifra explica cómo se calcula")

    def test_cada_garantia_y_perdida_es_una_fila_con_los_cuatro_papeles_y_el_porque(self):
        h = self.html
        n_casos = len(self.d["atribucion"]["casos"])
        self.assertGreater(n_casos, 0)
        self.assertEqual(h.count('class="gc-caso gc-item"'), n_casos)
        for texto in ("¿Quién lo asume?", "¿Qué es y por qué?", "¿Quién lo autorizó?", "<dt>Originó</dt>", "<dt>Autorizó</dt>",
                      "<dt>Lleva</dt>", "<dt>Ejecutó</dt>", "Gerente Demo", "Garantía de fábrica: falló la tarjeta"):
            self.assertIn(texto, h)
        self.assertIn("Cargado al centro de costo: <b>", h)

    def test_lo_que_falta_clasificar_sale_en_ambar_suave_con_su_texto(self):
        h = self.html
        self.assertIn('class="gc-sc"', h)
        self.assertIn("Sin clasificar", h)
        self.assertIn("Sin autorización registrada", h)
        css = A.leer("static", "ot_gerencial.css")
        regla = re.search(r"\.gc-sc\{[^}]*\}", css).group(0)
        self.assertNotIn("dc2626", regla.lower())
        self.assertNotIn("var(--rojo)", regla)
        self.assertIn("var(--ambar-s)", regla)

    def test_todo_es_clickeable_hacia_la_ot(self):
        h = self.html
        ids_casos = {c["id"] for c in self.d["atribucion"]["casos"]}
        for i in ids_casos:
            self.assertIn('href="/ot/%s"' % i, h)
        self.assertIn('data-href="/ot/', h, "la tarjeta entera lleva a la OT")
        for g in self.d["tecnicos"]["grupos"][:1]:
            for o in g["ots"]:
                self.assertIn('href="/ot/%s"' % o["id"], h)
        for e in self.d["garantias"]["por_cliente"][:1]:
            for o in e["ots"]:
                self.assertIn('href="/ot/%s"' % o["id"], h)
        for c in self.d["reincidencia"]["casos"]:
            self.assertIn('href="/ot/%s"' % c["x_id"], h)
            self.assertIn('href="/ot/%s"' % c["y_id"], h)

    def test_nombres_completos_nada_truncado(self):
        largo = "Corporación Internacional de Centros Deportivos y Bienestar de la Región Metropolitana Sociedad por Acciones"
        d = D.tablero_demo()
        d["atribucion"]["casos"][0]["cliente"] = largo
        d["garantias"]["por_cliente"][0]["cliente"] = largo
        html = self.render(d)
        self.assertGreaterEqual(html.count(largo), 2)
        css = A.leer("static", "ot_gerencial.css")
        for malo in ("text-overflow", "line-clamp", "-webkit-line-clamp", "ellipsis"):
            self.assertNotIn(malo, css, "REGLA #15: nunca cortar un nombre")

    def test_dice_confidencial_y_no_muestra_fechas_iso_ni_alert_nativo(self):
        h = self.html
        self.assertGreaterEqual(h.count("Confidencial · solo gerencia"), 3)
        secciones = re.findall(r'<section class="gc-sec"[^>]*>\s*<header class="gc-sec-h">.*?</header>', h, flags=re.S)
        self.assertGreaterEqual(len(secciones), 7)
        for s in secciones:
            self.assertIn("gc-clasif", s, "cada sección lleva su etiqueta de clasificación")
        visible = self.texto()
        self.assertIsNone(re.search(r"\d{4}-\d{2}-\d{2}", visible), "ninguna fecha ISO a la vista")
        self.assertIsNone(re.search(r"\d{4}-\d{2}-\d{2}T", h))
        js = A.leer("static", "ot_gerencial.js")
        for malo in ("alert(", "confirm(", "prompt("):
            self.assertNotIn(malo, js)
            self.assertNotIn(malo, h)

    def test_las_listas_largas_se_paginan_como_etiquetas(self):
        h = self.html
        self.assertGreaterEqual(h.count("data-gc-pag"), 5)
        js = A.leer("static", "ot_gerencial.js")
        for pista in ("Mostrando ", "Anterior", "Siguiente", "Página ", "Por página"):
            self.assertIn(pista, js)

    def test_el_periodo_ofrece_semana_mes_trimestre_y_rango_y_navega(self):
        h = self.html
        for t in ("Semana", "Mes", "Trimestre", "Rango"):
            self.assertIn(">%s</a>" % t, h)
        self.assertIn("aria-label=\"Período anterior\"", h)
        self.assertNotIn("aria-label=\"Período siguiente\"", h, "la semana en curso no tiene siguiente")
        rango = self.render(D.tablero_demo("rango", None))
        self.assertNotIn('name="desde"', rango, "sin ser modo rango no hay formulario de fechas")

    def test_la_tendencia_trae_las_barras_del_neto_y_la_linea_del_margen(self):
        h = self.html
        self.assertIn('aria-label="Neto semanal de las últimas 12 semanas"', h)
        self.assertIn("Margen por semana", h)
        self.assertIn('aria-label="Margen por semana, últimas 12 semanas"', h)
        self.assertIn("Ver los números semana a semana", h)

    def test_la_reincidencia_es_un_indicador_para_revisar_no_una_sentencia(self):
        t = self.texto()
        self.assertIn("indicador para revisar, no una sentencia", t)
        self.assertIn("conviene revisarlo", t)
        self.assertIn("mismo equipo", t)
        for n in (30, 60, 90):
            self.assertIn(f"{n} días", t)

    def test_con_las_columnas_de_extension_vacias_explica_que_aun_no_se_registran(self):
        t = self.texto()
        self.assertIn("aún no se registra", t)

    def test_un_periodo_sin_datos_se_ve_con_mensajes_claros(self):
        d = G.armar_tablero(G.rango_periodo(HOY), [], agregar(), HOY, reinc={"cerradas": [], "candidatas": [], "dias": 60, "hoy_utc": datetime(2026, 10, 23)},
                            meta={"dias_reinc": 60, "fallos": [], "truncado": False, "generado": "x"})
        t = self.texto(self.render(d))
        self.assertIn("Todavía no hay OT medidas en este período", t)
        self.assertIn("En este período no hubo OT sin cobrar", t)
        self.assertIn("Sin garantías, coberturas ni OT cobradas a pérdida en este período", t)
        self.assertIn("Ningún técnico externo cerró OT en los últimos 6 meses", t)

    def test_los_avisos_de_lectura_incompleta_y_ot_sin_medir_se_muestran(self):
        d = D.tablero_demo()
        d["meta"]["fallos"] = ["reincidencia"]
        d["meta"]["truncado"] = True
        t = self.texto(self.render(d))
        self.assertIn("No se pudo leer: reincidencia", t)
        self.assertIn("Se leyeron las 6000 OT cerradas más recientes", t)
        self.assertIn("OT cerradas no están en las cifras", self.texto())

    def test_css_cumple_44px_16px_y_paleta_ilus(self):
        css = A.leer("static", "ot_gerencial.css")
        for pista in ("min-height:44px", "font-size:16px", "#dc2626", "#0a0a0a", "prefers-reduced-motion", "focus-visible"):
            self.assertIn(pista, css)


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestAppBloque(unittest.TestCase):
    """El bloque de app.py: contiguo, solo lectura, solo superadmin, con sus puntos de extensión."""

    @classmethod
    def setUpClass(cls):
        cls.src = A.app_src()
        cls.bloque = A.bloque_gerencial()

    def test_el_bloque_es_uno_solo_y_va_justo_antes_del_hub(self):
        self.assertEqual(self.src.count("#  📊 COSTEO DE ÓRDENES DE TRABAJO — el tablero gerencial"), 1)
        fin = self.src.index("#  HUB TOP-LEVEL — Dashboard / Órdenes de Trabajo / Plantillas")
        ini = self.src.index("#  📊 COSTEO DE ÓRDENES DE TRABAJO — el tablero gerencial")
        self.assertLess(ini, fin)
        self.assertLess(fin - ini, 40000)
        self.assertEqual(len(re.findall(r"^def _ger_", self.bloque, re.M)) >= 12, True)

    def test_las_rutas_exigen_gestion_y_superadmin_en_este_orden(self):
        for ruta, fn in (('@app.route("/ot/gerencial", methods=["GET"])', "ot_gerencial_pagina"),
                         ('@app.route("/ot/gerencial/excel", methods=["GET"])', "ot_gerencial_excel")):
            i = self.bloque.index(ruta)
            cab = self.bloque[i:self.bloque.index("def " + fn, i)]
            self.assertIn("@_mant_required\n@_no_tecnico\n@_ger_solo_superadmin\n", cab, fn)
        self.assertEqual(self.bloque.count("@app.route("), 2, "el costeo solo agrega estas dos rutas")

    def test_la_puerta_de_superadmin_se_resuelve_al_llamar_porque_el_decorador_vive_mas_abajo(self):
        self.assertLess(self.src.index("def _ger_solo_superadmin"), self.src.index("def _ot_solo_superadmin"),
                        "el bloque va ANTES de _ot_solo_superadmin: por eso lo usa de forma perezosa")
        self.assertIn("return _ot_solo_superadmin(view)(*a, **kw)", self.bloque)

    def test_el_bloque_no_escribe_nada_en_ninguna_parte(self):
        sin_comentarios = "\n".join(l for l in self.bloque.split("\n") if not l.strip().startswith("#"))
        for verbo in ("INSERT", "UPDATE", "DELETE", "ALTER", "CREATE", "DROP", "TRUNCATE", "REPLACE", "GRANT", "MERGE"):
            self.assertIsNone(re.search(r"\b%s\b" % verbo, sin_comentarios), "el costeo es solo lectura: " + verbo)
        for llamada in ("mysql_execute", "mysql_execute_returning_rowcount", "conn.commit", ".commit(", "requests.post", "requests.put",
                        "_random_sql_query", "_checkwms_get", "erp_engine", "fetch_document", "_send_ilus_email", "_mant_log"):
            self.assertNotIn(llamada, self.bloque, llamada)

    def test_no_toca_lo_que_el_anexo_unico_esta_editando(self):
        for prohibido in ("_ot_anexo_", "ot2_api_anexo", "_anexo_bloquea_ot", "mant_ot_aprobar_cierre"):
            self.assertNotIn(prohibido, self.bloque)

    def test_cada_columna_que_lee_existe_en_su_tabla(self):
        """REGLA #5: nada de adivinar nombres de columna. Cada `v.<col>` del bloque debe estar declarada en mant_visitas."""
        usadas = set(re.findall(r"\bv\.(\w+)", self.bloque)) - {"id", "lower"}
        creada = self.src[self.src.index("CREATE TABLE IF NOT EXISTS mant_visitas ("):][:1800]
        declaradas = set(re.findall(r"\n\s+(\w+)\s+[A-Z]", creada))                 # columnas del CREATE TABLE original
        declaradas |= set(re.findall(r"ALTER TABLE mant_visitas ADD COLUMN (\w+)", self.src))   # y de sus ALTER
        i_cc = self.src.index('("centro_costo",  "VARCHAR(20) NULL')
        declaradas |= set(re.findall(r'\("(\w+)",\s+"[A-Z]', self.src[i_cc - 1500:i_cc + 1500]))   # lista de columnas de _ensure_mant_ot2_tables
        for col in sorted(usadas):
            self.assertIn(col, declaradas, "mant_visitas no declara la columna " + col)
        tk = A.leer("tickets_module.py")
        creado = tk[tk.index("CREATE TABLE IF NOT EXISTS tk_tickets"):][:3600]
        for col in sorted(set(re.findall(r"\btk\.(\w+)", self.bloque))):
            self.assertRegex(creado, r"\n\s+%s\s" % col, "tk_tickets." + col)
        te = self.src[self.src.index("CREATE TABLE IF NOT EXISTS mant_tecnicos_externos ("):][:3200]
        for col in sorted(set(re.findall(r"\bte\.(\w+)", self.bloque)) - {"es_prueba"}):
            self.assertRegex(te, r"\n\s+%s\s" % col, "mant_tecnicos_externos." + col)
        self.assertRegex(self.src, r"ALTER TABLE mant_tecnicos_externos ADD COLUMN es_prueba", "te.es_prueba")
        autor = self.src[self.src.index("CREATE TABLE IF NOT EXISTS mant_ot_autorizaciones"):][:2400]
        for col in ("visita_id", "visita_creada_id", "tipo", "estado", "motivo", "argumento", "solicitado_por_nombre", "solicitado_at",
                    "resuelto_por_nombre", "resuelto_at", "entidad"):
            self.assertRegex(autor, r"\n\s+%s\s" % col, "mant_ot_autorizaciones." + col)
        fp = self.src[self.src.index("CREATE TABLE IF NOT EXISTS mant_facturas_proveedor ("):][:1800]
        for col in ("estado_pago", "monto_total", "pagada_at"):
            self.assertRegex(fp, r"\n\s+%s\s" % col, "mant_facturas_proveedor." + col)
        eq = self.src[self.src.index("CREATE TABLE IF NOT EXISTS mant_visita_equipos"):][:900]
        for col in ("visita_id", "maquina_id"):
            self.assertRegex(eq, r"\n\s+%s\s" % col, "mant_visita_equipos." + col)
        self.assertIn("responsable_nombre", self.src[self.src.index("def _ensure_ot_responsable"):][:1200])

    def test_el_panel_ofrece_el_boton_solo_al_superadmin(self):
        panel = A.leer("templates", "ot2", "panel.html")
        i = panel.index("{% if es_superadmin_ot %}")
        j = panel.index("</nav>", i)
        self.assertIn("url_for('ot_gerencial_pagina')", panel[i:j])
        self.assertIn("Costeo de OT", panel[i:j])
        menu = A.leer("templates", "base.html") if os.path.exists(os.path.join(A.RAIZ, "templates", "base.html")) else ""
        self.assertNotIn("ot_gerencial", menu, "no va en el menú lateral")


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
class TestAppEjecutado(unittest.TestCase):
    """Ejecuta el bloque REAL de app.py con una base de mentira: todo SELECT, el costeo arma, la caché sirve y un técnico no entra."""

    def montar(self, cols_ext=(), filas=None, ctxs=None, es_superadmin=True, hoy=HOY):
        src = A.bloque_gerencial()
        vs, cx = D.dataset(hoy) if filas is None else (filas, ctxs or {})
        auts = [dict(a, visita_id=None, visita_creada_id=vid) for vid, c in cx.items() for a in c.get("autorizaciones", [])]
        clasif = {vid: c["clasif"] for vid, c in cx.items() if c.get("clasif")}
        cuenta = A.cuenta_real()

        def responder(sql, params):
            if "information_schema.COLUMNS" in sql:
                return [{"COLUMN_NAME": c} for c in cols_ext]
            if "SELECT username, nombre FROM app_users" in sql:
                return [{"username": k, "nombre": v} for k, v in D.USUARIOS.items()]
            if "FROM mant_ot_autorizaciones" in sql:
                return auts
            if "FROM mant_visita_equipos" in sql:
                return [{"visita_id": 1, "maquina_id": 1}]
            if "FROM mant_facturas_proveedor" in sql:
                return [{"n": 2, "monto": 500000}]
            if "tk.id AS tk_id" in sql:
                ini, fin = params[0], params[1]
                return [v for v in vs if ini <= v["cerrada_at"] < fin]
            if "v.fecha_programada >= %s" in sql:
                return []
            if "v.cerrada_at >= %s AND v.cliente_id IS NOT NULL" in sql:
                return [dict(v) for v in vs if v["proveedor_tipo"] == "externo"][:5]
            if "v.created_at >= %s" in sql:
                return []
            if "v.estado NOT IN ('cerrada','cancelada','anulada')" in sql:
                return [A.ot_cobrada(900, 100000, 60000)]
            raise AssertionError("SQL no previsto: " + sql[:160])

        db = A.BaseFalsa(responder)
        llamadas = {"lote": 0}

        def fin_lote(rows):
            llamadas["lote"] += 1
            return {int(r["id"]): cuenta["_ot_finanzas"](r) for r in rows}

        def clasif_lote(ids):
            return {i: clasif[i] for i in ids if i in clasif}

        amb = {
            "__name__": "app_bloque", "app": types.SimpleNamespace(route=lambda *a, **k: (lambda f: f)), "wraps": functools.wraps, "time": time,
            "io": io, "datetime": datetime, "timedelta": timedelta, "timezone": timezone, "date": date,
            "_mant_required": lambda f: f, "_no_tecnico": lambda f: f,
            "mysql_fetchall": db.fetchall, "mysql_fetchone": db.fetchone, "_ot_resp_ok": lambda: True,
            "_ot_fin_cols_sql": lambda alias="v": ", ".join(f"{alias}.{c}" for c in (
                "modalidad_cobro", "cubierto_por", "tipo", "cliente_id", "contrato_id", "costo", "zz_monto", "cobro_cero_motivo", "zz_codigo",
                "zz_envio_monto", "valor_origen", "costo_proveedor", "costo_despacho", "proveedor_tipo", "valorizado_clp", "valorizado_fuente",
                "zz_motivo_manual", "ot_madre_id", "ot_relacion")) + f", 0 AS contrato_real, 0 AS madre_valida",
            "_ot_fin_lote": fin_lote, "_ot_finanzas": cuenta["_ot_finanzas"], "_ot_fin_agregado": cuenta["_ot_fin_agregado"],
            "_ot_fin_clasif_lote": clasif_lote,
            "_now_chile": lambda: datetime(hoy.year, hoy.month, hoy.day, 15, 0, tzinfo=timezone(timedelta(hours=-3))),
            "_TIPO_OT_LABEL": {"instalacion": "Instalación", "correctiva": "Mantención correctiva"},
            "_OT2_CENTROS_COSTO": (("sstt", "Servicio Técnico"), ("logistica", "Logística"), ("comercial", "Comercial"), ("marketing", "Marketing")),
            "_ot_aut_es_superadmin": lambda: es_superadmin,
            "_friendly_error_page": lambda t, m, c: ("error", t, c),
            "render_template": lambda plantilla, **kw: ("render", plantilla, kw),
            "request": types.SimpleNamespace(args={}, path="/ot/gerencial"),
            "flash": lambda *a, **k: amb["_avisos"].append(a), "_avisos": [],
            "redirect": lambda u: ("redirect", u), "url_for": lambda e, **k: "/" + e,
            "jsonify": lambda d: ("json", d), "send_file": lambda *a, **k: ("archivo", k),
        }
        # _ot_solo_superadmin REAL (se mira cómo reacciona a un técnico)
        exec(compile(A.bloque(A.app_src(), "_ot_solo_superadmin"), "<app:_ot_solo_superadmin>", "exec"), amb)
        exec(compile(src, "<app:bloque_gerencial>", "exec"), amb)
        return amb, db, llamadas

    def p(self, **kw):
        base = {"modo": "semana", "ancla": None, "desde": None, "hasta": None, "dias": 60}
        base.update(kw)
        return base

    def test_arma_el_tablero_leyendo_solo_con_select(self):
        amb, db, _ = self.montar()
        d = amb["_ger_datos"](self.p())
        self.assertTrue(db.sqls, "tuvo que leer algo")
        self.assertTrue(all(s.split(" ", 1)[0].upper() == "SELECT" for s in db.sqls))
        self.assertEqual(d["periodo"]["desde"], date(2026, 10, 19))
        self.assertGreater(d["resumen"]["n_ot"], 0)
        self.assertEqual([t["titulo"] for t in d["tarjetas"]], ["Cobré", "Me cobraron", "Gané", "Perdí", "Neto"])
        self.assertEqual(d["resumen"]["neto"], d["resumen"]["ag"]["queda"])
        self.assertIsNotNone(d["cxp"])
        self.assertEqual(d["cxp"]["pendiente_n"], 2)

    def test_los_limites_del_periodo_viajan_en_utc_como_parametros(self):
        amb, db, _ = self.montar()
        capturado = []
        base = db.responder

        def espia(sql, params):
            if "tk.id AS tk_id" in sql:
                capturado.append(params)
            return base(sql, params)

        db.responder = espia
        amb["_ger_cargar"]("semana", None, None, None, 60)
        ini, fin, limite = capturado[0]
        self.assertLessEqual(ini, datetime(2026, 5, 1, 4, 0), "lee 6 meses hacia atrás")
        self.assertEqual(fin, datetime(2026, 10, 26, 3, 0), "hasta el fin de la semana (domingo 23:59 Chile)")
        self.assertEqual(limite, 6000)

    def test_los_datos_viajan_por_parametros_no_pegados_en_el_sql(self):
        amb, db, _ = self.montar()
        amb["_ger_datos"](self.p(modo="rango", desde=date(2026, 10, 1), hasta=date(2026, 10, 10)))
        for s in db.sqls:
            self.assertNotIn("2026-10", s, "ninguna fecha pegada en el texto del SQL")

    def test_la_cache_de_90_segundos_evita_releer(self):
        amb, db, llamadas = self.montar()
        amb["_ger_datos"](self.p())
        n = len(db.sqls)
        amb["_ger_datos"](self.p())
        self.assertEqual(len(db.sqls), n, "la segunda vez sale de la memoria")
        amb["_ger_datos"](self.p(), forzar=True)
        self.assertGreater(len(db.sqls), n, "Actualizar fuerza la lectura")
        amb["_ger_datos"](self.p(modo="mes"))
        self.assertGreater(len(db.sqls), n)
        self.assertEqual(amb["_GER_CACHE_SEG"], 90)

    def test_sin_columnas_de_extension_los_dos_puntos_devuelven_none(self):
        amb, db, _ = self.montar(cols_ext=())
        self.assertIsNone(amb["_ger_columna_intencion"]())
        self.assertIsNone(amb["_ger_columna_asume"]())
        sql = amb["_ger_select_cerradas"]()
        self.assertIn("NULL AS intencion", sql)
        self.assertIn("NULL AS asume_area", sql)
        d = amb["_ger_cargar"]("semana", None, None, None, 60)
        self.assertEqual(d["garantias"]["clasificadas"], 0)
        self.assertIsNone(d["meta"]["col_intencion"])

    def test_cuando_las_columnas_existan_las_usa_sin_romper_nada(self):
        amb, db, _ = self.montar(cols_ext=("cobro_cero_intencion", "cobro_cero_asume_area", "numero_ot"))
        self.assertEqual(amb["_ger_columna_intencion"](), "cobro_cero_intencion")
        self.assertEqual(amb["_ger_columna_asume"](), "cobro_cero_asume_area")
        sql = amb["_ger_select_cerradas"]()
        self.assertIn("v.cobro_cero_intencion AS intencion", sql)
        self.assertIn("v.cobro_cero_asume_area AS asume_area", sql)

    def test_el_codigo_de_un_centro_de_costo_se_muestra_con_su_nombre(self):
        amb, _, _ = self.montar()
        self.assertEqual(amb["_ger_etiqueta_area"]("marketing"), "Marketing")
        self.assertEqual(amb["_ger_etiqueta_area"]("Comercial"), "Comercial")
        self.assertEqual(amb["_ger_etiqueta_area"]("Gerencia General"), "Gerencia General")
        self.assertIsNone(amb["_ger_etiqueta_area"]("  "))
        self.assertIsNone(amb["_ger_etiqueta_area"](None))

    def test_un_tecnico_no_entra_ni_a_la_pagina_ni_al_excel(self):
        amb, db, _ = self.montar(es_superadmin=False)
        for vista in ("ot_gerencial_pagina", "ot_gerencial_excel"):
            n = len(db.sqls)
            r = amb[vista]()
            self.assertEqual(r, ("redirect", "/ot2_panel"), vista)
            self.assertEqual(len(db.sqls), n, "ni siquiera lee la base: " + vista)
        self.assertTrue(amb["_avisos"])
        # y una ruta de API respondería JSON 403
        amb["request"].path = "/ot/api/gerencial"
        r = amb["_ger_solo_superadmin"](lambda: "datos")()
        self.assertEqual(r[1], 403, "una ruta de API responde JSON 403 a un técnico")
        self.assertEqual(r[0][0], "json")
        self.assertFalse(r[0][1]["ok"])

    def test_el_superadmin_si_entra_y_recibe_la_plantilla_con_el_tablero(self):
        amb, db, _ = self.montar(es_superadmin=True)
        r = amb["ot_gerencial_pagina"]()
        self.assertEqual(r[0], "render")
        self.assertEqual(r[1], "ot2/gerencial.html")
        self.assertIn("d", r[2])
        self.assertEqual(r[2]["d"]["nombre"], "Costeo de órdenes de trabajo")
        x = amb["ot_gerencial_excel"]()
        self.assertEqual(x[0], "archivo")
        self.assertTrue(x[1]["download_name"].startswith("Costeo_OT_19-10-2026_al_25-10-2026"))

    def test_una_falla_de_lectura_se_dice_y_no_se_inventa(self):
        amb, db, _ = self.montar()
        base = db.responder

        def rota(sql, params):
            if "FROM mant_ot_autorizaciones" in sql:
                raise RuntimeError("se cayó")
            return base(sql, params)

        db.responder = rota
        d = amb["_ger_cargar"]("semana", None, None, None, 60)
        self.assertIn("autorizaciones", d["meta"]["fallos"])

    def test_una_ot_con_dos_tickets_se_cuenta_una_sola_vez(self):
        vs = [A.ot_cobrada(1, 500000, 300000, **en_semana(0)), A.ot_cobrada(1, 500000, 300000, **en_semana(0))]
        amb, db, _ = self.montar(filas=vs, ctxs={})
        d = amb["_ger_cargar"]("semana", None, None, None, 60)
        self.assertEqual(d["resumen"]["n_ot"], 1)
        self.assertEqual(d["resumen"]["cobre"], 500000)


if __name__ == "__main__":
    unittest.main()
