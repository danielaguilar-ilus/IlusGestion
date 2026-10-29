# -*- coding: utf-8 -*-
"""🔎 «Costeo de OT · Explorar» — el tablero para disertar ante la gerencia (FASE 1, Daniel 2026-10-29).

Pedido de Daniel: «En el mes tuvimos 10 millones de pesos: lo pueda desglosar en centro de costo, técnico, rango de tiempo,
si es trabajo interno, si es trabajo externo, quiénes son los proveedores (Mili, Felca, Daniel Pulgar = DAP…). Multifiltro,
algo súper estructurado, estilo Power BI: presionas un objeto y se filtra todo. Filtrar por día, por semana, por número de
semana, por mes y por año. Quiero saber cuánto cobra un técnico en un tiempo determinado y quién ha gastado más en cada
centro de costo (qué técnico).»

Cómo funciona: el servidor arma UNA fila plana por OT con la cuenta YA HECHA (la de su ficha, `_ot_finanzas`, normalizada
por `ot_gerencial.armar_fila`) y el navegador (static/ot_costeo_explorar.js) cruza los filtros. Este módulo es el lado
Python de ese contrato:
  · `leer_parametros`   → rango de fechas, qué fecha manda y si entran las OT abiertas (nada se confía del navegador).
  · `fila`              → la fila plana de una OT (fechas en hora Chile, dimensiones con su clave y su nombre, montos).
  · `documento_hoy`     → con qué documento se respalda HOY lo cobrado (no cómo nació la OT: eso es `nac`, aparte).
  · `armar_dataset`     → {cols, rows, dims, meta} (por columnas: pesa poco y viaja comprimido).
  · `kpis` / `filtrar`  → ESPEJO del navegador (misma suma que `_ot_fin_agregado`; mismo filtro), para el Excel de lo
                          filtrado y para las pruebas que comparan Python con JS.
  · `excel_bytes`       → el Excel de lo filtrado, con fechas reales de Excel (dd/mm/aaaa) y Año / Mes / N° semana / Día.

MÓDULO PURO (REGLAS #9 y #18): sin Flask, sin base, sin Random ni Check. La plata NO se recalcula aquí: cada monto sale de
`_ot_finanzas` tal cual; la suma de varias OT sigue la regla única del proyecto (interno y contrato aparte, las que no se
pueden medir fuera y contadas, nunca un $0 ni un 100 % inventado). La política de margen (REGLA #57) solo pinta.

Confidencial · solo gerencia (REGLAS #19 y #26): todo esto es plata y proveedores; la ruta exige superadmin.
"""
from __future__ import annotations

import io
import json
from datetime import date, datetime, timedelta

import ot_gerencial as G
import ot_margen_politica as POL
import ot_saldo_servicio as SALDO

NOMBRE = "Explorar"
CLASIFICACION = G.CLASIFICACION

# ── Qué fecha manda ────────────────────────────────────────────────────────────────────────────────────────────────
FECHAS = ("cierre", "ejecucion", "creacion", "programada")
FECHA_TXT = {"cierre": "Fecha de cierre", "ejecucion": "Fecha de ejecución", "creacion": "Fecha de creación",
             "programada": "Fecha programada"}
# Columna de mant_visitas de cada una (nombres fijos de código, nunca del navegador).
FECHA_COLUMNA = {"cierre": "cerrada_at", "ejecucion": "fecha_realizada", "creacion": "created_at",
                 "programada": "fecha_programada"}
# cerrada_at y created_at son DATETIME en UTC (NOW() de MySQL): se pasan a hora Chile. Las otras dos son DATE.
FECHA_ES_INSTANTE = {"cierre": True, "creacion": True, "ejecucion": False, "programada": False}
ALCANCES = ("cerradas", "todas")
ALCANCE_TXT = {"cerradas": "Solo OT cerradas", "todas": "Todas las OT (también las abiertas)"}
MAX_RANGO_DIAS = 1096          # tres años
MAX_FUTURO_DIAS = 366          # la fecha programada puede estar adelante
LIMITE_FILAS = 8000            # techo de OT que se leen (período + período anterior); si se toca, la pantalla lo dice

# ── Columnas del dataset (una fila = una OT) ───────────────────────────────────────────────────────────────────────
DIMS = ("estado", "cli", "reg", "tec", "ttipo", "emp", "cc", "tipo", "cob", "int", "doc", "nac", "banda", "clase")
DIMS_TIEMPO = ("anio", "mes", "sem", "nsem", "dia")
FILTRABLES = DIMS + DIMS_TIEMPO
DIM_TXT = {
    "estado": "Estado", "cli": "Cliente", "reg": "Región", "tec": "Técnico", "ttipo": "Interno o externo",
    "emp": "Empresa o proveedor", "cc": "Centro de costo", "tipo": "Tipo de OT", "cob": "Cobertura", "int": "Intención",
    "doc": "Documento hoy", "nac": "Cómo nació la OT", "banda": "Margen según la política", "clase": "En la cuenta",
    "anio": "Año", "mes": "Mes", "sem": "Semana", "nsem": "N° de semana", "dia": "Día",
}
COLS = ("id", "numero", "titulo", "estado", "f_creada", "f_programada", "f_realizada", "f_cerrada",
        "dia", "sem", "nsem", "mes", "anio", "prev") + DIMS[1:] + (
        "cobre_serv", "cobre_desp", "cobre", "me_tec", "me_desp", "me_rep", "me_total", "a_pagar", "queda", "pct",
        "val", "cobra", "falta_tec")

# ── Documento que respalda HOY lo cobrado ───────────────────────────────────────────────────────────────────────────
FACTURAS = ("FCV", "FCE")
BOLETAS = ("BLV", "BLE")
DOC_ORDEN = ("factura", "boleta", "nota_venta", "cotizacion", "otro_doc", "autorizado_sin_doc", "sin_documento",
             "no_aplica")
DOC_TXT = {
    "factura": "Factura", "boleta": "Boleta", "nota_venta": "Nota de venta (falta la factura)",
    "cotizacion": "Cotización (falta la factura)", "otro_doc": "Guía u otro documento (falta la factura)",
    "autorizado_sin_doc": "Sin documento, autorizado", "sin_documento": "Sin documento",
    "no_aplica": "No aplica (no se cobra)",
}
AUT_SIN_DOC = ("crear_sin_documento", "cerrar_sin_documento")

# ── Otras etiquetas ─────────────────────────────────────────────────────────────────────────────────────────────────
BANDA_TXT = dict(POL.BANDA_TXT, sin_medir="Sin medir (falta un costo o el cobro)", no_cobra="No se cobra (va aparte)")
BANDA_ORDEN = POL.BANDAS + ("sin_medir", "no_cobra")
CLASE_TXT = {"ganada": "Se cobra y dejó margen", "perdida": "Se cobra y quedó a pérdida",
             "cobertura": "Garantía, regalía, arriendo o cortesía", "cubierta": "Desinstalación cubierta",
             "aparte": "Trabajo interno o contrato (aparte)", "sin_medir": "Sin medir (falta un costo)"}
ESTADO_TXT = {"cerrada": "Cerrada", "completada": "Completada", "programada": "Programada", "en_curso": "En curso",
              "en_ejecucion": "En ejecución", "firmada_tecnico": "Firmada por el técnico",
              "pendiente_aprobacion": "Pendiente de aprobación", "reagendada": "Reagendada", "borrador": "Borrador",
              "pendiente": "Pendiente", "asignada": "Asignada"}
COBERTURAS_REGALADAS = G.COBERTURAS_REGALADAS
SIN = {"cli": "interno", "reg": "sin_region", "cc": "sin_centro", "tipo": "sin_tipo", "int": "sin_clasificar",
       "nac": "sin_clasif", "tec": "sin", "emp": "sin"}
DIAS_SEMANA = ("lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo")


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Parámetros (nada se confía del navegador)
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def leer_parametros(args, hoy):
    """?desde=&hasta=&fecha=&alcance= → el rango a mirar y el período anterior del mismo largo (para comparar).
    Defectos: los últimos 12 meses (desde el día 1 del mes de hace 11 meses hasta hoy), fecha de cierre, solo cerradas.
    Un rango al revés se da vuelta; uno de más de 3 años se acorta; `hasta` no pasa de un año hacia adelante."""
    fecha = (args.get("fecha") or "cierre").strip().lower()
    if fecha not in FECHAS:
        fecha = "cierre"
    alcance = (args.get("alcance") or "cerradas").strip().lower()
    if alcance not in ALCANCES:
        alcance = "cerradas"
    desde, hasta = G._fecha(args.get("desde")), G._fecha(args.get("hasta"))
    tope = hoy + timedelta(days=MAX_FUTURO_DIAS)
    if not hasta:
        hasta = hoy
    if hasta > tope:
        hasta = tope
    if not desde:
        a, m = G._mes_mas(hasta.year, hasta.month, -11)
        desde = date(a, m, 1)
    if desde > hasta:
        desde, hasta = hasta, desde
    if (hasta - desde).days + 1 > MAX_RANGO_DIAS:
        desde = hasta - timedelta(days=MAX_RANGO_DIAS - 1)
    n = (hasta - desde).days + 1
    return {"desde": desde, "hasta": hasta, "fecha": fecha, "alcance": alcance, "dias": n, "hoy": hoy,
            "carga_desde": desde - timedelta(days=n), "previo_hasta": desde - timedelta(days=1)}


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Fechas en hora Chile
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def fecha_local(valor, instante):
    """La fecha (date) en hora Chile. `instante`=True: DATETIME naive en UTC (cerrada_at, created_at). False: DATE."""
    if valor is None or valor == "":
        return None
    if instante:
        if isinstance(valor, str):
            try:
                valor = datetime.fromisoformat(valor.strip().replace("Z", ""))
            except ValueError:
                return None
        c = G.a_chile(valor)
        return c.date() if c else None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return G._fecha(str(valor))


def iso(d):
    return d.isoformat() if d else ""


def derivar(d):
    """Día, semana (su lunes), N° de semana ISO, mes y año de una fecha. Todo texto ordenable (claves del filtro)."""
    if not d:
        return {"dia": "", "sem": "", "nsem": "", "mes": "", "anio": ""}
    lunes = d - timedelta(days=d.weekday())
    return {"dia": d.isoformat(), "sem": lunes.isoformat(), "nsem": "%02d" % d.isocalendar()[1],
            "mes": "%04d-%02d" % (d.year, d.month), "anio": "%04d" % d.year}


def etiqueta_tiempo(dim, clave):
    """El nombre legible de una clave de tiempo (el navegador tiene el mismo: etiquetaTiempo)."""
    if not clave:
        return "Sin fecha"
    if dim == "anio":
        return clave
    if dim == "mes":
        a, m = clave.split("-")
        return f"{G.MESES_ES[int(m) - 1].capitalize()} {a}"
    if dim == "nsem":
        return f"Semana {int(clave)}"
    d = G._fecha(clave)
    if not d:
        return clave
    if dim == "sem":
        dom = d + timedelta(days=6)
        return f"Semana {d.isocalendar()[1]} · {G._etiqueta_semana(d)} de {dom.year}"
    return f"{DIAS_SEMANA[d.weekday()]} {d.strftime('%d/%m/%Y')}"


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Documento de HOY (no cómo nació)
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def tipo_doc(d):
    """factura | boleta | nota_venta | cotizacion | otro_doc | None (sin tipo) de un documento {origen, tido}."""
    if (d.get("origen") or "erp") == "cotizacion":
        return "cotizacion"
    t = str(d.get("tido") or "").strip().upper()
    if not t:
        return None
    if t in FACTURAS:
        return "factura"
    if t in BOLETAS:
        return "boleta"
    if t in SALDO.NOTAS_VENTA:
        return "nota_venta"
    if t in SALDO.COTIZACIONES:
        return "cotizacion"
    return "otro_doc"


def docs_con_principal(docs, v):
    """Los documentos de la tabla puente + el principal de mant_visitas SOLO si se validó contra Random (asociar-factura:
    factura_asociada_por o factura_rut_ok), igual que _ot_puerta_docs_lote. Sin repetir el mismo documento."""
    out = [dict(d) for d in (docs or [])]
    t = str(v.get("factura_tido") or "").strip().upper()
    n = str(v.get("factura_nudo") or "").strip()
    validado = bool(str(v.get("factura_asociada_por") or "").strip()) or v.get("factura_rut_ok") is not None
    if t and n and validado:
        k = (t, n.lstrip("0"))
        if not any((str(d.get("tido") or "").strip().upper(), str(d.get("nudo") or "").strip().lstrip("0")) == k
                   for d in out):
            out.append({"origen": "erp", "tido": t, "nudo": n, "es_cobro": True})
    return out


def documento_hoy(cobra, docs, auts=()):
    """Con qué documento se respalda HOY lo cobrado. Manda el documento que lleva el cobro (es_cobro) y, entre ellos, el
    más firme: factura > boleta > nota de venta > cotización > otro. Una OT que nació con nota de venta y ya tiene su
    factura cuenta como FACTURA. Sin documento: «autorizado» si hay una autorización aprobada de crear o cerrar sin
    documento. Una OT que no se cobra (garantía, interno, contrato…) → «no aplica»."""
    if not cobra:
        return "no_aplica"
    con_tipo = [d for d in (docs or []) if tipo_doc(d)]
    de_cobro = [d for d in con_tipo if d.get("es_cobro", True) not in (False, 0, "0")]
    for grupo in (de_cobro, con_tipo):
        tipos = {tipo_doc(d) for d in grupo}
        for k in ("factura", "boleta", "nota_venta", "cotizacion", "otro_doc"):
            if k in tipos:
                return k
    for a in (auts or []):
        if str(a.get("tipo") or "") in AUT_SIN_DOC and str(a.get("estado") or "aprobada").lower() == "aprobada":
            return "autorizado_sin_doc"
    return "sin_documento"


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Técnico y empresa
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def tecnico_de(f):
    """(clave, nombre) de la PERSONA que ejecutó: el usuario asignado. Un externo sin usuario queda con su empresa."""
    e = f["ejecutor"]
    if e["tipo"] == "sin_asignar":
        return SIN["tec"], "Sin técnico asignado"
    nombre = (e.get("persona") or e.get("empresa") or e.get("nombre") or "Técnico sin nombre").strip()
    if e.get("user_id"):
        return "u:%s" % e["user_id"], nombre
    return "n:" + nombre.lower(), nombre


def empresa_de(f):
    """(clave, nombre) de la EMPRESA: los externos se agrupan por empresa (Mili, Felca, DAP…), igual que
    ot_gerencial.armar_ejecutor; los técnicos propios van juntos (igual que «¿A dónde se va la plata?» del Resumen)."""
    e = f["ejecutor"]
    if e["tipo"] == "externo":
        return e["clave"], e["nombre"]
    if e["tipo"] == "sin_asignar":
        return SIN["emp"], "Sin técnico asignado"
    return "propios", "Técnicos propios"


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Una fila = una OT
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def _r2(x):
    return None if x is None else round(float(x), 2)


def banda_de(clase, pct, queda):
    """La banda de margen de UNA OT. Solo se pinta lo que se cobra y se puede medir; el resto se dice, no se inventa."""
    if clase == "sin_medir":
        return "sin_medir"
    if clase in ("aparte", "cobertura", "cubierta"):
        return "no_cobra"
    return POL.banda(pct, queda) or "sin_medir"


def fila(v, f, docs=(), auts=(), p=None):
    """La fila plana de una OT. `v`: la fila de mant_visitas con sus joins; `f`: ot_gerencial.armar_fila(v, _ot_finanzas);
    `docs`: sus documentos (tabla puente); `auts`: sus autorizaciones APROBADAS; `p`: leer_parametros. Devuelve un dict con
    las claves de COLS y `_lab` = {dimensión: nombre legible}."""
    p = p or {"fecha": "cierre", "desde": None}
    fin = f["fin"]
    me = fin.get("me_cobraron") or {}
    cob_d = fin.get("cobre") or {}
    fechas = {k: fecha_local(v.get(col), FECHA_ES_INSTANTE[k]) for k, col in FECHA_COLUMNA.items()}
    elegida = fechas[p["fecha"]]
    tiempo = derivar(elegida)
    tec_k, tec_n = tecnico_de(f)
    emp_k, emp_n = empresa_de(f)
    e = f["ejecutor"]
    cliente_id = v.get("cliente_id")
    region = str(v.get("cliente_region") or "").strip()
    cl = (f.get("clasif_codigo") or "").strip()
    intencion = (f.get("intencion") or "").strip()
    doc = documento_hoy(f["cobra"], docs_con_principal(docs, v), auts)
    clase = f["clase_neto"]
    r = {
        "id": int(f["id"]), "numero": f["numero"], "titulo": f["titulo"], "estado": f["estado"] or "sin_estado",
        "f_creada": iso(fechas["creacion"]), "f_programada": iso(fechas["programada"]),
        "f_realizada": iso(fechas["ejecucion"]), "f_cerrada": iso(fechas["cierre"]),
        **tiempo,
        "prev": 1 if (elegida and p.get("desde") and elegida < p["desde"]) else 0,
        "cli": str(cliente_id) if cliente_id else SIN["cli"],
        "reg": region.lower() if region else SIN["reg"],
        "tec": tec_k, "ttipo": e["tipo"], "emp": emp_k,
        "cc": f["centro_costo"] or SIN["cc"], "tipo": f["tipo"] or SIN["tipo"], "cob": f["cobertura"],
        "int": intencion.lower() if intencion else SIN["int"], "doc": doc, "nac": cl or SIN["nac"],
        "banda": banda_de(clase, f["pct"], f["queda"]), "clase": clase,
        "cobre_serv": _r2(cob_d.get("servicio")) or 0.0, "cobre_desp": _r2(cob_d.get("despacho")) or 0.0,
        "cobre": _r2(f["cobre"]), "me_tec": _r2(me.get("tecnico")), "me_desp": _r2(me.get("despacho")),
        "me_rep": _r2(me.get("repuestos")) or 0.0, "me_total": _r2(f["me_total"]), "a_pagar": _r2(f["a_pagar"]),
        # Una OT a la que le falta un costo (o que va aparte) NO tiene margen: _ot_finanzas igual calcula un % con lo que
        # hay (un técnico sin declarar da «100 %»), así que aquí se deja vacío. Nunca un 100 % ni un $0 inventado.
        "queda": None if clase in ("sin_medir", "aparte") else _r2(f["queda"]),
        "pct": f["pct"] if clase in ("ganada", "perdida") else None, "val": f["valorizado"],
        "cobra": 1 if f["cobra"] else 0, "falta_tec": 1 if f["falta_tecnico"] else 0,
    }
    r["_lab"] = {
        "estado": ESTADO_TXT.get(r["estado"], r["estado"].replace("_", " ").capitalize()),
        "cli": f["cliente"] if cliente_id else "Trabajo interno (sin cliente)",
        "reg": region or "Sin región",
        "tec": tec_n, "ttipo": G.TIPO_EJECUTOR_TXT.get(e["tipo"], e["tipo"]), "emp": emp_n,
        "cc": f["centro_txt"] if f["centro_costo"] else "Sin centro de costo",
        "tipo": f["tipo_txt"] or "Sin tipo", "cob": G.COBERTURA_TXT.get(f["cobertura"], f["cobertura"]),
        "int": intencion or G.SIN_CLASIFICAR, "doc": DOC_TXT[doc],
        "nac": SALDO.CLASIF_TXT.get(cl, f.get("clasif_texto") or cl) if cl else "Sin clasificación",
        "banda": BANDA_TXT[r["banda"]], "clase": CLASE_TXT.get(clase, clase),
    }
    return r


def actuales(filas):
    """Las OT del rango elegido (sin las del período anterior, que solo sirven para comparar)."""
    return [r for r in filas if not r.get("prev")]


def armar_dataset(filas, p, meta=None):
    """{cols, rows, dims, tec_info, meta}: por columnas (cada fila es una lista en el orden de COLS) y los nombres de cada
    dimensión una sola vez. Las fechas viajan como AAAA-MM-DD (el navegador las muestra dd/mm/aaaa, REGLA #6)."""
    dims = {d: {} for d in DIMS}
    tec_info = {}
    rows = []
    for r in filas:
        for d in DIMS:
            dims[d].setdefault(r[d], r["_lab"][d])
        tec_info.setdefault(r["tec"], {"ttipo": r["ttipo"], "emp": r["emp"]})
        rows.append([r[c] for c in COLS])
    m = dict(meta or {})
    m.update({
        "desde": iso(p["desde"]), "hasta": iso(p["hasta"]), "carga_desde": iso(p["carga_desde"]),
        "previo_hasta": iso(p["previo_hasta"]), "dias": p["dias"], "fecha": p["fecha"], "alcance": p["alcance"],
        "hoy": iso(p.get("hoy")),
        "fecha_txt": FECHA_TXT[p["fecha"]], "fechas_txt": dict(FECHA_TXT), "alcances_txt": dict(ALCANCE_TXT),
        "politica": POL.como_json(), "dims_txt": dict(DIM_TXT), "dims": list(DIMS), "dims_tiempo": list(DIMS_TIEMPO),
        "regaladas": list(COBERTURAS_REGALADAS), "orden": {"doc": list(DOC_ORDEN), "banda": list(BANDA_ORDEN),
                                                           "cob": list(G.COBERTURA_ORDEN) + ["cobra"]},
        "clasificacion": CLASIFICACION, "n": len(rows), "n_actuales": sum(1 for r in filas if not r["prev"]),
        "max_rango_dias": MAX_RANGO_DIAS,
    })
    return {"cols": list(COLS), "rows": rows, "dims": dims, "tec_info": tec_info, "meta": m}


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Espejo del navegador: filtro y suma (las pruebas comparan uno con otro)
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def validar_filtros(raw):
    """{dimensión: [valores]} desde el navegador (texto JSON o dict). Solo dimensiones conocidas; valores como texto
    (a lo más 500 por dimensión y 200 caracteres cada uno). Lo demás se descarta sin reventar."""
    if isinstance(raw, (bytes, str)):
        try:
            raw = json.loads(raw or "{}")
        except (TypeError, ValueError):
            return {}
    if not isinstance(raw, dict):
        return {}
    out = {}
    for d, vs in raw.items():
        if d not in FILTRABLES or not isinstance(vs, (list, tuple)):
            continue
        limpios = []
        for x in vs[:500]:
            if isinstance(x, (str, int, float)) and not isinstance(x, bool):
                s = str(x)[:200]
                if s not in limpios:
                    limpios.append(s)
        if limpios:
            out[d] = limpios
    return out


def leer_ids(raw, tope=5000):
    """«12,15,20» → {12, 15, 20} (solo enteros, a lo más `tope`)."""
    out = set()
    for x in str(raw or "").split(",")[:tope]:
        x = x.strip()
        if x.isdigit():
            out.add(int(x))
    return out


def filtrar(filas, filtros, excepto=()):
    """Las filas que cumplen TODAS las dimensiones filtradas (dentro de una dimensión, cualquiera de sus valores).
    `excepto`: dimensiones que no se aplican (cada visual se calcula con los filtros de las OTRAS dimensiones)."""
    activos = [(d, set(str(x) for x in vs)) for d, vs in (filtros or {}).items() if vs and d not in (excepto or ())]
    if not activos:
        return list(filas)

    def _txt(x):
        return "" if x is None else str(x)       # igual que el navegador: null → ''
    return [r for r in filas if all(_txt(r.get(d)) in vs for d, vs in activos)]


def kpis(filas):
    """La suma de varias OT con la REGLA ÚNICA del proyecto (espejo de `_ot_fin_agregado`, por `clase`):
    interno y contrato APARTE; las que no se pueden medir FUERA (y contadas); garantías y coberturas ADENTRO (es plata
    que se puso sin cobrar) y además informadas aparte. Más lo que sirve al tablero: a_pagar y costo total de TODAS."""
    t = {"n": 0, "cobre": 0.0, "me_cobraron": 0.0, "n_fuera": 0, "cobre_fuera": 0.0, "n_no_cobra": 0,
         "costo_no_cobra": 0.0, "aparte": {}, "n_ot": len(filas), "a_pagar": 0.0, "costo_total": 0.0, "val_reg": 0.0,
         "n_perdidas": 0, "perdido": 0.0, "costo_cubiertas": 0.0}
    for r in filas:
        t["a_pagar"] += r.get("a_pagar") or 0.0
        t["costo_total"] += r.get("me_total") or 0.0
        if r.get("cob") in COBERTURAS_REGALADAS:
            t["val_reg"] += r.get("val") or 0.0
        c = r.get("clase")
        if c == "aparte":
            a = t["aparte"].setdefault(r.get("cob"), {"n": 0, "costo": 0.0})
            a["n"] += 1
            a["costo"] = round(a["costo"] + (r.get("me_total") or 0.0), 2)
            continue
        if c == "sin_medir":
            t["n_fuera"] += 1
            t["cobre_fuera"] += r.get("cobre") or 0.0
            continue
        t["n"] += 1
        t["cobre"] += r.get("cobre") or 0.0
        t["me_cobraron"] += r.get("me_total") or 0.0
        if not r.get("cobra"):
            t["n_no_cobra"] += 1
            t["costo_no_cobra"] += r.get("me_total") or 0.0
            if c == "cubierta":
                t["costo_cubiertas"] += r.get("me_total") or 0.0
        elif c == "perdida":
            t["n_perdidas"] += 1
            t["perdido"] += -(r.get("queda") or 0.0)
    for k in ("cobre", "me_cobraron", "cobre_fuera", "costo_no_cobra", "a_pagar", "costo_total", "val_reg", "perdido",
              "costo_cubiertas"):
        t[k] = round(t[k], 2)
    t["queda"] = round(t["cobre"] - t["me_cobraron"], 2)
    t["pct"] = round(t["queda"] / t["cobre"] * 100, 1) if t["cobre"] > 0 else None
    t["aparte_costo"] = round(sum(a["costo"] for a in t["aparte"].values()), 2)
    t["n_aparte"] = sum(a["n"] for a in t["aparte"].values())
    # Rojo sin % solo si hay algo QUE SE COBRA a pérdida: un grupo de puras garantías no «pierde», cuesta (va en gris).
    t["banda"] = POL.banda(t["pct"], t["queda"] if (t["n"] and t["n_no_cobra"] < t["n"]) else None)
    return t


# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
#  Excel de lo filtrado (solo superadmin)
# ════════════════════════════════════════════════════════════════════════════════════════════════════════════════
def texto_filtros(filtros, dims):
    """«Técnico: Pedro Ejemplo, Marta Demo · Centro de costo: Servicio Técnico» con los nombres (no las claves)."""
    partes = []
    for d in FILTRABLES:
        vs = (filtros or {}).get(d)
        if not vs:
            continue
        if d in DIMS_TIEMPO:
            nombres = [etiqueta_tiempo(d, x) for x in vs]
        else:
            nombres = [(dims.get(d) or {}).get(x, x) for x in vs]
        partes.append(f"{DIM_TXT[d]}: " + ", ".join(nombres))
    return " · ".join(partes) if partes else "Sin filtros: todas las OT del rango"


def excel_bytes(filas, p, filtros=None, dims=None, generado="", truncado=False, seleccion=0):
    """Las OT filtradas en un .xlsx: hoja «OT filtradas» (una fila por OT, fechas REALES de Excel con formato dd/mm/aaaa,
    columnas Año / Mes / N° semana / Día / Día de la semana de la fecha que manda, montos como números) y hoja «Resumen»
    (las cifras con la misma suma de la pantalla). Cada hoja dice «Confidencial · solo gerencia»."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    dims = dims or {}
    NEGRO, GRIS = "0A0A0A", "F3F4F6"
    titulo = (f"Costeo de OT · Explorar · Del {G.fecha_txt(p['desde'])} al {G.fecha_txt(p['hasta'])} · "
              f"{FECHA_TXT[p['fecha']]} · {ALCANCE_TXT[p['alcance']]} · {CLASIFICACION}")
    nota_filtros = "Filtros: " + texto_filtros(filtros, dims)
    if seleccion:
        nota_filtros += f" · Solo las {seleccion} OT seleccionadas"
    wb = openpyxl.Workbook()

    def encabezado(ws, encabezados, fila0):
        for ci, h in enumerate(encabezados, 1):
            c = ws.cell(row=fila0, column=ci, value=h)
            c.font = Font(bold=True, color="FFFFFF", size=10)
            c.fill = PatternFill("solid", fgColor=NEGRO)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # ── Hoja 1: una fila por OT ──
    ws = wb.active
    ws.title = "OT filtradas"
    ws.cell(row=1, column=1, value=titulo).font = Font(bold=True, size=11)
    ws.cell(row=2, column=1, value=nota_filtros).font = Font(italic=True, size=9, color="6B7280")
    aviso = f"Generado el {generado} (hora de Chile)." if generado else ""
    if truncado:
        aviso += f" Se leyeron las {LIMITE_FILAS} OT más recientes: puede haber más en el rango."
    ws.cell(row=3, column=1, value=aviso.strip()).font = Font(italic=True, size=9, color="6B7280")
    cols = [("N° OT", 16, "t"), ("Título", 40, "t"), ("Estado", 16, "t"), ("Cliente", 38, "t"), ("Región", 18, "t"),
            ("Técnico", 28, "t"), ("Interno o externo", 18, "t"), ("Empresa o proveedor", 28, "t"),
            ("Centro de costo", 18, "t"), ("Tipo de OT", 22, "t"), ("Cobertura", 22, "t"), ("Intención", 22, "t"),
            ("Documento hoy", 26, "t"), ("Cómo nació la OT", 30, "t"),
            (FECHA_TXT[p["fecha"]], 14, "f"), ("Año", 8, "n"), ("Mes", 12, "t"), ("N° semana", 10, "n"), ("Día", 7, "n"),
            ("Día de la semana", 13, "t"), ("Creada", 12, "f"), ("Programada", 12, "f"), ("Realizada", 12, "f"),
            ("Cerrada", 12, "f"),
            ("Cobré servicio ($)", 15, "$"), ("Cobré despacho ($)", 15, "$"), ("Cobré ($)", 14, "$"),
            ("Me cobró el técnico ($)", 16, "$"), ("Me cobró el despacho ($)", 16, "$"), ("Repuestos ($)", 13, "$"),
            ("Me cobraron ($)", 15, "$"), ("A pagar al técnico o proveedor ($)", 18, "$"), ("Margen ($)", 14, "$"),
            ("Margen %", 10, "%"), ("Margen según la política", 30, "t"), ("En la cuenta", 30, "t"),
            ("Valorizado ($)", 14, "$")]
    fila0 = 5
    encabezado(ws, [c[0] for c in cols], fila0)
    for ri, r in enumerate(filas, fila0 + 1):
        lab = r.get("_lab") or {}
        d = G._fecha(r.get("dia"))
        valores = [r["numero"], r["titulo"], lab.get("estado"), lab.get("cli"), lab.get("reg"), lab.get("tec"),
                   lab.get("ttipo"), lab.get("emp"), lab.get("cc"), lab.get("tipo"), lab.get("cob"), lab.get("int"),
                   lab.get("doc"), lab.get("nac"),
                   d, (d.year if d else None), (G.MESES_ES[d.month - 1].capitalize() if d else None),
                   (d.isocalendar()[1] if d else None), (d.day if d else None), (DIAS_SEMANA[d.weekday()] if d else None),
                   G._fecha(r.get("f_creada")), G._fecha(r.get("f_programada")), G._fecha(r.get("f_realizada")),
                   G._fecha(r.get("f_cerrada")),
                   r["cobre_serv"], r["cobre_desp"], r["cobre"], r["me_tec"], r["me_desp"], r["me_rep"], r["me_total"],
                   r["a_pagar"], (r["queda"] if r["clase"] not in ("sin_medir", "aparte") else None), r["pct"],
                   lab.get("banda"), lab.get("clase"), r["val"]]
        for ci, (val, (_h, _w, tipo)) in enumerate(zip(valores, cols), 1):
            c = ws.cell(row=ri, column=ci, value=val)
            if tipo == "f" and val is not None:
                c.number_format = "dd/mm/yyyy"
            elif tipo == "$" and val is not None:
                c.number_format = '"$"#,##0'
            elif tipo == "%" and val is not None:
                c.number_format = "0.0"
            if (ri - fila0) % 2 == 0:
                c.fill = PatternFill("solid", fgColor=GRIS)
    for ci, (_h, w, _t) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = ws.cell(row=fila0 + 1, column=2)
    if filas:
        ws.auto_filter.ref = f"A{fila0}:{get_column_letter(len(cols))}{fila0 + len(filas)}"

    # ── Hoja 2: las cifras ──
    k = kpis(filas)
    rs = wb.create_sheet("Resumen")
    rs.cell(row=1, column=1, value=titulo).font = Font(bold=True, size=11)
    rs.cell(row=2, column=1, value=nota_filtros).font = Font(italic=True, size=9, color="6B7280")
    encabezado(rs, ["Cifra", "Valor", "Cómo se calcula"], 4)
    lineas = [
        ("Cobré ($)", k["cobre"], "Lo cobrado al cliente (servicio + despacho) en las OT medidas. Misma suma que la ficha de cada OT."),
        ("Me cobraron ($)", k["me_cobraron"], "Técnico, despacho y repuestos de las OT medidas (garantías y coberturas incluidas)."),
        ("Margen ($)", k["queda"], "Cobré − Me cobraron."),
        ("Margen %", k["pct"], "Margen sobre lo cobrado. Política: objetivo 30 %, tolerancia desde 20 %, bajo 20 % amarillo, bajo 0 % rojo."),
        ("Margen según la política", POL.texto(k["banda"]), ""),
        ("N° de OT", k["n_ot"], f"{k['n']} medidas · {k['n_fuera']} sin medir · {k['n_aparte']} aparte (interno o contrato)."),
        ("OT sin medir", k["n_fuera"], f"Les falta un costo: no están en las cifras. Cobrado en ellas: {G.clp(k['cobre_fuera'])}."),
        ("Garantías y coberturas: costo ($)", k["costo_no_cobra"], f"{k['n_no_cobra']} OT que no se cobran; ya restan en el margen."),
        ("Trabajo interno y contratos: costo ($)", k["aparte_costo"], "Van aparte: no son negocio con nadie."),
        ("Valor del servicio regalado ($)", k["val_reg"], "Valorizado de garantías, regalías, arriendo o leasing y cortesías (referencia)."),
        ("Nos cobraron técnicos y proveedores ($)", k["a_pagar"], "Técnico + despacho de TODAS las OT filtradas (sin repuestos)."),
    ]
    for ri, (a, b, c_) in enumerate(lineas, 5):
        rs.cell(row=ri, column=1, value=a).font = Font(bold=True)
        cel = rs.cell(row=ri, column=2, value=b)
        if isinstance(b, (int, float)) and not isinstance(b, bool):
            cel.number_format = "0.0" if "%" in a and "$" not in a else ('"$"#,##0' if "$" in a else "#,##0")
        rs.cell(row=ri, column=3, value=c_).alignment = Alignment(wrap_text=True, vertical="top")
    rs.column_dimensions["A"].width = 40
    rs.column_dimensions["B"].width = 18
    rs.column_dimensions["C"].width = 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
